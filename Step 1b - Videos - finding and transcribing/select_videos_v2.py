"""select_videos_v2.py: choose which YouTube videos to transcribe.
Step 1 (videos). Version 2, 15 Oct 2026.

Reads harrison_v7_VIDEO_LINKS.xlsx (video_links_v2.py 2.1, with V IDs) and writes harrison_v7_VIDEO_SELECTION.xlsx,
and beside it the download lists for get_audio.sh: video_audio_list_all.txt (every chosen video, priority 1 first)
and video_audio_list_pilot.txt (the videos of one act, PILOT_ACT, to try the download on first).
The rule:
  1. Window: no limit before the nomination; after it, anything up to 1999.
  2. Thin acts: at least 10 videos each.
  3. Dating is the weak point, so the chosen rows were checked by hand. The decisions are in checks.csv, keyed by
     the V IDs of the links file they were made on.
  4. Interview and Q&A heavy, not banter: chat-show and music-show slots are left out unless an act needs them.
Usage: python select_videos_v2.py [links.xlsx] [out.xlsx]
  By default both files are in HARRISON_DIR (in Colab, /content/drive/MyDrive/harrison_v7). No key and no network
  are needed. select_videos.py (version 1, for the speaker check and the shared helpers), video_links_v2.py and
  checks.csv must sit beside this script.
"""
import csv, importlib.util, os, re, sys
from datetime import date
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
DIR = os.environ.get("HARRISON_DIR", "/content/drive/MyDrive/harrison_v7")
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DIR, "harrison_v7_VIDEO_LINKS.xlsx")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(DIR, "harrison_v7_VIDEO_SELECTION.xlsx")
CHECKS = os.path.join(HERE, "checks.csv")
LAST_YEAR = 1999      # latest year kept
MIN_PER_ACT = 10      # thin acts are topped up to this many videos
NEAR = 3              # priority 1 = within this many years of the nomination
PILOT_ACT = "Big Country"   # the act whose videos make the pilot download list
MIN_ONE, MAX_ONE = 1.0, 120.0
SPOT_CHECK = ("A blind check of 60 chosen videos (25, 20 and 15 by priority) by a separate AI agent, reading only "
              "the titles and descriptions, found 2 problems (a chat-show slot, and a later interview dated a year "
              "late) and 3 doubtful rows; all five are now fixed or confirmed.")

_sv1 = os.path.join(HERE, "select_videos.py")
_spec = importlib.util.spec_from_file_location("sv1", _sv1)
sv1 = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(sv1)
vl = sv1.vl

TALK = ["Interview", "Q&A"]
# Slots on chat shows and music or variety shows: mostly host banter, or a quick chat between songs.
BANTER = re.compile(
    r"American Bandstand|\bBandstand\b|Soul Train|Solid Gold|Top of the Pops|\bTOTP\b|\bTop ?Pop\b|Tiswas|Razzmatazz|"
    r"Dance Party USA|Saturday Night Live|\bSNL\b|Saturday Superstore|Swap Shop|Wide Awake Club|Going Live|Motormouth|"
    r"Big Breakfast|Blue Peter|\bWogan\b|Letterman|Johnny Carson|Tonight Show|\bCarson\b|Arsenio|Oprah|Donahue|"
    r"Joan Rivers|Aspel|Russell Harty|Des O['’]Connor|Pebble Mill|Howard Stern|Stern Show|Midday Show|Ray Martin|"
    r"Mike Walsh|\bRove\b|Conan O['’]Brien|Team Coco|Graham Norton|Jonathan Ross|Loose Women|Richard (?:&|and) Judy|"
    r"Kelly Clarkson Show|Tamron Hall|Drew Barrymore|Jay Leno|Jimmy Kimmel|Jimmy Fallon|Rosie O['’]Donnell|"
    r"\bRosie Show\b|\bRegis\b|Parkinson|Merv Griffin|Byron Allen|Dini Petty|Jane Hawtin|Sally (?:Jessy )?Raphael|"
    r"Jennifer Hudson Show|Diana Dors", re.I)
BANTER_CS = re.compile(r"\bLate Show\b|\bLate Night\b|\bCONAN\b|\bThe View\b|\bEllen\b|\bMerv\b|\bThe Word\b")
# Two-digit years in titles: '3-15-89', "'85", "91'"; not decades ("'80s", "80's").
TWO_DIGIT = re.compile(r"(?<!\d)\d{1,2}[-/.]\d{1,2}[-/.](\d{2})(?!\d)|['’‘](\d{2})(?![\ds])|(?<![\d'’‘])(\d{2})['’‘](?![\ds])")
DECISIONS = {"keep", "redate", "undated", "banter", "drop", "add", "include"}


def year_guess(r):
    """Best-guess year and where it came from (title, description, two-digit year in the title, or upload)."""
    y, src = r["Year (best guess)"], str(r["Year from"])
    if src.startswith("upload"):
        m = TWO_DIGIT.search(str(r["Title"]))
        if m:
            yy = int(next(g for g in m.groups() if g))
            up = int(str(r["Upload date"])[:4]) if str(r["Upload date"])[:4].isdigit() else 2100
            if 70 <= yy <= 99 and 1900 + yy <= up:
                return 1900 + yy, "two-digit year in title"
    return y, src


def load_checks():
    """ID -> (decision, year, reason) from checks.csv, the log of rows checked by hand.
    keep: checked, fine. redate: use Year. undated: the year found is not the recording date.
    banter: a chat-show or music-show slot, or a performance with a chat. drop: leave out (reason given).
    add: use although it is not typed as an interview (an era documentary for a thin act); counts like banter.
    include: an interview the automatic checks missed (for example the act is named after 'INTERVIEW:')."""
    out = {}
    if os.path.exists(CHECKS):
        with open(CHECKS, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                dec = r["Decision"].strip()
                assert dec in DECISIONS, (r["ID"], dec)
                y = int(r["Year"]) if r.get("Year", "").strip() else None
                out[r["ID"].strip()] = (dec, y, r["Reason"].strip())
    return out


def eligible(r):
    """'' if the row is a usable interview or Q&A with the act talking, else the reason it is not."""
    ev = str(r["Type evidence"] or "")
    lang = str(r["Language"]) if pd.notna(r["Language"]) else ""
    if r["Type"] not in TALK:
        return f"not an interview or Q&A ({r['Type']})"
    if not ev.startswith("title"):
        return "typed from a show name or the description only"
    if lang[:2] not in ("", "en", "zx", "na"):
        return f"not in English ({lang})"
    if r["Minutes"] < MIN_ONE:
        return "under a minute"
    if r["Minutes"] > MAX_ONE:
        return f"over {MAX_ONE:.0f} minutes"
    if pd.notna(r["Other-era members named"]) and str(r["Other-era members named"]).strip() \
            and not (pd.notna(r["Members named"]) and str(r["Members named"]).strip()):
        return "only an other-era member is named"
    if sv1.TRIBUTE.search(str(r["Title"])):
        return "about the act's legacy or a death, not the act talking"
    ok, why = sv1.speaker_ok(r["Act"], str(r["Title"]), ev)
    return "" if ok else why


def main():
    L = pd.read_excel(SRC, "Links")
    checks = load_checks()
    unknown = set(checks) - set(L["ID"].astype(str))
    assert not unknown, f"checks.csv has IDs not in the links file: {sorted(unknown)[:5]}"
    rows = []
    for _, r in L.iterrows():
        y, ysrc = year_guess(r)
        reason = eligible(r)
        tc = f"{r['Title']}\n{r['Channel']}"
        banter = bool(BANTER.search(tc) or BANTER_CS.search(tc))
        note = ""
        chk = checks.get(str(r["ID"]))
        if chk:
            dec, cy, why = chk
            note = f"{dec}: {why}" if why else dec
            if cy:
                y, ysrc = cy, "checked by hand"
            if dec == "undated":
                up = str(r["Upload date"])[:4]
                y, ysrc = (int(up) if up.isdigit() else None), "upload year (the year found is not the recording date)"
            elif dec == "banter":
                banter = True
            elif dec == "add":
                reason, banter = "", True
            elif dec == "include":
                reason = ""
            elif dec == "drop":
                reason = f"dropped on checking: {why}"
        dated = not str(ysrc).startswith("upload") and pd.notna(y)
        rows.append({**r.to_dict(), "Year (best guess)": y, "Year from": ysrc,
                     "Years from nomination": (y - r["Grammy year"]) if pd.notna(y) else None,
                     "_year": y, "_dated": dated, "_banter": banter, "_reason": reason, "Checked by hand": note})
    D = pd.DataFrame(rows)
    D["_rank"] = D["Best rank"].fillna(999)
    D["_era"] = D["_dated"] & (D["_year"] <= LAST_YEAR)
    D["_ok"] = D["_reason"].eq("")

    # Selection: up to 1999 and not banter; then top thin acts up to MIN_PER_ACT.
    D["Priority"] = pd.NA
    D["Why chosen"] = ""
    core = D["_ok"] & D["_era"] & ~D["_banter"]
    near = core & ((D["_year"] - D["Grammy year"]).abs() <= NEAR)
    D.loc[near, "Priority"] = 1
    D.loc[near, "Why chosen"] = f"up to {LAST_YEAR}, within {NEAR} years of the nomination"
    D.loc[core & ~near, "Priority"] = 2
    D.loc[core & ~near, "Why chosen"] = f"up to {LAST_YEAR}"
    D["_eff"] = D["_year"].fillna(2100)   # undated rows carry their upload year, the latest they can date from
    for act in vl.ACTS:
        have = int((core & D["Act"].eq(act)).sum())
        if have >= MIN_PER_ACT:
            continue
        pools = [
            (D["_ok"] & D["_era"] & D["_banter"], 2,
             f"act has under {MIN_PER_ACT}: chat-show or music-show slot, or performance with a chat, up to {LAST_YEAR}"),
            (D["_ok"] & ~D["_era"] & ~D["_banter"], 3, f"act has under {MIN_PER_ACT}: earliest later or undated interview"),
            (D["_ok"] & ~D["_era"] & D["_banter"], 3, f"act has under {MIN_PER_ACT}: later chat-show slot"),
        ]
        for mask, pri, why in pools:
            P = D[mask & D["Act"].eq(act) & D["Priority"].isna()]
            if pri == 2:
                P = P.assign(_d=(P["_year"] - P["Grammy year"]).abs()).sort_values(["_d", "_rank"])
            else:
                P = P.sort_values(["_eff", "_rank"])
            for i in P.index:
                if have >= MIN_PER_ACT:
                    break
                D.at[i, "Priority"], D.at[i, "Why chosen"] = pri, why
                have += 1
    D["Selected"] = D["Priority"].notna()
    D["Why not"] = ""
    for i, r in D[~D["Selected"]].iterrows():
        if r["_reason"]:
            D.at[i, "Why not"] = r["_reason"]
        elif r["_era"] and r["_banter"]:
            D.at[i, "Why not"] = "chat-show or music-show slot (banter)"
        elif not r["_era"]:
            D.at[i, "Why not"] = "after 1999" if r["_dated"] else "undated (upload year only)"
        else:
            D.at[i, "Why not"] = "not needed"
    return D


COLS = ["ID", "Act", "Grammy year", "Priority", "Why chosen", "Type", "Title", "Channel", "URL", "Minutes",
        "Year (best guess)", "Year from", "Years from nomination", "Checked by hand", "Show or event", "Members named",
        "Upload date", "Captions by uploader", "Search source", "Best rank", "Notes", "YouTube ID"]
OCOLS = ["ID", "Act", "Grammy year", "Why not", "Type", "Title", "Channel", "URL", "Minutes", "Year (best guess)",
         "Year from", "Years from nomination", "Checked by hand", "Show or event", "Members named", "Upload date",
         "Search source", "Best rank", "Notes", "YouTube ID"]


def write_book(D, path):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    order = {a: k for k, a in enumerate(vl.ACTS)}
    D = D.copy()
    D["_o"] = D["Act"].map(order)
    S = D[D["Selected"]].sort_values(["_o", "Priority", "_eff", "_rank"])
    # One row per video: a clip chosen for two acts is listed once, with both acts.
    acts_of = S.groupby("ID")["Act"].apply(lambda s: "; ".join(dict.fromkeys(s)))
    T = S.drop_duplicates("ID").copy()
    T["Act"] = T["ID"].map(acts_of)
    R = D[~D["Selected"]].sort_values(["_o", "Why not", "_rank"])
    by = []
    for a in vl.ACTS:
        s_ = S[S["Act"] == a]
        by.append({"Act": a, "Grammy year": vl.ACTS[a]["year"], "Videos": len(s_),
                   "Minutes": round(s_["Minutes"].sum(), 1),
                   f"Up to 1999, within {NEAR} years (videos)": int((s_["Priority"] == 1).sum()),
                   "Up to 1999, other (videos)": int(((s_["Priority"] == 2) & ~s_["_banter"]).sum()),
                   "Banter slots added (videos)": int(((s_["Priority"] == 2) & s_["_banter"]).sum()),
                   "Later or undated added (videos)": int((s_["Priority"] == 3).sum()),
                   "Short of 10": max(0, MIN_PER_ACT - len(s_))})
    B = pd.DataFrame(by)
    hrs = lambda m: T[m]["Minutes"].sum() / 60
    short = B[B["Short of 10"] > 0]
    checks = D[D["Checked by hand"].ne("")]
    readme = [
        ["Video selection", f"Version 2, 15 October 2026, from harrison_v7_VIDEO_LINKS.xlsx (V IDs). {len(T)} videos, "
                            f"{T['Minutes'].sum() / 60:.1f} hours, all {len(vl.ACTS)} acts. Replaces version 1, which "
                            f"capped the minutes per act."],
        ["Rule", f"Interviews and Q&As in which the act or a nomination-era member is talking, in English, "
                 f"{MIN_ONE:.0f} to {MAX_ONE:.0f} minutes long, recorded in {LAST_YEAR} or earlier (no earlier limit). "
                 f"A video counts as dated only when its title or description gives the year, or it comes from a show "
                 f"that ended before 2000; the upload year alone does not date it."],
        ["Banter", "Chat-show and children's-show slots (Tonight Show, Letterman, Arsenio, Wogan, Swap Shop and the "
                   "like), quick chats on music shows (American Bandstand, Soul Train) and performances with a short "
                   f"chat are left out, unless an act needs them to reach {MIN_PER_ACT}."],
        ["Thin acts", f"Acts with fewer than {MIN_PER_ACT} are topped up to {MIN_PER_ACT}: first with their own banter "
                      f"slots up to {LAST_YEAR}, then with their earliest later or undated interviews (an undated video "
                      f"counts as its upload year, the latest it can date from). "
                      + ("Still short, because little else exists: "
                         + ", ".join(f"{a} ({v})" for a, v in zip(short['Act'], short['Videos'])) + "."
                         if len(short) else "")],
        ["Hand checks", f"Every chosen title was read, and the description wherever the date or the speaker was in "
                        f"doubt. {checks['ID'].nunique()} rows were changed or confirmed by hand (Checks sheet, and the 'Checked "
                        f"by hand' column): wrong dates (an album or song year, a later look back, a year in a "
                        f"biography note), dates found from the description or the show, re-uploads of the same clip, "
                        f"someone else talking, reaction videos, and banter. " + SPOT_CHECK],
        ["Priority", f"1 = up to {LAST_YEAR} and within {NEAR} years of the nomination ({hrs(T['Priority'] == 1):.1f} h); "
                     f"2 = other videos up to {LAST_YEAR}, and banter slots for thin acts "
                     f"({hrs(T['Priority'] == 2):.1f} h); 3 = later or undated interviews for thin acts "
                     f"({hrs(T['Priority'] == 3):.1f} h). Transcribe in this order."],
        ["To change it", "The transcription step reads the To transcribe sheet. Delete a row to drop a video; copy a "
                         "row in from Other candidates to add one. Re-uploads that slipped through will show up when "
                         "the transcripts are compared."],
    ]
    head = PatternFill("solid", fgColor="1F4E5F")
    widths = {"ID": 8, "Act": 18, "Grammy year": 8, "Priority": 8, "Why chosen": 30, "Why not": 30, "Type": 11,
              "Title": 60, "Channel": 22, "URL": 34, "Minutes": 8, "Year (best guess)": 9, "Year from": 20,
              "Years from nomination": 10, "Checked by hand": 30, "Show or event": 16, "Members named": 22,
              "Upload date": 11, "Captions by uploader": 10, "Search source": 12, "Best rank": 7, "Notes": 30,
              "YouTube ID": 13, "Decision": 10, "Reason": 70}
    DC = checks.drop_duplicates("ID")[["ID", "Act", "Title", "Checked by hand", "Year (best guess)", "Selected"]].copy()
    DC["Selected"] = DC["Selected"].map({True: "yes", False: "no"})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(readme).to_excel(xw, sheet_name="Read me", index=False, header=False)
        T[COLS].to_excel(xw, sheet_name="To transcribe", index=False)
        B.to_excel(xw, sheet_name="By act", index=False)
        R[OCOLS].to_excel(xw, sheet_name="Other candidates", index=False)
        DC.to_excel(xw, sheet_name="Checks", index=False)
        for name, ws in xw.sheets.items():
            if name == "Read me":
                ws.column_dimensions["A"].width = 22
                ws.column_dimensions["B"].width = 120
                for row in ws.iter_rows():
                    for c in row:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                    row[0].font = Font(bold=True)
                continue
            for j, c in enumerate(ws[1], start=1):
                c.fill, c.font = head, Font(bold=True, color="FFFFFF")
                c.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(j)].width = widths.get(c.value, 14)
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions
            heads = [c.value for c in ws[1]]
            col = {h: heads.index(h) + 1 for h in heads}
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top", wrap_text=heads[c.column - 1] in
                                            ("Title", "Notes", "Why chosen", "Why not", "Checked by hand"))
                if "URL" in col:
                    u = row[col["URL"] - 1]
                    if isinstance(u.value, str) and u.value.startswith("http"):
                        u.hyperlink, u.font = u.value, Font(color="1155CC", underline="single")
                if "Upload date" in col:
                    row[col["Upload date"] - 1].number_format = "yyyy-mm-dd"
                if "Minutes" in col:
                    row[col["Minutes"] - 1].number_format = "0.0"
    return T, B


def write_audio_lists(T, out_dir):
    """The download lists for get_audio.sh, from the To transcribe rows: every chosen video, priority 1 first (in
    the order of the sheet), and a pilot list of one act's videos. Each link follows a note line (#), which yt-dlp
    ignores."""
    made = date.today()
    made = f"{made.day} {made:%B %Y}"
    T = T.sort_values("Priority", kind="stable")
    pilot = T[T["Act"].astype(str).str.split("; ").map(lambda acts: PILOT_ACT in acts)]
    paths = []
    for which, rows, title in (("all", T, "All chosen videos, priority 1 first"),
                               ("pilot", pilot, f"Pilot: {PILOT_ACT}")):
        lines = [f"# {title}",
                 "# One YouTube link per line; lines starting with # are notes for people, ignored by yt-dlp.",
                 f"# {len(rows)} videos, {rows['Minutes'].sum() / 60:.1f} hours of audio. Made {made} from "
                 f"{os.path.basename(OUT)} (To transcribe)."]
        for _, r in rows.iterrows():
            lines.append(f"# {r['ID']} | {r['Act']} | priority {int(r['Priority'])} | {r['Minutes']:.1f} min")
            lines.append(str(r["URL"]))
        path = os.path.join(out_dir, f"video_audio_list_{which}.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
        paths.append(path)
    return paths


if __name__ == "__main__":
    D = main()
    T, B = write_book(D, OUT)
    lists = write_audio_lists(T, os.path.dirname(os.path.abspath(OUT)))
    print(f"{len(T)} videos, {T['Minutes'].sum() / 60:.1f} h")
    print(T.groupby("Priority")["Minutes"].agg(["count", "sum"]).assign(h=lambda x: (x["sum"] / 60).round(1)))
    print("short of 10:", B[B["Short of 10"] > 0][["Act", "Videos"]].values.tolist())
    print("Wrote:\n  " + "\n  ".join([OUT] + lists))
