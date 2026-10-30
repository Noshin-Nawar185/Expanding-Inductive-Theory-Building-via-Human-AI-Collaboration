"""merge_video_turns.py: join the two transcripts of each video into speaker turns.
Step 2 (videos). Version 1.3, 18 October 2026.

Joins the two transcripts of each video into one list of speaker turns:
  WhisperX (transcribe_videos.py)   the words and their times, and its own grouping of the voices
  Gemini (gemini_speakers.py)       the turns: who speaks (label, role, a name only on evidence) and the language

How a video is joined:
1. Gemini's words are lined up with WhisperX's words (the same words in the same order).
2. Each Gemini turn takes the WhisperX words lined up with it, and their times. Short stretches of WhisperX words
   that Gemini does not have go to the nearer turn; "thank you"-type lines WhisperX made up over silence or music,
   lines where it repeated its prompt (the names list) and transcription loops are left out unless Gemini has them.
3. Gemini's own text is used instead when the turn is not in English, or when WhisperX's words for it are mostly
   different (an interpreter talking over a band member: WhisperX takes down the interpreter, in English).
   Such a turn's times come from the nearest lined-up words, because Gemini's own clock can run fast.
4. Where Gemini has nothing (it refused a stretch, or left it out), WhisperX's words are kept, with the speaker
   taken from WhisperX's voices where a voice clearly belongs to one of Gemini's labels.
5. Videos sent to Gemini in parts: when a label clearly moves to another person's voice in a later part (a swap),
   it is put back, trusting the part where that person was named. Labels that still move are marked "check".
6. Names: kept only on evidence, that is said in the video (and found in the words) or shown in an on-screen
   caption, or the same person as someone so named. A name from the video's title or the member list alone is not
   kept; that person gets a plain label ("band member, unspecified (B)", "artist, unconfirmed (A)", "unknown (C)").
7. Band voice: a turn by a band member (bands, duos, groups) or the artist (solo artists); song lyrics are not.

Input, in the folder this script is in (Drive: MyDrive/harrison_v7/Video for Harrison):
  gemini_speakers/              Gemini's replies (gemini_speakers.py)
  video_transcripts/            the WhisperX transcripts (transcribe_videos.py)
  harrison_v7_VIDEO_SELECTION.xlsx, acts_frame.json (the acts' types, line-ups and name lists)
Output, in the same folder:
  harrison_v7_VIDEO_TURNS.xlsx  sheets Read me, Turns, Videos, People, Label repairs
  harrison_v7_VIDEO_TURNS.json  the Turns sheet in full ('columns', 'data'), for build_video_oral_histories.py
  video_turns/<ID>.json         the joined turns, the people, and what was left out or set aside, per video

In Colab (any runtime; no GPU and no key needed), after Drive is mounted:
  import os; B = "/content/drive/MyDrive/harrison_v7/Video for Harrison"
  print(len(os.listdir(B + "/gemini_speakers")), len(os.listdir(B + "/video_transcripts")))
  %run "/content/drive/MyDrive/harrison_v7/Video for Harrison/merge_video_turns.py"
The listing first makes Drive show the whole folders. It takes a few minutes and can be run again at any time
(after more videos are sent to Gemini, say): it rebuilds everything from the two folders.
Options: --gemini, --whisperx, --selection, --frame, --out (other places), --ids V1158,V1319 (only these),
--no-xlsx (the per-video files only).
"""
import argparse, collections, json, math, os, re, sys, unicodedata
import difflib

VERSION = "merge_video_turns 1.3 (18 Oct 2026)"

SHORT_RUN_WORDS = 20        # unlined-up WhisperX words fewer than this, and ...
SHORT_RUN_SECONDS = 15.0    # ... shorter than this, go to the nearer Gemini turn
MATCH_SHARE = 0.5           # a turn keeps WhisperX's words when at least this share of its own words line up
ISOLATED = 4                # a lined-up stretch of 1-2 words with this many unlined words on both sides is chance
VOICE_MIN_WORDS = 30        # a WhisperX voice is given a Gemini label when it has this many lined-up words ...
VOICE_PURITY = 0.7          # ... and this share of them carry that label
REPAIR_MIN_WORDS = 40       # label repair across parts: words of the label in the part ...
REPAIR_PURITY = 0.75        # ... and the share of them in one voice
NAME_OK = {"said in the video", "on-screen caption", "same person as one named earlier"}
ROLES = ["band member", "artist", "interviewer", "narrator", "lyrics", "other"]
VOICE_ROLES = {"band member", "artist"}
HALLU = re.compile(r"^(thank you( (so much|very much|for watching))?|thanks( for watching)?|you|bye|"
                   r"please subscribe.*|subtitles? .*|.*amara\.org.*|transcription by.*|"
                   r"(music|applause|laughter)|i'll see you (in the )?next (time|video).*)[.!?]*$", re.I)
TIME_RE = re.compile(r"\b(\d{1,2}):(\d{2})(?::(\d{2}))?\b")
NAME_STOP = {"the", "and", "of", "mr", "mrs", "ms", "dr", "jr", "sr", "de", "la", "le", "van", "von", "der", "di", "da", "del"}


# ------------------------------------------------------------------ small helpers
def norm(w):
    return re.sub(r"[^a-z0-9']", "", fold(w).lower())


def fold(s):
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()


def secs(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    parts = str(x).strip().split(":")
    try:
        nums = [float(p) for p in parts]
    except ValueError:
        return None
    total = 0.0
    for n in nums:
        total = total * 60 + n
    return total


def mmss(s):
    s = int(round(s or 0))
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60}:{s % 60:02d}"


def parse_json(text):
    """The JSON object in a reply (as in gemini_speakers.py), and, failing that, with two common slips mended:
    an escaped apostrophe (\\') and a comma before a closing bracket."""
    out = parse_json_strict(text)
    if out is None and text:
        mended = re.sub(r",(\s*[}\]])", r"\1", text.replace("\\'", "'"))
        if mended != text:
            out = parse_json_strict(mended)
    return out


def parse_json_strict(text):
    if not text:
        return None
    t = text.strip()
    m = re.search(r"```(?:json)?\s*(.*?)```", t, re.S)
    if m:
        t = m.group(1).strip()
    starts = [i for i in (t.find("{"), t.find("[")) if i >= 0]
    if not starts:
        return None
    t = t[min(starts):]
    try:
        return json.loads(t)
    except ValueError:
        pass
    tried = 0
    for end in range(len(t), 0, -1):
        if t[end - 1] in "}]":
            tried += 1
            if tried > 200:
                break
            try:
                return json.loads(t[:end])
            except ValueError:
                continue
    return None


def clean(x):
    return "" if x is None else str(x).strip()


def role_of(r, kind):
    r = clean(r).lower()
    if r in ROLES:
        return r
    if "band" in r or "member" in r:
        return "band member" if kind == "band" else "other"
    if "artist" in r:
        return "artist" if kind == "solo" else "other"
    if "interview" in r or "host" in r:
        return "interviewer"
    if "narrat" in r or "announc" in r:
        return "narrator"
    if "lyric" in r or "sing" in r or "song" in r:
        return "lyrics"
    return "other"


# ------------------------------------------------------------------ WhisperX
def whisperx_words(wx):
    """Words in time order, each with its voice, whether sung, and why it may be left out (junk)."""
    words = [w for seg in wx["segments"] for w in seg.get("words", []) if clean(w.get("word"))]
    # untimed words take the times of their neighbours
    for k, w in enumerate(words):
        if w.get("start") is None or w.get("end") is None:
            prev = next((words[j] for j in range(k - 1, -1, -1) if words[j].get("end") is not None), None)
            nxt = next((words[j] for j in range(k + 1, len(words)) if words[j].get("start") is not None), None)
            a = prev["end"] if prev else (nxt["start"] if nxt else 0.0)
            b = nxt["start"] if nxt else a
            w["start"], w["end"] = a, max(a, b)
    order = sorted(range(len(words)), key=lambda k: words[k]["start"])
    W = []
    for k in order:
        w = words[k]
        W.append({"w": clean(w["word"]), "n": norm(w["word"]), "s": float(w["start"]), "e": float(w["end"]),
                  "spk": w.get("speaker") or "?", "sung": bool(w.get("sung")), "junk": None, "lowconf": False})
    W = [w for w in W if w["n"]]
    # line flags (lines are runs of one voice): made-up lines, prompt echoes, loops, low confidence
    by_spk = collections.defaultdict(list)
    for idx, w in enumerate(W):
        by_spk[w["spk"]].append(idx)
    for ln in wx.get("lines", []):
        flags = set(ln.get("flags") or [])
        hallu = bool(HALLU.match(norm_line(ln.get("text", ""))))
        if not (flags & {"prompt echo", "repetitive", "low confidence"} or hallu):
            continue
        for idx in by_spk.get(ln.get("speaker"), []):
            if ln["start"] - 0.02 <= W[idx]["s"] <= ln["end"] + 0.02:
                if hallu:
                    W[idx]["junk"] = "made-up line"
                elif "prompt echo" in flags:
                    W[idx]["junk"] = "prompt echo"
                elif "repetitive" in flags:
                    W[idx]["junk"] = "transcription loop"
                if "low confidence" in flags:
                    W[idx]["lowconf"] = True
    return W


def norm_line(t):
    return re.sub(r"[^a-z0-9' ]", "", fold(t).lower()).strip()


# ------------------------------------------------------------------ Gemini
def gemini_turns(doc, kind):
    """Gemini's turns with the part each came from (read back from the stored replies)."""
    plan = doc.get("plan") or [{"start": 0.0, "end": (doc.get("minutes") or 0) * 60}]
    keys = collections.defaultdict(list)
    for p in doc.get("parts", []):
        if "prompt" not in p or not p.get("ok"):
            continue
        k = next((i for i, q in enumerate(plan) if abs(q["start"] - p["start"]) < 0.05 and abs(q["end"] - p["end"]) < 0.05), None)
        parsed = parse_json(p.get("raw"))
        rows = (parsed.get("turns") if isinstance(parsed, dict) else parsed if isinstance(parsed, list) else []) or []
        shift = p.get("shift") or 0.0
        for t in rows:
            if isinstance(t, dict):
                s = secs(t.get("start"))
                keys[(clean(t.get("speaker")) or "unknown", clean(t.get("text")),
                      None if s is None else round(s + shift, 1))].append(k)

    def part_by_time(s):
        for i, q in enumerate(plan):
            if s is not None and q["start"] - 0.01 <= s <= q["end"] + 0.01:
                return i
        return len(plan) - 1
    raw_turns = [(t, None) for t in doc.get("turns", [])]
    extra_speakers = []
    ok_ranges = {(round(p["start"], 1), round(p["end"], 1)) for p in doc.get("parts", []) if "prompt" in p and p.get("ok")}
    for p in doc.get("parts", []):
        rng = (round(p["start"], 1), round(p["end"], 1))
        if "prompt" not in p or p.get("ok") or not p.get("raw") or rng in ok_ranges:
            continue
        k = next((i for i, q in enumerate(plan) if abs(q["start"] - p["start"]) < 0.05 and abs(q["end"] - p["end"]) < 0.05), None)
        if k is None:
            continue                        # it was split in two and sent again
        parsed = parse_json(p["raw"])
        if not isinstance(parsed, dict):
            continue
        rows = [t for t in parsed.get("turns") or [] if isinstance(t, dict)]
        starts = [secs(t.get("start")) for t in rows if secs(t.get("start")) is not None]
        shift = 0.0
        if p.get("clip") and p["start"] >= 60 and starts and max(starts) <= p["end"] - p["start"] + 60 and min(starts) < p["start"] - 60:
            shift = p["start"]              # times from the start of the part, as gemini_speakers.py would move them
        for t in rows:
            t2 = dict(t)
            for key in ("start", "end"):
                v = secs(t.get(key))
                t2[key] = None if v is None else v + shift
            raw_turns.append((t2, k))
        extra_speakers += [sp for sp in parsed.get("speakers") or [] if isinstance(sp, dict)]
    turns = []
    for idx, (t, given_part) in enumerate(raw_turns):
        s, e = secs(t.get("start")), secs(t.get("end"))
        key = (clean(t.get("speaker")) or "unknown", clean(t.get("text")), None if s is None else round(s, 1))
        part = given_part if given_part is not None else (keys[key].pop(0) if keys.get(key) else None)
        if part is None:
            part = part_by_time(s)
        if s is not None and e is not None and e < s:
            s, e = e, s
        turns.append({"k": idx, "sg": s, "eg": e if e is not None else s, "label": clean(t.get("speaker")) or "unknown",
                      "role": role_of(t.get("role"), kind), "lang": (clean(t.get("language")) or "en").lower(),
                      "text": clean(t.get("text")), "on_screen": clean(t.get("on_screen")), "part": part,
                      "recovered": given_part is not None})
    return turns, plan, extra_speakers


# ------------------------------------------------------------------ lining up
def line_up(W, GW):
    """pos[i] = index in GW of the Gemini word lined up with WhisperX word i, or None."""
    sm = difflib.SequenceMatcher(a=[w["n"] for w in W], b=[g[0] for g in GW], autojunk=False)
    blocks = [b for b in sm.get_matching_blocks() if b.size]
    keep = []
    for x, b in enumerate(blocks):        # drop 1-2 word matches standing alone in a sea of unlined words
        if b.size <= 2:
            gap_before = b.a - (blocks[x - 1].a + blocks[x - 1].size) if x else b.a
            gap_after = (blocks[x + 1].a if x + 1 < len(blocks) else len(W)) - (b.a + b.size)
            if gap_before >= ISOLATED and gap_after >= ISOLATED:
                continue
        keep.append(b)
    pos = [None] * len(W)
    for b in keep:
        for k in range(b.size):
            pos[b.a + k] = b.b + k
    return pos


def time_map(turns, W, pos, GW, audio_s):
    """Gemini time -> WhisperX time, from turns that line up well (Gemini's clock can drift)."""
    first, last = {}, {}
    for i, p in enumerate(pos):
        if p is None:
            continue
        k = GW[p][1]
        first.setdefault(k, W[i]["s"])
        last[k] = W[i]["e"]
    pts = []
    for t in turns:
        if t["k"] in first and t.get("share", 0) >= MATCH_SHARE and t["sg"] is not None:
            pts.append((t["sg"], first[t["k"]]))
            pts.append((t["eg"], last[t["k"]]))
    pts.sort()
    xs, ys = [], []
    for x, y in pts:
        if xs and y < ys[-1]:
            y = ys[-1]                      # keep it running forwards
        if xs and x == xs[-1]:
            ys[-1] = max(ys[-1], y)
            continue
        xs.append(x); ys.append(y)

    def f(t):
        if t is None:
            return None
        if not xs:
            return min(max(t, 0.0), audio_s)
        if t <= xs[0]:
            v = ys[0] - (xs[0] - t)
        elif t >= xs[-1]:
            v = ys[-1] + (t - xs[-1])
        else:
            j = next(j for j in range(1, len(xs)) if xs[j] >= t)
            x0, x1, y0, y1 = xs[j - 1], xs[j], ys[j - 1], ys[j]
            v = y0 + (y1 - y0) * (t - x0) / (x1 - x0) if x1 > x0 else y0
        return round(min(max(v, 0.0), audio_s), 2)
    return f


def covering(turns, s, slack=2.0):
    """The Gemini turn (on WhisperX's clock) that covers time s, or None."""
    best, bd = None, None
    for t in turns:
        if t["sw"] is None or t["ew"] is None:
            continue
        if t["sw"] - slack <= s <= t["ew"] + slack:
            d = 0.0 if t["sw"] <= s <= t["ew"] else min(abs(s - t["sw"]), abs(s - t["ew"]))
            if bd is None or d < bd:
                best, bd = t, d
    return best


# ------------------------------------------------------------------ voices and label repair
def voice_counts(turns, W, pos, GW):
    """(label, part) -> Counter of WhisperX voices over lined-up words (songs left out)."""
    vc = collections.defaultdict(collections.Counter)
    for i, p in enumerate(pos):
        if p is not None:
            t = turns[GW[p][1]]
            if t["role"] != "lyrics":
                vc[(t["label"], t["part"])][W[i]["spk"]] += 1
    return vc


def voice_shares(vc):
    """(voice, part) -> Counter of labels: whose words a WhisperX voice carries in a part."""
    vs = collections.defaultdict(collections.Counter)
    for (lab, p), c in vc.items():
        for v, n in c.items():
            vs[(v, p)][lab] += n
    return vs


def evidence_part(sp, plan, vc):
    m = TIME_RE.search(sp.get("evidence") or "")
    if m:
        h_or_m, mm, ss = m.group(1), m.group(2), m.group(3)
        s = int(h_or_m) * 3600 + int(mm) * 60 + int(ss) if ss else int(h_or_m) * 60 + int(mm)
        for i, q in enumerate(plan):
            if q["start"] - 1 <= s <= q["end"] + 1:
                return i
    parts = {p: sum(c.values()) for (lab, p), c in vc.items() if lab == sp["label"]}
    return max(parts, key=parts.get) if parts else 0


def repair_labels(turns, plan, speakers, W, pos, GW):
    """Puts back labels that moved to another person's voice in a later part. Returns (repairs, labels to check)."""
    if len(plan) < 2:
        return [], []
    vc = voice_counts(turns, W, pos, GW)
    role = {s["label"]: s["role"] for s in speakers}
    ev = {s["label"]: evidence_part(s, plan, vc) for s in speakers}
    vs = voice_shares(vc)
    sig = {}
    for lab, p in ev.items():
        c = vc.get((lab, p))
        if c and sum(c.values()) >= VOICE_MIN_WORDS:
            v, n = c.most_common(1)[0]
            own = vs[(v, p)]
            # the person speaks mostly in that voice, and that voice is mostly that person (WhisperX can lump voices)
            if n / sum(c.values()) >= VOICE_PURITY and own[lab] / sum(own.values()) >= VOICE_PURITY:
                sig[lab] = v
    owners = collections.defaultdict(list)
    for lab, v in sig.items():
        owners[v].append(lab)
    repairs = []
    for j in range(len(plan)):
        props = {}
        present = {lab: sum(c.values()) for (lab, p), c in vc.items() if p == j}
        for lab, n in present.items():
            if n < REPAIR_MIN_WORDS or ev.get(lab) == j:
                continue
            v, nv = vc[(lab, j)].most_common(1)[0]
            if nv / n < REPAIR_PURITY or len(owners.get(v, [])) != 1:
                continue
            o = owners[v][0]
            if o != lab and role.get(o) == role.get(lab) and ev.get(o) != j:
                props[lab] = o
        # a consistent set only: no two labels onto one, and a target still speaking in its own voice blocks it
        targets = collections.Counter(props.values())
        ok = {}
        for lab, o in props.items():
            if targets[o] > 1:
                continue
            if present.get(o, 0) >= 20 and o not in props:
                continue
            ok[lab] = o
        if ok:
            for t in turns:
                if t["part"] == j and t["label"] in ok:
                    t["was"] = t["label"]
                    t["label"] = ok[t["label"]]
            for lab, o in ok.items():
                repairs.append({"part": j + 1, "from": lab, "to": o, "words": present[lab]})
    # labels whose main voice still differs between parts
    vc = voice_counts(turns, W, pos, GW)
    main = collections.defaultdict(set)
    for (lab, p), c in vc.items():
        n = sum(c.values())
        if n >= REPAIR_MIN_WORDS:
            v, nv = c.most_common(1)[0]
            if nv / n >= 0.6:
                main[lab].add(v)
    check = sorted(lab for lab, vs in main.items() if len(vs) > 1)
    return repairs, check


def voice_labels(turns, W, pos, GW):
    """WhisperX voice -> (label, role) where the voice clearly belongs to one Gemini label."""
    c = collections.defaultdict(collections.Counter)
    for i, p in enumerate(pos):
        if p is not None:
            t = turns[GW[p][1]]
            c[W[i]["spk"]][(t["label"], t["role"])] += 1
    out = {}
    for v, cc in c.items():
        by_label = collections.Counter()
        for (lab, r), n in cc.items():
            if r != "lyrics":
                by_label[lab] += n
        if not by_label:
            continue
        lab, n = by_label.most_common(1)[0]
        tot = sum(by_label.values())
        if tot >= VOICE_MIN_WORDS and n / tot >= VOICE_PURITY:
            out[v] = lab
    return out


# ------------------------------------------------------------------ names
def name_found(name, text_words, caps):
    toks = [t for t in re.split(r"[\s\-\.]+", norm_line(name)) if len(t) >= 3 and t not in NAME_STOP]
    for t in toks:
        t2 = t.replace("'", "")
        if t2 in text_words:
            return True
        # a short form ("Phil" for Philip), written with a capital
        if any(len(c) >= 3 and t2.startswith(c) for c in caps):
            return True
    return False


def resolve_people(speakers, turns, W, kind):
    """label -> {person, name_source, guess, final label}; names kept only on evidence."""
    words = set()
    caps = set()
    for w in W:
        words.add(w["n"].replace("'", ""))
        if w["w"][:1].isupper():
            caps.add(w["n"].replace("'", ""))
    for t in turns:
        for x in t["text"].split():
            n = norm(x).replace("'", "")
            if n:
                words.add(n)
                if fold(x)[:1].isupper():
                    caps.add(n)
    used = {lab for lab in (s["label"] for s in speakers)}
    letters = collections.defaultdict(int)
    for lab in used:
        m = re.match(r"^(band member, unspecified|artist, unconfirmed|unknown) \(([A-Z])\)$", lab)
        if m:
            letters[m.group(1)] = max(letters[m.group(1)], ord(m.group(2)) - 64)
    out = {}
    for s in speakers:
        name, src = s.get("name"), (s.get("name_source") or "").lower() or None
        info = {"person": "", "name_source": src or "", "guess": "", "label": s["label"], "note": ""}
        if name:
            if src in NAME_OK and (src != "said in the video" or name_found(name, words, caps)):
                info["person"] = name
            else:
                info["guess"] = name
                info["note"] = ("name said in the video, but not found in either transcript" if src == "said in the video"
                                else "name from the title or the member list only")
                if norm_line(name) and norm_line(name) in norm_line(s["label"]):     # the label is the name itself
                    base = ("band member, unspecified" if s["role"] == "band member" else
                            "artist, unconfirmed" if s["role"] == "artist" else "unknown")
                    letters[base] += 1
                    info["label"] = f"{base} ({chr(64 + letters[base])})"
        out[s["label"]] = info
    return out


def guess_voices(wx, W, kind):
    """For a video Gemini refused as a whole: the voice with the most words is taken to be the artist (or a band
    member), the voice that asks most questions the interviewer, and anyone else unknown. Marked as guessed."""
    words = collections.Counter(w["spk"] for w in W if not w["junk"])
    lines = collections.defaultdict(lambda: [0, 0])
    for ln in wx.get("lines", []):
        if ln.get("words", 0) >= 3:
            lines[ln.get("speaker")][0] += 1
            lines[ln.get("speaker")][1] += ln.get("text", "").rstrip().endswith("?")
    order = [v for v, n in words.most_common() if n >= 5]
    vlab, sps = {}, []
    if not order:
        return vlab, sps
    main = order[0]
    lab = "artist, unconfirmed (A)" if kind == "solo" else "band member, unspecified (A)"
    vlab[main] = lab
    sps.append({"label": lab, "name": None, "name_source": None, "role": "artist" if kind == "solo" else "band member",
                "description": "the voice with the most words (guessed)", "evidence": ""})
    rest = [v for v in order[1:]]
    q = {v: lines[v][1] / lines[v][0] for v in rest if lines[v][0]}
    asker = max(q, key=q.get) if q and max(q.values()) >= 0.2 else None
    n = 0
    for v in rest:
        if v == asker:
            vlab[v] = "interviewer (guessed)"
            sps.append({"label": "interviewer (guessed)", "name": None, "name_source": None, "role": "interviewer",
                        "description": "the voice that asks most questions (guessed)", "evidence": ""})
        else:
            n += 1
            vlab[v] = f"unknown ({chr(64 + n)})"
            sps.append({"label": vlab[v], "name": None, "name_source": None, "role": "other", "description": "", "evidence": ""})
    return vlab, sps


# ------------------------------------------------------------------ one video
def merge_video(doc, wx, kind):
    audio_s = float(wx.get("audio_seconds") or (doc.get("minutes") or 0) * 60)
    W = whisperx_words(wx)
    turns, plan, extra_speakers = gemini_turns(doc, kind)
    speakers = []
    seen = set()
    for s in list(doc.get("speakers", [])) + extra_speakers:
        lab = clean(s.get("label") or s.get("name"))
        if lab and lab not in seen:
            seen.add(lab)
            speakers.append({"label": lab, "name": clean(s.get("name")) or None, "name_source": clean(s.get("name_source")) or None,
                             "role": role_of(s.get("role"), kind), "description": clean(s.get("description")),
                             "evidence": clean(s.get("evidence"))})
    for t in turns:                       # labels used in turns but missing from the list
        if t["label"] not in seen:
            seen.add(t["label"])
            speakers.append({"label": t["label"], "name": None, "name_source": None, "role": t["role"] if t["role"] != "lyrics" else "other",
                             "description": t["on_screen"], "evidence": ""})
    GW = [(norm(x), t["k"]) for t in turns for x in t["text"].split() if norm(x)]
    pos = line_up(W, GW)
    n_g = collections.Counter(k for _, k in GW)
    n_m = collections.Counter(GW[p][1] for p in pos if p is not None)
    for t in turns:
        t["n"] = n_g.get(t["k"], 0)
        t["share"] = n_m.get(t["k"], 0) / t["n"] if t["n"] else 0.0
        t["english"] = t["lang"].startswith("en")
    g2w = time_map(turns, W, pos, GW, audio_s)
    for t in turns:
        t["sw"], t["ew"] = g2w(t["sg"]), g2w(t["eg"])
    repairs, check = repair_labels(turns, plan, speakers, W, pos, GW)

    # which turn each WhisperX word goes to
    owner = [None] * len(W)            # turn index
    aside = [None] * len(W)            # set aside for a turn whose Gemini text is used
    fallback = [False] * len(W)
    dropped = collections.Counter()
    for i, p in enumerate(pos):
        if p is not None:
            owner[i] = GW[p][1]
    i = 0
    while i < len(W):
        if pos[i] is not None:
            i += 1
            continue
        j = i
        while j < len(W) and pos[j] is None:
            j += 1
        run = list(range(i, j))
        keep = []
        for k in run:
            if W[k]["junk"]:
                dropped[W[k]["junk"]] += 1
            else:
                keep.append(k)
        left = owner[i - 1] if i > 0 else None
        right = owner[j] if j < len(W) else None
        if keep:
            dur = W[keep[-1]]["e"] - W[keep[0]]["s"]
            short = len(keep) < SHORT_RUN_WORDS and dur < SHORT_RUN_SECONDS
            # next to a turn whose words mostly differ (it will use Gemini's text), the stray words are part of it
            weak_l = left is not None and (turns[left]["share"] < MATCH_SHARE or not turns[left]["english"])
            weak_r = right is not None and (turns[right]["share"] < MATCH_SHARE or not turns[right]["english"])
            for k in keep:
                dl = W[k]["s"] - W[i - 1]["e"] if left is not None else 1e9
                dr = W[j]["s"] - W[k]["e"] if right is not None else 1e9
                if short and min(dl, dr) <= SHORT_RUN_SECONDS:
                    vl = W[i - 1]["spk"] if left is not None else None
                    vr = W[j]["spk"] if right is not None else None
                    if weak_l != weak_r:
                        owner[k] = left if weak_l else right
                    elif vl != vr and W[k]["spk"] in (vl, vr):     # the neighbour in the same WhisperX voice
                        owner[k] = left if W[k]["spk"] == vl else right
                    else:
                        owner[k] = left if dl <= dr else right     # a few words between lined-up ones: the nearer turn
                    continue
                t = covering(turns, W[k]["s"])                 # otherwise: the Gemini turn at that time, if any
                if t is None:
                    if W[k]["lowconf"]:          # WhisperX unsure and Gemini heard nothing: most likely noise
                        dropped["low confidence, Gemini silent"] += 1
                    else:
                        fallback[k] = True
                elif t["share"] >= MATCH_SHARE and t["english"]:
                    owner[k] = t["k"]
                else:
                    aside[k] = t["k"]
        i = j

    # a few WhisperX-only words right next to a turn belong to it (Gemini's times are only to the second)
    k = 0
    while k < len(W):
        if not fallback[k]:
            k += 1
            continue
        m = k
        while m < len(W) and fallback[m]:
            m += 1
        if m - k < 5:
            lo = next((x for x in range(k - 1, -1, -1) if owner[x] is not None), None)
            hi = next((x for x in range(m, len(W)) if owner[x] is not None), None)
            for x in range(k, m):
                dl = W[x]["s"] - W[lo]["e"] if lo is not None else 1e9
                dr = W[hi]["s"] - W[x]["e"] if hi is not None else 1e9
                if min(dl, dr) <= 5:
                    owner[x] = owner[lo] if dl <= dr else owner[hi]
                    fallback[x] = False
        k = m

    # the joined turns
    vlab = voice_labels(turns, W, pos, GW)
    guessed = False
    if not turns and W:                    # Gemini gave nothing for the whole video: guess from WhisperX's voices
        vlab, guessed_speakers = guess_voices(wx, W, kind)
        speakers += guessed_speakers
        guessed = True
    role = {s["label"]: s["role"] for s in speakers}
    people = resolve_people(speakers, turns, W, kind)
    words_of = collections.defaultdict(list)
    aside_of = collections.defaultdict(list)
    for k in range(len(W)):
        if owner[k] is not None:
            words_of[owner[k]].append(k)
        elif aside[k] is not None:
            aside_of[aside[k]].append(k)
    out = []
    n_aside = 0
    for t in turns:
        ws = words_of.get(t["k"], [])
        use_wx = t["english"] and t["share"] >= MATCH_SHARE and bool(ws)
        if not use_wx:
            aside_of[t["k"]] = sorted(aside_of.get(t["k"], []) + ws)
            n_aside += len(aside_of[t["k"]])
            ws = []
        if not t["text"] and not ws:
            continue
        flags = []
        if t.get("was"):
            flags.append(f"label put back (Gemini had {t['was']})")
        if t["label"] in check:
            flags.append("label to check across parts")
        if use_wx:
            text = " ".join(W[k]["w"] for k in ws)
            s, e = W[ws[0]]["s"], W[ws[-1]]["e"]
            src = "WhisperX"
            sung = sum(W[k]["sung"] for k in ws) / len(ws)
            if sung >= 0.5 and t["role"] != "lyrics":
                flags.append("sung")
        else:
            text = t["text"]
            s, e = t["sw"], t["ew"]
            src = "Gemini"
            flags.append("times from the lined-up words nearby")
        info = people.get(t["label"], {"person": "", "name_source": "", "guess": "", "label": t["label"], "note": ""})
        r = t["role"]
        out.append({"start": s, "end": e, "speaker": info["label"], "person": info["person"],
                    "name_source": info["name_source"] if info["person"] else "", "role": r,
                    "band_voice": r in VOICE_ROLES, "language": t["lang"], "text": text, "text_from": src,
                    "words": len(text.split()), "part": (t["part"] or 0) + 1, "flags": flags,
                    "gemini_label": t.get("was") or t["label"], "gemini_text": t["text"] if src == "WhisperX" else "",
                    "whisperx_text": " ".join(W[k]["w"] for k in aside_of.get(t["k"], [])) if src == "Gemini" else ""})
    # WhisperX only, where Gemini has nothing: one turn per run of one voice
    cur = None
    for k in range(len(W)):
        if not fallback[k]:
            cur = None
            continue
        lab = vlab.get(W[k]["spk"])
        if cur is None or cur["voice"] != W[k]["spk"] or W[k]["s"] - cur["end"] > 10:
            info = people.get(lab, {}) if lab else {}
            r = role.get(lab, "other") if lab else "other"
            if W[k]["sung"]:
                r = "lyrics"
            cur = {"start": W[k]["s"], "end": W[k]["e"], "speaker": info.get("label") or (lab or f"voice {W[k]['spk']} (WhisperX)"),
                   "person": info.get("person", ""), "name_source": info.get("name_source", "") if info.get("person") else "",
                   "role": r, "band_voice": r in VOICE_ROLES, "language": "en", "text": W[k]["w"], "text_from": "WhisperX only",
                   "words": 1, "part": 0,
                   "flags": ["speaker guessed from WhisperX's voices (Gemini gave nothing for this video)" if guessed else
                             "speaker from WhisperX's voices" if lab else "speaker unknown (WhisperX voice only)"],
                   "gemini_label": "", "gemini_text": "", "whisperx_text": "", "voice": W[k]["spk"]}
            out.append(cur)
        else:
            cur["text"] += " " + W[k]["w"]
            cur["end"] = W[k]["e"]
            cur["words"] += 1
    for t in out:
        t.pop("voice", None)
    out.sort(key=lambda t: (t["start"] if t["start"] is not None else 1e9))
    for n, t in enumerate(out, 1):
        t["turn"] = n
    rec_parts = {t["part"] for t in turns if t.get("recovered")}
    ok_r = {(round(p["start"], 1), round(p["end"], 1)) for p in doc.get("parts", []) if "prompt" in p and p.get("ok")}
    refused = [q for i, q in enumerate(plan) if (round(q["start"], 1), round(q["end"], 1)) not in ok_r and i not in rec_parts]
    stats = {"WhisperX words": len(W), "lined up": sum(p is not None for p in pos),
             "dropped": dict(dropped), "set aside": n_aside, "guessed": guessed,
             "replies read here": len(rec_parts),
             "WhisperX only": sum(fallback), "refused minutes": round(sum(p["end"] - p["start"] for p in refused) / 60, 1),
             "label repairs": repairs, "labels to check": check}
    plist = []
    words_by = collections.Counter()
    turns_by = collections.Counter()
    for t in out:
        words_by[t["speaker"]] += t["words"]
        turns_by[t["speaker"]] += 1
    for s in speakers:
        info = people[s["label"]]
        plist.append({"label": info["label"], "gemini_label": s["label"], "person": info["person"],
                      "name_source": info["name_source"] if info["person"] else "", "name_not_kept": info["guess"],
                      "why_not_kept": info["note"], "role": s["role"], "description": s["description"], "evidence": s["evidence"],
                      "turns": turns_by.get(info["label"], 0), "words": words_by.get(info["label"], 0)})
    return out, plist, stats


# ------------------------------------------------------------------ the text corpus's layout
CORPUS_COLS = ["Record ID", "Turn ID", "Turn order", "Turn type", "Speaker type", "Speaker name", "Speaker candidates",
               "Reported member", "Band voice", "Main set", "Confidence", "Attribution rule", "Attribution evidence",
               "Embedded quote", "Question", "Holds an undecided quote", "Turn text", "Words", "Repeat of", "Span start",
               "Span end", "Interview date", "Interview date source", "Interview date precision", "Interview date evidence",
               "Route", "Quotation convention", "Focal Act", "Grammy Year", "Outlet / Source Name", "Publication Date",
               "Date Confidence", "Date Precision", "Content Era", "Temporal Class", "Rules version", "Human verified"]
VIDEO_COLS = ["Video speaker label", "Start", "Start (s)", "End (s)", "Language", "Text from", "Video flags", "Video title",
              "Video URL", "Kind"]
BAND_TYPES = {"band member", "band member, unspecified", "band collective"}


def name_tokens(n):
    n = norm_line(n).replace("'", " ")
    toks = [{"jr": "junior", "sr": "senior"}.get(t, t) for t in n.split()]
    return [t for t in toks if t and t not in {"mr", "mrs", "ms", "dr", "the"}]


def canonical_member(name, frames):
    """A name as Gemini gives it -> (line-up name, 'nomination' or 'other'), (None, candidates) when ambiguous,
    or (None, None) when it is not on the act's name lists."""
    want = name_tokens(name)
    if not want:
        return None, None
    best, scores = None, {}
    for f in frames:
        for variant, canon, era, _ in f.get("names", []):
            vt = name_tokens(variant)
            if vt == want:
                return canon, era
            overlap = len(set(want) & set(vt))
            if overlap:
                sc = overlap - 0.5 * len(set(vt) - set(want)) - 0.5 * len(set(want) - set(vt))
                key = (canon, era)
                scores[key] = max(scores.get(key, -9), sc)
    if not scores:
        return None, None
    top = max(scores.values())
    winners = [k for k, v in scores.items() if v == top]
    if top <= -1:
        return None, None
    if len({c for c, _ in winners}) == 1:
        return winners[0]
    nom = [w for w in winners if w[1] == "nomination"]
    if len({c for c, _ in nom}) == 1:          # a nomination-era member over someone who joined later or left earlier
        return nom[0]
    return None, sorted({c for c, _ in winners})


def temporal_class(years_from):
    if years_from is None:
        return ""
    y = float(years_from)
    return "pre-nomination" if y < 0 else "contemporaneous" if y <= 3 else "near-term" if y <= 10 else "retrospective"


def corpus_row(meta, sel_row, t, person, frames, acts, kind):
    """One turn in the text corpus's layout (harrison_v7_turns_final), plus the video's own columns."""
    role, label, name = t["role"], t["speaker"], t["person"]
    flags = list(t["flags"])
    turn_type = "song" if role == "lyrics" else "narration" if role == "narrator" else "speech"
    stype, sname, cands, ev_note = "", name or "", "", ""
    acts_n = {norm_line(a).replace("the ", "", 1) for a in acts}
    collective = kind == "band" and norm_line(name or label).replace("the ", "", 1) in acts_n   # a solo artist is not a group
    if role in VOICE_ROLES or (role == "lyrics" and (person or {}).get("role") in VOICE_ROLES):
        if collective:
            stype, sname = "band collective", "; ".join(acts)
        elif name:
            canon, era = canonical_member(name, frames)
            if canon and era == "nomination":
                stype, sname = "band member", canon
            elif canon:
                stype, sname = "other-era member", canon
            elif isinstance(era, list):
                stype, sname, cands = "ambiguous", "", " / ".join(era)
            else:
                ev = norm_line(" ".join([(person or {}).get("evidence", ""), (person or {}).get("description", "")]))
                if any(a and a in ev for a in acts_n):     # with the act, but not a name on its lists (a nickname?)
                    stype, sname, cands = "band member, unspecified", "band member, unspecified", name
                    ev_note = "named in the video, but not a name on the act's line-up list"
                else:                                      # someone else (another band's member, a guest)
                    stype, sname = "other person", name
                    ev_note = "not on the act's line-up list"
        elif kind == "solo":
            stype, sname, cands = "band member", "", "; ".join(acts)
        else:
            stype, sname = "band member, unspecified", "band member, unspecified"
    elif role == "interviewer":
        stype = "interviewer"
    elif role == "narrator":
        stype = "journalist"
    elif t["text_from"] == "WhisperX only" and label.startswith("voice "):
        stype = "unknown"
    else:
        stype = "other person"
    band_voice = turn_type == "speech" and stype in BAND_TYPES
    words = t["words"]
    if t["text_from"] == "WhisperX only" or any("to check" in f for f in flags) or words <= 3:
        conf = "low"
    elif any("put back" in f for f in flags) or t["text_from"] == "Gemini" or stype == "ambiguous":
        conf = "medium"
    else:
        conf = "high" if words >= 20 else "medium"
    if any("guessed" in f for f in flags):
        rule = "video: guessed from WhisperX's voices (Gemini refused the video)"
    elif t["text_from"] == "WhisperX only":
        rule = "video: WhisperX's voice (Gemini has nothing here)"
    elif any("put back" in f for f in flags):
        rule = "video: Gemini's label, put back across parts"
    else:
        rule = "video: Gemini's label"
    p = person or {}
    if name:
        evidence = f"{p.get('name_source', '')}: {p.get('evidence', '')}".strip(": ")
    else:
        evidence = p.get("description", "")
    if ev_note:
        evidence = (evidence + "; " if evidence else "") + ev_note
    year = meta.get("Year")
    upload = None if sel_row is None else sel_row.get("Upload date")
    upload = None if upload is None or (isinstance(upload, float) and math.isnan(upload)) else str(upload)[:10]
    show = None if sel_row is None else sel_row.get("Show or event")
    channel = None if sel_row is None else sel_row.get("Channel")
    outlet = show if isinstance(show, str) and show.strip() else (f"YouTube: {channel}" if isinstance(channel, str) else "YouTube")
    gy = None if sel_row is None else sel_row.get("Grammy year")
    row = {"Record ID": meta["ID"], "Turn ID": f"{meta['ID']}-{t['turn']}", "Turn order": t["turn"], "Turn type": turn_type,
           "Speaker type": stype, "Speaker name": sname, "Speaker candidates": cands, "Reported member": "",
           "Band voice": "yes" if band_voice else "no",
           "Main set": ("yes" if conf in ("high", "medium") else "no") if band_voice else "",
           "Confidence": conf, "Attribution rule": rule, "Attribution evidence": evidence, "Embedded quote": "",
           "Question": "yes" if role == "interviewer" and t["text"].rstrip().endswith("?") else "no",
           "Holds an undecided quote": "no", "Turn text": t["text"], "Words": words, "Repeat of": "",
           "Span start": "", "Span end": "",
           "Interview date": "" if year is None else str(int(year)), "Interview date source": f"video: {meta.get('Year from') or ''}".strip(),
           "Interview date precision": "year", "Interview date evidence": meta.get("Title", ""), "Route": "video",
           "Quotation convention": "", "Focal Act": "; ".join(acts), "Grammy Year": None if gy is None else int(gy),
           "Outlet / Source Name": outlet, "Publication Date": upload or "", "Date Confidence": "high", "Date Precision": "day",
           "Content Era": None if year is None else int(year), "Temporal Class": temporal_class(meta.get("Years from nomination")),
           "Rules version": VERSION, "Human verified": "no",
           "Video speaker label": label, "Start": mmss(t["start"]), "Start (s)": t["start"], "End (s)": t["end"],
           "Language": t["language"], "Text from": t["text_from"], "Video flags": "; ".join(flags),
           "Video title": meta.get("Title", ""), "Video URL": meta.get("URL", ""), "Kind": meta.get("Kind", "")}
    return row


# ------------------------------------------------------------------ the corpus
# Videos dropped in the WhisperX audit of 18 October 2026 (never sent to Gemini); only for the report below.
DROPPED = ["V314", "V1225", "V1970", "V1980", "V1986", "V1987", "V2226", "V2232", "V2240", "V2695", "V2822", "V3148",
           "V3152", "V3485"]


def script_folder():
    f = globals().get("__file__")
    return os.path.dirname(os.path.abspath(f)) if f else os.getcwd()


def main(argv=None):
    base = script_folder()
    ap = argparse.ArgumentParser(prog="merge_video_turns.py")
    ap.add_argument("--gemini", default=os.path.join(base, "gemini_speakers"))
    ap.add_argument("--whisperx", default=os.path.join(base, "video_transcripts"))
    ap.add_argument("--selection", default=os.path.join(base, "harrison_v7_VIDEO_SELECTION.xlsx"))
    ap.add_argument("--frame", default=os.path.join(base, "acts_frame.json"))
    ap.add_argument("--out", default=base)
    ap.add_argument("--ids", default="")
    ap.add_argument("--no-xlsx", action="store_true")
    a = ap.parse_args(argv)
    print(VERSION)
    missing = [f"{what} ({path})" for what, path in [("the Gemini folder", a.gemini), ("the WhisperX folder", a.whisperx),
                                                     ("the selection workbook", a.selection), ("acts_frame.json", a.frame)]
               if not os.path.exists(path)]
    if missing:
        print("STOPPED: not found: " + "; ".join(missing) + ". Upload them next to this script, or give their places.")
        return [], [], [], []
    import pandas as pd
    frame_all = json.load(open(a.frame, encoding="utf-8"))
    types = {act: f["type"] for act, f in frame_all.items()}
    sel = pd.read_excel(a.selection, "To transcribe")
    sel = {str(r["ID"]): r for _, r in sel.iterrows()}
    ids = [x.strip().upper() for x in a.ids.split(",") if x.strip()]
    files = sorted((f for f in os.listdir(a.gemini) if re.match(r"V\d+\.json$", f)), key=lambda f: int(f[1:-5]))
    wx_files = {f for f in os.listdir(a.whisperx) if re.match(r"V\d+\.json$", f)}
    no_gemini = sorted((f[:-5] for f in wx_files - set(files) if f[:-5] not in DROPPED), key=lambda v: int(v[1:]))
    print(f"Found {len(files)} Gemini files and {len(wx_files)} WhisperX transcripts. "
          + (f"WhisperX transcripts with no Gemini file (apart from the {len(DROPPED)} dropped): {len(no_gemini)} "
             f"{no_gemini[:10]}" if no_gemini else "Every WhisperX transcript (apart from the dropped ones) has a Gemini file."))
    if ids:
        files = [f for f in files if f[:-5] in ids]
    os.makedirs(os.path.join(a.out, "video_turns"), exist_ok=True)
    T, V, P, R = [], [], [], []
    for f in files:
        vid = f[:-5]
        doc = json.load(open(os.path.join(a.gemini, f), encoding="utf-8"))
        wp = os.path.join(a.whisperx, f)
        if not os.path.exists(wp):
            print(f"{vid}: no WhisperX transcript, skipped")
            continue
        wx = json.load(open(wp, encoding="utf-8"))
        acts = [x.strip() for x in str(doc.get("act", "")).split(";") if x.strip()]
        kind = "band" if any(types.get(x, "?") != "Solo" for x in acts) else "solo"
        turns, people, stats = merge_video(doc, wx, kind)
        row = sel.get(vid)
        frames = [frame_all[x] for x in acts if x in frame_all]
        by_label = {p["label"]: p for p in people}
        meta = {"ID": vid, "Act": doc.get("act", ""), "Kind": "solo artist" if kind == "solo" else "band, duo or group",
                "Year": None if row is None else row.get("Year (best guess)"),
                "Year from": None if row is None else row.get("Year from"),
                "Years from nomination": None if row is None else row.get("Years from nomination"),
                "Show or event": None if row is None else row.get("Show or event"),
                "Title": doc.get("title", ""), "URL": doc.get("url", "")}
        meta = {k: (None if (isinstance(v, float) and math.isnan(v)) else v) for k, v in meta.items()}
        rows = []
        for t in turns:
            r = corpus_row(meta, row, t, by_label.get(t["speaker"]), frames, acts, kind)
            assert list(r) == CORPUS_COLS + VIDEO_COLS, list(r)
            rows.append(r)
            t.update({"speaker_type": r["Speaker type"], "speaker_name": r["Speaker name"],
                      "speaker_candidates": r["Speaker candidates"], "band_voice": r["Band voice"] == "yes",
                      "confidence": r["Confidence"], "main_set": r["Main set"]})
        with open(os.path.join(a.out, "video_turns", f), "w", encoding="utf-8") as fh:
            json.dump({"version": VERSION, "video": meta, "gemini": {"status": doc.get("status"), "version": doc.get("version"),
                       "model": doc.get("model_version")}, "whisperx": {"version": wx.get("version")}, "stats": stats,
                       "people": people, "turns": turns}, fh, ensure_ascii=False, indent=1)
        bv = [t for t in turns if t["band_voice"]]
        src = collections.Counter()
        for t in turns:
            src[t["text_from"]] += t["words"]
        flags = []
        if stats["guessed"]:
            flags.append("Gemini refused the whole video: speakers guessed from WhisperX's voices")
        elif stats["refused minutes"] > 0:
            flags.append(f"Gemini refused {stats['refused minutes']} min (WhisperX only there)")
        if stats["replies read here"]:
            flags.append("a Gemini reply its script could not read was read here")
        if stats["label repairs"]:
            flags.append(f"{len(stats['label repairs'])} label(s) put back across parts")
        if stats["labels to check"]:
            flags.append("labels to check across parts: " + "; ".join(stats["labels to check"]))
        if any(not t["language"].startswith("en") for t in turns):
            flags.append("turns not in English")
        V.append({**meta, "Minutes": round(float(wx.get("audio_seconds") or 0) / 60, 2), "Turns": len(turns),
                  "Band-voice turns": len(bv), "Band-voice words": sum(t["words"] for t in bv),
                  "Band-voice words in the main set": sum(t["words"] for t in bv if t["main_set"] == "yes"),
                  "Words": sum(t["words"] for t in turns), "Words from WhisperX": src["WhisperX"],
                  "Words from Gemini": src["Gemini"], "Words from WhisperX only": src["WhisperX only"],
                  "WhisperX words set aside": stats["set aside"], "WhisperX words left out": sum(stats["dropped"].values()),
                  "People named": sum(1 for p in people if p["person"]), "Gemini status": doc.get("status"),
                  "Flags": "; ".join(flags)})
        T.extend(rows)
        for p in people:
            P.append({"ID": vid, "Act": meta["Act"], "Label": p["label"], "Person": p["person"], "Name source": p["name_source"],
                      "Name not kept": p["name_not_kept"], "Why not kept": p["why_not_kept"], "Role": p["role"],
                      "Turns": p["turns"], "Words": p["words"], "Description": p["description"], "Evidence (Gemini)": p["evidence"],
                      "Gemini label": p["gemini_label"]})
        for r in stats["label repairs"]:
            R.append({"ID": vid, "Act": meta["Act"], "Part": r["part"], "Gemini label": r["from"], "Put back as": r["to"],
                      "Words": r["words"]})
        n_done = len(V)
        if flags and any(("refused" in x or "check" in x or "read here" in x) for x in flags):
            print(f"  {vid}: {'; '.join(x for x in flags if 'not in English' not in x)}", flush=True)
        if n_done % 100 == 0:
            print(f"{n_done} of {len(files)} videos joined", flush=True)
    if V:
        cols = CORPUS_COLS + VIDEO_COLS
        with open(os.path.join(a.out, "harrison_v7_VIDEO_TURNS.json"), "w", encoding="utf-8") as fh:
            json.dump({"about": {"version": VERSION, "notes": dict(README)}, "columns": cols,
                       "data": [[r[c] for c in cols] for r in T]}, fh, ensure_ascii=False)
    if not a.no_xlsx and V:
        print("Writing the workbook ...", flush=True)
        write_xlsx(os.path.join(a.out, "harrison_v7_VIDEO_TURNS.xlsx"), T, V, P, R)
    bvt = [t for t in T if t["Band voice"] == "yes"]
    print(f"Done: {len(V)} videos, {len(T)} turns, {sum(t['Words'] for t in T)} words; band voice {len(bvt)} turns, "
          f"{sum(t['Words'] for t in bvt)} words ({sum(t['Words'] for t in bvt if t['Main set'] == 'yes')} in the main set). "
          + ("Files: " + ("" if a.no_xlsx or not V else os.path.join(a.out, "harrison_v7_VIDEO_TURNS.xlsx") + ", ")
             + os.path.join(a.out, "harrison_v7_VIDEO_TURNS.json") + " and " + os.path.join(a.out, "video_turns") + "/"))
    return T, V, P, R


README = [
    ("What this is", "Every video's speaker turns, joining WhisperX (the words and their times) and Gemini (who speaks, the "
                     "role, a name only on evidence, the language). Made by " + VERSION + ". Not yet checked by a person: "
                     "on the two check videos (V1158, V1319) the speaker was right for 84% of sentences, and band voice "
                     "was 94% precise and 98% complete."),
    ("Same layout as the text corpus", "The Turns sheet's first 37 columns are those of harrison_v7_turns_final, with the "
                     "same meanings, so the two tables can be stacked. Route = video. What differs for videos is below; "
                     "the last 10 columns are the video's own."),
    ("Record ID, Turn ID", "The video's ID (V1158) and the turn's number in it (V1158-12), in time order."),
    ("Turn type", "speech; narration (a TV narrator or announcer); song (singing, never band voice)."),
    ("Speaker type", "band member: named on evidence in the video and on the act's nomination-era line-up (for a solo "
                     "artist, the artist); band member, unspecified: a member the video never names (Video speaker label "
                     "tells them apart within a video); band collective: the act as a group; other-era member; "
                     "interviewer; journalist (the programme's narrator); other person; unknown (a WhisperX voice that "
                     "belongs to no Gemini label); ambiguous (Speaker candidates lists who it could be)."),
    ("Speaker name", "The line-up name for members; other people as the video names them; blank when nothing in the video "
                     "names them. A solo artist's turns with no name in the video have the artist in Speaker candidates."),
    ("Names", "Kept only on evidence: said in the video (and found in the words), shown in an on-screen caption, or "
              "visibly the same person as someone so named. A name from the video's title or the member list alone is "
              "not kept; the People sheet shows it under Name not kept."),
    ("Confidence", "Provisional until the spot check. low: WhisperX-only stretches, labels to check across parts, and turns "
                   "of three words or fewer; medium: turns with Gemini's text, labels put back, and other turns under 20 "
                   "words; high: turns of 20 words or more whose words line up. Main set and Band voice as in the text corpus."),
    ("Dates", "Interview date and Content Era: the selection's best-guess year (Interview date source says how it was "
              "dated). Temporal Class from the years since the nomination, as in the text corpus. Publication Date: the "
              "upload date. Outlet: the show where known, otherwise YouTube and the channel."),
    ("Not filled for videos", "Span start and Span end (the times are in Start (s) and End (s)), Reported member, "
                              "Embedded quote, Repeat of, Quotation convention."),
    ("Text from", "WhisperX: WhisperX's words, lined up with Gemini's. Gemini: Gemini's text, for turns not in English and "
                  "where WhisperX's words are mostly different (often an interpreter talking over a band member). "
                  "WhisperX only: Gemini has nothing there (it refused the stretch or left it out); the speaker comes "
                  "from WhisperX's voices."),
    ("Video flags", "label put back: Gemini moved the label to another voice in a later part and it was put back; label to "
                    "check across parts: the label still moves between voices; sung: mostly sung; times from the lined-up "
                    "words nearby: turns with Gemini's text, whose times are estimated."),
    ("Gemini refused", "Gemini refuses to write out text it takes for published material (song lyrics, a published "
                       "interview). Where it refused part of a video, that stretch keeps WhisperX's words and the "
                       "speakers come from WhisperX's voices; where it refused a whole video, the voice with the most "
                       "words is guessed to be the artist (or a band member) and the voice asking most questions the "
                       "interviewer, at low confidence, so outside the main set."),
    ("Left out", "WhisperX lines made up over silence or music ('Thank you.'), lines repeating its prompt, and "
                 "transcription loops, unless Gemini heard the same words. Counted per video in Videos."),
    ("Other sheets", "Videos: one row per video, with its counts and flags. People: every label in every video, with the "
                     "name kept or not kept and why. Label repairs: the labels put back across parts."),
]


def write_xlsx(path, T, V, P, R):
    import pandas as pd
    from openpyxl.styles import Font, Alignment
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def frame(rows, cols=None):
        df = pd.DataFrame(rows, columns=cols)
        for c in df.columns:
            if df[c].dtype == object or str(df[c].dtype) in ("str", "string"):
                df[c] = df[c].map(lambda x: ILLEGAL_CHARACTERS_RE.sub("", x) if isinstance(x, str) else x)
        return df
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        frame(README, ["Item", "Explanation"]).to_excel(xw, sheet_name="Read me", index=False)
        frame(T).to_excel(xw, sheet_name="Turns", index=False)
        frame(V).to_excel(xw, sheet_name="Videos", index=False)
        frame(P).to_excel(xw, sheet_name="People", index=False)
        frame(R if R else [{"ID": "", "Act": "", "Part": "", "Gemini label": "", "Put back as": "", "Words": ""}]).to_excel(
            xw, sheet_name="Label repairs", index=False)
        widths = {"Read me": {"A": 26, "B": 120}, "Turns": {"E": 22, "F": 24, "L": 30, "M": 40, "Q": 90, "AL": 28, "AQ": 40},
                  "Videos": {"H": 50, "V": 60}, "People": {"C": 30, "D": 24, "K": 50, "L": 50}}
        for name, ws in xw.sheets.items():
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True)
            for col, w in widths.get(name, {}).items():
                ws.column_dimensions[col].width = w
            if name == "Read me":
                for row in ws.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical="top")


if __name__ == "__main__":
    main()
