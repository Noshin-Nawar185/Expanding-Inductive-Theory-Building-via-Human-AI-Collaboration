"""step3_open_coding.py, version 1.9 (27 October 2026)

What it does: runs Step 3, open coding with in-vivo labels, with three AI agents, in five rounds:
  1. reading: each agent reads each act in full (every turn, without Before), one act per call, and writes a few
     notes on it, with the data problems it finds listed apart (Prompt 1). The reading goes to each provider as a
     batch at half price (READ_BATCH). Each act's notes go into the agent's memory just before it codes that act.
     make_exclusion_list() then lists the turns that two or more agents marked as not the member's words, in
     step3_inputs/excluded_turns.csv; round 2 leaves those turns out (the corpus files are never changed);
  2. coding: each agent codes the acts in nomination-year order, one batch of whole interviews per call, turn by
     turn, giving each code a label, a one-sentence description, the evidence and whether it is unsure, saying for
     each turn whether it looks back and any data problem, defining each new label, and adding to its memory after
     each turn (Prompt 2). The script works out whether each label is the members' words. After each act, the agent
     reviews it (Prompt 2b): it renames labels that do not make sense on their own, writes their definitions, writes
     a memo on the act and keeps the notes it still needs; from then on its memory carries that memo and those notes
     instead of all its notes on the act. Within an act, the memory and the labels list stay as they were at its
     start, so that both are read from the cache; what the act adds comes after them;
  3. memo part 1: each agent writes part 1 of its memo, before seeing anyone else's codes (Prompt 3), from its memos
     on the acts, reading its own codes in its folder as it needs (up to READ_BUDGET characters in all);
  4. merge: the script combines the three agents' codes into one sheet, turn by turn, without changing or resolving
     any code, and works out the agreement figures and the sample of turns for round 5;
  5. memo parts 2 and 3: each agent reads the figures and the sample and writes parts 2 and 3 (Prompt 4).
The three agents work at the same time, each on its own, and never see each other's work before round 5.

Agents (see MODELS below): A = Claude Opus 5.5 (the lead, at low effort), B = Gemini 4 Argon and C = GPT-6.1 Sol
(both at medium effort).

Reads, in MyDrive/harrison_v7 (it never changes them):
  band_voice_by_act/*.jsonl and harrison_v7_band_voice_acts.json   the corpus, from build_band_voice_jsonl.py
  step3_inputs/prompt1.txt, prompt2.txt, prompt2b.txt, prompt3.txt, prompt4.txt
                                                                    the five prompts as plain text, one file each,
                                                                    with the {placeholders} as in the prompt doc
  step3_inputs/<GUIDE_FILE>                                         the best practices guide (.docx, .txt or .md)
  step3_inputs/excluded_turns.csv (if there is one)                 turns round 2 leaves out (make_exclusion_list)
Writes, in MyDrive/harrison_v7/step3_<RUN>/:
  agent_A, agent_B, agent_C   each agent's work: state.json, memory.txt (what it carries) and notes_all.txt (every
                              note), reading/ (its notes and data problems on each act, and the batch jobs),
                              batches/<act>/ (each accepted reply, one file per batch, the first with the reading
                              notes that went into memory), reviews/<act>.json (its review of each act), codes/ (its
                              codes by act, as CSV, with the review's new names), calls/ (every call: the changing
                              part of the prompt, the reply, tokens and cost), prompts/ (the fixed part of the
                              prompts), flags.jsonl (problems that remained, and data problems), memo_<letter>.md
  merged/                     round 4: step3_merged_codes.xlsx, the agreement figures and the sample of turns
  reports/                    act_reports.jsonl (after each act), pilot_report.md, data_problems.csv and
                              exclusion_candidates.csv
  run_settings.json, run_log.txt

Saved as it goes: each accepted batch is saved as its own file, and whenever round 2 starts, each agent's memory,
labels and progress are rebuilt from those files. So a run that stops (Colab disconnects, you press Stop, the budget
runs out) carries on where it left off when you run the same cell again, and nothing is coded twice or left out.
Every prompt and reply is saved too, so any code can be traced to the call that made it.

HOW TO RUN, in Colab (any runtime; no GPU):
  0. Upload this script to MyDrive/harrison_v7, and the five prompts and the guide to MyDrive/harrison_v7/step3_inputs.
     Add three keys to Colab Secrets (the key icon on the left), with notebook access on:
     ANTHROPIC_API_KEY, OPENAI_API_KEY and GEMINI_API_KEY. The script reads them there and never prints them.
  1. !pip install -q "anthropic==2.3.0" "openai==3.37.0" "google-genai==2.33.0" tiktoken python-docx openpyxl
     If Colab asks you to restart the session, do so (Runtime > Restart session), then carry on from step 2.
  2. from google.colab import drive; drive.mount("/content/drive")
     %run -i "/content/drive/MyDrive/harrison_v7/step3_open_coding.py"
     Run this cell again after you change a prompt or a setting. For the full run, add a line under the %run line:
     RUN = "full"; ACTS = "all"
  3. check_inputs()      checks the corpus, the prompts, the guide, the keys and the SDK versions (no model calls)
     dry_run()           writes the filled prompts to step3_<RUN>/dry_run, to read before spending anything
     test_models()       a few small calls per model: the keys, the JSON format, the cache and the folder tool
  4. run_round1()        reading: sends the reading as a batch and waits up to READ_WAIT_MINUTES; if the batch is
                         not back by then, the agent says "waiting for its batch": run it again later to collect it
                         (batches finish within 24 hours)
     make_exclusion_list()   lists the turns that two or more agents marked as not the member's words in
                         step3_inputs/excluded_turns.csv (open it to check), and all marked turns in
                         reports/exclusion_candidates.csv; it waits until every agent has read every act. To change
                         the list, use remove_from_exclusion_list("32-2") or add_to_exclusion_list("113-32",
                         reason="..."), not a spreadsheet, which can turn Turn IDs into dates
     run_round2()        coding (the acts in ACTS: the pilot act, or all), with the review after each act; for the
                         full run, in parts: run_round2(part=1) and so on (show_parts() lists them)
     pilot_report()      unique labels, codes per turn, the other checks, and the cost and hours round 2 will take
                         in the full run
     run_round3()        memo part 1
     run_round4()        merge, agreement figures and sample
     run_round5()        memo parts 2 and 3
     status()            where each agent is, its cost and its flags, at any time
  Each round can be run again: it skips what is done and carries on. To stop, press Stop once and wait: the calls in
  progress finish and are saved first (a new round will not start until they have).
  If an agent stops on the same batch every time (for example a reply the model refuses), skip_batch(...) leaves that
  batch's turns uncoded, marked as skipped and flagged, so that the agent can carry on.
The full run starts fresh, in nomination-year order, with its own memory, in its own folder (step3_full).

THE FULL RUN, in this order (the same cells as the pilot's, with RUN = "full" and ACTS = "all" under the %run line,
and BUDGET_USD set from the pilot report's projection, with some room):
  1. check_inputs() and dry_run(), as for the pilot; check_inputs() also lists the parts of round 2.
  2. run_round1(): the reading of all 52 acts, as batches. Run it again until all three agents say "round 1 done".
  3. make_exclusion_list(), then check step3_inputs/excluded_turns.csv before round 2 (the turns already on it, from
     the pilot, stay on it).
  4. run_round2(part=1), later run_round2(part=2), then run_round2(part=3): round 2 in PARTS parts of about a third
     of the turns each, one Colab session each. A part ends after an act, and the next part carries on from there
     with the same memory and labels (a new act writes its cache afresh anyway, so a break between parts costs
     nothing). run_round2(up_to=4) stops after the fourth act instead, for a first look.
  5. run_round3(), run_round4() and run_round5(): the memos and the merge.
If Colab disconnects, everything done is on Drive: run the setup cells again with the same settings, then the same
round or part again, and it carries on where it stopped; nothing is coded twice or left out. All a disconnect can
cost is the calls in progress at that moment (one per agent, well under a dollar in all); if Drive had not yet saved
an agent's last batch, that batch is coded again; and after a break of over an hour in the middle of an act, Claude
and GPT write their caches again (a few dollars at most, late in the run). If the budget stopped the run, raise
BUDGET_USD first. Through a provider's outage, an agent waits and tries again for about 45 minutes before it stops;
if one stops while the others go on, press Stop once, wait for the calls in progress to finish, and run the same
round or part again: all three carry on.
"""
import collections, csv, datetime, hashlib, json, os, random, re, shutil, threading, time, traceback, unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed

VERSION = "step3_open_coding 1.9 (27 Oct 2026)"

if any(t.name.startswith("step3-agent") and t.is_alive() for t in threading.enumerate()):
    raise SystemExit("Agents from the last run are still finishing their calls in this session, so the script was "
                     "not loaded again. Wait a minute (status() shows where they are), then run this cell again.")

# ==================================================================================================== settings
BASE = "/content/drive/MyDrive/harrison_v7"
RUN = "pilot"                      # the run folder is step3_<RUN>; a new name starts a new run
ACTS = "pilot"                     # "pilot" (PILOT_ACTS), "all" (all 52, in nomination-year order) or a list of names
PILOT_ACTS = ["Big Country"]
INPUT_DIR = "step3_inputs"
PROMPT_FILES = {1: "prompt1.txt", 2: "prompt2.txt", "2b": "prompt2b.txt", 3: "prompt3.txt", 4: "prompt4.txt"}
GUIDE_FILE = "best_practices_guide.docx"
EXCLUDE_FILE = "excluded_turns.csv"   # in INPUT_DIR: turns round 2 leaves out (Turn ID, Act, Reason)
EXCLUDE_MIN_AGENTS = 2             # make_exclusion_list(): a turn goes on the list when this many agents marked it

MODELS = {   # which model is each agent; effort is the thinking level (Claude at low, the other two at medium)
    "A": {"provider": "anthropic", "model": "claude-opus-5-5", "effort": "low", "max_output": 128000},
    "B": {"provider": "google", "model": "gemini-4-argon", "effort": "medium", "max_output": 65536},
    "C": {"provider": "openai", "model": "gpt-6.1-sol", "effort": "medium", "max_output": 128000},
}
PRICES = {   # US dollars per million tokens, as published on 25 October 2026
    "claude-opus-5-5": {"input": 4.00, "output": 20.00, "cache_read": 0.20, "cache_write_5m": 5.00,
                        "cache_write_1h": 8.00},
    "gemini-4-argon": {"input": 2.00, "output": 10.00, "cache_read": 0.10,
                       "from": {"2026-12-01": {"input": 4.00, "output": 20.00, "cache_read": 0.20}}},
    "gpt-6.1-sol": {"input": 2.00, "output": 10.00, "cache_read": 0.10, "cache_write": 2.50,
                    "long_input": 272000, "long_input_x": 2.0, "long_output_x": 1.5},
}
CONTEXT = {  # each model's context window, in its own tokens (the prompt plus room for max_output), 25 October 2026
    "claude-opus-5-5": 1000000, "gemini-4-argon": 1048576, "gpt-6.1-sol": 1050000}
SECRETS = {"anthropic": "ANTHROPIC_API_KEY", "google": "GEMINI_API_KEY", "openai": "OPENAI_API_KEY"}
BUDGET_USD = 800.0                 # the run stops when the three agents' total cost in this run passes this; for the
                                   # full run, set it from the pilot report's projection, with some room

SEED = 20261026                    # the random half of the round 5 sample
TARGET_WORDS = 3000                # whole interviews per call, small ones grouped up to about this many words
SPLIT_WORDS = 8000                 # an interview longer than this is coded in parts of consecutive turns, and no
MAX_TURNS = 150                    # call has more than MAX_TURNS turns, so that a reply is never cut off
SAMPLE_SIZE = 60                   # turns in the round 5 sample: half at random, half where the agents differed most
RETRIES_UNUSABLE = 2               # re-asks when a reply cannot be used (not JSON, records missing, numbering wrong)
RETRIES_PROBLEMS = 1               # re-asks when a reply can be used but a check finds problems; then it is flagged
TRANSIENT_RETRIES = 8              # after network or server errors (once the SDKs' own retries are used up), waits
FIRST_WAIT, MAX_WAIT = 30, 600     # 30 seconds, then twice as long each time up to 10 minutes, and tries again: about
                                   # 45 minutes in all, so that a provider's outage does not stop an agent overnight
OTHER_RETRIES = 4                  # other errors that are not fatal: tried again only 4 times (about 7 minutes)
MAX_TOOL_STEPS = 60                # memo rounds: most times an agent may use its folder in one memo,
READ_BUDGET = 800000               # and most characters it may read from it in all (about 200,000 tokens)
READ_LINES, READ_CHARS = 300, 60000                    # and in one read
GPT_CHUNK_TOKENS = 4000            # GPT caches the memory in chunks of at least this many tokens, each with its own
GPT_MAX_CHUNKS = 40                # cache breakpoint; past this many chunks, the chunks double in size
READ_ACTS = True                   # round 1: each agent reads each whole act (without Before) and writes a few notes,
                                   # which go into its memory just before it codes that act
READ_BATCH = True                  # the reading goes to each provider as a batch, at half price, back within 24 hours
READ_WAIT_MINUTES = 30             # run_round1() waits this long for the batches; running it again collects the rest
BATCH_POLL_SECONDS = 60
GEMINI_BATCH_BYTES = 8000000       # Gemini takes batch requests inline, so they go in parts of about this size
GEMINI_BATCH_MAX_OUTPUT = 32768    # Gemini refuses a batch request that allows more output tokens than this
MEMO_POINTS = 5                    # the review: most points in the memo on an act
REVIEW_EXAMPLES = 3                # the review: codes shown for each label the act created
PARTS = 3                          # the full run's round 2 in this many parts, each about a third of the turns:
                                   # run_round2(part=1), then part=2 and part=3 (show_parts() lists them)
TOKEN_RATIO = {"anthropic": 1.5, "google": 1.1, "openai": 1.0}   # each model's tokens for each token counted here
LIST_ROOM = 0.85                   # if, with the labels list in full, an act's calls could fill more than this share of
NEW_LABELS_ROOM = 40000            # the context window (allowing this many tokens for the act's own labels and notes),
SHORT_DEF_WORDS = 12               # the earlier labels are listed with only this many words of each definition, or,
                                   # if that is still too long, without definitions (labels_level)

# ==================================================================================================== basics
_G = globals()                     # with %run -i, loading the script again keeps the same locks and stop signal
PRINT_LOCK = _G.get("PRINT_LOCK") or threading.Lock()
COST_LOCK = _G.get("COST_LOCK") or threading.Lock()
FILE_LOCK = _G.get("FILE_LOCK") or threading.Lock()
STOP = _G.get("STOP") or threading.Event()
STOP_WHY = _G.get("STOP_WHY") or ["stopped by hand"]      # why STOP was set: by hand, or the budget
THREAD_PREFIX = "step3-agent"
COSTS, CACHED_CALLS = {}, {}
LETTERS = list(MODELS)
PERIOD = {"before": "before the nomination year", "nomination year": "in the nomination year",
          "after": "after the nomination year"}


class AgentStop(Exception):
    """Stops one agent's run, with a message saying why; the other agents carry on."""


def stop_reason():
    return AgentStop(STOP_WHY[0])


def now():
    return datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


_LOCAL = threading.local()         # each agent thread keeps the run folder it started with, whatever RUN is set to later


def rpath(*p):
    base, run = getattr(_LOCAL, "base", None) or BASE, getattr(_LOCAL, "run", None) or RUN
    return os.path.join(base, f"step3_{run}", *p)


def say(*a):
    """Prints a line, and keeps it in step3_<RUN>/run_log.txt, so the log survives a lost Colab page."""
    line = " ".join(str(x) for x in a)
    with PRINT_LOCK:
        print(line, flush=True)
        try:
            if os.path.isdir(rpath()):
                with open(rpath("run_log.txt"), "a", encoding="utf-8") as f:
                    f.write(f"{now()}  {line}\n")
        except Exception:
            pass


def ensure(d):
    os.makedirs(d, exist_ok=True)
    return d


def write_text(path, s):
    ensure(os.path.dirname(path))
    tmp = f"{path}.{os.getpid()}.{threading.get_ident()}.tmp"     # one per thread: agents never share one
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(s)
    os.replace(tmp, path)


def write_json(path, obj):
    write_text(path, json.dumps(obj, ensure_ascii=False, indent=1))


def read_json(path, default=None):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def append_jsonl(path, obj):
    ensure(os.path.dirname(path))
    with FILE_LOCK, open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False) + "\n")


def read_jsonl(path):
    out = []
    if os.path.exists(path):
        with FILE_LOCK, open(path, encoding="utf-8") as f:
            for line in f:
                try:
                    out.append(json.loads(line))
                except Exception:
                    pass
    return out


def sha(s):
    return hashlib.sha256(s.encode("utf-8")).hexdigest()


def short(s, n=160):
    s = re.sub(r"\s+", " ", str(s)).strip()
    return s if len(s) <= n else s[:n - 1] + "…"


_ENC = None


def est_tokens(s):
    """Tokens counted before a call: o200k_base where tiktoken can load it, otherwise characters / 4."""
    global _ENC
    if _ENC is None:
        try:
            import tiktoken
            _ENC = tiktoken.get_encoding("o200k_base")
        except Exception:
            _ENC = False
    if _ENC:
        try:
            return len(_ENC.encode(s, disallowed_special=()))
        except Exception:
            pass
    return len(s) // 4


def secret(name):
    """A key from Colab Secrets. Quotation marks, spaces and invisible characters copied with it are dropped (a key
    never has them); a key that still has other characters is refused with a clear message, and never printed."""
    v = None
    try:
        from google.colab import userdata
        v = userdata.get(name)
    except Exception:
        pass
    v = v or os.environ.get(name)
    if not v:
        raise AgentStop(f"the key {name} is not in Colab Secrets (or notebook access is off)")
    v = re.sub(r"[\s ​-‏⁠﻿\"'`‘’“”]", "", v)
    if not v or not v.isascii():
        raise AgentStop(f"the key {name} in Colab Secrets has characters a key never has, perhaps copied with it from "
                        f"a document; paste it again, straight from the provider's page")
    return v


def busy():
    """Agent threads still running, for example after Stop while their calls finish."""
    return sorted(t.name for t in threading.enumerate() if t.name.startswith(THREAD_PREFIX) and t.is_alive())

# ==================================================================================================== corpus


CORPUS = {"acts": None, "turns": {}, "all_turns": {}, "by_id": {}, "stem_of": {}, "excluded": {}, "excluded_unknown": []}


def act_stem(a):
    return a["File"][:-6] if a["File"].endswith(".jsonl") else a["File"]


def exclude_path():
    return os.path.join(BASE, INPUT_DIR, EXCLUDE_FILE)


def read_exclusions():
    """The turns round 2 leaves out: step3_inputs/excluded_turns.csv, a Turn ID per row (other columns are notes)."""
    p = exclude_path()
    if not os.path.exists(p):
        return {}
    out = {}
    with open(p, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            tid = str(row.get("Turn ID") or "").strip()
            if tid:
                out[tid] = {k: (v or "") for k, v in row.items() if k}
    return out


def load_corpus():
    """The corpus, act by act in nomination-year order. turns: what round 2 codes (without the excluded turns);
    all_turns: every turn, as the reading round shows them; by_id: every turn by its Turn ID."""
    if CORPUS["acts"] is not None:
        return CORPUS
    j = read_json(os.path.join(BASE, "harrison_v7_band_voice_acts.json"))
    if j is None:
        raise SystemExit("STOPPED: harrison_v7_band_voice_acts.json not found in " + BASE)
    acts = j["acts"]
    acts = [acts[i] for i in sorted(range(len(acts)), key=lambda i: (int(acts[i]["Nominated"]), i))]
    excl = read_exclusions()
    turns, all_turns, by_id, stem_of = {}, {}, {}, {}
    for a in acts:
        p = os.path.join(BASE, "band_voice_by_act", a["File"])
        with open(p, encoding="utf-8") as f:
            rows = [json.loads(line) for line in f if line.strip()]
        rows.sort(key=lambda r: r["Order"])
        if len(rows) != int(a["Turns"]):
            raise SystemExit(f"STOPPED: {a['File']} has {len(rows)} turns; the act list says {a['Turns']}.")
        for r in rows:
            if r["Turn ID"] in by_id:
                raise SystemExit(f"STOPPED: Turn ID {r['Turn ID']} appears twice in the corpus.")
            if not (r.get("Text") or "").strip():
                raise SystemExit(f"STOPPED: Turn {r['Turn ID']} has no text.")
            by_id[r["Turn ID"]] = r
            stem_of[r["Turn ID"]] = act_stem(a)
        all_turns[act_stem(a)] = rows
        kept = [r for r in rows if r["Turn ID"] not in excl]
        if not kept:
            raise SystemExit(f"STOPPED: {EXCLUDE_FILE} leaves out every turn of {a['Act']}.")
        turns[act_stem(a)] = kept
    unknown = sorted(t for t in excl if t not in by_id)
    if unknown:
        print(f"PROBLEM: {len(unknown)} Turn ID(s) in {EXCLUDE_FILE} are not in the corpus, e.g. "
              f"{', '.join(unknown[:5])}. A spreadsheet may have turned them into dates; round 2 will not start until "
              f"they are put right.")
    CORPUS.update(acts=acts, turns=turns, all_turns=all_turns, by_id=by_id, stem_of=stem_of, excluded_unknown=unknown,
                  excluded={t: v for t, v in excl.items() if t in by_id})
    return CORPUS


def fresh_corpus():
    """Reads the corpus and the list of excluded turns again (unless agents are still running), so that a change to
    step3_inputs/excluded_turns.csv is seen without loading the script again."""
    if not busy():
        CORPUS["acts"] = None
    return load_corpus()


def corpus_fingerprint(coded=True, stems=None):
    """A fingerprint of the corpus: with coded, of the turns round 2 codes (so a change to the excluded turns counts as
    another corpus); without, of every turn, as the reading round reads them. With stems, of those acts only (so that
    turns excluded in other acts do not change a run of these)."""
    c = load_corpus()
    h = hashlib.sha256()
    for a in c["acts"]:
        if stems is not None and act_stem(a) not in stems:
            continue
        for t in c["turns" if coded else "all_turns"][act_stem(a)]:
            h.update((t["Turn ID"] + "\t" + t["Text"] + "\n").encode("utf-8"))
    return h.hexdigest()


def act_of_turn(tid):
    """The stem of the act a turn belongs to."""
    return load_corpus()["stem_of"].get(tid)


def select_acts(acts=None):
    c = load_corpus()
    acts = ACTS if acts is None else acts
    if acts == "all":
        return list(c["acts"])
    names = PILOT_ACTS if acts == "pilot" else ([acts] if isinstance(acts, str) else list(acts))
    by_name = {a["Act"].casefold(): a for a in c["acts"]}
    out = []
    for n in names:
        if n.casefold() not in by_name:
            raise SystemExit(f"STOPPED: no act called {n!r} in the corpus.")
        out.append(by_name[n.casefold()])
    order = {act_stem(a): i for i, a in enumerate(c["acts"])}
    return sorted(out, key=lambda a: order[act_stem(a)])


def kind(a):
    return "solo artist" if str(a.get("Solo or group", "")).lower().startswith("solo") else "group"


def fmt_date(t):
    d, p = (t.get("Date") or "").strip(), (t.get("Date precision") or "").lower()
    if not d:
        return "unknown"
    return d[:4] if p == "year" else d[:7] if p == "month" else d[:10]


def render_turn(t, with_act=False, with_before=True):
    src = ["article" if t.get("Source") == "text" else "video"]
    src += [str(x).strip() for x in (t.get("Outlet"), t.get("Title")) if x and str(x).strip()]
    before = (t.get("Before") or "").strip()
    lines = [f"Turn ID: {t['Turn ID']}"] + ([f"Act: {t['Act']}"] if with_act else []) + [
        f"Speaker: {t.get('Speaker') or 'unknown'}", f"Date: {fmt_date(t)}",
        f"Period: {PERIOD.get(t.get('Period'), 'unknown')}", "Source: " + "; ".join(src)] + (
        [f"Before: {before or '(nothing)'}"] if with_before else []) + [f"Text: {t['Text']}"]
    return "\n".join(lines)


def render_turns(ts, with_act=False):
    return "\n\n".join(render_turn(t, with_act) for t in ts)


def render_turns_reading(ts):
    """The turns as the reading round shows them: everything but Before (the turns of an act are read in order, so
    what each one follows is there to see)."""
    return "\n\n".join(render_turn(t, with_before=False) for t in ts)


def n_interviews(turns):
    return len({t["Record ID"] for t in turns})


def words(t):
    return t.get("Words") or len(t["Text"].split())


def make_batches(turns):
    """Whole interviews per batch, in date order, small ones grouped up to TARGET_WORDS (and MAX_TURNS); a very long
    interview is split into parts of consecutive turns. Every turn is in exactly one batch."""
    groups = []
    for t in turns:
        if groups and groups[-1][-1]["Record ID"] == t["Record ID"]:
            groups[-1].append(t)
        else:
            groups.append([t])
    units = []
    for g in groups:
        if sum(map(words, g)) <= SPLIT_WORDS and len(g) <= MAX_TURNS:
            units.append(g)
            continue
        cur, w = [], 0
        for t in g:
            if cur and (w + words(t) > SPLIT_WORDS or len(cur) >= MAX_TURNS):
                units.append(cur)
                cur, w = [], 0
            cur.append(t)
            w += words(t)
        units.append(cur)
    batches, cur, w, nt = [], [], 0, 0
    for u in units:
        uw = sum(map(words, u))
        if cur and (w + uw > TARGET_WORDS or nt + len(u) > MAX_TURNS):
            batches.append(cur)
            cur, w, nt = [], 0, 0
        cur.append(u)
        w += uw
        nt += len(u)
    if cur:
        batches.append(cur)
    out = [{"k": len(b), "turns": [t for u in b for t in u]} for b in batches]
    assert [t["Turn ID"] for b in out for t in b["turns"]] == [t["Turn ID"] for t in turns]
    return out

# ==================================================================================================== prompts and guide


PROMPTS, GUIDE = {}, {"text": None}
PH = re.compile(r"\{([^{}\"\n]+)\}")
REQUIRED = {1: {"GUIDE", "LETTER", "ACT", "KIND", "N_INTERVIEWS", "N_TURNS", "TURNS"},
            2: {"GUIDE", "LETTER", "MEMORY", "NEXT", "LABELS", "ACT", "KIND", "BATCH_N", "BATCH_TOTAL",
                "N_INTERVIEWS", "N_TURNS", "TURNS"},
            "2b": {"ACT", "KIND", "ACT_LABELS"},
            3: {"GUIDE", "LETTER", "MEMORY"},
            4: {"GUIDE", "LETTER", "MEMO1", "FIGURES", "P4SAMPLE"}}


def ph_key(inner):
    s = inner.strip()
    low = s.lower()
    if "best practices" in low:
        return "GUIDE"
    if "definitions and codes" in low:
        return "ACT_LABELS"
    if re.fullmatch(r"a\s*[,/]\s*b\s*(,\s*)?(or|/)\s*c", low):
        return "LETTER"
    if "100 turns" in low:
        return "SAMPLE"
    if "memory so far" in low:
        return "MEMORY"
    if "labels so far" in low:
        return "LABELS"
    if "memo part 1" in low:
        return "MEMO1"
    if "agreement figures" in low:
        return "FIGURES"
    if "sample of turns" in low:
        return "P4SAMPLE"
    if low == "the turns":
        return "TURNS"
    if low == "next":
        return "NEXT"
    if low == "name":
        return "ACT"
    if "group or solo" in low or "solo artist" in low:
        return "KIND"
    return {"n": "BATCH_N", "N": "BATCH_TOTAL", "k": "N_INTERVIEWS", "m": "N_TURNS"}.get(s)


def keys_in(text):
    return [ph_key(m.group(1)) for m in PH.finditer(text)]


def fill(template, values):
    """Fills the {placeholders} in one pass, so that braces inside the filled text are never touched."""
    def sub(m):
        k = ph_key(m.group(1))
        return str(values[k]) if k in values else m.group(0)
    return PH.sub(sub, template)


def split_at_memory(template):
    ms = [m for m in PH.finditer(template) if ph_key(m.group(1)) == "MEMORY"]
    return template[:ms[0].start()], template[ms[0].end():]


def read_prompt(n):
    p = os.path.join(BASE, INPUT_DIR, PROMPT_FILES[n])
    if not os.path.exists(p):
        raise SystemExit(f"STOPPED: {p} not found.")
    with open(p, encoding="utf-8-sig") as f:
        s = f.read().replace("\r\n", "\n").replace("\r", "\n")
    lines = [ln.rstrip() for ln in s.split("\n")]
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and re.match(r"\s*Prompt\s*\d+[a-z]?\s*:", lines[0], re.I):
        lines.pop(0)
    return "\n".join(lines).strip("\n")


def check_prompt(n, s):
    problems = []
    found = [m.group(0) for m in PH.finditer(s)]
    ks = keys_in(s)
    if n == 1 and "SAMPLE" in ks:
        return ["this is the old Prompt 1 (100 random turns). From version 1.6, round 1 reads each act in full, so "
                "Prompt 1 is the reading prompt: upload the new prompt1.txt"]
    if n == 1 and '"data_problems"' not in s:
        return ["this is the Prompt 1 of version 1.6, without the list of data problems: upload the v7 prompt1.txt"]
    if n == 2 and ('"label_from"' in s or '"confidence"' in s or '"new_labels"' not in s):
        return ["this is the Prompt 2 of version 1.6 (with label_from and confidence): upload the v7 prompt2.txt"]
    paste = [f for f in found if f[1:].lower().startswith("paste here")]
    if paste:
        return [f"paste the passage into {paste[0][:60]}… (the slot is still in the prompt)"]
    unknown = [f for f, k in zip(found, ks) if k is None]
    if unknown:
        problems.append(f"placeholders the script does not know: {unknown}")
    missing = REQUIRED[n] - set(ks)
    if missing:
        problems.append(f"placeholders missing: {sorted(missing)}")
    extra = set(k for k in ks if k) - REQUIRED[n]
    if extra:
        problems.append(f"placeholders that belong to other prompts: {sorted(extra)}")
    for k in ("GUIDE", "MEMORY", "TURNS", "SAMPLE", "LABELS", "MEMO1", "FIGURES", "P4SAMPLE", "ACT_LABELS"):
        if ks.count(k) > 1:
            problems.append(f"{k} appears {ks.count(k)} times")
    if n == 2 and ks.count("MEMORY") == 1:
        pre, _ = split_at_memory(s)
        bad = set(k for k in keys_in(pre) if k) - {"GUIDE", "LETTER"}
        if bad:
            problems.append(f"these change in every call, so they must come after the memory: {sorted(bad)}")
    if n in (1, 2, "2b") and "JSON" not in s:
        problems.append("no JSON format in the prompt")
    return problems


def load_prompts():
    if not PROMPTS:
        for n in PROMPT_FILES:
            s = read_prompt(n)
            pr = check_prompt(n, s)
            if pr:
                raise SystemExit(f"STOPPED: {PROMPT_FILES[n]}: " + "; ".join(pr))
            PROMPTS[n] = s
    return PROMPTS


def docx_text(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") or style == "title":
                lvl = re.findall(r"\d+", style)
                out.append("#" * min(int(lvl[0]) if lvl else 1, 6) + " " + t)
            elif "list" in style:
                out.append("- " + t)
            else:
                out.append(t)
        elif tag == "tbl":
            for row in Table(el, d).rows:
                cells = []
                for c in row.cells:
                    ct = " ".join(x.text.strip() for x in c.paragraphs if x.text.strip())
                    if not cells or ct != cells[-1]:
                        cells.append(ct)
                if any(cells):
                    out.append(" | ".join(cells))
    return "\n\n".join(out)


def load_guide():
    if GUIDE["text"] is None:
        p = os.path.join(BASE, INPUT_DIR, GUIDE_FILE)
        if not os.path.exists(p):
            raise SystemExit(f"STOPPED: the guide {p} not found (set GUIDE_FILE).")
        if p.lower().endswith(".docx"):
            g = docx_text(p)
        else:
            with open(p, encoding="utf-8-sig") as f:
                g = f.read().replace("\r\n", "\n").strip()
        if len(g.split()) < 500:
            raise SystemExit(f"STOPPED: the guide has only {len(g.split())} words; is it the right file?")
        GUIDE["text"] = g
    return GUIDE["text"]

# ==================================================================================================== run settings record


def round_settings(r, stems=None):
    load_prompts()
    s = {"script": VERSION, "seed": SEED}
    if r in (1, 2, 3, 5):
        s["models"] = MODELS
        s["guide_sha"] = sha(load_guide())
        s["prompt_sha"] = sha(PROMPTS[2] + "\n" + PROMPTS["2b"]) if r == 2 else sha(PROMPTS[{1: 1, 3: 3, 5: 4}[r]])
    if r in (1, 2):
        s["corpus"] = corpus_fingerprint(coded=(r == 2), stems=stems if r == 2 else None)
    if r == 2:
        s["excluded_turns"] = len([t for t in load_corpus()["excluded"]
                                   if stems is None or act_of_turn(t) in stems])
    if r == 2:
        s["batching"] = [TARGET_WORDS, SPLIT_WORDS, MAX_TURNS]
    if r in (4, 5):
        s["sample_size"] = SAMPLE_SIZE
    return s


def round_started(r):
    for L in LETTERS:
        st = load_state(L)
        if (r == 1 and st["round1"]) or (r == 2 and any(st["done"].values())) or (r == 3 and st["memo1"]) \
                or (r == 5 and st["memo23"]):
            return True
        if r == 1 and (os.path.exists(rpath(f"agent_{L}", "round1_memory.json"))
                       or os.path.isdir(rpath(f"agent_{L}", "reading"))):
            return True
        if r == 2 and any(fn.endswith(".json") for _, _, fns in os.walk(rpath(f"agent_{L}", "batches")) for fn in fns):
            return True
    return r == 4 and os.path.exists(rpath("merged", "done.json"))


def check_settings(r, extra=None, agents=None):
    """The settings a round started with are kept: once a round has work saved, it never carries on with other
    prompts, guide, models, batching or acts. Before that, the settings are simply recorded again. Each time the
    round is run, the script version and the agents are added to its "runs", so that the record shows which version
    did which part (the round's own "script" is the version it started with)."""
    path = rpath("run_settings.json")
    rec = read_json(path, {})
    cur = dict(round_settings(r, (extra or {}).get("acts")), **(extra or {}))
    old = rec.get(str(r))
    run = {"script": VERSION, "agents": list(agents or LETTERS) if r != 4 else [], "time": now()}
    if old is None or not round_started(r):
        rec[str(r)] = dict(cur, started=now(), runs=[run])
        write_json(path, rec)
        return
    diff = [k for k in cur if k != "script" and old.get(k) != cur[k]]
    if diff:
        names = {"seed": "SEED", "models": "MODELS", "guide_sha": f"guide ({GUIDE_FILE})",
                 "prompt_sha": "prompts (" + {1: "prompt1.txt", 2: "prompt2.txt and prompt2b.txt", 3: "prompt3.txt",
                                               5: "prompt4.txt"}.get(r, "") + ")",
                 "corpus": "corpus" if r != 2 else f"turns to code (the corpus, less the turns in {EXCLUDE_FILE})",
                 "excluded_turns": f"number of excluded turns ({EXCLUDE_FILE})",
                 "batching": "batching (TARGET_WORDS, SPLIT_WORDS, MAX_TURNS)", "sample_size": "SAMPLE_SIZE",
                 "acts": "acts"}
        raise SystemExit(f"STOPPED: round {r} of run {RUN!r} started with other settings: "
                         f"{'; '.join(names.get(k, k) for k in diff)}. Put back what it started with, or set RUN to a "
                         f"new name to start a new run.")
    old.setdefault("runs", []).append(run)
    write_json(path, rec)

# ==================================================================================================== agent state


def new_state(letter):
    """memory: every note, numbered M1 on; act_memos: the review of each act (its memo and the notes kept), which is
    what the agent carries into later acts; act_start: the act being coded and its labels as they were at its start,
    which is the labels list every call of the act shows (so that it is read from the cache)."""
    return {"agent": letter, "round1": False, "memory": [], "blocks": [], "next": 1, "labels": {}, "done": {},
            "acts_complete": [], "act_memos": {}, "act_start": {"stem": None, "labels": {}},
            "memo1": False, "memo23": False, "cost": 0.0, "calls": 0, "flags": 0, "stopped": None, "label_levels": {}}


def load_state(letter):
    path = rpath(f"agent_{letter}", "state.json")
    try:
        st = read_json(path)
    except Exception:                          # a damaged file: kept aside, and the state rebuilt
        st = None
        shutil.move(path, path[:-5] + f"_damaged_{datetime.datetime.now():%Y%m%d_%H%M%S}.json")
    if st is None:
        st = new_state(letter)
        if os.path.isdir(rpath(f"agent_{letter}")):
            recover_state(st)
            save_state(st)
    for k, v in new_state(letter).items():
        st.setdefault(k, v)
    return st


def logged_calls(letter):
    """Every call in the agent's call logs."""
    d = rpath(f"agent_{letter}", "calls")
    return [c for fn in (sorted(os.listdir(d)) if os.path.isdir(d) else []) if fn.endswith(".jsonl")
            for c in read_jsonl(os.path.join(d, fn))]


def call_numbers(letter, calls):
    return [int(str(c.get("call", "")).rsplit("-", 1)[-1]) for c in calls
            if re.fullmatch(rf"{letter}-\d+", str(c.get("call", "")))]


def sync_counters(st):
    """After a crash, Drive can keep an older state.json than the call logs and memo files. Where those are ahead,
    the number of calls, their cost and the memos written are taken from them, so that call IDs never repeat, the
    budget counts every call and no memo is paid for twice. Returns True if anything changed."""
    L = st["agent"]
    calls = logged_calls(L)
    n = max(call_numbers(L, calls) + [0])
    cost = round(sum(c.get("cost", 0) or 0 for c in calls), 6)
    new = {}
    if n > st["calls"]:
        new["calls"] = n
    if cost > st["cost"] + 0.005:
        new["cost"] = cost
    if not st["memo1"] and os.path.exists(rpath(f"agent_{L}", "memo_part1.md")):
        new["memo1"] = True
    if not st["memo23"] and os.path.exists(rpath(f"agent_{L}", "memo_parts2and3.md")):
        new["memo23"] = True
    if new:
        st.update(new)
        say(f"[{L}] state.json was behind the agent's call logs or memo files (Drive may have kept an older copy), so "
            f"its counts were taken from them: {st['calls']} calls, ${st['cost']:.2f}"
            + (", memo part 1 written" if new.get("memo1") else "")
            + (", memo parts 2 and 3 written" if new.get("memo23") else "") + ".")
    return bool(new)


def recover_state(st):
    """When an agent's state.json is lost or damaged (Drive can lose the last write in a crash), its state is rebuilt
    from its other files: the cost and the number of calls from its call logs, round 1 from its reading files, its
    coding (and the reading notes in its memory) from its batch files, and the memos from the memo files."""
    letter = st["agent"]
    calls = logged_calls(letter)
    st.update(cost=round(sum(c.get("cost", 0) or 0 for c in calls), 6),
              calls=max(call_numbers(letter, calls) + [len(calls)]))
    rs = read_json(rpath("run_settings.json"), {})
    r1 = read_json(rpath(f"agent_{letter}", "round1_memory.json"))      # runs made with versions before 1.6
    if r1 is not None:
        st.update(round1=True, memory=r1, blocks=[[0, len(r1)]] if r1 else [], next=len(r1) + 1)
    acts1 = rs.get("1", {}).get("acts") or []
    if acts1 and all(os.path.exists(reading_path(letter, s)) for s in acts1):
        st["round1"] = True
    stems = rs.get("2", {}).get("acts") or []
    if stems:
        by_stem = {act_stem(a): a for a in load_corpus()["acts"]}
        rebuild_state(st, [by_stem[s] for s in stems if s in by_stem], move=False)
    st["memo1"] = os.path.exists(rpath(f"agent_{letter}", "memo_part1.md"))
    st["memo23"] = os.path.exists(rpath(f"agent_{letter}", "memo_parts2and3.md"))
    st["flags"] = len(read_jsonl(rpath(f"agent_{letter}", "flags.jsonl")))
    say(f"[{letter}] state.json was missing or damaged, so it was rebuilt from the agent's other files: "
        f"{st['calls']} calls, ${st['cost']:.2f}, {sum(len(v) for v in st['done'].values())} batches coded.")


def save_state(st):
    write_text(rpath(f"agent_{st['agent']}", "state.json"), json.dumps(st, ensure_ascii=False, separators=(",", ":")))


def total_cost():
    with COST_LOCK:
        if COSTS:
            return sum(COSTS.get(L, 0.0) for L in LETTERS)
    return sum(load_state(L)["cost"] for L in LETTERS)


def flag(st, round_no, act, batch, items, call_id, kind="check"):
    """Keeps problems in flags.jsonl: kind "check" for a problem the checks found that remained, "data" for a data
    problem an agent reported (a turn that seems to be someone else's words, or garbled text)."""
    for p in items:
        append_jsonl(rpath(f"agent_{st['agent']}", "flags.jsonl"),
                     {"agent": st["agent"], "round": round_no, "act": act, "batch": batch, "kind": kind, "problem": p,
                      "call": call_id, "time": now()})
    st["flags"] += len(items)


def entry_act(m):
    """The act a memory entry belongs to (entries from version 1.6 carry it as their round)."""
    return m.get("act") or (m.get("round") if m.get("round") not in ("reading", "familiarisation") else None)


def act_entries(st, stem):
    return [m for m in st["memory"] if entry_act(m) == stem]


def act_name(stem):
    a = next((x for x in load_corpus()["acts"] if act_stem(x) == stem), None)
    return f"{a['Act']} ({a['Nominated']})" if a else stem


def act_block(st, stem):
    """What the agent carries from an act it has reviewed: its memo on the act and the notes it kept (from an act
    coded before version 1.7, without a review: all its notes on it)."""
    entries = act_entries(st, stem)
    am = st["act_memos"].get(stem)
    if am is None:
        return f"Notes on {act_name(stem)}:\n" + "\n".join(m["text"] for m in entries)
    keep = set(am.get("keep") or [])
    lines = [f"Memo on {act_name(stem)}:"] + [f"{i}. {memo_point(p)}" for i, p in enumerate(am.get("memo") or [], 1)]
    kept = [m["text"] for m in entries if m["n"] in keep]
    if kept:
        lines += [f"Notes kept from {act_name(stem)}:"] + kept
    return "\n".join(lines)


def memory_blocks(st):
    """The memory every call carries: one block for each act the agent has coded and reviewed, in order. A block
    never changes once written, so each call reads all of them but a new one from the cache."""
    out = []
    pre = [m for m in st["memory"] if m.get("round") == "familiarisation"]       # runs before version 1.6
    if pre:
        out.append("Notes from before coding:\n" + "\n".join(m["text"] for m in pre))
    out += [act_block(st, s) for s in st["acts_complete"]]
    if not out:
        return ["Your memory so far: none yet, as this is your first act."]
    return ["Your memory so far, from the acts you have coded:\n" + out[0]] + ["\n\n" + b for b in out[1:]]


def memory_text(st):
    return "".join(memory_blocks(st))


def act_memory_text(st, stem):
    """The notes of the act being coded: its reading notes and what its batches have added. They come after the
    memory and the labels list, and change from call to call."""
    es = act_entries(st, stem)
    if not es:
        return "\n\nYour notes from this act so far: none yet."
    return "\n\nYour notes from this act so far:\n" + "\n".join(m["text"] for m in es)


def all_notes_text(st):
    """Every note the agent has written, act by act, with the memo on each act it has reviewed (notes_all.txt)."""
    out = []
    for stem in dict.fromkeys(entry_act(m) for m in st["memory"]):
        am = st["act_memos"].get(stem)
        out.append(f"== {act_name(stem) if stem else 'before coding'}\n" + "\n".join(
            m["text"] for m in st["memory"] if entry_act(m) == stem)
            + (("\nMemo:\n" + "\n".join(f"{i}. {memo_point(p)}" for i, p in enumerate(am["memo"], 1))
                + f"\nKept: {', '.join(f'M{n}' for n in am.get('keep') or []) or 'none'}") if am else ""))
    return "\n\n".join(out) + "\n"


_SIZES = {}


def _size(text):
    """est_tokens, kept for texts seen before (the fixed part, the memory blocks, full chunks)."""
    k = hashlib.sha1(text.encode("utf-8")).hexdigest()
    if k not in _SIZES:
        _SIZES[k] = est_tokens(text)
    return _SIZES[k]


def stable_chunks(blocks, size=None, cap=None):
    """The memory blocks (one per act), joined in order into chunks of at least `size` tokens; the last chunk is
    the one still filling. The text is the same; only the cut points differ. As the memory only grows at the end, a
    full chunk is the same in every later call, and so is the place of the breakpoint after it. If there would be
    more than `cap` chunks, the size doubles: the cut points change once, and the memory is cached again once."""
    size, cap = size or GPT_CHUNK_TOKENS, cap or GPT_MAX_CHUNKS
    sizes = [_size(b) for b in blocks]
    while True:
        chunks, cur, n = [], [], 0
        for b, s in zip(blocks, sizes):
            cur.append(b)
            n += s
            if n >= size:
                chunks.append("".join(cur))
                cur, n = [], 0
        if cur:
            chunks.append("".join(cur))
        if len(chunks) <= cap:
            return chunks
        size *= 2


def label_key(label):
    return re.sub(r"\s+", " ", str(label).casefold()).strip(" .,;:!?\"'“”‘’")


def _one(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


POINT_MARK = re.compile(r"^\**(?:(?:[Mm]emo\s+)?[Pp]oint\s+)?\(?(?:[1-9]|1\d)\s*[.):]\)?\**(?:\s+|(?=[A-Z“\"']))"
                        r"|^[-*•]\s+")


def memo_point(p):
    """A point of a memo on an act, without the number or bullet the agent may have put before it, as in "1. …" or
    "(2) …" (the script numbers the points itself); a year such as "1984: …" is not taken for a number."""
    s = _one(p)
    for _ in range(3):
        t = POINT_MARK.sub("", s, count=1).strip()
        if t == s:
            break
        s = t
    return s


def label_line(x, level=0):
    """A label as the lists show it: label | times used | its definition (its first description until it has one);
    at level 1 only the first SHORT_DEF_WORDS words of the definition, at level 2 no definition."""
    if level >= 2:
        return f"{x['label']} | {x['count']}"
    d = _one(x.get("definition") or x.get("description"))
    if level == 1 and len(d.split()) > SHORT_DEF_WORDS:
        d = " ".join(d.split()[:SHORT_DEF_WORDS]) + " …"
    return f"{x['label']} | {x['count']} | {d}"


def _sorted_labels(labels):
    return sorted(labels.values(), key=lambda x: (-x["count"], x["label"].casefold()))


LIST_HEADS = {0: "(label | times used | definition)",
              1: "(label | times used | the first words of its definition: the list is shortened to fit the model's "
                 "limit on the length of a prompt)",
              2: "(label | times used; the definitions are left out to fit the model's limit on the length of a prompt)"}


def start_labels(st):
    """The labels from earlier acts, with their uses when the act being coded began. The act's start keeps only the
    uses (its labels, names and definitions do not change until the act's own review), or, from versions before 1.9,
    a whole copy of each label."""
    start, out = (st.get("act_start") or {}).get("labels") or {}, {}
    for k, v in start.items():
        if isinstance(v, dict):
            out[k] = v
        elif k in st["labels"]:
            out[k] = dict(st["labels"][k], count=v)
    return out


def frozen_labels_text(st, level=None):
    """The labels from earlier acts as they stood when the act being coded began: the same in every call of the act
    (it is read from the cache), with the uses then; level (labels_level) is 0 unless the list must be shortened."""
    labs = start_labels(st)
    if not labs:
        return "\n\nYour labels from earlier acts: none yet."
    lvl = (st.get("act_start") or {}).get("level", 0) if level is None else level
    return (f"\n\nYour labels from earlier acts, as they stood at the start of this act {LIST_HEADS[lvl]}:\n"
            + "\n".join(label_line(x, lvl) for x in _sorted_labels(labs)))


def labels_level(st, stem, fixed):
    """How the labels from earlier acts are listed in an act's calls: 0 in full; 1 with each definition cut to its
    first words; 2 with names and uses only. A shorter form is used only when, in full, the act's biggest call could
    fill more than LIST_ROOM of what the model's context window leaves for the prompt (with room for the notes and
    the labels the act adds). It is worked out from the act's start, so it is the same for every call of the act."""
    cfg = MODELS[st["agent"]]
    room = (CONTEXT.get(cfg["model"], 10 ** 9) - cfg["max_output"]) * LIST_ROOM
    ratio = TOKEN_RATIO.get(cfg["provider"], 1.5)
    bs = make_batches(CORPUS["turns"][stem])
    rest = (est_tokens(fixed) + est_tokens(memory_text(st)) + NEW_LABELS_ROOM
            + max(est_tokens(render_turns(b["turns"])) for b in bs))
    for lvl in (0, 1):
        if (rest + est_tokens(frozen_labels_text(st, lvl))) * ratio <= room:
            return lvl
    return 2


def new_labels_text(st, stem):
    """The labels the agent has created in the act it is coding, with their uses so far."""
    items = {k: x for k, x in st["labels"].items() if x.get("first_act") == stem}
    if not items:
        return "Your labels so far from this act: none yet."
    return ("Your labels so far from this act (label | times used | definition):\n"
            + "\n".join(label_line(x) for x in _sorted_labels(items)))


def labels_text(st):
    """All the agent's labels (for reports and estimates)."""
    if not st["labels"]:
        return "Your labels so far: none yet."
    return ("Your labels so far (label | times used | definition):\n"
            + "\n".join(label_line(x) for x in _sorted_labels(st["labels"])))

# ==================================================================================================== checks


QUOTES = dict.fromkeys(map(ord, "\"'“”‘’`´«»"), None)
ELLIPSIS = re.compile(r"[\[(]?\s*(?:\.\s*\.\s*\.|…)\s*[\])]?")
MEM_RE = re.compile(r"^\s*M(\d+)\s*\(([^)]*)\)\s*:\s*(\S.*)$", re.S)
REF_RE = re.compile(r"\bM(\d+)\b")


def norm_words(s):
    """The words of a text, for comparing: accents and letter forms normalised, lower case, quotation marks and
    apostrophes dropped, other punctuation treated as a space."""
    s = unicodedata.normalize("NFKC", str(s)).casefold().translate(QUOTES)
    return re.findall(r"\w+", s)


BRACKETED = re.compile(r"\[[^\[\]]*\]")


def turn_words(text):
    """The words of a turn, as norm_words gives them, each with whether it lies inside square brackets: an editor's
    insertion, such as '[Adamson – guitar]' or '[laughs]', which a quotation may leave out."""
    s = unicodedata.normalize("NFKC", str(text)).casefold().translate(QUOTES)
    inside = bytearray(len(s) + 1)
    for m in BRACKETED.finditer(s):
        inside[m.start():m.end()] = b"\x01" * (m.end() - m.start())
    return [(m.group(0), bool(inside[m.start()])) for m in re.finditer(r"\w+", s)]


def _match_from(frag, tw, i):
    """Where the words of frag end if they match the turn's words from position i, words in square brackets being
    optional (they may be quoted or left out); -1 if they do not match."""
    todo, seen = [(0, i)], set()
    while todo:
        k, j = todo.pop()
        if k == len(frag):
            return j
        if j >= len(tw) or (k, j) in seen:
            continue
        seen.add((k, j))
        word, optional = tw[j]
        if optional:
            todo.append((k, j + 1))            # the quotation leaves this bracketed word out
        if word == frag[k]:
            todo.append((k + 1, j + 1))        # tried first
    return -1


def _seq_at(frag, tw, start=0):
    """Where the words of frag end, the first time they appear in a row in tw (from turn_words) from start, bracketed
    words being optional; -1 if they do not."""
    for i in range(start, len(tw)):
        if tw[i][0] == frag[0]:
            end = _match_from(frag, tw, i)
            if end >= 0:
                return end
    return -1


def evidence_in(ev, text):
    """The evidence is the turn's own words, in order: spacing, capitals, punctuation and quotation marks aside;
    '...' (also '[...]' or '(...)') may join parts of the turn, and an editor's insertion in square brackets in the
    turn may be left out."""
    tw = turn_words(text)
    frags = [f for f in (norm_words(x) for x in ELLIPSIS.split(ev)) if f]
    pos = 0
    for f in frags:
        pos = _seq_at(f, tw, pos)
        if pos < 0:
            return False
    return bool(frags)


def label_in(label, text):
    """A label from the data: its words are the turn's own words, checked as the evidence is ('...' may join parts
    of the turn, and bracketed insertions may be left out)."""
    return evidence_in(label, text)


def label_from(label, text, first_turn=None):
    """Worked out by the script (the agents no longer say it): "data" when the label's words are word for word in
    this turn, or in the turn where the label was first used; otherwise "own"."""
    if label_in(label, text):
        return "data"
    first = CORPUS["by_id"].get(first_turn, {}).get("Text", "") if first_turn else ""
    return "data" if first and label_in(label, first) else "own"


LABEL_STOP = set("a an the and or but of to in on at for with by from as is was were be been being it its this that "
                 "these those i we you he she they me us him her them my our your his their not no so do did just "
                 "about into over out up than then there what when where who why how all any some more most very can "
                 "could would should will one".split())
SOURCES = ("word for word", "built on their words", "own words")


def label_source(label, text, first_turn=None):
    """For the reports only: where a label's words come from, worked out by the script. "word for word" is "data" in
    the agents' sheets; of the rest, "built on their words" when at least half of the label's main words are in this
    turn or the turn where the label was first used, as with "never-planned look" for "we never planned it", and
    otherwise "own words"."""
    if label_from(label, text, first_turn) == "data":
        return SOURCES[0]
    first = CORPUS["by_id"].get(first_turn, {}).get("Text", "") if first_turn else ""
    have = set(norm_words(text)) | set(norm_words(first))
    main = [w for w in norm_words(label) if w not in LABEL_STOP and len(w) > 2]
    return SOURCES[1] if main and sum(w in have for w in main) / len(main) >= 0.5 else SOURCES[2]


def source_counts(st, recs):
    """How many of these records' codes have labels word for word, built on the members' words, or in the agent's own
    words (label_source)."""
    n = collections.Counter()
    for r in recs:
        text = CORPUS["by_id"].get(r["turn_id"], {}).get("Text", "")
        for c in r["codes"]:
            info = st["labels"].get(label_key(c["label"])) or {}
            n[label_source(c["label"], text, info.get("first_turn"))] += 1
    return n


def source_text(n):
    tot = max(sum(n.values()), 1)
    return ", ".join(f"{s} {n[s]} ({n[s] / tot:.0%})" for s in SOURCES)


def _bool(v):
    """true/false as the agents give it (Gemini, without a schema, may write it as a word)."""
    return v is True or str(v).strip().lower() in ("true", "yes", "1")


TID_RE = re.compile(r"\b(?:V\d+-\d+|\d{1,6}-\d+)\b")


_RECORDS = {"key": None, "ids": set()}


def corpus_records():
    """The record IDs of the corpus: a Turn ID is its article's or video's record ID, a hyphen and the turn's place."""
    by = CORPUS["by_id"]
    if _RECORDS["key"] != (id(by), len(by)):
        _RECORDS.update(key=(id(by), len(by)), ids={t.rsplit("-", 1)[0] for t in by})
    return _RECORDS["ids"]


def looks_like_tid(x):
    """Whether a Turn ID-shaped string that is not in the corpus is worth a problem: only when its record part is a
    record of the corpus (so "100-150 turns" is not), and it is not a span of years such as "1984-85", "1985-6" or
    "1984-1986" (records with those numbers exist, so a made-up Turn ID that looks like a span of years goes
    unnoticed: a rare case)."""
    a, b = x.rsplit("-", 1)
    if a not in corpus_records():
        return False
    return not (a.isdigit() and len(a) == 4 and 1900 <= int(a) <= 2099
                and (len(b) <= 2 or (len(b) == 4 and 1900 <= int(b) <= 2099)))
QUOTE_RE = re.compile(r"“([^”\n]{1,500})”|\"([^\"\n]{1,500})\"")     # curly or straight, each pair on its own
NAME_RE = re.compile(r"\b((?:Mc|Mac|O')?[A-Z][a-z]+(?:[-'][A-Z]?[a-z]+)?(?:\s+(?:Mc|Mac|O')?[A-Z][a-z]+"
                     r"(?:[-'][A-Z]?[a-z]+)?){1,3})\b")
NAME_SKIP = set("The A An And But In On At Of For To From By With As This That These Those Their His Her Its Our My "
                "Your When After Before During While Both Each Every Some Many Most Part Point Memo Notes Note Labels "
                "Label Agent Agents Turn Turns Act Acts Step Steps Later Early Here There It We They He She I You".split())
NAME_STARTERS = set("Whereas Although Though Despite Unlike Following However Meanwhile Moreover Instead Rather "
                    "Initially Eventually Ultimately Throughout Within Across Among Beyond Since Until Once Upon Over "
                    "Regarding Overall Thus Hence Still Yet Even Only Also Just Then Now Again Already Often Keep "
                    "Treat Read Compare Like".split())


def _act_haystack(stems):
    """The words the agent saw for these acts (text, Before, speaker, source), for the name check: an editor's square
    brackets are taken out, so that "Tommy [Paxton]" in a turn counts as "Tommy Paxton"."""
    c = load_corpus()
    s = "\n".join(f"{t.get('Text', '')}\n{t.get('Before', '')}\n{t.get('Speaker', '')}\n{t.get('Title', '')}\n"
                  f"{t.get('Outlet', '')}" for s in stems for t in c["all_turns"].get(s, []))
    return re.sub(r"[ \t ]+", " ", s.replace("[", "").replace("]", "").replace("’", "'"))


def _sentence_start(text, i):
    pre = text[:i].rstrip(" \t\"'“‘(*•-–—")
    return not pre or pre[-1] in ".!?:;\n"


def _in_hay(s, haystack):
    return re.search(r"(?<![A-Za-z])" + re.escape(s) + r"(?![a-z])", haystack) is not None


_HAY_NORM = {}


def _hay_norm(haystack):
    """The haystack as lower-case words with single spaces, so that a name is found whatever the punctuation between
    its words, as in "Countdown (Australia)" (kept for the few haystacks in use)."""
    k = hash(haystack)
    v = _HAY_NORM.get(k)
    if v is None:
        if len(_HAY_NORM) > 8:
            _HAY_NORM.clear()
        v = _HAY_NORM[k] = " " + " ".join(norm_words(haystack)) + " "
    return v


def _name_in(name, haystack):
    return name in haystack or f" {' '.join(norm_words(name))} " in _hay_norm(haystack)


HYPHEN_WORD = re.compile(r"[A-Z][a-z]+-[a-z]")       # "London-accented": a word of description, not part of a name


def _ordinary(w, haystack):
    """Whether a capitalised word at the start of a sentence is an ordinary word, not a name: one that often starts a
    sentence; or, if the turns never put it before another capitalised word as they would a first name ("Mark
    Brzezicki"), a word in -ing or one the turns use in lower case (as a word of its own: "don't" is not "don")."""
    if w in NAME_SKIP or w in NAME_STARTERS:
        return True
    if re.search(r"(?<![A-Za-z'])" + re.escape(w) + r"\s+(?:Mc|Mac|O')?[A-Z][a-z]", haystack):
        return False
    return ((len(w) >= 6 and w.endswith("ing"))
            or re.search(r"(?<![A-Za-z'])" + re.escape(w.lower()) + r"(?![A-Za-z'])", haystack) is not None)


def unknown_names(text, haystack):
    """Names of two or more capitalised words that appear nowhere in the act's turns: often from outside knowledge
    (only flagged, as the check is rough). A possessive's "'s" is left out, and at the start of a sentence, an ordinary
    word before a name the turns give is not taken for part of the name, as in "Keep Mark's …" or "Whereas Stuart …"."""
    text = str(text).replace("’", "'")
    found = []
    for m in NAME_RE.finditer(text):
        segs, cur = [], []                     # a possessive ends a name: "Stuart's Big Country" is two
        for w in m.group(1).split():
            if w.endswith("'s"):
                segs.append(cur + [w[:-2]])
                cur = []
            else:
                cur.append(w)
        if cur:
            segs.append(cur)
        for k, words in enumerate(segs):
            start = k == 0 and _sentence_start(text, m.start())
            while words and (words[0] in NAME_SKIP or HYPHEN_WORD.match(words[0])):
                words, start = words[1:], False
            while words and words[-1] in NAME_SKIP:
                words = words[:-1]
            if len(words) < 2:
                continue
            name, rest = " ".join(words), " ".join(words[1:])
            if _name_in(name, haystack):
                continue
            if (start and (_in_hay(rest, haystack) or (len(words) > 2 and _name_in(rest, haystack)))
                    and _ordinary(words[0], haystack)):
                continue
            found.append(name)
    return sorted(set(found))


_NORM_CACHE = {}


def _norm_text(tid):
    if tid not in _NORM_CACHE:
        _NORM_CACHE[tid] = " ".join(norm_words(CORPUS["by_id"][tid]["Text"]))
    return _NORM_CACHE[tid]


def quote_problems(text, scope_ids, labels_ok=()):
    """The quotes in a memo, point by point: each must be word for word in a turn the point cites (or, if it cites
    none, in a turn of the acts in scope), unless it is a label; each Turn ID must be in the corpus."""
    problems, ok_labels = [], {label_key(x) for x in labels_ok}
    points = re.split(r"\n\s*(?=(?:\d+[.)]|[-*•])\s)", str(text))
    for pt in points:
        ids = TID_RE.findall(pt)
        bad = sorted(set(i for i in ids if i not in CORPUS["by_id"] and looks_like_tid(i)))
        if bad:
            problems.append(f"Turn ID(s) that are not in the corpus: {', '.join(bad[:5])}.")
        cited = [i for i in dict.fromkeys(ids) if i in CORPUS["by_id"]]
        for mq in QUOTE_RE.finditer(pt):
            q = (mq.group(1) or mq.group(2) or "").strip(" .,;:")
            words = norm_words(q)
            if len(q) < 6 or len(words) < 3 or label_key(q) in ok_labels:
                continue
            first = " ".join(norm_words(ELLIPSIS.split(q)[0])) or " ".join(words)
            pool = cited or [i for i in scope_ids if first in _norm_text(i)]
            if any(evidence_in(q, CORPUS["by_id"][i]["Text"]) for i in pool):
                continue
            problems.append(f'The quote "{short(q, 90)}" is not word for word in '
                            + (f"the turn(s) it cites ({', '.join(cited[:4])})." if cited else "any turn of the acts."))
    return problems


def parse_json(text):
    if text is None or not str(text).strip():
        return None, "the reply is empty"
    s = str(text).strip()
    if s.startswith("```"):
        s = re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", s).strip()
    try:
        return json.loads(s), None
    except Exception as e:
        i, j = s.find("{"), s.rfind("}")
        if 0 <= i < j:
            try:
                return json.loads(s[i:j + 1]), None
            except Exception as e2:
                return None, short(e2, 200)
        return None, short(e, 200)


def parse_entry(e):
    if not isinstance(e, str):
        return None
    m = MEM_RE.match(e)
    return (int(m.group(1)), m.group(2).strip(), m.group(3)) if m else None


ID_SEP = re.compile(r"\s*(?:,|/|;|&|\band\b)\s*", re.I)


def entry_ids(tid):
    """The Turn IDs an entry names in its brackets: one, or several joined by commas, slashes, semicolons, '&' or
    'and', as in 'M7 (3153-25 / 3004-28)'."""
    return [x.strip().strip("\"'‘’“”") for x in ID_SEP.split(tid) if x.strip().strip("\"'‘’“”")]


def check_refs(n, body):
    return [f"M{n} names M{r}, which does not exist (if it is meant as a reference)."
            for r in map(int, REF_RE.findall(body)) if not 1 <= r < n]


# Each check returns (what to keep, unusable, problems, notes). A reply that cannot be used is asked again up to
# RETRIES_UNUSABLE times; problems are asked again RETRIES_PROBLEMS times, then flagged; notes are only flagged.
NOTE_RE = re.compile(r"^\s*(?:[-*•]|\d+[.)])?\s*\(([^()]*)\)\s*:?\s*(\S.*)$", re.S)


RANGE_RE = re.compile(r"\b(V\d+-\d+|\d{1,6}-\d+)\s*(?:to|through|until|–|—|-)\s*(V\d+-\d+|\d{1,6}-\d+)\b")


def turn_ids_in(s, order):
    """The Turn IDs of the act (its turns in order) that a text names: one, several, or a range such as "113-32 to
    113-40", which stands for every turn of the act from the one to the other."""
    pos = {t: i for i, t in enumerate(order)}
    out = []
    for m in RANGE_RE.finditer(str(s)):
        a, b = m.group(1), m.group(2)
        if a in pos and b in pos and 0 <= pos[b] - pos[a] < 200:
            out += order[pos[a]:pos[b] + 1]
    out += [t for t in TID_RE.findall(str(s)) if t in pos]
    return list(dict.fromkeys(out))


def check_reading(obj, act_ids, haystack=None):
    """The checks after a reading call: a list of notes, each "(Turn ID or general): note", and a list of data
    problems, each with its Turn ID and whether the turn is not the member's words at all. A note that names no Turn
    ID is kept as general; a Turn ID from outside the act, and names that appear nowhere in the act, are only
    flagged."""
    soft, notes, out, dps = [], [], [], []
    order = sorted(act_ids) if isinstance(act_ids, (set, frozenset)) else list(act_ids)
    act_ids = set(order)
    if not isinstance(obj, dict) or not isinstance(obj.get("notes"), list):
        return None, ['The reply is not a JSON object with a "notes" list.'], [], []
    for x in obj["notes"]:
        if not isinstance(x, str) or not x.strip():
            continue
        m = NOTE_RE.match(x)
        if m:
            tid, body = m.group(1).strip() or "general", m.group(2).strip()
        else:
            tid, body = "general", x.strip()
            notes.append(f'A note does not start with a Turn ID or "general" in brackets; kept as general: '
                         f'"{short(x, 80)}".')
        ids = entry_ids(tid)
        if not ids or not all(i.lower() == "general" or i in act_ids for i in ids):
            notes.append(f'A note names "{tid}", which is not a Turn ID of this act, or "general".')
        out.append({"turn": tid, "note": body})
    for d in obj.get("data_problems") or []:
        if not isinstance(d, dict):
            continue
        tid, prob = str(d.get("turn_id") or "").strip(), _one(d.get("problem"))
        if not tid and not prob:
            continue
        found = [tid] if tid in act_ids else turn_ids_in(tid, order)     # "113-32, 113-33", "113-32 to 113-40"
        if not found:
            notes.append(f'A data problem names "{tid}", which is not a Turn ID of this act.')
            found = [tid]
        for t1 in found:
            dps.append({"turn_id": t1, "problem": prob, "not_members_words": _bool(d.get("not_members_words"))})
    if haystack is not None:
        names = unknown_names("\n".join(o["note"] for o in out), haystack)
        if names:
            notes.append(f"The notes name people or things that appear nowhere in the act's turns: {', '.join(names)}.")
    if not out:
        soft.append("The reply has no notes.")
    return {"notes": out, "data_problems": dps}, [], soft, notes


def check_round2(obj, batch_turns, st, haystack=None):
    """The checks after every coding call. Returns what to keep: the records (turn by turn: whether it looks back,
    any data problem, the codes, each with whether its label is the members' words, worked out here, and the memory
    entries) and the definitions of the labels first used in the reply."""
    hard, soft, notes = [], [], []
    if not isinstance(obj, dict) or not isinstance(obj.get("turns"), list):
        return None, ['The reply is not a JSON object with a "turns" list.'], [], []
    recs, ids = obj["turns"], [t["Turn ID"] for t in batch_turns]
    got = [str(r.get("turn_id", "")).strip() if isinstance(r, dict) else "?" for r in recs]
    if got != ids:
        miss = [i for i in ids if i not in got]
        extra = [g for g in got if g not in ids]
        hard.append(f"The reply must have exactly one record per turn, in the order of the turns: {len(ids)} records, "
                    f"from {ids[0]} to {ids[-1]}. It has {len(recs)}"
                    + (f"; missing: {', '.join(miss[:10])}" if miss else "")
                    + (f"; not in this batch: {', '.join(extra[:10])}" if extra else "")
                    + ("; out of order" if not miss and not extra else "") + ".")
        return None, hard, soft, notes
    defs = {}
    for d in obj.get("new_labels") or []:
        if isinstance(d, dict) and _one(d.get("label")) and _one(d.get("definition")):
            defs[label_key(_one(d.get("label")))] = _one(d.get("definition"))
    n, labels, new_labels, out = st["next"], st["labels"], {}, []
    old_names = {}                             # a label's name before an earlier act's review renamed it
    for v in labels.values():
        if v.get("renamed_from"):
            old_names.setdefault(label_key(v["renamed_from"]), v["label"])
    for r, t in zip(recs, batch_turns):
        tid = t["Turn ID"]
        codes, mem = r.get("codes"), r.get("memory")
        if not isinstance(codes, list) or not isinstance(mem, list):
            hard.append(f"{tid}: the record needs a list of codes and a list of memory entries.")
            continue
        rc = []
        for c in codes:
            if not isinstance(c, dict):
                hard.append(f"{tid}: a code is not an object with label, description, evidence and unsure.")
                continue
            label = _one(c.get("label"))
            desc = _one(c.get("description"))
            ev = str(c.get("evidence") or "").strip()
            if not label:
                hard.append(f"{tid}: a code has no label.")
                continue
            if not desc:
                soft.append(f'{tid}: the code "{label}" has no description.')
            if not ev:
                soft.append(f'{tid}: the code "{label}" has no evidence.')
            elif not evidence_in(ev, t["Text"]):
                soft.append(f'{tid}: the evidence for "{label}" is not word for word in this turn: "{short(ev, 120)}".')
            key, as_coded = label_key(label), None
            if key not in labels and key not in new_labels and key in old_names:
                as_coded, label = label, old_names[key]      # the old name: the code goes to the renamed label
                key = label_key(label)
            info = labels.get(key) or new_labels.get(key)
            lf = label_from(label, t["Text"], info["first_turn"] if info else None)
            if not info:
                new_labels[key] = {"first_turn": tid, "label": label}
            code = {"label": label, "label_from": lf, "description": desc, "evidence": ev,
                    "unsure": _bool(c.get("unsure"))}
            if as_coded:
                code["original_label"] = as_coded
            rc.append(code)
        rm = []
        for e in mem:
            p = parse_entry(e)
            if p is None:
                hard.append(f'{tid}: a memory entry is not written as "M{n} ({tid}): …": "{short(e)}".')
                n += 1
                continue
            num, etid, body = p
            if num != n:
                hard.append(f"{tid}: memory entry M{num} should be M{n}. New entries are numbered in order from "
                            f"M{st['next']}, across the whole reply.")
            if tid not in entry_ids(etid):
                notes.append(f'{tid}: memory entry M{num} carries "{etid}" instead of this turn\'s ID, {tid}.')
            notes += [f"{tid}: {x}" for x in check_refs(num, body)]
            if haystack is not None:
                names = unknown_names(body, haystack)
                if names:
                    notes.append(f"{tid}: memory entry M{num} names people or things that appear nowhere in the act's "
                                 f"turns: {', '.join(names)}.")
            rm.append({"n": num, "turn": tid, "text": e.strip()})
            n += 1
        out.append({"turn_id": tid, "looks_back": _bool(r.get("looks_back")), "data_problem": _one(r.get("data_problem")),
                    "codes": rc, "memory": rm})
    missing = [v["label"] for k, v in new_labels.items() if not defs.get(k)]
    if missing:
        notes.append(f"No definition was given for the new label(s) {', '.join(repr(x) for x in missing[:8])}; the "
                     f"first description stands in until the act's review.")
    return {"records": out, "new_labels": {k: {"label": v["label"], "definition": defs[k]}
                                           for k, v in new_labels.items() if defs.get(k)}}, hard, soft, notes


def check_review(obj, st, stem, haystack=None):
    """The checks after the review of an act. Renaming: only the labels created in the act, and never to a name that
    is already a label (renaming is not merging). A definition for every label created in the act (the one given
    when it was coded stands where the review gives none). The memo: one to MEMO_POINTS points, each with a Turn ID
    of the act, its quotes word for word. Keep: entries of this act's notes."""
    hard, soft, notes = [], [], []
    if not isinstance(obj, dict) or not isinstance(obj.get("memo"), list):
        return None, ['The reply is not a JSON object with "renames", "definitions", "memo" and "keep".'], [], []
    created = {k: v for k, v in st["labels"].items() if v.get("first_act") == stem}
    renames, taken, new_key = [], set(), {}
    for x in obj.get("renames") or []:
        if not isinstance(x, dict):
            continue
        old, new = _one(x.get("label")), _one(x.get("new_label"))
        ko, kn = label_key(old), label_key(new)
        if ko not in created:
            notes.append(f'"{old}" is not a label created in this act, so it was not renamed.')
            continue
        if not new or kn == ko or ko in new_key:
            continue
        if kn in st["labels"] or kn in taken:
            soft.append(f'"{old}" cannot be renamed "{new}": that is already one of your labels, and renaming is not '
                        f'merging. Choose another name, or keep the label as it is.')
            continue
        if len(new.split()) > 8:
            soft.append(f'The new name for "{old}" has {len(new.split())} words: a label is a short name, not a '
                        f'quotation.')
            continue
        renames.append({"label": created[ko]["label"], "new_label": new})
        taken.add(kn)
        new_key[ko] = kn
    back = {kn: ko for ko, kn in new_key.items()}
    defs = {}
    for d in obj.get("definitions") or []:
        if not isinstance(d, dict):
            continue
        k, df = label_key(_one(d.get("label"))), _one(d.get("definition"))
        ko = back.get(k, k)
        if ko in created and df:
            defs[ko] = df
        elif df:
            notes.append(f'A definition is given for "{_one(d.get("label"))}", which is not a label created in this act.')
    missing = [v["label"] for k, v in created.items() if not defs.get(k) and not v.get("definition")]
    if missing:
        soft.append(f"Give a definition for each label you created in this act; these have none: "
                    f"{', '.join(repr(x) for x in missing[:12])}" + (" and more." if len(missing) > 12 else "."))
    final = {ko: (next(r["new_label"] for r in renames if label_key(r["label"]) == ko) if ko in new_key
                  else created[ko]["label"]) for ko in created}
    definitions = [{"label": final[ko], "definition": defs.get(ko) or created[ko].get("definition") or ""}
                   for ko in created if defs.get(ko) or created[ko].get("definition")]
    memo = [memo_point(p) for p in obj.get("memo") or [] if memo_point(p)]
    act_ids = [t["Turn ID"] for t in CORPUS["turns"].get(stem, [])]
    act_set = set(act_ids)
    if not memo:
        soft.append("The memo on the act has no points.")
    elif len(memo) > MEMO_POINTS:
        soft.append(f"The memo has {len(memo)} points: write up to {MEMO_POINTS}, keeping the ones with the most "
                    f"substance.")
    for i, p in enumerate(memo, 1):
        if not any(x in act_set for x in TID_RE.findall(p)):
            soft.append(f"Point {i} of the memo gives no Turn ID of this act.")
        soft += [f"Point {i}: {q}" for q in quote_problems(p, act_ids, [v["label"] for v in st["labels"].values()]
                                                           + [r["new_label"] for r in renames])]
    if haystack is not None:
        names = unknown_names("\n".join(memo), haystack)
        if names:
            notes.append(f"The memo names people or things that appear nowhere in the act's turns: {', '.join(names)}.")
    nums = {m["n"] for m in act_entries(st, stem)}
    keep = []
    for k in obj.get("keep") or []:
        m = re.match(r"\s*M?(\d+)\b", str(k))
        if m and int(m.group(1)) in nums:
            keep.append(int(m.group(1)))
        else:
            notes.append(f'"{short(k, 40)}" in the notes to keep is not one of this act\'s notes.')
    return {"renames": renames, "definitions": definitions, "memo": memo, "keep": sorted(set(keep))}, hard, soft, notes


def _obj(props, required=None):
    return {"type": "object", "additionalProperties": False, "required": required or list(props), "properties": props}


_S, _B = {"type": "string"}, {"type": "boolean"}
_LIST = {"type": "array", "items": _S}
_CODE = _obj({"label": _S, "description": _S, "evidence": _S, "unsure": _B})
_NEWLAB = {"type": "array", "items": _obj({"label": _S, "definition": _S})}
SCHEMA_READ = _obj({"notes": _LIST, "data_problems": {"type": "array", "items": _obj(
    {"turn_id": _S, "problem": _S, "not_members_words": _B})}})
SCHEMA1 = _obj({"memory": _LIST})
SCHEMA2 = _obj({"turns": {"type": "array", "items": _obj({
    "turn_id": _S, "looks_back": _B, "data_problem": _S, "codes": {"type": "array", "items": _CODE},
    "memory": _LIST})}, "new_labels": _NEWLAB})
SCHEMA_REVIEW = _obj({"renames": {"type": "array", "items": _obj({"label": _S, "new_label": _S})},
                      "definitions": _NEWLAB, "memo": _LIST, "keep": _LIST})

# ==================================================================================================== the three models


def _status(e):
    for a in ("status_code", "code", "status"):
        v = getattr(e, a, None)
        try:
            return int(v)
        except (TypeError, ValueError):
            continue
    return None


def is_fatal(e):
    return _status(e) in (400, 401, 403, 404, 413, 422) or isinstance(e, (TypeError, UnicodeError, AgentStop))


OUTAGE_NAMES = re.compile(r"Connection|Timeout|Overloaded|RateLimit|InternalServer|ServiceUnavailable|ServerError|"
                          r"Protocol|ReadError|WriteError|NetworkError|Unavailable|DeadlineExceeded")


def is_outage(e):
    """A network or server error such as a provider's outage causes, worth waiting through for TRANSIENT_RETRIES;
    any other error that is not fatal is tried again only OTHER_RETRIES times, as waiting will not cure it."""
    if _status(e) in (408, 409, 425, 429, 500, 502, 503, 504, 520, 522, 524, 529):
        return True
    if isinstance(e, (ConnectionError, TimeoutError)):
        return True
    if isinstance(e, RuntimeError) and re.search(r"stream ended|failed on OpenAI's side", str(e)):
        return True
    return any(OUTAGE_NAMES.search(c.__name__) for c in type(e).__mro__)


def price(model, day=None):
    """The prices that apply on a day, today unless given (a model's prices can change on a set date, as Gemini's do
    on 1 December)."""
    day = day or datetime.date.today().isoformat()
    p = dict(PRICES.get(model) or {})
    for d, newp in sorted((p.pop("from", None) or {}).items()):
        if day >= d:
            p.update(newp)
    return p


def price_changes(model):
    """The dates after today on which a model's prices change."""
    return sorted(d for d in (PRICES.get(model) or {}).get("from", {}) if d > datetime.date.today().isoformat())


def nice_date(d):
    dt = datetime.date.fromisoformat(d)
    return f"{dt.day} {dt:%B %Y}"


def cost_of(model, u, batch=False):
    """The cost of one response, from its usage; a response from a batch costs half."""
    p = price(model)
    if not p:
        return 0.0
    if "cache_write_1h" in p:                          # Claude
        c = (u["input"] * p["input"] + u["cache_read"] * p["cache_read"] + u["cache_write_5m"] * p["cache_write_5m"]
             + u["cache_write_1h"] * p["cache_write_1h"] + u["output"] * p["output"])
    elif "cache_write" in p:                           # GPT: the long-input rates apply to each response on its own
        long = u["input_total"] > p["long_input"]
        xi, xo = (p["long_input_x"], p["long_output_x"]) if long else (1.0, 1.0)
        c = (u["input"] * p["input"] + u["cache_read"] * p["cache_read"] + u["cache_write"] * p["cache_write"]) * xi \
            + u["output"] * p["output"] * xo
    else:                                              # Gemini
        c = u["input"] * p["input"] + u["cache_read"] * p["cache_read"] + u["output"] * p["output"]
    return c / 1e6 * (0.5 if batch else 1.0)


def empty_usage():
    return {"input_total": 0, "input": 0, "cache_read": 0, "cache_write": 0, "cache_write_5m": 0,
            "cache_write_1h": 0, "output": 0, "thinking": 0}


def sum_usage(us):
    tot = empty_usage()
    for u in us:
        for k in tot:
            tot[k] += u.get(k, 0) or 0
    return tot


def own_tokens(letter, n):
    """Roughly how many of its own tokens a model counts for n tokens counted here (o200k_base): in the pilot,
    Claude's tokenizer counted about 1.45 times as many, Gemini's and GPT's about 1.05 times."""
    return int(n * (1.45 if MODELS[letter]["provider"] == "anthropic" else 1.05))


def reading_cost(letter, toks, batch):
    """The estimated cost of reading acts whose prompts have toks tokens each (counted here), with about 3,000 tokens
    out per act (the notes and the thinking)."""
    model = MODELS[letter]["model"]
    return sum(cost_of(model, dict(empty_usage(), input=own_tokens(letter, t), input_total=own_tokens(letter, t),
                                   output=3000), batch) for t in toks)


def _int(x):
    try:
        return int(x or 0)
    except (TypeError, ValueError):
        return 0


def full_text(parts):
    """The whole prompt of a call, as one text (for the providers that take it so, and for counting)."""
    return (parts["fixed"] + "".join(parts.get("memory") or []) + parts.get("labels", "") + parts.get("act_memory", "")
            + parts.get("tail", ""))


TOOL_DOCS = {
    "list_files": "List the files in your folder, with the number of lines in each.",
    "read_file": f"Read a file in your folder from start_line (1 for the start of the file): whole lines, up to "
                 f"{READ_LINES} at a time. The reply says which lines it shows, how many the file has, and where to "
                 f"read on. One memo can read about {READ_BUDGET:,} characters in all, in up to {MAX_TOOL_STEPS} "
                 f"uses of the tools, so choose what to read.",
}

# Each provider's json_call and tool_call add the usage of every response to the list `usages` as it arrives, so that
# responses already paid for are counted even if a later step fails. They return the reply's text, whether it is
# complete (ok), why the model stopped, and the model version.


class ClaudeAgent:
    def __init__(self, cfg, letter):
        import anthropic
        self.cfg, self.letter = cfg, letter
        self.client = anthropic.Anthropic(api_key=secret(SECRETS["anthropic"]), max_retries=6, timeout=1800)

    def _usage(self, msg):
        u = msg.usage
        cc = getattr(u, "cache_creation", None)
        w1 = _int(getattr(cc, "ephemeral_1h_input_tokens", 0)) if cc else 0
        w5 = _int(getattr(cc, "ephemeral_5m_input_tokens", 0)) if cc else 0
        cw = _int(u.cache_creation_input_tokens)
        if cc is None:
            w5 = cw
        cr, inp = _int(u.cache_read_input_tokens), _int(u.input_tokens)
        return {"input_total": inp + cr + cw, "input": inp, "cache_read": cr, "cache_write": cw, "cache_write_5m": w5,
                "cache_write_1h": w1, "output": _int(u.output_tokens), "thinking": 0}

    def _content(self, parts, cache):
        """The fixed part, the memory (a block per act), the labels list as it was at the act's start, then what
        changes: the act's notes and the batch. Cache breakpoints, each kept for an hour (four at most): the fixed
        part, the last two memory blocks (where the last act's cache entry ends, and this act's) and the labels."""
        content = [{"type": "text", "text": parts["fixed"]}]
        if cache:
            content[0]["cache_control"] = {"type": "ephemeral", "ttl": "1h"}
        mb = parts["memory"]
        for i, b in enumerate(mb):
            blk = {"type": "text", "text": b}
            if cache and i >= len(mb) - 2:
                blk["cache_control"] = {"type": "ephemeral", "ttl": "1h"}
            content.append(blk)
        if parts.get("labels", "").strip():
            content.append({"type": "text", "text": parts["labels"]})
            if cache:
                content[-1]["cache_control"] = {"type": "ephemeral", "ttl": "1h"}
        for k in ("act_memory", "tail"):
            if parts.get(k, "").strip():
                content.append({"type": "text", "text": parts[k]})
        return content

    def _stream(self, **kw):
        with self.client.messages.stream(model=self.cfg["model"], max_tokens=self.cfg["max_output"], **kw) as s:
            return s.get_final_message()

    def json_call(self, parts, schema, cache, usages=None):
        usages = [] if usages is None else usages
        oc = {"format": {"type": "json_schema", "schema": schema}}
        if self.cfg.get("effort"):
            oc["effort"] = self.cfg["effort"]
        msg = self._stream(messages=[{"role": "user", "content": self._content(parts, cache)}], output_config=oc)
        usages.append(self._usage(msg))
        text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
        return {"text": text, "usages": usages, "stop": msg.stop_reason,
                "ok": msg.stop_reason in ("end_turn", "stop_sequence"), "model": getattr(msg, "model", "")}

    def batch_submit(self, items, schema, keep):
        """Sends (custom_id, prompt) pairs as one Message Batch, at half price; keep(job) records the job."""
        oc = {"format": {"type": "json_schema", "schema": schema}}
        if self.cfg.get("effort"):
            oc["effort"] = self.cfg["effort"]
        reqs = [{"custom_id": cid, "params": {"model": self.cfg["model"], "max_tokens": self.cfg["max_output"],
                                              "messages": [{"role": "user", "content": text}], "output_config": oc}}
                for cid, text in items]
        b = self.client.messages.batches.create(requests=reqs)
        keep({"provider": "anthropic", "id": b.id, "custom_ids": [cid for cid, _ in items]})

    def batch_ended(self, job):
        return self.client.messages.batches.retrieve(job["id"]).processing_status == "ended"

    def batch_results(self, job):
        out = {}
        for r in self.client.messages.batches.results(job["id"]):
            res = r.result
            if getattr(res, "type", "") == "succeeded":
                msg = res.message
                text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
                out[r.custom_id] = {"text": text, "usages": [self._usage(msg)], "stop": msg.stop_reason,
                                    "ok": msg.stop_reason in ("end_turn", "stop_sequence"),
                                    "model": getattr(msg, "model", "")}
            else:
                out[r.custom_id] = {"error": f"{getattr(res, 'type', 'failed')}: {short(getattr(res, 'error', ''), 300)}"}
        return out

    def tool_call(self, parts, tools, usages=None):
        usages = [] if usages is None else usages
        tdefs = [{"name": "list_files", "description": TOOL_DOCS["list_files"],
                  "input_schema": {"type": "object", "properties": {}}},
                 {"name": "read_file", "description": TOOL_DOCS["read_file"],
                  "input_schema": {"type": "object", "properties": {"name": {"type": "string"},
                                                                    "start_line": {"type": "integer"}},
                                   "required": ["name"]}}]
        # one breakpoint only, set by the API at the end of each request (top-level cache_control): a block with its
        # own one-hour breakpoint as well would clash with it, which the API refuses
        msgs = [{"role": "user", "content": self._content(parts, False)}]
        kw = {"output_config": {"effort": self.cfg["effort"]}} if self.cfg.get("effort") else {}
        steps = 0
        while True:
            msg = self._stream(messages=msgs, tools=tdefs, cache_control={"type": "ephemeral"}, **kw)
            usages.append(self._usage(msg))
            calls = [b for b in msg.content if getattr(b, "type", "") == "tool_use"]
            if msg.stop_reason != "tool_use" or not calls:
                text = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
                return {"text": text, "usages": usages, "stop": msg.stop_reason, "tool_steps": steps,
                        "ok": msg.stop_reason in ("end_turn", "stop_sequence"), "model": getattr(msg, "model", "")}
            if STOP.is_set():
                raise stop_reason()
            steps += 1
            if steps > MAX_TOOL_STEPS + 10:            # past the limit, every tool use is told to write now
                raise AgentStop("the memo kept using its folder after the reading limit")
            msgs.append({"role": "assistant", "content": msg.content})
            msgs.append({"role": "user", "content": [
                {"type": "tool_result", "tool_use_id": b.id, "content": tools(b.name, dict(b.input or {}))}
                for b in calls]})


class GPTAgent:
    def __init__(self, cfg, letter):
        import openai
        self.cfg, self.letter, self.openai, self.streaming = cfg, letter, openai, True
        self.client = openai.OpenAI(api_key=secret(SECRETS["openai"]), max_retries=6, timeout=1800)

    def _usage(self, r):
        u = getattr(r, "usage", None)
        if u is None:
            return empty_usage()
        d = getattr(u, "input_tokens_details", None)
        total = _int(u.input_tokens)
        cr = _int(getattr(d, "cached_tokens", 0)) if d else 0
        cw = _int(getattr(d, "cache_write_tokens", 0)) if d else 0
        od = getattr(u, "output_tokens_details", None)
        return {"input_total": total, "input": max(total - cr - cw, 0), "cache_read": cr, "cache_write": cw,
                "cache_write_5m": 0, "cache_write_1h": 0, "output": _int(u.output_tokens),
                "thinking": _int(getattr(od, "reasoning_tokens", 0)) if od else 0}

    def _stream(self, kw):
        """Streams the response, so that a long call never sits silent on the connection, and returns the final
        response, complete or stopped early. If OpenAI asks the organisation to verify before it streams this model,
        the calls go on without streaming. A response that failed on OpenAI's side is raised as an error, so that it is
        tried again after a wait, like any other server error."""
        final = None
        if self.streaming:
            try:
                with self.client.responses.stream(**kw) as s:
                    for ev in s:
                        if getattr(ev, "type", "") in ("response.completed", "response.incomplete", "response.failed"):
                            final = ev.response
                if final is None:
                    raise RuntimeError("the response stream ended without a final response")
            except self.openai.BadRequestError as e:
                if not re.search(r"verif", str(e), re.I):
                    raise
                self.streaming = False
                say(f"[{self.letter}] OpenAI asks for organisation verification to stream {self.cfg['model']}; "
                    f"carrying on without streaming.")
        if not self.streaming:
            final = self.client.responses.create(**kw)
        if getattr(final, "status", "") == "failed":
            raise RuntimeError(f"the response failed on OpenAI's side: {short(getattr(final, 'error', None), 200)}")
        return final

    def _result(self, r, usages, steps=0):
        inc = getattr(r, "incomplete_details", None)
        status = getattr(r, "status", "") or ""
        return {"text": r.output_text or "", "usages": usages, "ok": status == "completed", "tool_steps": steps,
                "stop": status + (f" ({getattr(inc, 'reason', '')})" if inc else ""), "model": getattr(r, "model", "")}

    def json_call(self, parts, schema, cache, usages=None):
        """With cache, a breakpoint after the fixed part and after every chunk of the memory (stable_chunks). GPT
        reuses a cached prefix only when the breakpoints before it are where they were (in the pilot, breakpoints
        that moved on with each batch meant that the memory was written to the cache again on every call), so the
        breakpoints stay put: a chunk, once full, never changes, and each call adds to the last chunk or starts a
        new one. Each call writes at most the last chunk or two to the cache and reads the rest."""
        usages = [] if usages is None else usages
        content = [{"type": "input_text", "text": parts["fixed"]}]
        mb = stable_chunks(parts["memory"]) if cache else parts["memory"]
        for b in mb:
            content.append({"type": "input_text", "text": b})
        n_cached = len(content)
        if parts.get("labels", "").strip():           # the labels list as it was at the act's start: cached too
            content.append({"type": "input_text", "text": parts["labels"]})
            n_cached += 1
        for k in ("act_memory", "tail"):
            if parts.get(k, "").strip():
                content.append({"type": "input_text", "text": parts[k]})
        if cache:
            for i in range(n_cached):
                content[i]["prompt_cache_breakpoint"] = {"mode": "explicit"}
        kw = dict(model=self.cfg["model"], input=[{"role": "user", "content": content}],
                  max_output_tokens=self.cfg["max_output"], store=False,
                  prompt_cache_options={"mode": "explicit"}, prompt_cache_key=f"step3-{RUN}-{self.letter}",
                  text={"format": {"type": "json_schema", "name": "reply", "schema": schema, "strict": True}})
        if self.cfg.get("effort"):
            kw["reasoning"] = {"effort": self.cfg["effort"]}
        r = self._stream(kw)
        usages.append(self._usage(r))
        return self._result(r, usages)

    def batch_submit(self, items, schema, keep):
        """Sends (custom_id, prompt) pairs as one Batch (a JSONL file of Responses requests), at half price; keep(job)
        records the job."""
        lines = []
        for cid, text in items:
            body = dict(model=self.cfg["model"], input=[{"role": "user", "content": [{"type": "input_text", "text": text}]}],
                        max_output_tokens=self.cfg["max_output"], store=False,
                        text={"format": {"type": "json_schema", "name": "reply", "schema": schema, "strict": True}})
            if self.cfg.get("effort"):
                body["reasoning"] = {"effort": self.cfg["effort"]}
            lines.append(json.dumps({"custom_id": cid, "method": "POST", "url": "/v1/responses", "body": body}))
        f = self.client.files.create(file=("step3_reading.jsonl", ("\n".join(lines) + "\n").encode("utf-8")),
                                     purpose="batch")
        b = self.client.batches.create(input_file_id=f.id, endpoint="/v1/responses", completion_window="24h")
        keep({"provider": "openai", "id": b.id, "custom_ids": [cid for cid, _ in items]})

    def batch_ended(self, job):
        return self.client.batches.retrieve(job["id"]).status in ("completed", "failed", "expired", "cancelled")

    @staticmethod
    def _usage_body(u):
        if not u:
            return empty_usage()
        d, od = u.get("input_tokens_details") or {}, u.get("output_tokens_details") or {}
        total, cr, cw = _int(u.get("input_tokens")), _int(d.get("cached_tokens")), _int(d.get("cache_write_tokens"))
        return {"input_total": total, "input": max(total - cr - cw, 0), "cache_read": cr, "cache_write": cw,
                "cache_write_5m": 0, "cache_write_1h": 0, "output": _int(u.get("output_tokens")),
                "thinking": _int(od.get("reasoning_tokens"))}

    def batch_results(self, job):
        b = self.client.batches.retrieve(job["id"])
        out = {}
        for fid in (b.output_file_id, b.error_file_id):
            if not fid:
                continue
            for line in self.client.files.content(fid).text.splitlines():
                if not line.strip():
                    continue
                d = json.loads(line)
                resp = d.get("response") or {}
                body = resp.get("body")
                if d.get("error") or resp.get("status_code") != 200 or not isinstance(body, dict):
                    out[d.get("custom_id")] = {"error": short(d.get("error") or body or resp, 300)}
                    continue
                status, inc = body.get("status") or "", body.get("incomplete_details") or {}
                text = "".join(c.get("text", "") for o in body.get("output") or [] if o.get("type") == "message"
                               for c in o.get("content") or [] if c.get("type") == "output_text")
                out[d["custom_id"]] = {"text": text, "usages": [self._usage_body(body.get("usage"))],
                                       "ok": status == "completed", "model": body.get("model", ""),
                                       "stop": status + (f" ({inc.get('reason', '')})" if inc else "")}
        return out

    def tool_call(self, parts, tools, usages=None):
        usages = [] if usages is None else usages
        tdefs = [{"type": "function", "name": "list_files", "description": TOOL_DOCS["list_files"], "strict": True,
                  "parameters": {"type": "object", "properties": {}, "required": [], "additionalProperties": False}},
                 {"type": "function", "name": "read_file", "description": TOOL_DOCS["read_file"], "strict": True,
                  "parameters": {"type": "object", "properties": {"name": {"type": "string"},
                                                                  "start_line": {"type": "integer"}},
                                 "required": ["name", "start_line"], "additionalProperties": False}}]
        text = full_text(parts)
        base = dict(model=self.cfg["model"], max_output_tokens=self.cfg["max_output"], tools=tdefs, store=True,
                    prompt_cache_key=f"step3-{RUN}-{self.letter}")
        if self.cfg.get("effort"):
            base["reasoning"] = {"effort": self.cfg["effort"]}
        r = self._stream(dict(base, input=[{"role": "user", "content": [{"type": "input_text", "text": text}]}]))
        usages.append(self._usage(r))
        steps = 0
        while True:
            calls = [o for o in (r.output or []) if getattr(o, "type", "") == "function_call"]
            if not calls or getattr(r, "status", "") != "completed":
                return self._result(r, usages, steps)
            if STOP.is_set():
                raise stop_reason()
            steps += 1
            if steps > MAX_TOOL_STEPS + 10:            # past the limit, every tool use is told to write now
                raise AgentStop("the memo kept using its folder after the reading limit")
            outs = []
            for c in calls:
                try:
                    args = json.loads(c.arguments or "{}")
                except Exception:
                    args = {}
                outs.append({"type": "function_call_output", "call_id": c.call_id, "output": tools(c.name, args)})
            r = self._stream(dict(base, previous_response_id=r.id, input=outs))
            usages.append(self._usage(r))


class GeminiAgent:
    def __init__(self, cfg, letter):
        from google import genai
        from google.genai import types
        self.cfg, self.letter, self.types, self.json_mode = cfg, letter, types, True
        self.client = genai.Client(api_key=secret(SECRETS["google"]), http_options=types.HttpOptions(
            timeout=30 * 60 * 1000, retry_options=types.HttpRetryOptions(attempts=6)))

    def _usage(self, um):
        if um is None:
            return empty_usage()
        total, cr = _int(um.prompt_token_count), _int(um.cached_content_token_count)
        th = _int(um.thoughts_token_count)
        return {"input_total": total, "input": max(total - cr, 0), "cache_read": cr, "cache_write": 0,
                "cache_write_5m": 0, "cache_write_1h": 0, "output": _int(um.candidates_token_count) + th,
                "thinking": th}

    def _config(self, **kw):
        t = self.types
        if self.cfg.get("effort"):
            kw["thinking_config"] = t.ThinkingConfig(thinking_level=self.cfg["effort"])
        return t.GenerateContentConfig(**kw)

    def json_call(self, parts, schema, cache, usages=None):
        usages = [] if usages is None else usages
        text = full_text(parts)
        for attempt in (1, 2):
            kw = {"max_output_tokens": self.cfg["max_output"]}
            if self.json_mode:
                kw["response_mime_type"] = "application/json"
            try:
                chunks, um, finish, model = [], None, None, ""
                for ch in self.client.models.generate_content_stream(model=self.cfg["model"], contents=text,
                                                                     config=self._config(**kw)):
                    if ch.candidates and ch.candidates[0].content and ch.candidates[0].content.parts:
                        chunks += [p.text for p in ch.candidates[0].content.parts
                                   if getattr(p, "text", None) and not getattr(p, "thought", False)]
                    if ch.candidates and ch.candidates[0].finish_reason:
                        finish = ch.candidates[0].finish_reason
                    um = ch.usage_metadata or um
                    model = getattr(ch, "model_version", None) or model
                usages.append(self._usage(um))
                fin = getattr(finish, "name", str(finish))
                return {"text": "".join(chunks), "usages": usages, "stop": fin, "ok": fin == "STOP", "model": model}
            except Exception as e:
                if attempt == 1 and self.json_mode and _status(e) == 400 and re.search(r"mime|json", str(e), re.I):
                    self.json_mode = False
                    say(f"[{self.letter}] JSON mode refused by {self.cfg['model']}; carrying on without it.")
                    continue
                raise

    def batch_submit(self, items, schema, keep):
        """Sends (custom_id, prompt) pairs as batch jobs with the requests inline, in parts of up to
        GEMINI_BATCH_BYTES, at half price; keep(job) records each job as it is made."""
        t = self.types
        kw = {"max_output_tokens": min(self.cfg["max_output"], GEMINI_BATCH_MAX_OUTPUT)}   # the most a batch allows
        if self.json_mode:
            kw["response_mime_type"] = "application/json"
        cfg = self._config(**kw)
        part, size = [], 0

        def send(part):
            src = [t.InlinedRequest(contents=text, config=cfg, metadata={"key": cid}) for cid, text in part]
            j = self.client.batches.create(model=self.cfg["model"], src=src,
                                           config={"display_name": f"step3-{RUN}-{self.letter}-reading"})
            keep({"provider": "google", "id": j.name, "custom_ids": [cid for cid, _ in part]})
        for cid, text in items:
            n = len(text.encode("utf-8")) + 2000
            if part and size + n > GEMINI_BATCH_BYTES:
                send(part)
                part, size = [], 0
            part.append((cid, text))
            size += n
        if part:
            send(part)

    def batch_ended(self, job):
        j = self.client.batches.get(name=job["id"])
        state = getattr(j.state, "name", None) or str(j.state or "")
        return state in ("JOB_STATE_SUCCEEDED", "JOB_STATE_FAILED", "JOB_STATE_CANCELLED", "JOB_STATE_EXPIRED",
                         "JOB_STATE_PARTIALLY_SUCCEEDED")

    def batch_results(self, job):
        j = self.client.batches.get(name=job["id"])
        out = {}
        for cid, r in zip(job["custom_ids"], (j.dest.inlined_responses if j.dest else None) or []):
            key = (r.metadata or {}).get("key") or cid
            if r.error or not r.response:
                out[key] = {"error": short(r.error or "no response", 300)}
                continue
            resp = r.response
            cand = resp.candidates[0] if resp.candidates else None
            ps = cand.content.parts if cand and cand.content and cand.content.parts else []
            text = "".join(x.text for x in ps if getattr(x, "text", None) and not getattr(x, "thought", False))
            fin = getattr(getattr(cand, "finish_reason", None), "name", "") if cand else ""
            out[key] = {"text": text, "usages": [self._usage(resp.usage_metadata)], "ok": fin == "STOP", "stop": fin,
                        "model": getattr(resp, "model_version", "") or ""}
        return out

    def tool_call(self, parts, tools, usages=None):
        usages = [] if usages is None else usages
        t = self.types
        decls = [t.FunctionDeclaration(name="list_files", description=TOOL_DOCS["list_files"],
                                       parameters_json_schema={"type": "object", "properties": {}}),
                 t.FunctionDeclaration(name="read_file", description=TOOL_DOCS["read_file"],
                                       parameters_json_schema={"type": "object", "properties": {
                                           "name": {"type": "string"}, "start_line": {"type": "integer"}},
                                           "required": ["name"]})]
        # one request without streaming returns at most 32,768 tokens; past that, the SDK continues by itself up to
        # max_output_tokens (automatic continuation), so a long memo still comes back whole
        config = self._config(max_output_tokens=self.cfg["max_output"], tools=[t.Tool(function_declarations=decls)],
                              automatic_function_calling=t.AutomaticFunctionCallingConfig(disable=True))
        text = full_text(parts)
        contents = [t.Content(role="user", parts=[t.Part(text=text)])]
        steps = 0
        while True:
            r = self.client.models.generate_content(model=self.cfg["model"], contents=contents, config=config)
            usages.append(self._usage(r.usage_metadata))
            calls = r.function_calls or []
            cand = r.candidates[0] if r.candidates else None
            if not calls:
                fin = getattr(getattr(cand, "finish_reason", None), "name", "")
                body = "".join(p.text for p in (cand.content.parts if cand and cand.content and cand.content.parts
                                                else []) if getattr(p, "text", None) and not getattr(p, "thought", False))
                return {"text": body, "usages": usages, "stop": fin, "ok": fin == "STOP", "tool_steps": steps,
                        "model": getattr(r, "model_version", "")}
            if STOP.is_set():
                raise stop_reason()
            steps += 1
            if steps > MAX_TOOL_STEPS + 10:            # past the limit, every tool use is told to write now
                raise AgentStop("the memo kept using its folder after the reading limit")
            contents.append(cand.content)
            contents.append(t.Content(role="user", parts=[
                t.Part(function_response=t.FunctionResponse(id=c.id, name=c.name,
                                                            response={"result": tools(c.name, dict(c.args or {}))}))
                for c in calls]))


PROVIDERS = {"anthropic": ClaudeAgent, "openai": GPTAgent, "google": GeminiAgent}
_CLIENTS = {}


def provider(letter):
    if letter not in _CLIENTS:
        cfg = MODELS[letter]
        _CLIENTS[letter] = PROVIDERS[cfg["provider"]](cfg, letter)
    return _CLIENTS[letter]


def _book(st, kind_, parts, rec, usages, est, secs, batch=False):
    """Adds a call's cost to the agent's, saves the state, and keeps the record of the call: in calls/, one file per
    round (round 2: one per act), with the changing part of the prompt and the reply; the fixed part in prompts/.
    batch_api is True for a reply that came back from a provider's batch, at half price (round 2's "batch" is the
    number of the batch of turns)."""
    letter = st["agent"]
    model = MODELS[letter]["model"]
    c = sum(cost_of(model, u, batch) for u in usages)
    with COST_LOCK:
        st["cost"] = round(st["cost"] + c, 6)
        st["calls"] += 1
        COSTS[letter] = st["cost"]
        call_id = f"{letter}-{st['calls']:05d}"
    save_state(st)
    rec = dict(rec, call=call_id, agent=letter, kind=kind_, time=now(), seconds=secs, script=VERSION, model=model,
               effort=MODELS[letter].get("effort"), tokens_counted_before=est, usage=sum_usage(usages),
               responses=len(usages), cost=round(c, 6), batch_api=batch, fixed_sha=sha(parts["fixed"]), tail=parts["tail"],
               memory_blocks=len(parts["memory"]))
    append_jsonl(rpath(f"agent_{letter}", "calls", f"{rec.get('log', kind_)}.jsonl"), rec)
    fixed_path = rpath(f"agent_{letter}", "prompts", f"fixed_{rec['fixed_sha'][:12]}.txt")
    if not os.path.exists(fixed_path):
        write_text(fixed_path, parts["fixed"])
    return call_id, c


def call_model(st, kind_, parts, schema=None, tools=None, cache=False, meta=None):
    """One model call: the budget check, waits and retries after network or server errors (which Stop cuts short),
    the cost, the log line and the record of the call. Returns the result."""
    letter, meta = st["agent"], dict(meta or {})
    prov = provider(letter)
    if total_cost() >= BUDGET_USD:
        STOP_WHY[0] = f"the budget of ${BUDGET_USD:,.0f} is spent"
        STOP.set()                             # stops the other agents too
        raise stop_reason()
    est = est_tokens(full_text(parts))
    wait = FIRST_WAIT
    for attempt in range(TRANSIENT_RETRIES + 1):
        if STOP.is_set():
            raise stop_reason()
        t0, usages = time.time(), []
        try:
            res = prov.tool_call(parts, tools, usages) if tools else prov.json_call(parts, schema, cache, usages)
            break
        except Exception as e:
            if usages:                 # responses received before the failure (in a memo's tool loop) are paid for
                _book(st, kind_, parts, dict(meta, ok=False, error=short(e, 400)), usages, est,
                      round(time.time() - t0, 1))
            if isinstance(e, AgentStop):
                raise
            if is_fatal(e) or attempt >= (TRANSIENT_RETRIES if is_outage(e) else OTHER_RETRIES):
                raise AgentStop(f"{MODELS[letter]['model']} returned an error: {short(e, 400)}")
            say(f"[{letter}] {type(e).__name__}: {short(e, 160)}; trying again in {wait} s")
            if STOP.wait(wait):
                raise stop_reason()
            wait = min(wait * 2, MAX_WAIT)
    secs = round(time.time() - t0, 1)
    u = sum_usage(res["usages"])
    miss = False
    if cache and MODELS[letter]["provider"] != "google":   # Gemini caches by itself, when it can
        with COST_LOCK:
            CACHED_CALLS[letter] = CACHED_CALLS.get(letter, 0) + 1
            first = CACHED_CALLS[letter] == 1
        # what the call should have read from the cache at least: the fixed part and the memory, but for its newest
        # part (Claude: the last batch's block; GPT: the chunk still filling, and one that may have just filled)
        mem = parts["memory"]
        older = stable_chunks(mem)[:-2] if MODELS[letter]["provider"] == "openai" else mem[:-1]
        miss = not first and u["cache_read"] < 0.8 * (_size(parts["fixed"]) + sum(_size(b) for b in older))
    call_id, c = _book(st, kind_, parts, dict(meta, model_version=res.get("model", ""), stop=res["stop"],
                                               ok=res["ok"], cache_miss=miss, tool_steps=res.get("tool_steps", 0),
                                               reply=res["text"]), res["usages"], est, secs)
    res.update(call_id=call_id, cost=c, usage=u)
    n = len(res["usages"])
    say(f"[{letter}] {meta.get('label', kind_)}: ~{est:,} tokens in" + (f", {n} responses (file reads)" if n > 1 else "")
        + f", {u['output']:,} out, {u['cache_read']:,} from the cache, {secs:.0f} s, ${c:.3f} (agent "
          f"${st['cost']:.2f}, all ${total_cost():.2f})" + ("  CACHE MISS" if miss else ""))
    return res


def problem_note(problems, unusable, tail_hint):
    head = ("Your previous reply to this prompt could not be used, for these reasons:" if unusable
            else "Your previous reply to this prompt had these problems:")
    lst = "\n".join(f"{i}. {p}" for i, p in enumerate(problems[:20], 1))
    more = f"\n(and {len(problems) - 20} more)" if len(problems) > 20 else ""
    return f"\n\n{head}\n{lst}{more}\nReply again in full, {tail_hint}, correcting them."


def ask_until_usable(st, kind_, parts, schema, check, meta, hint, soft_retries=None):
    """Asks, checks, and asks again: up to RETRIES_UNUSABLE times when the reply cannot be used, and RETRIES_PROBLEMS
    times (or soft_retries) when it can be used but has problems. Returns the best usable reply (what to keep,
    unusable, problems, notes, result), or stops the agent. After Stop, a usable reply already received is kept
    rather than lost. (Round 2 asks again about problems itself, for the turns that have them only: fix_turns.)"""
    tail0, tries = parts["tail"], []
    left_u, left_p = RETRIES_UNUSABLE, (RETRIES_PROBLEMS if soft_retries is None else soft_retries)
    while not STOP.is_set():
        try:
            res = call_model(st, kind_, parts, schema=schema, cache=meta.get("cache", False),
                             meta=dict(meta, attempt=len(tries) + 1))
        except AgentStop:
            if STOP.is_set() and any(not x[1] for x in tries):
                break
            raise
        obj, err = parse_json(res["text"])
        if not res["ok"]:
            out, hard, soft, notes = None, [f"The reply was cut off or stopped early ({res['stop']})."], [], []
        elif obj is None:
            out, hard, soft, notes = None, [f"The reply is not valid JSON ({err})."], [], []
        else:
            out, hard, soft, notes = check(obj)
        tries.append((out, hard, soft, notes, res))
        if hard:
            say(f"[{st['agent']}] {meta['label']}: reply not usable: {short(hard[0], 200)}")
        elif soft:
            say(f"[{st['agent']}] {meta['label']}: {len(soft)} problem(s): {short(soft[0], 200)}")
        if not hard and not soft:
            break
        if (hard and left_u == 0) or (not hard and left_p == 0):
            break
        if hard:
            left_u -= 1
        else:
            left_p -= 1
        parts = dict(parts, tail=tail0 + problem_note(hard or soft, bool(hard), hint))
    usable = [x for x in tries if not x[1]]
    if not usable:
        if STOP.is_set():
            raise stop_reason()
        raise AgentStop(f"{meta['label']}: no usable reply after {len(tries)} tries; last problem: "
                        f"{short(tries[-1][1][0], 300)}. Run the round again; if it stops here every time, see "
                        f"skip_batch().")
    return min(usable, key=lambda x: len(x[2]))


def run_agents(fn, letters, *args):
    """Runs fn(letter, *args) for the agents at the same time. Stop (in Colab) lets the calls in progress finish."""
    left = busy()
    if left:
        say(f"Not started: agents from the last run are still finishing their calls ({', '.join(left)}). Wait a "
            f"minute and run this again; status() shows where each agent is.")
        return {}
    STOP_WHY[0] = "stopped by hand"
    STOP.clear()
    ensure(rpath())
    letters, results = [L for L in (letters or LETTERS)], {}
    states = {}
    for L in LETTERS:                          # counts that a stale state.json would get wrong, from the logs
        st = load_state(L)
        if os.path.isdir(rpath(f"agent_{L}")) and sync_counters(st):
            save_state(st)
        states[L] = st
    with COST_LOCK:
        COSTS.clear()
        CACHED_CALLS.clear()
        for L in LETTERS:
            COSTS[L] = states[L]["cost"]
    for L in list(letters):                    # keys and clients first, here in the main thread (Colab Secrets)
        try:
            provider(L)
        except Exception as e:
            say(f"[{L}] NOT STARTED: {e}")
            results[L] = f"not started: {e}"
            letters.remove(L)
    if not letters:
        return results
    ex = ThreadPoolExecutor(max_workers=len(letters), thread_name_prefix=THREAD_PREFIX)
    futs = {ex.submit(_guard, fn, L, (BASE, RUN), *args): L for L in letters}
    try:
        for f in as_completed(futs):
            results[futs[f]] = f.result()
    except KeyboardInterrupt:
        STOP.set()
        say("Stopping: the calls in progress will finish and be saved first. Please wait ...")
        for f in futs:
            try:
                results[futs[f]] = f.result()
            except Exception:
                pass
    finally:
        ex.shutdown(wait=True)
    return results


def _guard(fn, letter, where, *args):
    _LOCAL.base, _LOCAL.run = where
    try:
        return _guarded(fn, letter, *args)
    finally:
        _LOCAL.base = _LOCAL.run = None


def _guarded(fn, letter, *args):
    try:
        r = fn(letter, *args)
        return r if isinstance(r, str) else "done"     # a round may say what it is waiting for instead
    except AgentStop as e:
        st = load_state(letter)
        st["stopped"] = f"{now()}: {e}"
        save_state(st)
        say(f"[{letter}] STOPPED: {e}")
        return f"stopped: {e}"
    except Exception as e:
        tb = traceback.format_exc()
        append_jsonl(rpath(f"agent_{letter}", "errors.jsonl"), {"time": now(), "error": repr(e), "traceback": tb})
        st = load_state(letter)
        st["stopped"] = f"{now()}: error {e!r}"
        save_state(st)
        say(f"[{letter}] ERROR: {e!r} (the traceback is in agent_{letter}/errors.jsonl)")
        return f"error: {e!r}"

# ==================================================================================================== round 1


def reading_path(letter, stem):
    return rpath(f"agent_{letter}", "reading", f"{stem}.json")


def reading_jobs_path(letter):
    return rpath(f"agent_{letter}", "reading", "batch_jobs.json")


def batch_id(letter, a):
    return re.sub(r"[^A-Za-z0-9_-]", "_", f"{letter}_{act_stem(a)}")[:64]


def reading_parts(letter, a):
    """The reading prompt: every turn of the act, the excluded ones too (the reading is where they are found)."""
    turns = CORPUS["all_turns"][act_stem(a)]
    text = fill(PROMPTS[1], {"GUIDE": load_guide(), "LETTER": letter, "ACT": a["Act"], "KIND": kind(a),
                             "N_INTERVIEWS": n_interviews(turns), "N_TURNS": len(turns),
                             "TURNS": render_turns_reading(turns)})
    return {"fixed": text, "memory": [], "tail": ""}


def check_reading_act(obj, stem):
    return check_reading(obj, [t["Turn ID"] for t in CORPUS["all_turns"][stem]], _act_haystack([stem]))


def _save_reading(st, a, out, problems, call_id, via):
    """Keeps an act's reading: its notes (which go into the memory just before the act is coded) and the data
    problems the agent found (which go to the flags, and to make_exclusion_list)."""
    stem = act_stem(a)
    notes_, dps = out["notes"], out["data_problems"]
    write_json(reading_path(st["agent"], stem), {"agent": st["agent"], "act": a["Act"], "stem": stem, "notes": notes_,
                                                "data_problems": dps, "problems": problems, "call": call_id,
                                                "via": via, "time": now()})
    flag(st, 1, stem, 0, problems, call_id)
    flag(st, 1, stem, 0, [f"{d['turn_id']}: " + ("not the member's words: " if d["not_members_words"] else "")
                          + d["problem"] for d in dps], call_id, kind="data")
    save_state(st)
    nm = sum(1 for d in dps if d["not_members_words"])
    say(f"[{st['agent']}] {a['Act']}: read ({'in the batch' if via == 'batch' else 'in a normal call'}), "
        f"{len(notes_)} notes, {len(dps)} data problems ({nm} turns not the member's words)"
        + (f", {len(problems)} flagged" if problems else ""))


def read_act_live(st, a):
    """Reads one act in a normal call, with the usual re-asks: when READ_BATCH is off, or the batch could not read
    the act."""
    stem = act_stem(a)
    out, hard, soft, notes_, res = ask_until_usable(
        st, "round1", reading_parts(st["agent"], a), SCHEMA_READ, lambda o: check_reading_act(o, stem),
        {"round": 1, "act": stem, "label": f"{a['Act']} (reading)", "log": "round1"}, "in the JSON format given above")
    _save_reading(st, a, out, soft + notes_, res["call_id"], "live")


def _collect(st, job, results, by_cid, failed):
    """Takes a finished batch job's results: each usable reply becomes that act's reading; an act the batch could not
    read is read again in a normal call."""
    letter = st["agent"]
    for cid in job["custom_ids"]:
        a = by_cid.get(cid)
        if a is None or os.path.exists(reading_path(letter, act_stem(a))):
            continue
        stem, r = act_stem(a), results.get(cid)
        if not r or r.get("error"):
            failed.append(a)
            say(f"[{letter}] {a['Act']}: the batch did not read it ({short((r or {}).get('error') or 'no result', 160)});"
                f" it is read in a normal call.")
            continue
        parts = reading_parts(letter, a)
        call_id, _ = _book(st, "round1", parts, {"round": 1, "act": stem, "label": f"{a['Act']} (reading)",
                                                 "log": "round1", "attempt": 1, "model_version": r.get("model", ""),
                                                 "stop": r["stop"], "ok": r["ok"], "reply": r["text"],
                                                 "batch_job": job["id"]},
                           r["usages"], est_tokens(parts["fixed"]), 0, batch=True)
        obj, err = parse_json(r["text"]) if r["ok"] else (None, f"the reply stopped early ({r['stop']})")
        if obj is None:
            out, hard, soft, notes_ = None, [f"The reply could not be used: {err}."], [], []
        else:
            out, hard, soft, notes_ = check_reading_act(obj, stem)
        if hard or soft:
            failed.append(a)
            say(f"[{letter}] {a['Act']}: the batch reply could not be used ({short((hard or soft)[0], 160)}); it is "
                f"read in a normal call.")
            continue
        _save_reading(st, a, out, notes_, call_id, "batch")


def reading_pending(letter):
    """The custom IDs of acts still in a batch job that has not been collected."""
    return {cid for j in read_json(reading_jobs_path(letter), []) if not j.get("collected") for cid in j["custom_ids"]}


def _round1(letter, acts, wait_minutes):
    st = load_state(letter)
    st["stopped"] = None

    def todo():
        return [a for a in acts if not os.path.exists(reading_path(letter, act_stem(a)))]
    by_cid = {batch_id(letter, a): a for a in acts}
    failed = []
    if todo() and READ_BATCH:
        jobs = read_json(reading_jobs_path(letter), [])
        pending = reading_pending(letter)
        new = [a for a in todo() if batch_id(letter, a) not in pending]
        if new:
            items = [(batch_id(letter, a), reading_parts(letter, a)["fixed"]) for a in new]
            est = reading_cost(letter, [est_tokens(t) for _, t in items], True)
            if total_cost() + est > BUDGET_USD:
                raise AgentStop(f"the budget of ${BUDGET_USD:,.0f} would be passed by this agent's reading (about "
                                f"${est:,.2f}, with ${total_cost():,.2f} spent so far)")
            made = []

            def keep(job):                     # saved the moment it is made, so that no job is ever lost track of
                job.update(submitted=now(), collected=False)
                made.append(job)
                jobs.append(job)
                write_json(reading_jobs_path(letter), jobs)
            try:
                provider(letter).batch_submit(items, SCHEMA_READ, keep)
            except Exception as e:
                sent = {c for j in made for c in j["custom_ids"]}
                rest = [a for a in new if batch_id(letter, a) not in sent]
                say(f"[{letter}] the reading batch could not be sent ({short(e, 200)}); {len(rest)} act(s) are read "
                    f"in normal calls instead.")
                failed += rest
            if made:
                n_sent = sum(len(j["custom_ids"]) for j in made)
                say(f"[{letter}] reading of {n_sent} act(s) sent as {len(made)} batch job(s), at half price (about "
                    f"${est * n_sent / len(new):,.2f}); batches come back within 24 hours, often much sooner.")
        deadline = time.time() + wait_minutes * 60
        while True:
            for j in jobs:
                if j.get("collected"):
                    continue
                try:
                    ended = provider(letter).batch_ended(j)
                except Exception as e:
                    if not is_fatal(e):
                        say(f"[{letter}] could not check batch {j['id']} ({short(e, 160)}); trying again later.")
                        continue
                    say(f"[{letter}] batch {j['id']} cannot be found or read ({short(e, 160)}); its acts are read in "
                        f"normal calls.")
                    ended, j["lost"] = True, True
                if ended:
                    if j.get("lost"):
                        failed += [by_cid[c] for c in j["custom_ids"] if c in by_cid]
                    else:
                        try:
                            results = provider(letter).batch_results(j)
                        except Exception as e:
                            if not is_fatal(e):        # the batch is done and paid for: try the download again
                                say(f"[{letter}] could not download the results of batch {j['id']} ({short(e, 160)});"
                                    f" trying again later.")
                                continue
                            say(f"[{letter}] the results of batch {j['id']} cannot be read ({short(e, 160)}); its "
                                f"acts are read in normal calls.")
                            results = {}
                        _collect(st, j, results, by_cid, failed)
                    j.update(collected=True, collected_at=now())
                    write_json(reading_jobs_path(letter), jobs)
            if all(j.get("collected") for j in jobs) or time.time() >= deadline:
                break
            if STOP.wait(BATCH_POLL_SECONDS):
                raise stop_reason()
    elif todo():
        failed = todo()
    for a in failed:
        if not os.path.exists(reading_path(letter, act_stem(a))):
            read_act_live(st, a)
    left = todo()
    if left:
        st["round1"] = False
        save_state(st)
        say(f"[{letter}] reading: {len(acts) - len(left)} of {len(acts)} act(s) read; the batch is still running for "
            f"{len(left)}. Run run_round1() again later to collect it.")
        return f"waiting for its batch ({len(left)} act(s) still to come): run run_round1() again later"
    st["round1"] = True
    save_state(st)
    say(f"[{letter}] round 1 done: {len(acts)} act(s) read.")


def run_round1(acts=None, agents=None, wait_minutes=None):
    """Round 1: each agent reads each act in full, without Before, and writes a few notes on it, which go into its
    memory just before it codes that act. With READ_BATCH, the reading is sent as a batch at half price: this sends
    it, waits up to READ_WAIT_MINUTES (or wait_minutes), and can be run again to collect what is still to come."""
    load_corpus(), load_prompts(), load_guide()
    sel = select_acts(acts)
    if not READ_ACTS:
        print("READ_ACTS is off, so there is no reading round: start with round 2.")
        return {}
    check_settings(1, {"acts": [act_stem(a) for a in sel]}, agents=agents)
    return run_agents(_round1, agents, sel, READ_WAIT_MINUTES if wait_minutes is None else wait_minutes)


# ==================================================================================================== excluded turns, data problems


def exclusion_candidates():
    """The turns the agents marked, in their reading, as not the member's words, with which agents marked each and
    why (from the reading of this run)."""
    c = load_corpus()
    marks = collections.OrderedDict()
    for a in c["acts"]:
        stem = act_stem(a)
        ids = {t["Turn ID"] for t in c["all_turns"][stem]}
        for L in LETTERS:
            p = reading_path(L, stem)
            if not os.path.exists(p):
                continue
            for d in read_json(p).get("data_problems") or []:
                tid = d.get("turn_id")
                if d.get("not_members_words") and tid in ids:
                    m = marks.setdefault(tid, {"act": a["Act"], "agents": [], "reasons": []})
                    if L not in m["agents"]:
                        m["agents"].append(L)
                        m["reasons"].append(f"{L}: {d.get('problem', '')}")
    return marks


EXCLUDE_COLS = ["Turn ID", "Act", "Speaker", "Reason", "Text starts"]


def _extra_columns(old):
    """Columns the team added to the list, which are kept as they are."""
    out = []
    for v in old.values():
        for k in v:
            if k not in EXCLUDE_COLS and k not in out:
                out.append(k)
    return out


def _write_exclusions(rows, extra=()):
    cols = EXCLUDE_COLS + [k for k in extra if k not in EXCLUDE_COLS]
    ensure(os.path.dirname(exclude_path()))
    with open(exclude_path(), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in cols})


def _exclusion_row(t, reason):
    return {"Turn ID": t["Turn ID"], "Act": t.get("Act", ""), "Speaker": t.get("Speaker", ""), "Reason": reason,
            "Text starts": short(t.get("Text", ""), 120)}


def make_exclusion_list(min_agents=None, force=False):
    """After the reading (round 1), before round 2: puts the turns that at least min_agents (EXCLUDE_MIN_AGENTS)
    agents marked as not the member's words into step3_inputs/excluded_turns.csv, keeping the turns already listed
    there (and any columns added to it), and writes every marked turn to reports/exclusion_candidates.csv. Round 2
    then leaves the listed turns out; the corpus files are not changed. It waits until every agent has read every
    act (force=True makes the list from the readings there are). To change the list afterwards, use
    remove_from_exclusion_list() and add_to_exclusion_list(), rather than a spreadsheet, which can turn Turn IDs
    such as 32-2 into dates."""
    if busy():
        print("Agents are still running; wait until they have stopped, then run this again.")
        return None
    k = EXCLUDE_MIN_AGENTS if min_agents is None else min_agents
    CORPUS["acts"] = None                      # read the corpus and the current list afresh
    c = load_corpus()
    missing = [(L, a) for a in select_acts() for L in LETTERS if not os.path.exists(reading_path(L, act_stem(a)))]
    if missing and not force:
        out_ = any(batch_id(L, a) in reading_pending(L) for L, a in missing)
        print(f"Not made yet: {len(missing)} reading(s) are not in, e.g. agent {missing[0][0]}'s reading of "
              f"{missing[0][1]['Act']}" + (" (a reading batch is still out: run run_round1() again to collect it)"
                                           if out_ else "")
              + ". The list should come from all three agents' readings, since making it again later would change "
                "it after round 2 has begun. To make it from the readings there are, run "
                "make_exclusion_list(force=True).")
        return None
    marks = exclusion_candidates()
    old = read_exclusions()
    rows = [dict(v, **{"Turn ID": t}) for t, v in old.items()]
    added = []
    for tid, m in marks.items():
        if len(m["agents"]) >= k and tid not in old:
            rows.append(_exclusion_row(c["by_id"][tid], f"marked by {len(m['agents'])} agents as not the member's "
                                                        f"words: " + " | ".join(m["reasons"])))
            added.append(tid)
    _write_exclusions(rows, _extra_columns(old))
    cand = rpath("reports", "exclusion_candidates.csv")
    ensure(os.path.dirname(cand))
    with open(cand, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Turn ID", "Act", "Agents that marked it", "How many", "On the list", "Reasons", "Text starts"])
        for tid, m in marks.items():
            w.writerow([tid, m["act"], ", ".join(m["agents"]), len(m["agents"]),
                        "yes" if tid in old or tid in added else "", " | ".join(m["reasons"]),
                        short(c["by_id"][tid].get("Text", ""), 120)])
    CORPUS["acts"] = None
    print(f"{len(marks)} turn(s) marked by at least one agent as not the member's words; {len(added)} added to "
          f"{EXCLUDE_FILE} (marked by {k} or more agents), which now lists {len(rows)}. All marked turns: "
          f"{cand}. Open {exclude_path()} to check the list before round 2 (to change it, use "
          f"remove_from_exclusion_list() or add_to_exclusion_list(), not a spreadsheet).")
    return added


def remove_from_exclusion_list(*turn_ids):
    """Takes turns off step3_inputs/excluded_turns.csv, so that round 2 codes them, without opening the file in a
    spreadsheet. Example: remove_from_exclusion_list("32-2", "1006-2")"""
    return _edit_exclusions(remove=turn_ids)


def add_to_exclusion_list(*turn_ids, reason="added by hand"):
    """Puts turns on step3_inputs/excluded_turns.csv, so that round 2 leaves them out.
    Example: add_to_exclusion_list("113-32", reason="the journalist's words")"""
    return _edit_exclusions(add=turn_ids, reason=reason)


def _edit_exclusions(remove=(), add=(), reason=""):
    if busy():
        print("Agents are still running; wait until they have stopped, then run this again.")
        return None
    CORPUS["acts"] = None
    c = load_corpus()
    old = read_exclusions()
    rows = {t: dict(v, **{"Turn ID": t}) for t, v in old.items()}
    done, skipped = [], []
    for t in (str(x).strip() for x in remove):
        if rows.pop(t, None) is not None:
            done.append(f"{t} taken off")
        else:
            skipped.append(f"{t} is not on the list")
    for t in (str(x).strip() for x in add):
        if t not in c["by_id"]:
            skipped.append(f"{t} is not a Turn ID of the corpus")
        elif t in rows:
            skipped.append(f"{t} is already on the list")
        else:
            rows[t] = _exclusion_row(c["by_id"][t], reason)
            done.append(f"{t} added")
    if done:
        _write_exclusions(list(rows.values()), _extra_columns(old))
    CORPUS["acts"] = None
    print(("; ".join(done) or "Nothing changed") + "." + (f" Not done: {'; '.join(skipped)}." if skipped else "")
          + f" The list has {len(rows)} turn(s)."
          + (" Round 2 of this run has begun: if these turns are in its acts, it will stop and say that the list "
             "changed (set RUN to a new name to start a new run)." if done and round_started(2) else ""))
    return done


def data_problems_report():
    """Every data problem the agents reported, in the reading and in the coding, in reports/data_problems.csv."""
    rows = []
    for L in LETTERS:
        for f in read_jsonl(rpath(f"agent_{L}", "flags.jsonl")):
            if f.get("kind") == "data":
                tid, _, prob = f["problem"].partition(": ")
                rows.append([L, f.get("round"), f.get("act"), tid, prob])
    p = rpath("reports", "data_problems.csv")
    ensure(os.path.dirname(p))
    with open(p, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["Agent", "Round", "Act", "Turn ID", "Problem"])
        w.writerows(rows)
    return p


def insert_reading(st, a):
    """Puts the act's reading notes into the agent's memory, numbered on from its last entry, as a block of their own,
    just before it codes the act. Returns the entries, which are kept in the act's first batch file. If the act has
    not been read (its batch failed and so did the normal call), it is read now."""
    letter, stem = st["agent"], act_stem(a)
    if not os.path.exists(reading_path(letter, stem)):
        read_act_live(st, a)
    rd = read_json(reading_path(letter, stem))
    entries, n = [], st["next"]
    for x in rd["notes"]:
        tid = x.get("turn") or "general"
        entries.append({"n": n, "turn": tid, "text": f"M{n} ({tid}, reading {a['Act']}): {x['note']}",
                        "round": "reading", "act": stem, "call": rd["call"]})
        n += 1
    if entries:
        add_block(st, entries)
        save_state(st)
    return entries


def add_block(st, entries):
    st["blocks"].append([len(st["memory"]), len(st["memory"]) + len(entries)])
    st["memory"].extend(entries)
    st["next"] = entries[-1]["n"] + 1

# ==================================================================================================== round 2


def apply_records(st, stem, recs, call_id, new_labels=None):
    """Adds a batch's codes to the agent's labels (a new label with its definition, if the batch gave one), and its
    memory entries to its memory."""
    defs = new_labels or {}
    for r in recs:
        for c in r["codes"]:
            key = label_key(c["label"])
            if key not in st["labels"]:
                st["labels"][key] = {"label": c["label"], "from": c["label_from"], "count": 0,
                                     "first_turn": r["turn_id"], "first_act": stem,
                                     "description": c.get("description", ""),
                                     "definition": (defs.get(key) or {}).get("definition", "")}
            st["labels"][key]["count"] += 1
    new = [dict(e, round=stem, act=stem, call=call_id) for r in recs for e in r["memory"]]
    if new:
        add_block(st, new)


def batch_path(letter, stem, bi):
    return rpath(f"agent_{letter}", "batches", stem, f"batch_{bi:03d}.json")


def review_path(letter, stem):
    return rpath(f"agent_{letter}", "reviews", f"{stem}.json")


def _accept(st, stem, bi, batch, out, problems, call_id, reading_memory=None, fix_call=None):
    """Keeps an accepted batch: its own file first (the record of what the agent coded, with the reading notes that
    went into its memory just before the act's first batch), then the agent's state. The data problems the agent
    reported go to the flags, not to its memory."""
    write_json(batch_path(st["agent"], stem, bi),
               {"agent": st["agent"], "act": stem, "batch": bi, "turn_ids": [t["Turn ID"] for t in batch["turns"]],
                "reading_memory": reading_memory or [], "records": out["records"],
                "new_labels": out.get("new_labels") or {}, "problems": problems, "call": call_id, "fix_call": fix_call,
                "time": now()})
    apply_records(st, stem, out["records"], call_id, out.get("new_labels"))
    flag(st, 2, stem, bi, problems, call_id)
    flag(st, 2, stem, bi, [f"{r['turn_id']}: {r['data_problem']}" for r in out["records"] if r.get("data_problem")],
         call_id, kind="data")
    st["done"].setdefault(stem, []).append(bi)
    save_state(st)


def apply_review(st, stem, rv):
    """The review of an act: its new names for labels (the batch files keep the labels as coded; act_records gives
    the new names), the definitions, and the memo and the notes kept, which the memory carries from now on."""
    for x in rv.get("renames") or []:
        ko, kn = label_key(x["label"]), label_key(x["new_label"])
        if ko in st["labels"] and kn not in st["labels"]:
            info = st["labels"].pop(ko)
            first = CORPUS["by_id"].get(info["first_turn"], {}).get("Text", "")
            info.update(renamed_from=info["label"], label=x["new_label"])
            info["from"] = "data" if first and label_in(x["new_label"], first) else "own"
            st["labels"][kn] = info
    for d in rv.get("definitions") or []:
        k = label_key(d.get("label", ""))
        if k in st["labels"] and d.get("definition"):
            st["labels"][k]["definition"] = d["definition"]
    st["act_memos"][stem] = {"memo": rv.get("memo") or [], "keep": rv.get("keep") or [], "call": rv.get("call")}


def rebuild_state(st, acts, move=True):
    """The batch and review files are the record of what an agent has coded. Whenever round 2 starts, the agent's
    memory, labels, reviews and progress are rebuilt from them, so that they always match, even if Drive kept one
    file and lost another in a crash. A file that does not follow on from the ones before it is set aside, and that
    batch (or review) is done again (with move=False, it is only left out)."""
    letter = st["agent"]
    r1 = read_json(rpath(f"agent_{letter}", "round1_memory.json"))
    if r1 is None:
        r1 = [m for m in st["memory"] if m.get("round") == "familiarisation"]
    tmp = dict(st, memory=[dict(m) for m in r1], blocks=[[0, len(r1)]] if r1 else [], next=len(r1) + 1, labels={},
               done={}, acts_complete=[], act_memos={}, act_start={"stem": None, "labels": {}})
    gap, current = False, None

    def set_aside(p, stem, what):
        dest = rpath(f"agent_{letter}", "batches_set_aside", stem,
                     f"{os.path.basename(p)[:-5]}_{datetime.datetime.now():%Y%m%d_%H%M%S}.json")
        ensure(os.path.dirname(dest))
        shutil.move(p, dest)
        say(f"[{letter}] {stem} {what}: its file does not follow on from the files before it; it is set aside (in "
            f"batches_set_aside) and will be done again.")
    for a in acts:
        stem = act_stem(a)
        start_labels = {k: v["count"] for k, v in tmp["labels"].items()}     # each earlier label's uses then
        batches = make_batches(CORPUS["turns"][stem])
        for bi, b in enumerate(batches, 1):
            p = batch_path(letter, stem, bi)
            if not os.path.exists(p):
                gap = True
                continue
            try:
                bf = read_json(p)
                nums = [e["n"] for e in bf.get("reading_memory") or []] + [e["n"] for r in bf["records"]
                                                                            for e in r["memory"]]
                ok = (not gap and bf.get("turn_ids") == [t["Turn ID"] for t in b["turns"]]
                      and [r["turn_id"] for r in bf["records"]] == bf["turn_ids"]
                      and nums == list(range(tmp["next"], tmp["next"] + len(nums))))
            except Exception:
                ok = False
            if not ok:
                gap = True
                if move:
                    set_aside(p, stem, f"batch {bi}")
                continue
            if bf.get("reading_memory"):
                add_block(tmp, [dict(e) for e in bf["reading_memory"]])
            apply_records(tmp, stem, bf["records"], bf["call"], bf.get("new_labels"))
            tmp["done"].setdefault(stem, []).append(bi)
        complete = len(tmp["done"].get(stem, [])) == len(batches)
        rp = review_path(letter, stem)
        if os.path.exists(rp):
            rv = None
            if complete and not gap:
                try:
                    rv = read_json(rp)
                    rv = rv if isinstance(rv.get("memo"), list) else None
                except Exception:
                    rv = None
            if rv is not None:
                apply_review(tmp, stem, rv)
                tmp["acts_complete"].append(stem)
            else:
                gap = True
                if move:
                    set_aside(rp, stem, "review")
        elif complete:
            gap = True                         # coded but not yet reviewed: the review comes next
        if stem not in tmp["acts_complete"] and current is None:
            current = {"stem": stem, "labels": start_labels}
    tmp["act_start"] = current or {"stem": None, "labels": {}}
    keys = ("memory", "blocks", "next", "labels", "done", "acts_complete", "act_memos", "act_start")
    if move and any(tmp[k] != st.get(k) for k in keys if k != "act_start"):
        say(f"[{letter}] the saved state did not match the batch and review files; it is rebuilt from them "
            f"({sum(len(v) for v in tmp['done'].values())} batches, {len(tmp['acts_complete'])} act(s) reviewed, "
            f"{len(tmp['memory'])} notes).")
    st.update({k: tmp[k] for k in keys})


def act_records(letter, stem, renamed=True):
    """The agent's records for an act, turn by turn, from its accepted batches. With renamed, each code carries the
    name the act's review gave its label (the label as coded is kept as original_label, and whether the new name is
    the members' words is worked out again)."""
    d = rpath(f"agent_{letter}", "batches", stem)
    recs = []
    if os.path.isdir(d):
        for fn in sorted(os.listdir(d)):
            if re.fullmatch(r"batch_\d{3}\.json", fn):
                recs += read_json(os.path.join(d, fn))["records"]
    rp = review_path(letter, stem)
    if not renamed or not os.path.exists(rp):
        return recs
    ren = {label_key(x["label"]): x["new_label"] for x in (read_json(rp) or {}).get("renames") or []}
    if not ren:
        return recs
    first = {}
    for r in recs:
        for c in r["codes"]:
            first.setdefault(label_key(c["label"]), r["turn_id"])
    out = []
    for r in recs:
        cs = []
        for c in r["codes"]:
            k = label_key(c["label"])
            if k in ren:
                text = CORPUS["by_id"].get(r["turn_id"], {}).get("Text", "")
                c = dict(c, label=ren[k], original_label=c["label"], label_from=label_from(ren[k], text, first[k]))
            cs.append(c)
        out.append(dict(r, codes=cs))
    return out


def count_codes(letter, stem):
    """Codes in an agent's accepted batch files for an act, counted file by file (a check on the merge)."""
    d, n = rpath(f"agent_{letter}", "batches", stem), 0
    for fn in (sorted(os.listdir(d)) if os.path.isdir(d) else []):
        if re.fullmatch(r"batch_\d{3}\.json", fn):
            n += sum(len(r["codes"]) for r in read_json(os.path.join(d, fn))["records"])
    return n


def reported(letter, act_name):
    return any(r.get("agent") == letter and r.get("act") == act_name
               for r in read_jsonl(rpath("reports", "act_reports.jsonl")))


def act_report(st, a):
    stem = act_stem(a)
    recs = act_records(st["agent"], stem)
    coded = [r for r in recs if r["codes"]]
    codes = [c for r in recs for c in r["codes"]]
    used = {label_key(c["label"]) for c in codes}
    new = sum(1 for v in st["labels"].values() if v.get("first_act") == stem)
    own = sum(1 for c in codes if c["label_from"] == "own")
    src = source_counts(st, recs)
    rv = read_json(review_path(st["agent"], stem)) or {}
    am = st["act_memos"].get(stem) or {}
    rep = {"agent": st["agent"], "act": a["Act"], "turns": len(recs), "coded": len(coded),
           "skipped": sum(1 for r in recs if r.get("skipped")), "codes": len(codes),
           "codes_per_coded_turn": round(len(codes) / max(len(coded), 1), 2), "labels_used": len(used),
           "new_labels": new, "renamed": len(rv.get("renames") or []), "own_share": round(own / max(len(codes), 1), 3),
           "label_sources": {s: src[s] for s in SOURCES},
           "unsure_share": round(sum(1 for c in codes if c.get("unsure")) / max(len(codes), 1), 3),
           "looks_back": sum(1 for r in recs if r.get("looks_back")),
           "data_problems": sum(1 for r in recs if r.get("data_problem")),
           "notes_written": len(act_entries(st, stem)), "notes_kept": len(am.get("keep") or []),
           "memo_points": len(am.get("memo") or []), "unique_labels_so_far": len(st["labels"]),
           "memory_tokens": est_tokens(memory_text(st)), "cost_so_far": round(st["cost"], 2), "time": now()}
    append_jsonl(rpath("reports", "act_reports.jsonl"), rep)
    say(f"[{st['agent']}] {a['Act']} done: {rep['coded']} of {rep['turns']} turns coded, "
        f"{rep['codes_per_coded_turn']} codes per coded turn, {rep['labels_used']} labels ({new} new, "
        f"{rep['renamed']} renamed in the review); labels {source_text(src)}; "
        f"{rep['unique_labels_so_far']} unique labels so far; {rep['notes_written']} notes, {rep['notes_kept']} kept, a "
        f"memo of {rep['memo_points']} points; memory about {rep['memory_tokens']:,} tokens.")


def act_parts(st, stem, fixed, tail):
    """A call of an act: the fixed part, the memory (the acts reviewed so far), the labels list as it was at the act's
    start (all three the same for every call of the act, so read from the cache), then the act's notes and the
    tail."""
    return {"fixed": fixed, "memory": memory_blocks(st), "labels": frozen_labels_text(st),
            "act_memory": act_memory_text(st, stem), "tail": tail}


FIX_TID = re.compile(r"^(\S+): ")
FIX_CODE = re.compile(r'^(\S+): the (?:code|evidence for) "(.+?)" (?:has no |is not word for word)')


def merge_codes(first, fix, flagged):
    """A turn's codes after a re-ask: the codes that had problems (their labels in flagged) give way to the
    answer's codes, in their place; the codes without problems stand, unless the answer gives a code of the same
    label, which then replaces it. With flagged empty (a problem that names no code), the answer replaces them all."""
    if not flagged:
        return list(fix)
    keys = {label_key(_one(c.get("label"))) for c in fix if isinstance(c, dict)}
    out, placed = [], False
    for c in first:
        k = label_key(_one(c.get("label"))) if isinstance(c, dict) else None
        if k in flagged:
            if not placed:
                out += list(fix)
                placed = True
        elif k not in keys:
            out.append(c)
    return out if placed else out + list(fix)


def fix_turns(st, parts, b, res, out, soft, notes, meta, hay):
    """Asks again about the problems of a usable reply, for the turns that have them only, and for the codes with
    problems only (the other codes and records, and every memory entry, stand). The answer's codes replace the codes
    with problems if they leave fewer problems; otherwise the first reply stands, and its problems are flagged. If
    the re-ask cannot be made, the first reply stands too, so that a reply already paid for is never lost."""
    ids = [t["Turn ID"] for t in b["turns"]]
    named, whole = collections.defaultdict(set), set()
    for p in soft:
        m = FIX_TID.match(p)
        if not m or m.group(1) not in ids:
            continue
        mc = FIX_CODE.match(p)
        if mc:
            named[m.group(1)].add(label_key(mc.group(2)))
        else:
            whole.add(m.group(1))
    bad = [t for t in ids if t in named or t in whole]
    if not bad:
        return out, soft, notes, None
    lst = "\n".join(f"{i}. {p}" for i, p in enumerate(soft[:20], 1))
    note = (f"\n\nYour reply to this prompt had these problems:\n{lst}\nReply again for these turns only: "
            f"{', '.join(bad)}, with the codes that had problems corrected, or left out if they do not hold, in the "
            "JSON format given above: a record for each of these turns, with its looks_back and data_problem as "
            "before and an empty list of memory entries, and a definition for any label you use here for the first "
            "time. Your other codes, the other turns and your memory entries stand as you gave them.")
    try:                                       # the same JSON schema as the batch, so the re-ask is read from the cache
        r2 = call_model(st, "round2", dict(parts, tail=parts["tail"] + note), schema=SCHEMA2, cache=True,
                        meta=dict(meta, attempt=2, fix_turns=bad))
    except AgentStop as e:
        if not STOP.is_set():
            say(f"[{st['agent']}] {meta['label']}: the re-ask could not be made ({short(e, 150)}); the first reply "
                f"stands.")
        return out, soft, notes, None
    obj2 = parse_json(r2["text"])[0] if r2["ok"] else None
    raw = parse_json(res["text"])[0]
    if not isinstance(obj2, dict) or not isinstance(obj2.get("turns"), list) or not isinstance(raw, dict):
        say(f"[{st['agent']}] {meta['label']}: the answer about the problems could not be used; the first reply stands.")
        return out, soft, notes, r2["call_id"]
    new_codes = {str(x.get("turn_id", "")).strip(): x["codes"] for x in obj2["turns"]
                 if isinstance(x, dict) and isinstance(x.get("codes"), list)}

    def fixed_record(r):
        tid = str(r.get("turn_id", "")).strip() if isinstance(r, dict) else None
        if tid not in bad or tid not in new_codes:
            return r
        return dict(r, codes=merge_codes(r.get("codes") or [], new_codes[tid],
                                         set() if tid in whole else named[tid]))
    taken = {label_key(_one(c.get("label"))) for tid in bad for c in new_codes.get(tid, []) if isinstance(c, dict)}
    had = {label_key(_one(d.get("label"))) for d in raw.get("new_labels") or [] if isinstance(d, dict)}
    defs2 = [d for d in obj2.get("new_labels") or [] if isinstance(d, dict)    # only for the codes taken from it,
             and label_key(_one(d.get("label"))) in taken - had]               # never over the first reply's
    merged = dict(raw, turns=[fixed_record(r) for r in raw["turns"]],
                  new_labels=list(raw.get("new_labels") or []) + defs2)
    out2, hard2, soft2, notes2 = check_round2(merged, b["turns"], st, hay)
    if hard2 or out2 is None or len(soft2) >= len(soft):
        say(f"[{st['agent']}] {meta['label']}: the answer about the problems was no better; the first reply stands.")
        return out, soft, notes, r2["call_id"]
    say(f"[{st['agent']}] {meta['label']}: asked again about {len(bad)} turn(s) only: {len(soft)} problem(s), now "
        f"{len(soft2)}.")
    return out2, soft2, notes2, r2["call_id"]


def act_labels_text(st, stem):
    """For the review: the labels the act created, each with its definition and up to REVIEW_EXAMPLES of its codes."""
    created = [x for x in st["labels"].values() if x.get("first_act") == stem]
    if not created:
        return "You created no new labels in this act."
    ex = collections.defaultdict(list)
    for r in act_records(st["agent"], stem, renamed=False):
        for c in r["codes"]:
            k = label_key(c["label"])
            if len(ex[k]) < REVIEW_EXAMPLES:
                ex[k].append(f"  {r['turn_id']}: {_one(c.get('description'))}")
    lines = [f"The labels you created in this act ({len(created)}), each as label | times used | definition, with up to "
             f"{REVIEW_EXAMPLES} of its codes (Turn ID: description):"]
    for x in _sorted_labels({label_key(x["label"]): x for x in created}):
        lines += [label_line(x)] + ex.get(label_key(x["label"]), [])
    return "\n".join(lines)


def review_act(st, a, fixed, hay):
    """The review of an act, after its last batch (Prompt 2b): sent after the same fixed part, memory and labels as
    the act's coding calls, so that they are read from the cache."""
    stem = act_stem(a)
    tail = "\n\n" + fill(PROMPTS["2b"], {"ACT": a["Act"], "KIND": kind(a), "ACT_LABELS": act_labels_text(st, stem)})
    out, hard, soft, notes_, res = ask_until_usable(
        st, "review", act_parts(st, stem, fixed, tail), SCHEMA_REVIEW, lambda o: check_review(o, st, stem, hay),
        {"round": 2, "act": stem, "batch": 0, "label": f"{a['Act']} (review)", "cache": True, "log": stem,
         "review": True, "memory_upto": len(st["memory"])},
        "in the JSON format given above, for the review of the act")
    rv = dict(out, agent=st["agent"], act=stem, problems=soft + notes_, call=res["call_id"], time=now())
    write_json(review_path(st["agent"], stem), rv)
    apply_review(st, stem, rv)
    flag(st, 2, stem, 0, soft + notes_, res["call_id"])
    if stem not in st["acts_complete"]:
        st["acts_complete"].append(stem)
    save_state(st)
    say(f"[{st['agent']}] {a['Act']} reviewed: {len(rv['renames'])} label(s) renamed, {len(rv['definitions'])} "
        f"definitions, a memo of {len(rv['memo'])} point(s), {len(rv['keep'])} of {len(act_entries(st, stem))} notes "
        f"kept.")


def write_memory_files(st):
    write_text(rpath(f"agent_{st['agent']}", "memory.txt"), memory_text(st) + "\n")
    write_text(rpath(f"agent_{st['agent']}", "notes_all.txt"), all_notes_text(st))


def _round2(letter, acts, stop_at=None):
    st = load_state(letter)
    if READ_ACTS:
        unread = [a for a in acts if not os.path.exists(reading_path(letter, act_stem(a)))]
        if unread and any(batch_id(letter, a) in reading_pending(letter) for a in unread):
            raise AgentStop(f"the reading of {len(unread)} act(s) is still in its batch: run run_round1() again to "
                            f"collect it, then round 2")
        if unread and not os.path.isdir(rpath(f"agent_{letter}", "reading")):
            raise AgentStop("round 1 (the reading) has not been run: run run_round1() first")
    st["stopped"] = None
    rebuild_state(st, acts)
    save_state(st)
    pre, post = split_at_memory(PROMPTS[2])
    fixed = fill(pre, {"GUIDE": load_guide(), "LETTER": letter})
    for a in acts:
        stem = act_stem(a)
        if stem not in st["acts_complete"]:
            if (st.get("act_start") or {}).get("stem") != stem:   # a new act: its labels list as it starts
                st["act_start"] = {"stem": stem, "labels": {k: v["count"] for k, v in st["labels"].items()}}
                save_state(st)
            lvl = labels_level(st, stem, fixed)        # the same for every call of the act, so it stays cached
            st["act_start"]["level"] = lvl
            if lvl and st["label_levels"].get(stem) != lvl:
                how = "with the first words of each definition" if lvl == 1 else "without their definitions"
                say(f"[{letter}] {a['Act']}: the labels from earlier acts are listed {how}, as in full the act's calls "
                    f"could pass {MODELS[letter]['model']}'s context window.")
                flag(st, 2, stem, 0, [f"In this act, the labels from earlier acts were listed {how}, to fit the "
                                      f"context window."], None)
            st["label_levels"][stem] = lvl
            save_state(st)
            batches = make_batches(CORPUS["turns"][stem])
            done = set(st["done"].get(stem, []))
            hay = _act_haystack([stem])
            for bi, b in enumerate(batches, 1):
                if bi in done:
                    continue
                if STOP.is_set():
                    raise stop_reason()
                rm = insert_reading(st, a) if READ_ACTS and bi == 1 else None
                vals = {"NEXT": st["next"], "LABELS": new_labels_text(st, stem), "ACT": a["Act"], "KIND": kind(a),
                        "BATCH_N": bi, "BATCH_TOTAL": len(batches), "N_INTERVIEWS": b["k"], "N_TURNS": len(b["turns"]),
                        "TURNS": render_turns(b["turns"])}
                parts = act_parts(st, stem, fixed, fill(post, vals))
                meta = {"round": 2, "act": stem, "batch": bi, "cache": True, "log": stem, "memory_upto": len(st["memory"]),
                        "label": f"{a['Act']} batch {bi}/{len(batches)} ({len(b['turns'])} turns)"}
                out, hard, soft, notes, res = ask_until_usable(
                    st, "round2", parts, SCHEMA2, lambda o: check_round2(o, b["turns"], st, hay), meta,
                    "in the JSON format given above, with one record per turn, in the order of the turns", soft_retries=0)
                fix_call = None
                if soft and RETRIES_PROBLEMS > 0 and not STOP.is_set():
                    out, soft, notes, fix_call = fix_turns(st, parts, b, res, out, soft, notes, meta, hay)
                _accept(st, stem, bi, b, out, soft + notes, res["call_id"], rm, fix_call)
            if STOP.is_set():
                raise stop_reason()
            review_act(st, a, fixed, hay)
            write_codes_csv(letter, a)
            write_memory_files(st)
        elif not os.path.exists(rpath(f"agent_{letter}", "codes", f"codes_{stem}.csv")):
            write_codes_csv(letter, a)
        if not reported(letter, a["Act"]):
            act_report(st, a)
        if stem == stop_at and stem != act_stem(acts[-1]):
            write_memory_files(st)
            n = sum(1 for x in acts if act_stem(x) in st["acts_complete"])
            say(f"[{letter}] stopped as asked, after {a['Act']}: {n} of {len(acts)} acts coded and reviewed.")
            return f"coded up to {a['Act']} ({n} of {len(acts)} acts)"
    write_memory_files(st)


def parts(acts=None):
    """The acts of round 2 in PARTS consecutive groups, in coding order, each with about the same number of turns
    (the cut falls between acts; it is worked out from all the turns, so the excluded turns do not move it)."""
    sel = select_acts(acts)
    k = max(1, min(PARTS, len(sel)))
    cum, s = [], 0
    for a in sel:
        s += len(load_corpus()["all_turns"][act_stem(a)])
        cum.append(s)
    cuts, lo = [], 0
    for j in range(1, k):
        best = min(range(lo, len(sel) - (k - j)), key=lambda i: abs(cum[i] - s * j / k))
        cuts.append(best)
        lo = best + 1
    bounds = [-1] + cuts + [len(sel) - 1]
    return [sel[bounds[i] + 1:bounds[i + 1] + 1] for i in range(k)]


def show_parts(acts=None):
    """The parts of round 2 (PARTS): their acts, turns and batches, and how far each agent has got in each."""
    load_corpus()
    states = {L: load_state(L) for L in LETTERS}
    for i, p in enumerate(parts(acts), 1):
        nt = sum(len(CORPUS["turns"][act_stem(a)]) for a in p)
        nb = sum(len(make_batches(CORPUS["turns"][act_stem(a)])) for a in p)
        n_done = {L: sum(1 for a in p if act_stem(a) in st["acts_complete"]) for L, st in states.items()}
        done = ", ".join(f"{L} {n}" for L, n in n_done.items())
        print(f"Part {i}: {len(p)} acts, {p[0]['Act']} ({p[0]['Nominated']}) to {p[-1]['Act']} ({p[-1]['Nominated']}), "
              f"{nt:,} turns in {nb} batches; acts done: {done} of {len(p)}"
              + (": done." if all(n == len(p) for n in n_done.values()) else f". Run run_round2(part={i})."))


def run_round2(acts=None, agents=None, part=None, up_to=None):
    """Round 2: each agent codes the acts (ACTS: the pilot act, or all 52 in nomination-year order), leaving out the
    turns in step3_inputs/excluded_turns.csv, which is read afresh each time the round is run. With part (1 to
    PARTS; show_parts() lists them), it codes the acts up to the end of that part and stops there; with up_to (an
    act's name, or a number of acts), up to that act. Either way it first finishes any act before it, so the acts are
    always coded in order, and the run's settings are those of the whole round."""
    fresh_corpus(), load_prompts(), load_guide()
    unk = CORPUS.get("excluded_unknown") or []
    if unk:
        raise SystemExit(f"STOPPED: {len(unk)} Turn ID(s) in {EXCLUDE_FILE} are not in the corpus, e.g. "
                         f"{', '.join(unk[:5])}. A spreadsheet may have turned them into dates. Put them right, or "
                         f"delete those rows, then run round 2 again. remove_from_exclusion_list() and "
                         f"add_to_exclusion_list() change the list without opening it in a spreadsheet.")
    sel = select_acts(acts)
    stop_at = None
    if part is not None:
        ps = parts(acts)
        if not isinstance(part, int) or not 1 <= part <= len(ps):
            raise SystemExit(f"STOPPED: part must be 1 to {len(ps)} (show_parts() lists them).")
        stop_at = act_stem(ps[part - 1][-1])
    elif up_to is not None:
        if isinstance(up_to, int) and not isinstance(up_to, bool):
            if up_to < 1:
                raise SystemExit("STOPPED: up_to must be 1 or more acts.")
            stop_at = act_stem(sel[min(up_to, len(sel)) - 1])
        else:
            hit = [a for a in sel if a["Act"].casefold() == str(up_to).casefold()]
            if not hit:
                raise SystemExit(f"STOPPED: no act called {up_to!r} among the acts of this run.")
            stop_at = act_stem(hit[0])
    check_settings(2, {"acts": [act_stem(a) for a in sel]}, agents=agents)
    return run_agents(_round2, agents, sel, stop_at)


def skip_batch(agent, act, batch, reason="skipped by hand"):
    """A last resort, when an agent stops on the same batch every time: that batch's turns are left uncoded, marked
    as skipped and flagged, and the agent carries on from the next batch when round 2 runs again.
    Example: skip_batch("B", "Big Country", 7, "Gemini refused the batch three times")"""
    load_corpus()
    if busy():
        print("Agents are still running; wait until they have stopped.")
        return
    a = select_acts([act])[0]
    stem = act_stem(a)
    run_acts = read_json(rpath("run_settings.json"), {}).get("2", {}).get("acts") or []
    if stem not in run_acts:
        print(f"{a['Act']} is not one of the acts of round 2 in run {RUN!r}; nothing skipped.")
        return
    st = load_state(agent)
    rebuild_state(st, [x for x in select_acts("all") if act_stem(x) in run_acts])
    batches = make_batches(CORPUS["turns"][stem])
    earlier = run_acts[:run_acts.index(stem)]
    if not 1 <= batch <= len(batches) or any(s not in st["acts_complete"] for s in earlier) \
            or sorted(st["done"].get(stem, [])) != list(range(1, batch)):
        print(f"Batch {batch} of {a['Act']} is not agent {agent}'s next batch; nothing skipped. status() shows where "
              f"each agent is.")
        return
    b = batches[batch - 1]
    recs = [{"turn_id": t["Turn ID"], "looks_back": False, "data_problem": "", "codes": [], "memory": [],
             "skipped": True} for t in b["turns"]]
    rm = insert_reading(st, a) if READ_ACTS and batch == 1 else None
    _accept(st, stem, batch, b, {"records": recs, "new_labels": {}},
            [f"batch skipped by hand, its {len(recs)} turns left uncoded: {reason}"], "skipped", rm)
    say(f"Agent {agent}: batch {batch} of {a['Act']} skipped ({len(recs)} turns, flagged). Run round 2 again.")


def write_codes_csv(letter, a, path=None):
    stem = act_stem(a)
    path = path or rpath(f"agent_{letter}", "codes", f"codes_{stem}.csv")
    ensure(os.path.dirname(path))
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Turn ID", "Label", "Label from", "Description", "Evidence", "Unsure", "Looks back",
                    "Data problem", "Label as coded"])
        for r in act_records(letter, stem):
            lb, dp = "yes" if r.get("looks_back") else "", r.get("data_problem") or ""
            if not r["codes"]:
                w.writerow([r["turn_id"], "(skipped)" if r.get("skipped") else "(no code)", "", "", "", "", lb, dp, ""])
            for c in r["codes"]:
                w.writerow([r["turn_id"], c["label"], c["label_from"], c.get("description", ""), c["evidence"],
                            "yes" if c.get("unsure") else "", lb, dp, c.get("original_label", "")])
    return path

# ==================================================================================================== rounds 3 and 5: the agent's folder


def folder_tools(folder):
    """The agent's folder tools for one memo. Past READ_BUDGET characters or MAX_TOOL_STEPS uses, every further use
    is answered with a note to write the memo now, so that a memo never stops half-way for reading too much."""
    files = sorted(f for f in os.listdir(folder) if not f.startswith("."))
    cache, used = {}, {"uses": 0, "chars": 0}

    def lines_of(name):
        if name not in cache:
            with open(os.path.join(folder, name), encoding="utf-8-sig") as fh:
                cache[name] = fh.read().splitlines()
        return cache[name]

    def run(name, args):
        used["uses"] += 1
        if used["uses"] > MAX_TOOL_STEPS or used["chars"] >= READ_BUDGET:
            return (f"You have reached this memo's reading limit (about {READ_BUDGET:,} characters, or "
                    f"{MAX_TOOL_STEPS} uses of the tools). Do not use the tools again: write the memo now, from what "
                    f"you have read and your memory.")
        out = _run(name, args)
        used["chars"] += len(out)
        return out

    def _run(name, args):
        try:
            if name == "list_files":
                return "\n".join(f"{f} ({len(lines_of(f))} lines)" for f in files)
            if name == "read_file":
                fname = str(args.get("name") or "").strip().split("/")[-1]
                if fname not in files:
                    return f"There is no file called {fname!r} in your folder. Use list_files to see the files."
                ls = lines_of(fname)
                s = max(1, _int(args.get("start_line") or 1))
                if s > len(ls):
                    return f"{fname} has {len(ls)} lines; there is nothing from line {s}."
                out, size, e = [], 0, s - 1
                for i in range(s - 1, min(len(ls), s - 1 + READ_LINES)):
                    line = ls[i]
                    if out and size + len(line) + 1 > READ_CHARS:
                        break
                    if len(line) > READ_CHARS:
                        line = line[:READ_CHARS] + " … [this line is longer; it is cut here]"
                    out.append(line)
                    size += len(line) + 1
                    e = i + 1
                more = f"\n(Read on from line {e + 1}.)" if e < len(ls) else "\n(End of the file.)"
                return f"{fname}, lines {s} to {e} of {len(ls)}:\n" + "\n".join(out) + more
            return f"Unknown tool {name!r}."
        except Exception as ex:
            return f"The file could not be read: {ex}"
    return run


def build_folder(letter, acts, with_merged):
    folder = rpath(f"agent_{letter}", "folder_round5" if with_merged else "folder_round3")
    ensure(folder)
    st = load_state(letter)
    for a in acts:
        write_codes_csv(letter, a, os.path.join(folder, f"codes_{act_stem(a)}.csv"))
    acts_of = collections.defaultdict(set)
    for a in acts:
        for r in act_records(letter, act_stem(a)):
            for c in r["codes"]:
                acts_of[label_key(c["label"])].add(a["Act"])
    with open(os.path.join(folder, "labels_summary.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Label", "Data or own", "Times used", "Acts", "First used in", "Definition"])
        for k, x in sorted(st["labels"].items(), key=lambda kv: (-kv[1]["count"], kv[1]["label"].casefold())):
            w.writerow([x["label"], x["from"], x["count"], "; ".join(sorted(acts_of.get(k, []))), x["first_turn"],
                        x.get("definition") or x.get("description", "")])
    if with_merged:
        for a in acts:
            shutil.copyfile(rpath("merged", "by_act", f"merged_{act_stem(a)}.csv"),
                            os.path.join(folder, f"merged_{act_stem(a)}.csv"))
        shutil.copyfile(rpath("merged", "agreement_figures.txt"), os.path.join(folder, "agreement_figures.txt"))
    return folder


def part_heading(n):
    """A line that starts with the heading Part n (with # or ** around it, or not)."""
    return re.compile(rf"^\s*(?:#+\s*)?\**\s*Part\s+{n}\b", re.I | re.M)


PART_RE = {n: part_heading(n) for n in (1, 2, 3)}
PART1_LINE = re.compile(r"^\s*(?:#+\s*)?\**\s*Part\s+1\s*\**\s*:?\s*\**\s*$", re.I | re.M)   # a bare Part 1 heading


def ask_text(st, kind_, parts, tools, check, meta):
    """Memo rounds: asks and checks the headings, and asks once more if a check fails. Keeps the better complete
    reply; a reply that was cut off is never kept (the agent stops, and running the round again asks again)."""
    tail0, tries = parts["tail"], []
    for attempt in (1, 2):
        if STOP.is_set():
            break
        res = call_model(st, kind_, parts, tools=tools, meta=dict(meta, attempt=attempt))
        probs = check(res["text"]) if res["ok"] else [f"The reply was cut off or stopped early ({res['stop']})."]
        tries.append((res, probs))
        if not probs:
            break
        say(f"[{st['agent']}] {meta['label']}: {probs[0]}")
        parts = dict(parts, tail=tail0 + problem_note(probs, False, "as the prompt asks"))
    good = [x for x in tries if x[0]["ok"] and x[0]["text"].strip()]
    if not good:
        raise AgentStop(f"{meta['label']}: no complete reply; " + (STOP_WHY[0] if STOP.is_set() else
                                                                   "run the round again"))
    return min(good, key=lambda x: len(x[1]))


def _round3(letter, acts):
    st = load_state(letter)
    stems = [act_stem(a) for a in acts]
    missing = [s for s in stems if s not in st["acts_complete"]]
    if missing:
        raise AgentStop(f"round 2 is not finished for {len(missing)} act(s), e.g. {missing[0]}")
    if st["memo1"]:
        say(f"[{letter}] round 3 is already done.")
        return
    st["stopped"] = None
    tools = folder_tools(build_folder(letter, acts, False))
    pre, post = split_at_memory(PROMPTS[3])
    vals = {"GUIDE": load_guide(), "LETTER": letter}
    parts = {"fixed": fill(pre, vals), "memory": [memory_text(st)], "tail": fill(post, vals)}

    scope = [t["Turn ID"] for s in stems for t in CORPUS["all_turns"][s]]     # every turn the agent has read
    own_labels = [v["label"] for v in st["labels"].values()]

    def check(text):
        first = next((ln for ln in text.splitlines() if ln.strip()), "")
        p = [] if PART_RE[1].match(first) else ['Part 1 must start with the heading "Part 1".']
        return p + quote_problems(text, scope, own_labels)
    res, probs = ask_text(st, "round3", parts, tools, check, {"round": 3, "label": "round 3 (memo part 1)",
                                                               "log": "round3"})
    write_text(rpath(f"agent_{letter}", "memo_part1.md"), res["text"].strip() + "\n")
    write_text(rpath(f"agent_{letter}", f"memo_{letter}.md"), res["text"].strip() + "\n")
    st["memo1"] = True
    save_state(st)                             # the memo, paid for, is kept before the slower name check
    names = unknown_names(res["text"], _act_haystack(stems))
    if names:
        probs = probs + [f"The memo names people or things that appear nowhere in the turns: {', '.join(names)}."]
    flag(st, 3, "", 0, probs, res["call_id"])
    save_state(st)
    say(f"[{letter}] round 3 done: memo part 1 has {len(res['text'].split()):,} words, "
        f"after {res.get('tool_steps', 0)} file reads.")


def run_round3(acts=None, agents=None):
    """Round 3: each agent writes part 1 of its memo, about its own coding only."""
    load_corpus(), load_prompts(), load_guide()
    check_settings(3, agents=agents)
    return run_agents(_round3, agents, select_acts(acts))

# ==================================================================================================== round 4: the merge


STOPWORDS = set("a an the and or but of to in on at for with by from as is was were be been being it its this that "
                "these those i we you he she they me us him her them my our your his their not no so do did".split())


def _codes_cell(rec):
    if rec.get("skipped"):
        return "(skipped)"
    if not rec["codes"]:
        return "(no code)"
    one = lambda x: re.sub(r"\s+", " ", str(x)).strip()

    def tag(c):                                # [data] or [own, unsure]; runs before 1.7 had a confidence instead
        extra = ", unsure" if c.get("unsure") else (f", {c['confidence']}" if c.get("confidence") else "")
        return f"[{c['label_from']}{extra}]"
    return "\n".join(f'{one(c["label"])} {tag(c)}: '
                     + (f'{one(c["description"])} ' if c.get("description") else "") + f'"{one(c["evidence"])}"'
                     for c in rec["codes"])


def _label_words(codes):
    return {w for c in codes for w in norm_words(c["label"]) if w not in STOPWORDS}


def run_round4(acts=None):
    """Round 4: the script merges the three agents' codes, turn by turn, without changing or resolving any; works out
    the agreement figures for each act and the sample of turns for round 5; and checks that nothing was lost."""
    load_corpus()
    acts = select_acts(acts)
    if busy():
        raise SystemExit("STOPPED: agents are still running; wait until they have stopped.")
    check_settings(4)
    states = {L: load_state(L) for L in LETTERS}
    for L, st in states.items():
        miss = [a["Act"] for a in acts if act_stem(a) not in st["acts_complete"]]
        if miss:
            raise SystemExit(f"STOPPED: agent {L} has not finished coding {len(miss)} act(s), e.g. {miss[0]}.")
    rows, figures, per_turn = [], [], []
    n_codes = {L: 0 for L in LETTERS}
    for a in acts:
        stem = act_stem(a)
        turns = CORPUS["turns"][stem]
        recs = {}
        for L in LETTERS:
            lst = act_records(L, stem)
            if [r["turn_id"] for r in lst] != [t["Turn ID"] for t in turns]:
                raise SystemExit(f"STOPPED: agent {L}'s records for {a['Act']} do not match the act's turns, one "
                                 f"record per turn in order.")
            recs[L] = {r["turn_id"]: r for r in lst}
            n_codes[L] += count_codes(L, stem)
        all3 = none = shared_turns = 0
        label_agents = collections.defaultdict(collections.Counter)
        skipped_by = collections.Counter()
        out_rows = []
        for t in turns:
            tid = t["Turn ID"]
            rr = {L: recs[L][tid] for L in LETTERS}
            cs = {L: rr[L]["codes"] for L in LETTERS}
            sets = {L: {label_key(c["label"]) for c in cs[L]} for L in LETTERS}
            skipped = [L for L in LETTERS if rr[L].get("skipped")]
            if skipped:                        # a turn an agent skipped is in the sheet, but not compared or sampled
                skipped_by.update(skipped)
                out_rows.append([tid, a["Act"], t["Order"], t.get("Speaker"), PERIOD.get(t.get("Period"), ""),
                                 t["Text"]] + [_codes_cell(rr[L]) for L in LETTERS] + [len(cs[L]) for L in LETTERS]
                                + [""])
                continue
            for L in LETTERS:
                for k in sets[L]:
                    label_agents[k][L] += 1
            coded = [L for L in LETTERS if cs[L]]
            all3 += len(coded) == len(LETTERS)
            none += not coded
            shared, sims = set(), []
            for i, x in enumerate(LETTERS):
                for y in LETTERS[i + 1:]:
                    shared |= sets[x] & sets[y]
                    if not cs[x] and not cs[y]:
                        sims.append(1.0)               # both left the turn uncoded: they agree
                    elif bool(cs[x]) != bool(cs[y]):
                        sims.append(0.0)               # one coded it, the other did not
                    else:                              # both coded it: the words their labels share
                        wx, wy = _label_words(cs[x]), _label_words(cs[y])
                        sims.append(len(wx & wy) / len(wx | wy) if wx | wy else float(sets[x] == sets[y]))
            shared_turns += bool(shared)
            per_turn.append({"tid": tid, "diff": round(1 - sum(sims) / len(sims), 3) if sims else 0.0})
            out_rows.append([tid, a["Act"], t["Order"], t.get("Speaker"), PERIOD.get(t.get("Period"), ""), t["Text"]]
                            + [_codes_cell(rr[L]) for L in LETTERS] + [len(cs[L]) for L in LETTERS]
                            + ["; ".join(sorted(shared))])
        rows += out_rows
        by2 = sorted((k for k, v in label_agents.items() if len(v) >= 2), key=lambda k: -sum(label_agents[k].values()))
        figures.append({"act": a["Act"], "turns": len(turns), "all_three_coded": all3, "none_coded": none,
                        "turns_with_a_shared_label": shared_turns,
                        "labels_shared_by_three": [k for k in by2 if len(label_agents[k]) == 3],
                        "labels_shared_by_two": [k for k in by2 if len(label_agents[k]) == 2],
                        "turns_skipped": dict(skipped_by)})
        p = rpath("merged", "by_act", f"merged_{stem}.csv")
        ensure(os.path.dirname(p))
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["Turn ID"] + [f"Agent {L}" for L in LETTERS])
            for r in out_rows:
                w.writerow([r[0]] + [x.replace("\n", " || ") for x in r[6:6 + len(LETTERS)]])
    # the check: every code of every agent is in the sheet, nothing dropped or changed
    in_sheet = {L: sum(0 if r[6 + i] in ("(no code)", "(skipped)") else r[6 + i].count("\n") + 1 for r in rows)
                for i, L in enumerate(LETTERS)}
    if in_sheet != n_codes:
        raise SystemExit(f"STOPPED: the merged sheet has {in_sheet} codes; the agents' batch files have {n_codes}.")
    # the figures, as the agents will read them
    lines = []
    for f in figures:
        sk = f["turns_skipped"]
        lines.append(f"{f['act']}: {f['turns']} turns"
                     + (" (left out of these figures, skipped by hand: " + ", ".join(f"Agent {L} {n}" for L, n in
                                                                                    sorted(sk.items())) + ")" if sk else "")
                     + f"; all three agents coded {f['all_three_coded']}, none coded {f['none_coded']}; on "
                       f"{f['turns_with_a_shared_label']} turns at least two agents gave the same label.")
        for key, what in (("labels_shared_by_three", "Labels all three used"),
                          ("labels_shared_by_two", "Labels two used")):
            if f[key]:
                lines.append(f"  {what}: " + "; ".join(f[key][:15])
                             + (f" (and {len(f[key]) - 15} more)" if len(f[key]) > 15 else ""))
    fig_text = ("Agreement figures for each act (labels compared after lower-casing and trimming punctuation):\n"
                + "\n".join(lines))
    write_text(rpath("merged", "agreement_figures.txt"), fig_text)
    write_json(rpath("merged", "agreement_figures.json"), figures)
    # the sample: half at random; half where the agents differed most (one coded a turn and another did not, or
    # their labels have few words in common)
    rnd = random.Random(SEED + 5)
    ids = [p["tid"] for p in per_turn]
    order = {tid: i for i, tid in enumerate(ids)}
    half = min(SAMPLE_SIZE // 2, len(ids))
    rand_half = rnd.sample(ids, half)
    chosen = set(rand_half)
    rest = sorted((p for p in per_turn if p["tid"] not in chosen), key=lambda p: (-p["diff"], rnd.random()))
    diff_half = [p["tid"] for p in rest[:max(0, min(SAMPLE_SIZE, len(ids)) - half)]]
    pick = sorted(rand_half + diff_half, key=order.get)
    by_tid = {r[0]: r for r in rows}
    blocks = []
    for tid in pick:
        r, t = by_tid[tid], CORPUS["by_id"][tid]
        b = [f"Turn ID: {tid}", f"Act: {r[1]}", f"Speaker: {r[3]}", f"Text: {t['Text']}"]
        for i, L in enumerate(LETTERS):
            b.append(f"Agent {L}'s codes:\n" + r[6 + i])
        blocks.append("\n".join(b))
    sample_text = f"A sample of {len(pick)} turns with all three agents' codes:\n\n" + "\n\n".join(blocks)
    write_text(rpath("merged", "sample_of_turns.txt"), sample_text)
    diffs = {p["tid"]: p["diff"] for p in per_turn}
    write_json(rpath("merged", "sample.json"), {"turn_ids": pick, "random_half": sorted(rand_half, key=order.get),
                                                "differed_most": sorted(diff_half, key=order.get), "seed": SEED + 5,
                                                "difference": {tid: diffs[tid] for tid in pick}})
    write_merged_xlsx(rows, figures)
    write_json(rpath("merged", "done.json"), {"acts": [act_stem(a) for a in acts], "time": now(),
                                              "codes": n_codes, "sample": len(pick)})
    say(f"Round 4 done: {len(rows):,} turns from {len(acts)} act(s) in merged/step3_merged_codes.xlsx; codes per "
        f"agent {n_codes}; a sample of {len(pick)} turns for round 5. Every code is in the sheet.")


def write_merged_xlsx(rows, figures):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def clean(x):
        return ILLEGAL_CHARACTERS_RE.sub("", x)[:32000] if isinstance(x, str) else x
    wb = Workbook()
    ws = wb.active
    ws.title = "Codes by turn"
    ws.append(["Turn ID", "Act", "Order", "Speaker", "Period", "Text"] + [f"Agent {L} codes" for L in LETTERS]
              + [f"Agent {L} n" for L in LETTERS] + ["Labels shared by two or three"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([clean(x) for x in r])
    widths = [10, 16, 7, 18, 20, 70] + [45] * len(LETTERS) + [6] * len(LETTERS) + [30]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = wdt
    ws.freeze_panes = "B2"
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    ws2 = wb.create_sheet("Agreement by act")
    ws2.append(["Act", "Turns", "All three coded", "None coded", "Turns with a shared label",
                "Labels all three used", "Labels two used", "Turns skipped by hand (left out of the figures)"])
    for f in figures:
        ws2.append([f["act"], f["turns"], f["all_three_coded"], f["none_coded"], f["turns_with_a_shared_label"],
                    clean("; ".join(f["labels_shared_by_three"])), clean("; ".join(f["labels_shared_by_two"])),
                    "; ".join(f"Agent {L}: {n}" for L, n in sorted(f["turns_skipped"].items()))])
    ws3 = wb.create_sheet("Read me")
    for line in [
        "Step 3 merged codes: what each agent coded, turn by turn. Nothing here is changed or resolved.",
        "Codes by turn: one row per turn, in each act's date order; each agent's codes as label [data or own, and "
        "unsure where the agent was unsure]: description \"evidence\", with each label as the act's review named it; "
        "(no code) where the agent left the turn uncoded; (skipped) where a batch was skipped by hand. Data or own is "
        "worked out by the script: data when the label's words are in the turn, or in the turn where it was first "
        "used.",
        "Agreement by act: turns all three agents coded, turns none coded, turns where at least two gave the same "
        "label (compared after lower-casing and trimming punctuation), and those labels.",
        f"Made by {VERSION}, run {RUN}, on {now()}."]:
        ws3.append([line])
    p = rpath("merged", "step3_merged_codes.xlsx")
    ensure(os.path.dirname(p))
    wb.save(p)

# ==================================================================================================== round 5


def _round5(letter, acts):
    st = load_state(letter)
    if not st["memo1"]:
        raise AgentStop("round 3 (memo part 1) is not done for this agent")
    done = read_json(rpath("merged", "done.json"))
    if not done or done["acts"] != [act_stem(a) for a in acts]:
        raise AgentStop("round 4 (the merge) has not been run for these acts")
    if st["memo23"]:
        say(f"[{letter}] round 5 is already done.")
        return
    st["stopped"] = None
    tools = folder_tools(build_folder(letter, acts, True))
    with open(rpath(f"agent_{letter}", "memo_part1.md"), encoding="utf-8") as f:
        memo1 = f.read().strip()
    with open(rpath("merged", "agreement_figures.txt"), encoding="utf-8") as f:
        figs = f.read()
    with open(rpath("merged", "sample_of_turns.txt"), encoding="utf-8") as f:
        sample = f.read()
    text = fill(PROMPTS[4], {"GUIDE": load_guide(), "LETTER": letter, "MEMO1": memo1, "FIGURES": figs,
                             "P4SAMPLE": sample})

    scope = [t["Turn ID"] for a in acts for t in CORPUS["all_turns"][act_stem(a)]]
    all_labels = [v["label"] for L in LETTERS for v in load_state(L)["labels"].values()]

    def check(t):
        p = []
        if not PART_RE[2].search(t) or not PART_RE[3].search(t):
            p.append('Parts 2 and 3 must start with the headings "Part 2" and "Part 3".')
        if PART1_LINE.search(t):
            p.append("Return parts 2 and 3 only; part 1 stays as it is.")
        return p + quote_problems(t, scope, all_labels)
    res, probs = ask_text(st, "round5", {"fixed": text, "memory": [], "tail": ""}, tools, check,
                          {"round": 5, "label": "round 5 (memo parts 2 and 3)", "log": "round5"})
    write_text(rpath(f"agent_{letter}", "memo_parts2and3.md"), res["text"].strip() + "\n")
    write_text(rpath(f"agent_{letter}", f"memo_{letter}.md"), memo1 + "\n\n" + res["text"].strip() + "\n")
    flag(st, 5, "", 0, probs, res["call_id"])
    st["memo23"] = True
    save_state(st)
    say(f"[{letter}] round 5 done: parts 2 and 3 appended to memo_{letter}.md ({len(res['text'].split()):,} words).")


def run_round5(acts=None, agents=None):
    """Round 5: each agent reads the agreement figures and the sample and writes memo parts 2 and 3."""
    load_corpus(), load_prompts(), load_guide()
    check_settings(5, agents=agents)
    return run_agents(_round5, agents, select_acts(acts))

# ==================================================================================================== checks, dry run, test, reports


def check_inputs():
    """Checks the corpus, the prompts, the guide, the keys and the SDK versions, without calling any model."""
    print(VERSION)
    for mod, want in (("anthropic", "2.3.0"), ("openai", "3.37.0"), ("google.genai", "2.33.0")):
        try:
            v = getattr(__import__(mod, fromlist=["_"]), "__version__", "?")
            print(f"{mod} {v}" + ("" if v == want else f"  (the script was tested with {want}: run the pip line in "
                                                         f"the notes at the top, then Runtime > Restart session)"))
        except Exception as e:
            print(f"{mod}: NOT INSTALLED ({e}); run the pip line in the notes at the top")
    c = fresh_corpus()
    n_coded = sum(len(v) for v in c["turns"].values())
    print(f"Corpus: {len(c['acts'])} acts, {len(c['by_id']):,} turns, "
          f"{sum(words(t) for t in c['by_id'].values()):,} words. Fingerprint {corpus_fingerprint(False)[:16]}.")
    print(f"Excluded turns ({EXCLUDE_FILE}): {len(c['excluded'])}" + (f", so round 2 codes {n_coded:,} turns"
                                                                     if c["excluded"] else ", so round 2 codes every turn")
          + (" (the list is made by make_exclusion_list() after round 1)" if not os.path.exists(exclude_path()) else ""))
    if c.get("excluded_unknown"):
        print(f"  PROBLEM: {len(c['excluded_unknown'])} Turn ID(s) in it are not in the corpus, e.g. "
              f"{', '.join(c['excluded_unknown'][:5])}: a spreadsheet may have turned them into dates. Round 2 will "
              f"not start until they are put right.")
    for n in PROMPT_FILES:
        s = read_prompt(n)
        pr = check_prompt(n, s)
        print(f"Prompt {n} ({PROMPT_FILES[n]}): {len(s.split()):,} words; placeholders "
              f"{sorted(set(k for k in keys_in(s) if k))}" + (f"; PROBLEMS: {pr}" if pr else "; fine"))
    g = load_guide()
    print(f"Guide ({GUIDE_FILE}): {len(g.split()):,} words, about {est_tokens(g):,} tokens. Starts: {short(g, 100)}")
    for a in select_acts():
        bs = make_batches(c["turns"][act_stem(a)])
        print(f"{a['Act']} ({kind(a)}, nominated {a['Nominated']}): {a['Turns']} turns in {len(bs)} batches, "
              f"{min(len(b['turns']) for b in bs)} to {max(len(b['turns']) for b in bs)} turns each")
    for L, cfg in MODELS.items():
        try:
            secret(SECRETS[cfg["provider"]])
            ok = "key found"
        except AgentStop as e:
            ok = f"PROBLEM: {e}"
        print(f"Agent {L}: {cfg['model']} ({cfg['provider']}), effort {cfg.get('effort')}: {ok}")
    load_prompts()
    print(f"Run {RUN!r}, acts {ACTS!r}; the run folder is {rpath()}")
    if len(select_acts()) > 1 and PARTS > 1:
        print(f"Round 2 in {len(parts())} parts:")
        show_parts()


def dry_run(acts=None):
    """Writes the filled prompts to step3_<RUN>/dry_run, so you can read exactly what each agent will get, and
    prints the token counts. No model is called."""
    fresh_corpus(), load_prompts()
    g = load_guide()
    d = ensure(rpath("dry_run"))
    sel = select_acts(acts)
    if READ_ACTS:
        p1 = reading_parts("A", sel[0])["fixed"]
        write_text(os.path.join(d, f"prompt1_agent_A_{act_stem(sel[0])}.txt"), p1)
        toks = [est_tokens(reading_parts("A", a)["fixed"]) for a in sel]
        cost = {L: reading_cost(L, toks, READ_BATCH) for L in LETTERS}
        print(f"Prompt 1 (reading), {sel[0]['Act']}: about {toks[0]:,} tokens (the guide is about {est_tokens(g):,}). "
              f"Reading all {len(sel)} act(s) {'as a batch at half price' if READ_BATCH else 'in normal calls'}: "
              f"about {sum(toks):,} tokens per agent, roughly "
              + ", ".join(f"Agent {L} ${c:,.2f}" for L, c in cost.items()) + f" (${sum(cost.values()):,.2f} in all).")
    pre, post = split_at_memory(PROMPTS[2])
    fixed = fill(pre, {"GUIDE": g, "LETTER": "A"})
    st = new_state("A")
    for a in sel:
        st.update(memory=[{"n": 1, "turn": "general", "act": act_stem(a), "round": "reading",
                           "text": "M1 (general, reading [the act]): [the notes from reading the act go here]"}],
                  blocks=[[0, 1]], next=2)
        bs = make_batches(CORPUS["turns"][act_stem(a)])
        with open(os.path.join(d, f"batches_{act_stem(a)}.csv"), "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["Batch", "Interviews or parts", "Turns", "Words", "First Turn ID", "Last Turn ID"])
            for i, b in enumerate(bs, 1):
                w.writerow([i, b["k"], len(b["turns"]), sum(map(words, b["turns"])), b["turns"][0]["Turn ID"],
                            b["turns"][-1]["Turn ID"]])
        b = bs[0]
        vals = {"NEXT": st["next"], "LABELS": new_labels_text(st, act_stem(a)), "ACT": a["Act"], "KIND": kind(a),
                "BATCH_N": 1, "BATCH_TOTAL": len(bs), "N_INTERVIEWS": b["k"], "N_TURNS": len(b["turns"]),
                "TURNS": render_turns(b["turns"])}
        write_text(os.path.join(d, f"prompt2_agent_A_{act_stem(a)}_batch1.txt"),
                   full_text(act_parts(st, act_stem(a), fixed, fill(post, vals))))
        rtail = "\n\n" + fill(PROMPTS["2b"], {"ACT": a["Act"], "KIND": kind(a),
                                              "ACT_LABELS": "[the labels the agent created in the act, with their "
                                                            "definitions and codes, go here]"})
        write_text(os.path.join(d, f"prompt2b_agent_A_{act_stem(a)}_review.txt"),
                   full_text(act_parts(st, act_stem(a), fixed, rtail)))
        tails = [est_tokens(render_turns(x["turns"])) for x in bs]
        nx = len(CORPUS["all_turns"][act_stem(a)]) - len(CORPUS["turns"][act_stem(a)])
        print(f"Prompt 2, {a['Act']}: {len(bs)} batches" + (f" ({nx} excluded turn(s) left out)" if nx else "")
              + f"; fixed part about {est_tokens(fixed):,} tokens (cached), turns about {min(tails):,} to "
              f"{max(tails):,} tokens per batch, plus the memory and the labels; then the review (Prompt 2b).")
    p3pre, p3post = split_at_memory(PROMPTS[3])
    write_text(os.path.join(d, "prompt3_agent_A.txt"), fill(p3pre, {"GUIDE": g, "LETTER": "A"})
               + "[the agent's whole memory goes here]" + fill(p3post, {"GUIDE": g, "LETTER": "A"}))
    write_text(os.path.join(d, "prompt4_agent_A.txt"), fill(PROMPTS[4], {
        "GUIDE": g, "LETTER": "A", "MEMO1": "[the agent's memo part 1 goes here]",
        "FIGURES": "[the agreement figures from round 4 go here]", "P4SAMPLE": "[the sample of turns from round 4]"}))
    print(f"The filled prompts are in {d}: read the prompt1, prompt2 and prompt2b files before running round 1.")


def test_models(agents=None):
    """A few small calls per model, with the same settings as the real calls: a JSON reply; the cached layout of
    round 2, three times with a memory that grows, to see each call read the earlier memory from the cache; and the
    folder tool of the memo rounds. Under a dollar in all; these calls are not counted in the run's budget."""
    global GPT_CHUNK_TOKENS
    if busy():
        print("Agents from the last run are still finishing their calls; wait a minute and run this again.")
        return
    STOP_WHY[0] = "stopped by hand"
    STOP.clear()                               # a Stop pressed during an earlier round must not stop these tests
    print(VERSION)
    g = load_guide()
    folder = ensure(rpath("test_folder"))
    with open(os.path.join(folder, "note.txt"), "w", encoding="utf-8") as f:
        f.write("The word to find is: lighthouse\n")
    for L in (agents or LETTERS):
        tools = folder_tools(folder)
        name = f"Agent {L} ({MODELS[L]['model']})"
        try:
            prov = provider(L)
            res = prov.json_call({"fixed": 'This is a test of the connection. Reply in this JSON format: {"memory": '
                                           '["M1 (general): connected"]}', "memory": [], "tail": ""}, SCHEMA1, False)
            obj, _ = parse_json(res["text"])
            print(f"{name}: JSON reply {'OK' if obj and res['ok'] else 'PROBLEM: ' + short(res['text'], 80)}; "
                  f"model version {res.get('model')}")
            # the cache as the memory grows, as in round 2: three calls, each with one more block of memory than the
            # last; the second and the third should read more from the cache each time
            text = (g + "\n") * (1 + 40000 // max(len(g), 1))
            fixed = ("This is a test of the cache. Here is a long text to cache:\n" + text[:16000]
                     + '\nReply in this JSON format: {"memory": ["M9 (general): cached"]}')
            blocks = ["Your memory so far:\n" + text[16000:24000], "\n" + text[24000:32000], "\n" + text[32000:40000]]
            keep, GPT_CHUNK_TOKENS = GPT_CHUNK_TOKENS, 1000    # so that each block of the test is a chunk of its own
            try:
                rs = [prov.json_call({"fixed": fixed, "memory": blocks[:k], "tail": "\nReply in the JSON format given "
                                      "above."}, SCHEMA1, True) for k in (1, 2, 3)]
            finally:
                GPT_CHUNK_TOKENS = keep
            us = [sum_usage(r["usages"]) for r in rs]
            grows = 0 < us[1]["cache_read"] < us[2]["cache_read"]
            if not all(r["ok"] for r in rs):
                verdict = "PROBLEM (a reply stopped early)"
            elif grows:
                verdict = "OK"
            elif MODELS[L]["provider"] == "google":
                verdict = "no read yet, which is fine: Gemini caches by itself, as it sees fit"
            else:
                verdict = "PROBLEM: the memory is not being read from the cache"
            print(f"{name}: cache as the memory grows: {verdict}. Read from the cache: "
                  + ", then ".join(f"{u['cache_read']:,} of {u['input_total']:,}" for u in us) + " tokens")
            r3 = prov.tool_call({"fixed": "This is a test of your folder tools. List the files in your folder, read "
                                          "note.txt, and reply with the one word it asks you to find, and nothing "
                                          "else.", "memory": [], "tail": ""}, tools)
            print(f"{name}: folder tool {'OK' if 'lighthouse' in r3['text'].lower() else 'PROBLEM'}: "
                  f"{r3.get('tool_steps', 0)} tool step(s); replied {short(r3['text'], 60)!r}")
            c = sum(cost_of(MODELS[L]["model"], u) for r in [res] + rs + [r3] for u in r["usages"])
            print(f"{name}: these tests cost about ${c:.3f}")
        except Exception as e:
            print(f"{name}: PROBLEM: {short(e, 400)}")


def status():
    """Where each agent is: rounds, acts and batches done, memory, labels, cost, flags, and any stop."""
    load_corpus()
    for L in LETTERS:
        st = load_state(L)
        nb = sum(len(v) for v in st["done"].values())
        d = rpath(f"agent_{L}", "reading")
        n_read = len([f for f in os.listdir(d) if f.endswith(".json") and f != "batch_jobs.json"]) if os.path.isdir(d) else 0
        in_batch = len(reading_pending(L))
        print(f"Agent {L} ({MODELS[L]['model']}): round 1 {'done' if st['round1'] else 'not done'} ({n_read} act(s) "
              f"read" + (f", {in_batch} still in a batch" if in_batch else "") + "); "
              f"{len(st['acts_complete'])} act(s) coded and reviewed, {nb} batches; {len(st['memory'])} notes, the "
              f"memory carrying about {est_tokens(memory_text(st)):,} tokens; "
              f"{len(st['labels'])} unique labels; memo part 1 {'done' if st['memo1'] else 'not done'}, parts 2 and 3 "
              f"{'done' if st['memo23'] else 'not done'}; {st['calls']} calls, ${st['cost']:.2f}; {st['flags']} flags"
              + (f"; STOPPED {st['stopped']}" if st.get("stopped") else ""))
    left = busy()
    print(f"Total cost of run {RUN!r}: ${sum(load_state(L)['cost'] for L in LETTERS):.2f} of the "
          f"${BUDGET_USD:,.0f} budget." + (f" Still running: {', '.join(left)}." if left else ""))
    if len(select_acts()) > 1 and PARTS > 1:
        show_parts()


def _project(L, st, stems, calls, n_done, fixed_est):
    """The full run of round 2 for one agent, act by act in nomination-year order, from the pilot: the cost of each
    turn's output and of the input that is never cached (the turns, the act's notes and labels, the instructions) as
    in the pilot; each call's fixed part, memory (which grows by the pilot's memo and kept notes per act) and labels
    list (which grows by the pilot's new labels per act, an upper bound: later acts should add fewer) read from the
    cache as often as the pilot's calls found them there (Gemini's cache is implicit, so it can miss), the rest at
    the full input price; the cache written again for each act; a review per act at the pilot's cost; GPT's higher
    prices past its long-input line; and, where a model's prices change on a set date, the cost at the later prices
    too. Returns the figures, or None."""
    model, p = MODELS[L]["model"], price(MODELS[L]["model"])
    coding = [c for c in calls if c.get("kind") == "round2"]
    reviews = [c for c in calls if c.get("kind") == "review"]
    if not coding or not n_done:
        return None
    ratios = sorted(c["usage"]["input_total"] / c["tokens_counted_before"] for c in coding
                    if c.get("tokens_counted_before") and c["usage"].get("input_total"))
    ratio = ratios[len(ratios) // 2] if ratios else 1.0
    reviewed = [s for s in stems if s in st["act_memos"]]
    mem_act = sum(est_tokens(act_block(st, s)) for s in reviewed) / max(len(reviewed), 1) if reviewed else 0.0
    lab_act = (sum(est_tokens(label_line(x)) for x in st["labels"].values() if x.get("first_act") in stems)
               / max(len(stems), 1))
    k_of = {s: i for i, s in enumerate(stems)}
    pre = [min(c["usage"].get("input_total", 0), (fixed_est + (mem_act + lab_act) * k_of.get(c.get("act"), 0)) * ratio)
           for c in coding]                    # the cached part of each pilot call, in model tokens
    firsts, seen = set(), set()
    for i, c in enumerate(coding):
        if c.get("act") not in seen:
            seen.add(c.get("act"))
            firsts.add(i)
    later = [i for i in range(len(coding)) if i not in firsts] or list(range(len(coding)))
    den = sum(pre[i] for i in later)
    hit = min(1.0, sum(min(coding[i]["usage"].get("cache_read", 0), pre[i]) for i in later) / den) if den else 1.0
    var_in = sum(max(c["usage"].get("input_total", 0) - q, 0) for c, q in zip(coding, pre))
    out_tok = sum(c["usage"]["output"] for c in coding)
    unc_call = var_in / len(coding)
    n_batches = sum(len(make_batches(CORPUS["turns"][s])) for s in stems)
    per_batch = len(coding) / max(n_batches, 1)
    rev_cost = sum(c["cost"] for c in reviews) / max(len(stems), 1) if reviews else 0.0
    secs = sum(c.get("seconds", 0) for c in coding + reviews) / n_done
    acts = load_corpus()["acts"]
    n_turns = sum(len(CORPUS["turns"][act_stem(a)]) for a in acts)

    def run(pp, rev):
        p_in, p_read = pp["input"], pp.get("cache_read", pp["input"])
        p_write = pp.get("cache_write_1h", pp.get("cache_write", 0.0))
        var = (out_tok * pp["output"] + var_in * p_in) / 1e6 / n_done
        total, cum, biggest = 0.0, [], 0.0
        for k, a in enumerate(acts):
            s = act_stem(a)
            t_k = len(CORPUS["turns"][s])
            calls_k = len(make_batches(CORPUS["turns"][s])) * per_batch
            prefix = (fixed_est + (mem_act + lab_act) * k) * ratio          # model tokens, cached as in the pilot
            biggest = max(biggest, prefix + unc_call)
            xi = xo = 1.0
            if "long_input" in pp and prefix + unc_call > pp["long_input"]:
                xi, xo = pp["long_input_x"], pp["long_output_x"]
            c_k = (t_k * var * (xi + xo) / 2 + calls_k * prefix * (hit * p_read + (1 - hit) * p_in) * xi / 1e6 + rev
                   + (mem_act + lab_act * (k + 1)) * ratio * p_write * xi / 1e6)
            total += c_k
            cum.append((total, sum(len(CORPUS["turns"][act_stem(x)]) for x in acts[:k + 1])))
        return total, cum, biggest
    total, cum, biggest = run(p, rev_cost)
    out = dict(cost=total, hours=secs * n_turns / 3600, ratio=ratio, biggest=biggest, cum=cum, n_turns=n_turns,
               mem_end=mem_act * len(acts) * ratio, lab_end=lab_act * len(acts) * ratio, mem_act=mem_act,
               lab_act=lab_act, review=rev_cost, hit=hit)
    ch = price_changes(model)
    if ch:
        pl = price(model, ch[-1])
        out.update(cost_late=run(pl, rev_cost * pl["output"] / max(p["output"], 1e-9))[0], late_date=ch[-1])
    return out


def new_label_rates(st, recs, stems, parts=4):
    """New labels per 100 turns in each of `parts` equal parts of the records, in the order coded, counting the
    labels as coded (a label from an act outside stems, under its old or new name, is not new)."""
    seen = set()
    for k, v in st["labels"].items():
        if v.get("first_act") not in stems:
            seen.add(k)
            if v.get("renamed_from"):
                seen.add(label_key(v["renamed_from"]))
    n, out = len(recs), []
    cuts = [round(n * i / parts) for i in range(parts + 1)]
    for i in range(parts):
        new = 0
        for r in recs[cuts[i]:cuts[i + 1]]:
            for c in r["codes"]:
                k = label_key(c["label"])
                if k not in seen:
                    seen.add(k)
                    new += 1
        out.append(100 * new / max(cuts[i + 1] - cuts[i], 1))
    return out


def pilot_report(acts=None):
    """The checks for the pilot: unique labels (published guidance puts a whole study at about 50 to 300), codes per
    turn (one is typical), where the labels' words come from (word for word, built on the members' words, or the
    agent's own, worked out by the script), labels
    of one word, renamed in the review, definitions, evidence length, unsure and looking back, data problems, notes
    written and kept, the memo on each act, the reading, re-asks, and a projection of the full run (cost against the
    budget, hours, and the size of the last call against each model's context window)."""
    load_corpus(), load_prompts()
    g = load_guide()
    acts = select_acts(acts)
    stems = [act_stem(a) for a in acts]
    fixed = est_tokens(fill(split_at_memory(PROMPTS[2])[0], {"GUIDE": g, "LETTER": "A"}))
    lines = [f"# Step 3 pilot report ({RUN}, {now()})", "", "Acts: " + ", ".join(a["Act"] for a in acts), ""]
    allsets, projected, reading_costs = {}, {}, {}
    for L in LETTERS:
        st = load_state(L)
        recs = [r for s in stems for r in act_records(L, s)]
        codes = [c for r in recs for c in r["codes"]]
        if not recs:
            lines.append(f"## Agent {L}: no codes yet\n")
            continue
        cnt = collections.Counter(label_key(c["label"]) for c in codes)
        allsets[L] = set(cnt)
        per = collections.Counter(min(len(r["codes"]), 4) for r in recs)
        src = source_counts(st, recs)
        once = sum(1 for v in cnt.values() if v == 1)
        lw = [len(c["label"].split()) for c in codes]
        one_word = sum(1 for k in cnt if len(k.split()) == 1)
        ew = [len(str(c.get("evidence") or "").split()) for c in codes]
        dw = [len(str(c.get("description") or "").split()) for c in codes]
        labs = [v for v in st["labels"].values() if v.get("first_act") in stems]
        n_def = sum(1 for v in labs if v.get("definition"))
        renamed = sum(len((read_json(review_path(L, s)) or {}).get("renames") or []) for s in stems)
        growth = new_label_rates(st, [r for s in stems for r in act_records(L, s, renamed=False)], stems)
        flags = read_jsonl(rpath(f"agent_{L}", "flags.jsonl"))
        n_data = sum(1 for f in flags if f.get("kind") == "data" and f.get("round") == 2)
        notes_w = [m for s in stems for m in act_entries(st, s)]
        kept = sum(len((st["act_memos"].get(s) or {}).get("keep") or []) for s in stems)
        memo_pts = [p for s in stems for p in (st["act_memos"].get(s) or {}).get("memo") or []]
        read_acts = [s for s in stems if os.path.exists(reading_path(L, s))]
        rd = [read_json(reading_path(L, s)) for s in read_acts]
        n_notes = sum(len(x["notes"]) for x in rd)
        n_dp = sum(len(x.get("data_problems") or []) for x in rd)
        n_nm = sum(1 for x in rd for d in x.get("data_problems") or [] if d.get("not_members_words"))
        rcalls = [c for c in read_jsonl(rpath(f"agent_{L}", "calls", "round1.jsonl")) if c.get("act") in stems]
        rcost = sum(c.get("cost", 0) or 0 for c in rcalls)
        if read_acts and len(read_acts) == len(stems):
            reading_costs[L] = rcost
        calls = [c for s in stems for c in read_jsonl(rpath(f"agent_{L}", "calls", f"{s}.jsonl"))]
        c2 = sum(c["cost"] for c in calls)
        again = [c for c in calls if (c.get("attempt") or 1) > 1]
        lines += [f"## Agent {L} ({MODELS[L]['model']})", "",
                  f"1. Unique labels: {len(cnt)} (published guidance puts a whole study at about 50 to 300, and never "
                  f"in the thousands); {once} used only once; {one_word} of one word; {renamed} renamed in the review. "
                  f"New labels per 100 turns, in four equal parts of the turns as coded: "
                  f"{', '.join(f'{x:.0f}' for x in growth)} (they should fall as the list grows).",
                  f"2. Turns: {len(recs)}; with no code {per[0]}, one {per[1]}, two {per[2]}, three {per[3]}, four or "
                  f"more {per[4]}; {len(codes) / max(len(recs) - per[0], 1):.2f} codes per coded turn.",
                  f"3. Codes: {len(codes)}; each label used {len(codes) / max(len(cnt), 1):.1f} times on average; label "
                  f"length {sum(lw) / max(len(lw), 1):.1f} words on average, {max(lw or [0])} at most.",
                  f"4. Where the labels' words come from, code by code (worked out by the script): {source_text(src)}. "
                  f"Word for word means the label's words are, as the member said them, in the turn or in the turn "
                  f"where the label was first used; built on their words, that at least half of its main words are "
                  f"theirs, as in \"never-planned look\" for \"we never planned it\". The agents' sheets count only "
                  f"word for word as the members' words (\"data\").",
                  f"5. Definitions: {n_def} of the {len(labs)} labels these acts created; descriptions "
                  f"{sum(dw) / max(sum(1 for x in dw if x), 1):.0f} words long on average; evidence "
                  f"{sum(ew) / max(len(ew), 1):.0f} words on average.",
                  f"6. Unsure: {sum(1 for c in codes if c.get('unsure'))} codes "
                  f"({sum(1 for c in codes if c.get('unsure')) / max(len(codes), 1):.0%}); turns that look back: "
                  f"{sum(1 for r in recs if r.get('looks_back'))} of {len(recs)}; data problems reported in coding: "
                  f"{n_data}.",
                  f"7. Notes: {len(notes_w)} written on these acts ({sum(1 for m in notes_w if m.get('round') == 'reading')}"
                  f" from the reading), {kept} kept in the review; memos on the acts: {len(memo_pts)} points, "
                  f"{sum(len(p.split()) for p in memo_pts) / max(len(memo_pts), 1):.0f} words each; the memory now "
                  f"carries about {est_tokens(memory_text(st)):,} tokens.",
                  f"8. Reading (round 1): {len(read_acts)} of {len(stems)} act(s) read, {n_notes} notes "
                  f"({n_notes / max(len(read_acts), 1):.1f} per act), {n_dp} data problems ({n_nm} turns not the "
                  f"member's words); cost ${rcost:.2f}"
                  + (", as a batch at half price." if any(c.get("batch_api") for c in rcalls) else "."),
                  f"9. Round 2 so far: {len(calls)} calls, ${c2:.2f}; asked again {len(again)} times "
                  f"(${sum(c['cost'] for c in again):.2f}, {sum(c['cost'] for c in again) / max(c2, 1e-9):.0%}); "
                  f"{len(flags)} flags; cost so far ${st['cost']:.2f}.", ""]
        pj = _project(L, st, stems, calls, len(recs), fixed)
        if pj:
            projected[L] = pj
            model = MODELS[L]["model"]
            room = CONTEXT.get(model, 0) - MODELS[L]["max_output"]
            warn = (f" By the last act, a call would carry about {pj['biggest']:,.0f} of {model}'s own tokens"
                    + (f" (it counts {pj['ratio']:.2f} for each token estimated here)" if abs(pj['ratio'] - 1) > 0.05
                       else ""))
            if 0 < room < pj["biggest"]:
                warn += f": more than the {room:,} its context window leaves for the prompt, so the run could not finish."
            elif room > 0:
                warn += f", {'close to' if pj['biggest'] > 0.8 * room else 'within'} the {room:,} its window leaves."
            p = price(model)
            if "long_input" in p and pj["biggest"] > p["long_input"]:
                warn += (f" Past {p['long_input']:,} tokens, {model}'s prices go up (input x{p['long_input_x']}, "
                         f"output x{p['long_output_x']}); the cost here allows for that.")
            lines += [f"Projection for round 2 on all {pj['n_turns']:,} turns: about ${pj['cost']:,.0f}, about "
                      f"{pj['hours']:,.0f} hours of calls. By the end, the memory carries about {pj['mem_end']:,.0f} "
                      f"tokens (a memo and kept notes of about {pj['mem_act']:,.0f} per act) and the labels list "
                      f"about {pj['lab_end']:,.0f} (as if every act added as many new labels as the pilot did: an "
                      f"upper bound), both read from the cache: in the pilot, its calls found about {pj['hit']:.0%} "
                      f"of that part in the cache, and the rest is priced in full. A review costs about "
                      f"${pj['review']:.2f} an act."
                      + (f" At {model}'s prices from {nice_date(pj['late_date'])}, the projection is about "
                         f"${pj['cost_late']:,.0f}." if pj.get("cost_late") is not None else "")
                      + warn, ""]
        lines += ["Most used labels: " + "; ".join(f"{st['labels'].get(k, {}).get('label', k)} ({v})"
                                                   for k, v in cnt.most_common(20)), ""]
    read_total = 0.0
    if reading_costs:                          # the reading of all the acts, scaled from the pilot's by their words
        w0 = len(fill(PROMPTS[1], {"GUIDE": g, "LETTER": "A", "ACT": "", "KIND": "", "N_INTERVIEWS": 0, "N_TURNS": 0,
                                   "TURNS": ""}).split())
        size = {act_stem(a): w0 + len(render_turns_reading(CORPUS["all_turns"][act_stem(a)]).split())
                for a in CORPUS["acts"]}
        read_total = sum(reading_costs.values()) * sum(size.values()) / max(sum(size[s] for s in stems), 1)
        who_r = "all three agents" if len(reading_costs) == len(LETTERS) else "agent " + " and ".join(sorted(reading_costs))
        lines += [f"Projected cost of round 1 (the reading) in the full run, {who_r}: about ${read_total:,.0f}"
                  + (" (as a batch, at half price)" if READ_BATCH else "") + f", from the pilot's reading, scaled by "
                  f"the words of all {len(CORPUS['acts'])} acts.", ""]
    avail = BUDGET_USD - read_total            # what the budget leaves for round 2 after the reading
    if projected:
        total = sum(x["cost"] for x in projected.values())
        if total > avail:
            n_all = max(x["n_turns"] for x in projected.values())
            steps = [sum(x["cum"][k][0] for x in projected.values()) for k in range(len(load_corpus()["acts"]))]
            k_stop = next((k for k, v in enumerate(steps) if v > avail), len(steps) - 1)
            done = next(iter(projected.values()))["cum"][k_stop][1]
            note = f": the budget would stop round 2 at about {done / max(n_all, 1):.0%} of the turns"
        else:
            note = ", within the budget"
        who = "all three agents" if len(projected) == len(LETTERS) else "agent " + " and ".join(sorted(projected))
        late = {L: x for L, x in projected.items() if x.get("cost_late") is not None}
        total_late = sum(x.get("cost_late", x["cost"]) for x in projected.values())
        late_note = ""
        if late:
            d_late = max(x["late_date"] for x in late.values())
            late_note = (f" If the run goes on past {nice_date(d_late)}, when "
                         f"{' and '.join(MODELS[L]['model'] for L in late)}'s prices change, the calls after that "
                         f"date cost more: all of round 2 at those prices would be about ${total_late:,.0f}.")
        lines += [f"Projected cost of round 2 (the coding and the reviews) in the full run, {who}: about "
                  f"${total:,.0f}, against BUDGET_USD = ${BUDGET_USD:,.0f}"
                  + (f" (about ${max(avail, 0):,.0f} after the reading)" if read_total else "") + f"{note}. The memo rounds "
                  f"come on top. The agents work side by side, so round 2 takes about as long as the slowest: about "
                  f"{max(x['hours'] for x in projected.values()):,.0f} hours. A rough guide: it assumes the later "
                  f"acts are coded much as the pilot was." + late_note, ""]
    if len(allsets) == len(LETTERS):
        inter = set.intersection(*allsets.values())
        lines += ["## Across agents", "", f"Labels all three used: {len(inter)}: " + "; ".join(sorted(inter)[:40]), ""]
    data_problems_report()
    rep = "\n".join(lines)
    write_text(rpath("reports", "pilot_report.md"), rep)
    print(rep)


if __name__ == "__main__":
    print(VERSION, "is loaded. Start with check_inputs(), then dry_run(); see the notes at the top of the script.")
