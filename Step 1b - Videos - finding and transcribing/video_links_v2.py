"""
===============================================================================
video_links_v2.py: FIND YOUTUBE VIDEOS IN WHICH THE ACTS SPEAK
Step 1 (videos). Version 2.1, 15 Oct 2026.
===============================================================================
Finds YouTube videos in which the 52 acts speak (interviews, Q&A, panel
discussions, workshops and documentaries) and writes
    harrison_v7_VIDEO_LINKS.xlsx   to read and check
    harrison_v7_VIDEO_LINKS.json   for the transcription step
Nothing is downloaded. The script collects links and details only.

HOW IT SEARCHES
Harrison et al. (2023, p. 104) "searched newspaper, magazine, and video
databases with the name of each band and band member". This script does the
same on YouTube:
  1. Serper video search (Google Videos), the same key as the text corpus:
       act      "{act}" interview, plus the same with the year before the
                ceremony and the ceremony year
                "{act}" documentary
                "{act}" Q&A OR panel OR workshop OR masterclass
       members  "{member}" "{act}" interview, for every nomination-era member
                "{member}" interview, and "{member}" Q&A OR panel OR workshop
                OR masterclass OR documentary, for members whose names are
                distinctive (a common name returns namesakes)
  2. YouTube's own search, once per act: "{act}" interview (50 results).
  3. YouTube details for every video found: the exact upload date, the full
     description, the duration, captions, views, and whether the video is
     still online.
Only YouTube videos are kept. Every raw response is cached in Drive, so a
rerun costs nothing and the search can be audited later.

WHAT GOES WHERE
  Links     Interview, Q&A, Panel discussion, Workshop or masterclass,
            Documentary, Award or speech, and Unclear (check)
  Left out  performances and music videos; covers, tributes, lessons and
            reactions; videos not about the act; videos naming only an
            other-era member; links that are not YouTube videos

COST
  All 52 acts: about 1,000 Serper credits, and about 5,300 of the 10,000 free
  YouTube units a day. A rerun answers from the cache and costs nothing.

COLAB (one cell; keys in Colab secrets named SERPER_API_KEY and YT_API_KEY,
with notebook access switched on; if a key is missing the script asks for it)
  from google.colab import drive
  drive.mount('/content/drive')
  %run /content/drive/MyDrive/harrison_v7/video_links_v2.py
===============================================================================
"""

import html
import json
import os
import random
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta

import pandas as pd
import requests

VERSION = "video_links_v2.1 (15 Oct 2026; simple video IDs)"

# ═════════════════════════════════════════════════════════════════════════════
# SETTINGS
# ═════════════════════════════════════════════════════════════════════════════
DIR = os.environ.get("HARRISON_DIR", "/content/drive/MyDrive/harrison_v7")
PILOT_ACTS = []              # e.g. ["Culture Club", "Amy Holland", "Asia"] to try a few acts first; [] = all 52
SERPER_PAGES = 2             # Serper pages per query (10 results a page)
YOUTUBE_SEARCH = True        # YouTube's own search, once per act (100 quota units each)
YT_RESULTS = 50              # results per YouTube search (one page holds at most 50)
DELAY = (0.4, 0.9)           # seconds between calls
SAVE_EVERY = 25              # save the cache every 25 queries
DESC_CHARS = 1500            # characters of each description shown in the workbook

OUT_XLSX = os.path.join(DIR, "harrison_v7_VIDEO_LINKS.xlsx")
OUT_JSON = os.path.join(DIR, "harrison_v7_VIDEO_LINKS.json")
CACHE = os.path.join(DIR, "video_search_cache.json")      # shared with video_links_v1

# ═════════════════════════════════════════════════════════════════════════════
# THE 52 ACTS
# From the final name lists (harrison_v7_NAME_LISTS.json, 5 Oct) and grammy_nominees.xlsx.
#   search         how the act's name is searched (ordinary-word names get a qualifier)
#   member_search  the act's name inside member queries
#   members        nomination-era members searched by name (bands only)
#   distinctive    members whose name alone is safe to search and to recognise
#   names          [name or alias, member, era, distinctive] used to recognise people in titles
#                  and descriptions (one-word nicknames and common aliases are not trusted alone)
# ═════════════════════════════════════════════════════════════════════════════
ACTS = {
    "Dire Straits": {
        "year": 1980, "type": "Band", "winner": False,
        "search": "\"Dire Straits\"", "member_search": "Dire Straits",
        "members": ["David Knopfler", "John Illsley", "Mark Knopfler", "Pick Withers"],
        "distinctive": ["David Knopfler", "John Illsley", "Mark Knopfler", "Pick Withers"],
        "other_era": ["Alan Clark", "Guy Fletcher", "Hal Lindes", "Jack Sonni", "Terry Williams"],
        "extra_queries": [],
        "names": [
            ["David Knopfler", "David Knopfler", "nomination", True],
            ["John Illsley", "John Illsley", "nomination", True],
            ["Mark Knopfler", "Mark Knopfler", "nomination", True],
            ["David 'Pick' Withers", "Pick Withers", "nomination", True],
            ["David Withers", "Pick Withers", "nomination", False],
            ["Pick Withers", "Pick Withers", "nomination", True],
            ["Alan Clark", "Alan Clark", "other", False],
            ["Guy Fletcher", "Guy Fletcher", "other", False],
            ["Hal Lindes", "Hal Lindes", "other", False],
            ["Jack Sonni", "Jack Sonni", "other", False],
            ["Terry Williams", "Terry Williams", "other", False],
        ],
    },
    "Rickie Lee Jones": {
        "year": 1980, "type": "Solo", "winner": True,
        "search": "\"Rickie Lee Jones\"", "member_search": "Rickie Lee Jones",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Rickie Lee Jones", "Rickie Lee Jones", "nomination", True],
        ],
    },
    "The Knack": {
        "year": 1980, "type": "Band", "winner": False,
        "search": "\"The Knack\" band", "member_search": "The Knack",
        "members": ["Berton Averre", "Bruce Gary", "Doug Fieger", "Prescott Niles"],
        "distinctive": ["Berton Averre", "Doug Fieger", "Prescott Niles"],
        "other_era": ["Billy Ward", "Pat Torpey", "Phil Jost", "Stephen 'Mac' McNally"],
        "extra_queries": [],
        "names": [
            ["Berton Averre", "Berton Averre", "nomination", True],
            ["Burton Averre", "Berton Averre", "nomination", True],
            ["Bruce Gary", "Bruce Gary", "nomination", False],
            ["Doug Fieger", "Doug Fieger", "nomination", True],
            ["Prescott Niles", "Prescott Niles", "nomination", True],
            ["Billy Ward", "Billy Ward", "other", False],
            ["Pat Torpey", "Pat Torpey", "other", False],
            ["Phil Jost", "Phil Jost", "other", False],
            ["Stephen 'Mac' McNally", "Stephen 'Mac' McNally", "other", False],
            ["Stephen McNally", "Stephen 'Mac' McNally", "other", False],
        ],
    },
    "Amy Holland": {
        "year": 1981, "type": "Solo", "winner": False,
        "search": "\"Amy Holland\"", "member_search": "Amy Holland",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Amy Celeste Boersma", "Amy Holland", "nomination", True],
            ["Amy Holland", "Amy Holland", "nomination", True],
        ],
    },
    "Christopher Cross": {
        "year": 1981, "type": "Solo", "winner": True,
        "search": "\"Christopher Cross\"", "member_search": "Christopher Cross",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Chris Cross", "Christopher Cross", "nomination", False],
            ["Christopher 'Sailing' Cross", "Christopher Cross", "nomination", True],
            ["Christopher Charles Geppert", "Christopher Cross", "nomination", True],
            ["Christopher Cross", "Christopher Cross", "nomination", True],
            ["Christopher Geppert", "Christopher Cross", "nomination", True],
        ],
    },
    "Irene Cara": {
        "year": 1981, "type": "Solo", "winner": False,
        "search": "\"Irene Cara\"", "member_search": "Irene Cara",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Irene Cara", "Irene Cara", "nomination", True],
            ["Irene Cara Escalera", "Irene Cara", "nomination", True],
        ],
    },
    "Robbie Dupree": {
        "year": 1981, "type": "Solo", "winner": False,
        "search": "\"Robbie Dupree\"", "member_search": "Robbie Dupree",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Robbie Dupree", "Robbie Dupree", "nomination", True],
            ["Robert Dupuis", "Robbie Dupree", "nomination", False],
        ],
    },
    "The Pretenders": {
        "year": 1981, "type": "Band", "winner": False,
        "search": "\"Pretenders\"", "member_search": "Pretenders",
        "members": ["Chrissie Hynde", "James Honeyman-Scott", "Martin Chambers", "Pete Farndon"],
        "distinctive": ["Chrissie Hynde", "James Honeyman-Scott", "Pete Farndon"],
        "other_era": ["Adam Seymour", "Andy Hobson", "Blair Cunningham", "Gerry Mcilduff", "Johnny Marr", "Malcolm Foster", "Robbie McIntosh", "T. M. Stevens"],
        "extra_queries": [],
        "names": [
            ["Chrissie Hynde", "Chrissie Hynde", "nomination", True],
            ["Chrissy Hynde", "Chrissie Hynde", "nomination", True],
            ["Christine Ellen Hynde", "Chrissie Hynde", "nomination", True],
            ["James Honeyman-Scott", "James Honeyman-Scott", "nomination", True],
            ["Jim Honeyman-Scott", "James Honeyman-Scott", "nomination", True],
            ["Jimmy Scott", "James Honeyman-Scott", "nomination", False],
            ["Martin Chambers", "Martin Chambers", "nomination", False],
            ["Martin Dale Chambers", "Martin Chambers", "nomination", False],
            ["Pete Farndon", "Pete Farndon", "nomination", True],
            ["Peter Farndon", "Pete Farndon", "nomination", True],
            ["Adam Seymour", "Adam Seymour", "other", False],
            ["Andy Hobson", "Andy Hobson", "other", False],
            ["Blair Cunningham", "Blair Cunningham", "other", False],
            ["Gerry Mcilduff", "Gerry Mcilduff", "other", False],
            ["Johnny Marr", "Johnny Marr", "other", False],
            ["Malcolm Foster", "Malcolm Foster", "other", False],
            ["Robbie McIntosh", "Robbie McIntosh", "other", False],
            ["T. M. Stevens", "T. M. Stevens", "other", False],
        ],
    },
    "Adam and the Ants": {
        "year": 1982, "type": "Band", "winner": False,
        "search": "\"Adam and the Ants\"", "member_search": "Adam and the Ants",
        "members": ["Adam Ant", "Chris Hughes", "Gary Tibbs", "Marco Pirroni", "Terry Lee Miall"],
        "distinctive": ["Adam Ant", "Gary Tibbs", "Marco Pirroni", "Terry Lee Miall"],
        "other_era": ["Andy Warren", "Dave Barbarossa", "Johnny Bivouac", "Jordan", "Kevin Mooney", "Leigh Gorman", "Lester Square", "Mark Ryan", "Matthew Ashman", "Paul Flanagan"],
        "extra_queries": [],
        "names": [
            ["Adam Ant", "Adam Ant", "nomination", True],
            ["Stuart Goddard", "Adam Ant", "nomination", True],
            ["Stuart Leslie Goddard", "Adam Ant", "nomination", True],
            ["Chris 'Merrick' Hughes", "Chris Hughes", "nomination", False],
            ["Chris Hughes", "Chris Hughes", "nomination", False],
            ["Christopher Merrick Hughes", "Chris Hughes", "nomination", False],
            ["Gary Brian Tibbs", "Gary Tibbs", "nomination", True],
            ["Gary Tibbs", "Gary Tibbs", "nomination", True],
            ["Marco Francesco Andrea Pirroni", "Marco Pirroni", "nomination", True],
            ["Marco Pirroni", "Marco Pirroni", "nomination", True],
            ["Terry Lee Miall", "Terry Lee Miall", "nomination", True],
            ["Andy Warren", "Andy Warren", "other", False],
            ["Dave 'Barbe' Barbarossa", "Dave Barbarossa", "other", False],
            ["Dave Barbarossa", "Dave Barbarossa", "other", False],
            ["Dave Barbe", "Dave Barbarossa", "other", False],
            ["David Barbarossa", "Dave Barbarossa", "other", False],
            ["Johnny Bivouac", "Johnny Bivouac", "other", False],
            ["Pamela Rooke", "Jordan", "other", False],
            ["Kevin Mooney", "Kevin Mooney", "other", False],
            ["Kevin Paul Mooney", "Kevin Mooney", "other", False],
            ["Leigh Gorman", "Leigh Gorman", "other", False],
            ["Lester Square", "Lester Square", "other", False],
            ["Thom Hardy", "Lester Square", "other", False],
            ["Thomas Hardy", "Lester Square", "other", False],
            ["Mark Ryan", "Mark Ryan", "other", False],
            ["Matthew Ashman", "Matthew Ashman", "other", False],
            ["Paul Flanagan", "Paul Flanagan", "other", False],
        ],
    },
    "James Ingram": {
        "year": 1982, "type": "Solo", "winner": False,
        "search": "\"James Ingram\"", "member_search": "James Ingram",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["James Ingram", "James Ingram", "nomination", True],
        ],
    },
    "Luther Vandross": {
        "year": 1982, "type": "Solo", "winner": False,
        "search": "\"Luther Vandross\"", "member_search": "Luther Vandross",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Luther Ronzoni Vandross", "Luther Vandross", "nomination", True],
            ["Luther Vandross", "Luther Vandross", "nomination", True],
        ],
    },
    "Sheena Easton": {
        "year": 1982, "type": "Solo", "winner": True,
        "search": "\"Sheena Easton\"", "member_search": "Sheena Easton",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Sheena Easton", "Sheena Easton", "nomination", True],
            ["Sheena Shirley Orr", "Sheena Easton", "nomination", True],
        ],
    },
    "The Go-Go's": {
        "year": 1982, "type": "Band", "winner": False,
        "search": "\"Go-Go's\"", "member_search": "Go-Go's",
        "members": ["Belinda Carlisle", "Charlotte Caffey", "Gina Schock", "Jane Wiedlin", "Kathy Valentine"],
        "distinctive": ["Belinda Carlisle", "Charlotte Caffey", "Gina Schock", "Jane Wiedlin", "Kathy Valentine"],
        "other_era": ["Elissa Bello", "Margot Olavarria", "Paula Jean Brown"],
        "extra_queries": [],
        "names": [
            ["Belinda Carlisle", "Belinda Carlisle", "nomination", True],
            ["Charlotte Caffey", "Charlotte Caffey", "nomination", True],
            ["Gina Schock", "Gina Schock", "nomination", True],
            ["Jane Wiedlin", "Jane Wiedlin", "nomination", True],
            ["Kathy Valentine", "Kathy Valentine", "nomination", True],
            ["Elissa Bello", "Elissa Bello", "other", False],
            ["Margot Olavarria", "Margot Olavarria", "other", False],
            ["Paula Brown", "Paula Jean Brown", "other", False],
            ["Paula Jean Brown", "Paula Jean Brown", "other", False],
        ],
    },
    "Asia": {
        "year": 1983, "type": "Band", "winner": False,
        "search": "\"Asia\" band", "member_search": "Asia",
        "members": ["Carl Palmer", "Geoff Downes", "John Wetton", "Steve Howe"],
        "distinctive": ["Carl Palmer", "Geoff Downes", "John Wetton", "Steve Howe"],
        "other_era": ["Al Pitrelli", "Greg Lake", "John Payne", "Mandy Meyer", "Pat Thrall"],
        "extra_queries": [],
        "names": [
            ["Carl Palmer", "Carl Palmer", "nomination", True],
            ["Geoff Downes", "Geoff Downes", "nomination", True],
            ["Geoffrey Downes", "Geoff Downes", "nomination", True],
            ["John 'Ken' Wetton", "John Wetton", "nomination", True],
            ["John Kenneth Wetton", "John Wetton", "nomination", True],
            ["John Wetton", "John Wetton", "nomination", True],
            ["Steve Howe", "Steve Howe", "nomination", True],
            ["Al Pitrelli", "Al Pitrelli", "other", False],
            ["Greg Lake", "Greg Lake", "other", False],
            ["John Payne", "John Payne", "other", False],
            ["Mandy Meyer", "Mandy Meyer", "other", False],
            ["Pat Thrall", "Pat Thrall", "other", False],
        ],
    },
    "Jennifer Holliday": {
        "year": 1983, "type": "Solo", "winner": False,
        "search": "\"Jennifer Holliday\"", "member_search": "Jennifer Holliday",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Jennifer Holliday", "Jennifer Holliday", "nomination", True],
        ],
    },
    "Men at Work": {
        "year": 1983, "type": "Band", "winner": True,
        "search": "\"Men at Work\" band", "member_search": "Men at Work",
        "members": ["Colin Hay", "Greg Ham", "Jerry Speiser", "John Rees", "Ron Strykert"],
        "distinctive": ["Colin Hay", "Jerry Speiser", "Ron Strykert"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Colin Hay", "Colin Hay", "nomination", True],
            ["Greg Ham", "Greg Ham", "nomination", False],
            ["Jerry Speiser", "Jerry Speiser", "nomination", True],
            ["John Rees", "John Rees", "nomination", False],
            ["Ron Strykert", "Ron Strykert", "nomination", True],
        ],
    },
    "Stray Cats": {
        "year": 1983, "type": "Band", "winner": False,
        "search": "\"Stray Cats\"", "member_search": "Stray Cats",
        "members": ["Brian Setzer", "Lee Rocker", "Slim Jim Phantom"],
        "distinctive": ["Brian Setzer", "Lee Rocker", "Slim Jim Phantom"],
        "other_era": ["Bob Beecher", "Gary Setzer", "Tommy Byrnes"],
        "extra_queries": [],
        "names": [
            ["Brian Setzer", "Brian Setzer", "nomination", True],
            ["Lee Drucker", "Lee Rocker", "nomination", False],
            ["Lee Rocker", "Lee Rocker", "nomination", True],
            ["Leon Drucker", "Lee Rocker", "nomination", False],
            ["James McDonnell", "Slim Jim Phantom", "nomination", False],
            ["Jim Phantom", "Slim Jim Phantom", "nomination", True],
            ["Slim Jim Phantom", "Slim Jim Phantom", "nomination", True],
            ["Bob Beecher", "Bob Beecher", "other", False],
            ["Gary Setzer", "Gary Setzer", "other", False],
            ["Tommy Byrnes", "Tommy Byrnes", "other", False],
        ],
    },
    "The Human League": {
        "year": 1983, "type": "Band", "winner": False,
        "search": "\"Human League\"", "member_search": "Human League",
        "members": ["Ian Burden", "Jo Callis", "Joanne Catherall", "Philip Adrian Wright", "Philip Oakey", "Susan Ann Sulley"],
        "distinctive": ["Jo Callis", "Joanne Catherall", "Philip Oakey", "Susan Ann Sulley"],
        "other_era": ["Ian Craig Marsh", "Jim Russell", "Martyn Ware", "Neil Sutton", "Russell Dennett"],
        "extra_queries": [],
        "names": [
            ["Ian Burden", "Ian Burden", "nomination", False],
            ["Jo Callis", "Jo Callis", "nomination", True],
            ["John William Callis", "Jo Callis", "nomination", True],
            ["Joanne Catherall", "Joanne Catherall", "nomination", True],
            ["Adrian Wright", "Philip Adrian Wright", "nomination", False],
            ["Philip Adrian Wright", "Philip Adrian Wright", "nomination", False],
            ["Philip Wright", "Philip Adrian Wright", "nomination", False],
            ["Phil Oakey", "Philip Oakey", "nomination", True],
            ["Philip Oakey", "Philip Oakey", "nomination", True],
            ["Phillip Oakey", "Philip Oakey", "nomination", True],
            ["Ann Sulley", "Susan Ann Sulley", "nomination", False],
            ["Susan Ann Sulley", "Susan Ann Sulley", "nomination", True],
            ["Susan Sulley", "Susan Ann Sulley", "nomination", False],
            ["Susanne Sulley", "Susan Ann Sulley", "nomination", True],
            ["Craig Marsh", "Ian Craig Marsh", "other", False],
            ["Ian Craig Marsh", "Ian Craig Marsh", "other", False],
            ["Ian Marsh", "Ian Craig Marsh", "other", False],
            ["Jim Russell", "Jim Russell", "other", False],
            ["Martin Ware", "Martyn Ware", "other", False],
            ["Martyn Ware", "Martyn Ware", "other", False],
            ["Neil Sutton", "Neil Sutton", "other", False],
            ["Russell Dennett", "Russell Dennett", "other", False],
        ],
    },
    "Big Country": {
        "year": 1984, "type": "Band", "winner": False,
        "search": "\"Big Country\" band", "member_search": "Big Country",
        "members": ["Bruce Watson", "Mark Brzezicki", "Stuart Adamson", "Tony Butler"],
        "distinctive": ["Mark Brzezicki", "Stuart Adamson"],
        "other_era": ["Alan Wishart", "Clive Parker", "Pat Ahern", "Pete Wishart", "Jamie Watson", "Mike Peters", "Derek Forbes", "Simon Hough", "Tommie Paxton", "Chris Stones", "Reece Dobbin"],
        "extra_queries": [],
        "names": [
            ["Bruce Watson", "Bruce Watson", "nomination", False],
            ["Mark Brzezicki", "Mark Brzezicki", "nomination", True],
            ["Stuart Adamson", "Stuart Adamson", "nomination", True],
            ["William Stuart Adamson", "Stuart Adamson", "nomination", True],
            ["Tony Butler", "Tony Butler", "nomination", False],
            ["Alan Wishart", "Alan Wishart", "other", False],
            ["Clive Parker", "Clive Parker", "other", False],
            ["Pat Ahern", "Pat Ahern", "other", False],
            ["Pete Wishart", "Pete Wishart", "other", False],
            ["Peter Wishart", "Pete Wishart", "other", False],
            ["Jamie Watson", "Jamie Watson", "other", False],
            ["Mike Peters", "Mike Peters", "other", False],
            ["Derek Forbes", "Derek Forbes", "other", False],
            ["Simon Hough", "Simon Hough", "other", False],
            ["Tommie Paxton", "Tommie Paxton", "other", False],
            ["Chris Stones", "Chris Stones", "other", False],
            ["Reece Dobbin", "Reece Dobbin", "other", False],
        ],
    },
    "Culture Club": {
        "year": 1984, "type": "Band", "winner": True,
        "search": "\"Culture Club\"", "member_search": "Culture Club",
        "members": ["Boy George", "Jon Moss", "Mikey Craig", "Roy Hay"],
        "distinctive": ["Boy George", "Mikey Craig"],
        "other_era": ["Jon Suede"],
        "extra_queries": [],
        "names": [
            ["Boy George", "Boy George", "nomination", True],
            ["George Alan O'Dowd", "Boy George", "nomination", True],
            ["George O'Dowd", "Boy George", "nomination", True],
            ["John Moss", "Jon Moss", "nomination", False],
            ["Jon Moss", "Jon Moss", "nomination", False],
            ["Jonathan Moss", "Jon Moss", "nomination", False],
            ["Michael Craig", "Mikey Craig", "nomination", False],
            ["Michael Emile Craig", "Mikey Craig", "nomination", True],
            ["Mikey Craig", "Mikey Craig", "nomination", True],
            ["Roy Hay", "Roy Hay", "nomination", False],
            ["Jon Suede", "Jon Suede", "other", False],
        ],
    },
    "Eurythmics": {
        "year": 1984, "type": "Duo", "winner": False,
        "search": "\"Eurythmics\"", "member_search": "Eurythmics",
        "members": ["Annie Lennox", "Dave Stewart"],
        "distinctive": ["Annie Lennox"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Ann Lennox", "Annie Lennox", "nomination", True],
            ["Annie Lennox", "Annie Lennox", "nomination", True],
            ["Dave Stewart", "Dave Stewart", "nomination", False],
            ["David A. Stewart", "Dave Stewart", "nomination", False],
            ["David Allan Stewart", "Dave Stewart", "nomination", False],
            ["David Stewart", "Dave Stewart", "nomination", False],
        ],
    },
    "Men Without Hats": {
        "year": 1984, "type": "Band", "winner": False,
        "search": "\"Men Without Hats\"", "member_search": "Men Without Hats",
        "members": ["Allan McCarthy", "Colin Doroschuk", "Ivan Doroschuk", "Stefan Doroschuk"],
        "distinctive": ["Colin Doroschuk", "Ivan Doroschuk", "Stefan Doroschuk"],
        "other_era": ["Al Gunn", "Bruce Murphy", "Dave Hill", "Denis Toupin", "Félix Matte", "Heidi Garcia", "Jean-Marc Pisapia", "John Gurrin", "John Kastner", "Johnny Souranis", "Jérémie Arrobas", "Lenny Pinkas", "Lysanne Thibodeau", "Marika Tjelios", "Michel Langevin", "Pete Seabrooke", "Pierre Taillefer", "Richard Samson", "Roman Martyn", "Tracy Howe"],
        "extra_queries": [],
        "names": [
            ["Alan McCarthy", "Allan McCarthy", "nomination", False],
            ["Allan McCarthy", "Allan McCarthy", "nomination", False],
            ["Colin Doroschuk", "Colin Doroschuk", "nomination", True],
            ["Ivan Doroschuk", "Ivan Doroschuk", "nomination", True],
            ["Stefan Doroschuk", "Stefan Doroschuk", "nomination", True],
            ["Al Gunn", "Al Gunn", "other", False],
            ["Alain Lefebvre", "Al Gunn", "other", False],
            ["Bruce Murphy", "Bruce Murphy", "other", False],
            ["Dave Hill", "Dave Hill", "other", False],
            ["Denis Toupin", "Denis Toupin", "other", False],
            ["Félix Matte", "Félix Matte", "other", False],
            ["Heidi Garcia", "Heidi Garcia", "other", False],
            ["Jean-Marc Pisapia", "Jean-Marc Pisapia", "other", False],
            ["John Gurrin", "John Gurrin", "other", False],
            ["John Kastner", "John Kastner", "other", False],
            ["Johnny Souranis", "Johnny Souranis", "other", False],
            ["Jeremie Arrobas", "Jérémie Arrobas", "other", False],
            ["Jérémie Arrobas", "Jérémie Arrobas", "other", False],
            ["Lenny Pinkas", "Lenny Pinkas", "other", False],
            ["Lysanne Thibodeau", "Lysanne Thibodeau", "other", False],
            ["Marika Tjelios", "Marika Tjelios", "other", False],
            ["Michel 'Away' Langevin", "Michel Langevin", "other", False],
            ["Michel Langevin", "Michel Langevin", "other", False],
            ["Pete Seabrooke", "Pete Seabrooke", "other", False],
            ["Pierre Taillefer", "Pierre Taillefer", "other", False],
            ["Richard Samson", "Richard Samson", "other", False],
            ["Roman Martyn", "Roman Martyn", "other", False],
            ["Tracy Howe", "Tracy Howe", "other", False],
        ],
    },
    "Musical Youth": {
        "year": 1984, "type": "Band", "winner": False,
        "search": "\"Musical Youth\" band", "member_search": "Musical Youth",
        "members": ["Dennis Seaton", "Frederick Waite", "Kelvin Grant", "Michael Grant", "Patrick Waite"],
        "distinctive": ["Frederick Waite", "Kelvin Grant"],
        "other_era": ["Frederick Waite Sr."],
        "extra_queries": [],
        "names": [
            ["Dennis Seaton", "Dennis Seaton", "nomination", False],
            ["Freddie 'Junior' Waite", "Frederick Waite", "nomination", True],
            ["Frederick Waite", "Frederick Waite", "nomination", True],
            ["Junior Waite", "Frederick Waite", "nomination", True],
            ["Kelvin Grant", "Kelvin Grant", "nomination", True],
            ["Michael Grant", "Michael Grant", "nomination", False],
            ["Mike Grant", "Michael Grant", "nomination", False],
            ["Patrick Waite", "Patrick Waite", "nomination", False],
            ["Frederick Waite Sr.", "Frederick Waite Sr.", "other", False],
        ],
    },
    "Corey Hart": {
        "year": 1985, "type": "Solo", "winner": False,
        "search": "\"Corey Hart\" singer", "member_search": "Corey Hart",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Corey Hart", "Corey Hart", "nomination", True],
            ["Corey Mitchell Hart", "Corey Hart", "nomination", True],
        ],
    },
    "Cyndi Lauper": {
        "year": 1985, "type": "Solo", "winner": True,
        "search": "\"Cyndi Lauper\"", "member_search": "Cyndi Lauper",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Cindy Lauper", "Cyndi Lauper", "nomination", True],
            ["Cyndi Lauper", "Cyndi Lauper", "nomination", True],
            ["Cynthia Ann Stephanie Lauper", "Cyndi Lauper", "nomination", True],
            ["Cynthia Lauper", "Cyndi Lauper", "nomination", True],
        ],
    },
    "Frankie Goes to Hollywood": {
        "year": 1985, "type": "Band", "winner": False,
        "search": "\"Frankie Goes to Hollywood\"", "member_search": "Frankie Goes to Hollywood",
        "members": ["Brian Nash", "Holly Johnson", "Mark O'Toole", "Paul Rutherford", "Peter Gill"],
        "distinctive": ["Holly Johnson"],
        "other_era": ["Ambrose Reynolds", "Jed O'Toole"],
        "extra_queries": [],
        "names": [
            ["Brian 'Nasher' Nash", "Brian Nash", "nomination", False],
            ["Brian Nash", "Brian Nash", "nomination", False],
            ["Holly Johnson", "Holly Johnson", "nomination", True],
            ["William 'Holly' Johnson", "Holly Johnson", "nomination", True],
            ["William Johnson", "Holly Johnson", "nomination", False],
            ["Mark O'Toole", "Mark O'Toole", "nomination", False],
            ["Paul Rutherford", "Paul Rutherford", "nomination", False],
            ["Peter 'Ped' Gill", "Peter Gill", "nomination", False],
            ["Peter Gill", "Peter Gill", "nomination", False],
            ["Ambrose Reynolds", "Ambrose Reynolds", "other", False],
            ["Jed O'Toole", "Jed O'Toole", "other", False],
        ],
    },
    "Sheila E.": {
        "year": 1985, "type": "Solo", "winner": False,
        "search": "\"Sheila E\"", "member_search": "Sheila E.",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Sheila Cecelia Escovedo", "Sheila E.", "nomination", True],
            ["Sheila Cecilia Escovedo", "Sheila E.", "nomination", True],
            ["Sheila E.", "Sheila E.", "nomination", True],
            ["Sheila Escovedo", "Sheila E.", "nomination", True],
        ],
    },
    "The Judds": {
        "year": 1985, "type": "Duo", "winner": False,
        "search": "\"The Judds\"", "member_search": "Judds",
        "members": ["Naomi Judd", "Wynonna Judd"],
        "distinctive": ["Naomi Judd", "Wynonna Judd"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Diana Ciminella", "Naomi Judd", "nomination", True],
            ["Diana Ellen Judd", "Naomi Judd", "nomination", True],
            ["Diana Judd", "Naomi Judd", "nomination", False],
            ["Naomi Judd", "Naomi Judd", "nomination", True],
            ["Christina Ciminella", "Wynonna Judd", "nomination", True],
            ["Christina Claire Ciminella", "Wynonna Judd", "nomination", True],
            ["Wynonna", "Wynonna Judd", "nomination", True],
            ["Wynonna Ellen Judd Moser", "Wynonna Judd", "nomination", True],
            ["Wynonna Judd", "Wynonna Judd", "nomination", True],
        ],
    },
    "a-ha": {
        "year": 1986, "type": "Band", "winner": False,
        "search": "\"a-ha\" band", "member_search": "a-ha",
        "members": ["Magne Furuholmen", "Morten Harket", "Paul Waaktaar-Savoy"],
        "distinctive": ["Magne Furuholmen", "Morten Harket", "Paul Waaktaar-Savoy"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Magne 'Mags' Furuholmen", "Magne Furuholmen", "nomination", True],
            ["Magne Furuholmen", "Magne Furuholmen", "nomination", True],
            ["Morten Harket", "Morten Harket", "nomination", True],
            ["Paul Waaktaar", "Paul Waaktaar-Savoy", "nomination", True],
            ["Paul Waaktaar-Savoy", "Paul Waaktaar-Savoy", "nomination", True],
            ["Pål Waaktaar", "Paul Waaktaar-Savoy", "nomination", True],
            ["Pål Waaktaar Gamst", "Paul Waaktaar-Savoy", "nomination", True],
            ["Pål Waaktaar-Savoy", "Paul Waaktaar-Savoy", "nomination", True],
        ],
    },
    "Freddie Jackson": {
        "year": 1986, "type": "Solo", "winner": False,
        "search": "\"Freddie Jackson\"", "member_search": "Freddie Jackson",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Freddie Jackson", "Freddie Jackson", "nomination", True],
        ],
    },
    "Julian Lennon": {
        "year": 1986, "type": "Solo", "winner": False,
        "search": "\"Julian Lennon\"", "member_search": "Julian Lennon",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["John Charles Julian Lennon", "Julian Lennon", "nomination", True],
            ["Julian Lennon", "Julian Lennon", "nomination", True],
        ],
    },
    "Katrina and the Waves": {
        "year": 1986, "type": "Band", "winner": False,
        "search": "\"Katrina and the Waves\"", "member_search": "Katrina and the Waves",
        "members": ["Alex Cooper", "Katrina Leskanich", "Kimberley Rew", "Vince de la Cruz"],
        "distinctive": ["Katrina Leskanich", "Kimberley Rew"],
        "other_era": ["Bob Jakins"],
        "extra_queries": [],
        "names": [
            ["Alex Cooper", "Alex Cooper", "nomination", False],
            ["Katrina Leskanich", "Katrina Leskanich", "nomination", True],
            ["Kimberley Rew", "Kimberley Rew", "nomination", True],
            ["Kimberly Rew", "Kimberley Rew", "nomination", True],
            ["Vince de la Cruz", "Vince de la Cruz", "nomination", False],
            ["Bob Jakins", "Bob Jakins", "other", False],
        ],
    },
    "Sade": {
        "year": 1986, "type": "Band", "winner": True,
        "search": "\"Sade\"", "member_search": "Sade",
        "members": ["Andrew Hale", "Paul S. Denman", "Sade Adu", "Stuart Matthewman"],
        "distinctive": ["Paul S. Denman", "Sade Adu", "Stuart Matthewman"],
        "other_era": ["Dave Early", "Paul Anthony Cooke"],
        "extra_queries": [],
        "names": [
            ["Andrew 'AJ' Hale", "Andrew Hale", "nomination", False],
            ["Andrew Hale", "Andrew Hale", "nomination", False],
            ["Andy Hale", "Andrew Hale", "nomination", False],
            ["Paul Denman", "Paul S. Denman", "nomination", True],
            ["Paul S. Denman", "Paul S. Denman", "nomination", True],
            ["Paul Spencer Denman", "Paul S. Denman", "nomination", True],
            ["Folasade Adu", "Sade Adu", "nomination", True],
            ["Helen Adu", "Sade Adu", "nomination", True],
            ["Helen Folasade Adu", "Sade Adu", "nomination", True],
            ["Sade Adu", "Sade Adu", "nomination", True],
            ["Stewart Matthewman", "Stuart Matthewman", "nomination", True],
            ["Stuart Matthewman", "Stuart Matthewman", "nomination", True],
            ["Dave Early", "Dave Early", "other", False],
            ["Paul Anthony Cook", "Paul Anthony Cooke", "other", False],
            ["Paul Anthony Cooke", "Paul Anthony Cooke", "other", False],
            ["Paul Cooke", "Paul Anthony Cooke", "other", False],
        ],
    },
    "Bruce Hornsby and the Range": {
        "year": 1987, "type": "Band", "winner": True,
        "search": "\"Bruce Hornsby\"", "member_search": "Bruce Hornsby",
        "members": ["Bruce Hornsby", "David Mansfield", "George Marinelli", "Joe Puerta", "John Molo", "Peter Harris"],
        "distinctive": ["Bruce Hornsby", "George Marinelli", "Joe Puerta", "John Molo"],
        "other_era": ["J. T. Thomas"],
        "extra_queries": [],
        "names": [
            ["Bruce Hornsby", "Bruce Hornsby", "nomination", True],
            ["David Mansfield", "David Mansfield", "nomination", False],
            ["George Marinelli", "George Marinelli", "nomination", True],
            ["Joe Puerta", "Joe Puerta", "nomination", True],
            ["Joseph Puerta", "Joe Puerta", "nomination", False],
            ["John Molo", "John Molo", "nomination", True],
            ["Pete Harris", "Peter Harris", "nomination", False],
            ["Peter Harris", "Peter Harris", "nomination", False],
            ["J. T. Thomas", "J. T. Thomas", "other", False],
        ],
    },
    "Glass Tiger": {
        "year": 1987, "type": "Band", "winner": False,
        "search": "\"Glass Tiger\"", "member_search": "Glass Tiger",
        "members": ["Al Connelly", "Alan Frew", "Michael Hanson", "Sam Reid", "Wayne Parker"],
        "distinctive": ["Alan Frew"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Al Connelly", "Al Connelly", "nomination", False],
            ["Alan Frew", "Alan Frew", "nomination", True],
            ["Michael Hanson", "Michael Hanson", "nomination", False],
            ["Sam Reid", "Sam Reid", "nomination", False],
            ["Wayne Parker", "Wayne Parker", "nomination", False],
        ],
    },
    "Nu Shooz": {
        "year": 1987, "type": "Duo", "winner": False,
        "search": "\"Nu Shooz\"", "member_search": "Nu Shooz",
        "members": ["John Smith", "Valerie Day"],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["John Smith", "John Smith", "nomination", False],
            ["Valerie Day", "Valerie Day", "nomination", False],
        ],
    },
    "Simply Red": {
        "year": 1987, "type": "Band", "winner": False,
        "search": "\"Simply Red\"", "member_search": "Simply Red",
        "members": ["Chris Joyce", "Fritz McIntyre", "Ian Kirkham", "Mick Hucknall", "Sylvan Richardson", "Tim Kellett", "Tony Bowers"],
        "distinctive": ["Fritz McIntyre", "Mick Hucknall", "Sylvan Richardson"],
        "other_era": ["Aziz Ibrahim", "David Fryman", "Gota Yashiki", "Heitor Pereira", "Shaun Ward"],
        "extra_queries": [],
        "names": [
            ["Chris Joyce", "Chris Joyce", "nomination", False],
            ["Fritz McIntyre", "Fritz McIntyre", "nomination", True],
            ["Ian Kirkham", "Ian Kirkham", "nomination", False],
            ["Michael Hucknall", "Mick Hucknall", "nomination", False],
            ["Michael James Hucknall", "Mick Hucknall", "nomination", True],
            ["Mick 'Hucko' Hucknall", "Mick Hucknall", "nomination", True],
            ["Mick 'Red' Hucknall", "Mick Hucknall", "nomination", True],
            ["Mick Hucknall", "Mick Hucknall", "nomination", True],
            ["Sylvan Richardson", "Sylvan Richardson", "nomination", True],
            ["Tim Kellett", "Tim Kellett", "nomination", False],
            ["Tony Bowers", "Tony Bowers", "nomination", False],
            ["Aziz Ibrahim", "Aziz Ibrahim", "other", False],
            ["David Fryman", "David Fryman", "other", False],
            ["Gota Yashiki", "Gota Yashiki", "other", False],
            ["Heitor Pereira", "Heitor Pereira", "other", False],
            ["Heitor TP", "Heitor Pereira", "other", False],
            ["Shaun Ward", "Shaun Ward", "other", False],
        ],
    },
    "Timbuk3": {
        "year": 1987, "type": "Duo", "winner": False,
        "search": "\"Timbuk 3\"", "member_search": "Timbuk 3",
        "members": ["Barbara K. MacDonald", "Pat MacDonald"],
        "distinctive": ["Barbara K. MacDonald"],
        "other_era": ["Courtney Audain", "Wally Ingram"],
        "extra_queries": [],
        "names": [
            ["Barbara K", "Barbara K. MacDonald", "nomination", True],
            ["Barbara K. MacDonald", "Barbara K. MacDonald", "nomination", True],
            ["Barbara Kooyman", "Barbara K. MacDonald", "nomination", True],
            ["Pat MacDonald", "Pat MacDonald", "nomination", False],
            ["Courtney Audain", "Courtney Audain", "other", False],
            ["Wally Ingram", "Wally Ingram", "other", False],
        ],
    },
    "Breakfast Club": {
        "year": 1988, "type": "Band", "winner": False,
        "search": "\"Breakfast Club\" band Gilroy", "member_search": "Breakfast Club",
        "members": ["Dan Gilroy", "Ed Gilroy", "Gary Burke", "Paul Kauk", "Randy Jackson", "Stephen Bray"],
        "distinctive": ["Paul Kauk"],
        "other_era": ["Angie Smit", "Madonna"],
        "extra_queries": [],
        "names": [
            ["Dan Gilroy", "Dan Gilroy", "nomination", False],
            ["Ed Gilroy", "Ed Gilroy", "nomination", False],
            ["Eddie Gilroy", "Ed Gilroy", "nomination", False],
            ["Gary Burke", "Gary Burke", "nomination", False],
            ["Paul Kauk", "Paul Kauk", "nomination", True],
            ["Randy Jackson", "Randy Jackson", "nomination", False],
            ["Stephen Bray", "Stephen Bray", "nomination", False],
            ["Steve Bray", "Stephen Bray", "nomination", False],
            ["Angie Schmit", "Angie Smit", "other", False],
            ["Angie Smit", "Angie Smit", "other", False],
            ["Madonna", "Madonna", "other", False],
        ],
    },
    "Cutting Crew": {
        "year": 1988, "type": "Band", "winner": False,
        "search": "\"Cutting Crew\"", "member_search": "Cutting Crew",
        "members": ["Colin Farley", "Kevin MacMichael", "Martin Beedle", "Nick Van Eede", "Tony Moore"],
        "distinctive": ["Kevin MacMichael", "Nick Van Eede"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Colin Farley", "Colin Farley", "nomination", False],
            ["Kevin MacMichael", "Kevin MacMichael", "nomination", True],
            ["Kevin Scott MacMichael", "Kevin MacMichael", "nomination", True],
            ["Frosty Beedle", "Martin Beedle", "nomination", False],
            ["Martin 'Frosty' Beedle", "Martin Beedle", "nomination", False],
            ["Martin Beedle", "Martin Beedle", "nomination", False],
            ["Martin Philip Beedle", "Martin Beedle", "nomination", False],
            ["Nicholas Eede", "Nick Van Eede", "nomination", True],
            ["Nicholas Van Eede", "Nick Van Eede", "nomination", True],
            ["Nick Eede", "Nick Van Eede", "nomination", False],
            ["Nick Van Eede", "Nick Van Eede", "nomination", True],
            ["Van Eede", "Nick Van Eede", "nomination", False],
            ["Tony Moore", "Tony Moore", "nomination", False],
        ],
    },
    "Jody Watley": {
        "year": 1988, "type": "Solo", "winner": True,
        "search": "\"Jody Watley\"", "member_search": "Jody Watley",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Jody Watley", "Jody Watley", "nomination", True],
        ],
    },
    "Swing Out Sister": {
        "year": 1988, "type": "Band/Duo", "winner": False,
        "search": "\"Swing Out Sister\"", "member_search": "Swing Out Sister",
        "members": ["Andy Connell", "Corinne Drewery", "Martin Jackson"],
        "distinctive": ["Corinne Drewery"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Andrew Connell", "Andy Connell", "nomination", False],
            ["Andy Connell", "Andy Connell", "nomination", False],
            ["Corinne Drewery", "Corinne Drewery", "nomination", True],
            ["Corrine Drewery", "Corinne Drewery", "nomination", True],
            ["Martin Jackson", "Martin Jackson", "nomination", False],
        ],
    },
    "Terence Trent D'Arby": {
        "year": 1988, "type": "Solo", "winner": False,
        "search": "\"Terence Trent D'Arby\"", "member_search": "Terence Trent D'Arby",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": ["\"Sananda Maitreya\" interview"],
        "names": [
            ["Francesco Maitreya", "Terence Trent D'Arby", "nomination", True],
            ["Sananda Francesco Maitreya", "Terence Trent D'Arby", "nomination", True],
            ["Sananda Maitreya", "Terence Trent D'Arby", "nomination", True],
            ["Terence Trent D'Arby", "Terence Trent D'Arby", "nomination", True],
            ["Terence Trent Howard", "Terence Trent D'Arby", "nomination", True],
            ["Terence Trent Susan Whitfield", "Terence Trent D'Arby", "nomination", True],
            ["Trent D'Arby", "Terence Trent D'Arby", "nomination", True],
        ],
    },
    "Rick Astley": {
        "year": 1989, "type": "Solo", "winner": False,
        "search": "\"Rick Astley\"", "member_search": "Rick Astley",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Richard Paul Astley", "Rick Astley", "nomination", True],
            ["Rick Astley", "Rick Astley", "nomination", True],
        ],
    },
    "Take 6": {
        "year": 1989, "type": "Group", "winner": False,
        "search": "\"Take 6\" a cappella", "member_search": "Take 6",
        "members": ["Alvin Chea", "Cedric Dent", "Claude McKnight", "David Thomas", "Mark Kibble", "Mervyn Warren"],
        "distinctive": ["Alvin Chea", "Claude McKnight", "Mark Kibble", "Mervyn Warren"],
        "other_era": ["Joel Kibble"],
        "extra_queries": [],
        "names": [
            ["Alvin 'Vinnie' Chea", "Alvin Chea", "nomination", True],
            ["Alvin 'Vinny' Chea", "Alvin Chea", "nomination", True],
            ["Alvin Chea", "Alvin Chea", "nomination", True],
            ["Cedric Dent", "Cedric Dent", "nomination", False],
            ["Claude McKnight", "Claude McKnight", "nomination", True],
            ["Claude V. McKnight III", "Claude McKnight", "nomination", True],
            ["Dave Thomas", "David Thomas", "nomination", False],
            ["David Thomas", "David Thomas", "nomination", False],
            ["Mark Kibble", "Mark Kibble", "nomination", True],
            ["Merv Warren", "Mervyn Warren", "nomination", False],
            ["Mervyn Warren", "Mervyn Warren", "nomination", True],
            ["Joel Kibble", "Joel Kibble", "other", False],
            ["Joey Kibble", "Joel Kibble", "other", False],
        ],
    },
    "Toni Childs": {
        "year": 1989, "type": "Solo", "winner": False,
        "search": "\"Toni Childs\"", "member_search": "Toni Childs",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Toni Childs", "Toni Childs", "nomination", True],
        ],
    },
    "Tracy Chapman": {
        "year": 1989, "type": "Solo", "winner": True,
        "search": "\"Tracy Chapman\"", "member_search": "Tracy Chapman",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Tracy Chapman", "Tracy Chapman", "nomination", True],
        ],
    },
    "Vanessa Williams": {
        "year": 1989, "type": "Solo", "winner": False,
        "search": "\"Vanessa Williams\" singer", "member_search": "Vanessa Williams",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Vanessa Lynn Williams", "Vanessa Williams", "nomination", True],
            ["Vanessa Williams", "Vanessa Williams", "nomination", True],
        ],
    },
    "Indigo Girls": {
        "year": 1990, "type": "Duo", "winner": False,
        "search": "\"Indigo Girls\"", "member_search": "Indigo Girls",
        "members": ["Amy Ray", "Emily Saliers"],
        "distinctive": ["Emily Saliers"],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Amy Ray", "Amy Ray", "nomination", False],
            ["Emily Saliers", "Emily Saliers", "nomination", True],
        ],
    },
    "Neneh Cherry": {
        "year": 1990, "type": "Solo", "winner": False,
        "search": "\"Neneh Cherry\"", "member_search": "Neneh Cherry",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Neneh Cherry", "Neneh Cherry", "nomination", True],
            ["Neneh Karlsson", "Neneh Cherry", "nomination", True],
            ["Neneh Mariann Karlsson", "Neneh Cherry", "nomination", True],
        ],
    },
    "Soul II Soul": {
        "year": 1990, "type": "Collective", "winner": False,
        "search": "\"Soul II Soul\"", "member_search": "Soul II Soul",
        "members": ["Aitch B", "Caron Wheeler", "Daddae Harvey", "Jazzie B", "Nellee Hooper", "Simon Law"],
        "distinctive": ["Aitch B", "Caron Wheeler", "Daddae Harvey", "Jazzie B", "Nellee Hooper"],
        "other_era": ["Doreen Waddell", "Kym Mazelle", "Lamya Al-Mugheiry", "Marcia Lewis", "Rose Windross", "Victoria Wilson-James"],
        "extra_queries": [],
        "names": [
            ["Aitch B", "Aitch B", "nomination", True],
            ["Aitch Bee", "Aitch B", "nomination", True],
            ["DJ Aitch B", "Aitch B", "nomination", True],
            ["Caron Melina Wheeler", "Caron Wheeler", "nomination", True],
            ["Caron Wheeler", "Caron Wheeler", "nomination", True],
            ["Daddae Harvey", "Daddae Harvey", "nomination", True],
            ["Philip 'Daddae' Harvey", "Daddae Harvey", "nomination", True],
            ["Philip Harvey", "Daddae Harvey", "nomination", False],
            ["Beresford Romeo", "Jazzie B", "nomination", True],
            ["Jazzie B", "Jazzie B", "nomination", True],
            ["Trevor Beresford Romeo", "Jazzie B", "nomination", True],
            ["Nellee Hooper", "Nellee Hooper", "nomination", True],
            ["Nellie Hooper", "Nellee Hooper", "nomination", True],
            ["Paul Andrew Hooper", "Nellee Hooper", "nomination", True],
            ["Simon Law", "Simon Law", "nomination", False],
            ["Doreen Tess Waddell", "Doreen Waddell", "other", False],
            ["Doreen Waddell", "Doreen Waddell", "other", False],
            ["Kym Mazelle", "Kym Mazelle", "other", False],
            ["Kymberly Grigsby", "Kym Mazelle", "other", False],
            ["Lamya Al-Mugheiry", "Lamya Al-Mugheiry", "other", False],
            ["Marcia Lewis", "Marcia Lewis", "other", False],
            ["Rose Windross", "Rose Windross", "other", False],
            ["Victoria Wilson-James", "Victoria Wilson-James", "other", False],
        ],
    },
    "Tone Loc": {
        "year": 1990, "type": "Solo", "winner": False,
        "search": "\"Tone Loc\"", "member_search": "Tone Loc",
        "members": [],
        "distinctive": [],
        "other_era": [],
        "extra_queries": [],
        "names": [
            ["Anthony Smith", "Tone Loc", "nomination", False],
            ["Anthony T. Smith", "Tone Loc", "nomination", False],
            ["Anthony Terrell Smith", "Tone Loc", "nomination", True],
            ["Tone Loc", "Tone Loc", "nomination", True],
        ],
    },
}

# ═════════════════════════════════════════════════════════════════════════════
# RECOGNISING THE ACT
# Names that are ordinary words must be capitalised ("Dire Straits", not "in dire straits").
# For Asia, Breakfast Club and Big Country the name alone proves nothing (Asia Argento, the
# Breakfast Club radio show, 'Big Country' Reeves), so a member, a song or a band reference
# must also appear. OTHER_SENSE lists uses of the name that are about something else.
# ═════════════════════════════════════════════════════════════════════════════
ACT_NAME = {
    "Dire Straits": (r"\b(?:Dire Straits|DIRE STRAITS)\b", True),
    "Rickie Lee Jones": (r"rickie lee jones", False),
    "The Knack": (r"\b(?:[Tt]he )?(?:Knack|KNACK)\b", True),
    "Christopher Cross": (r"christopher cross", False), "Irene Cara": (r"irene cara", False),
    "Robbie Dupree": (r"robbie dupree", False), "Amy Holland": (r"amy holland", False),
    "The Pretenders": (r"\b(?:Pretenders|PRETENDERS)\b", True),
    "Sheena Easton": (r"sheena easton", False),
    "Adam and the Ants": (r"adam (?:and|&|\+|n['’]?) ?the ants|adam ant\b", False),
    "The Go-Go's": (r"go[- ]?go['’]?s\b", False), "James Ingram": (r"james ingram", False),
    "Luther Vandross": (r"vandross", False),
    "Men at Work": (r"\b(?:Men [Aa]t Work|MEN AT WORK)\b", True),
    "Asia": (r"\b(?:Asia|ASIA)\b", True), "Jennifer Holliday": (r"jennifer holliday", False),
    "The Human League": (r"human league", False),
    "Stray Cats": (r"\b(?:Stray Cats?|STRAY CATS?)\b", True),
    "Culture Club": (r"\b(?:Culture Club|CULTURE CLUB)\b", True),
    "Big Country": (r"\b(?:Big Country|BIG COUNTRY)\b", True),
    "Eurythmics": (r"eurythmics", False), "Men Without Hats": (r"men without hats", False),
    "Musical Youth": (r"\b(?:Musical Youth|MUSICAL YOUTH)\b", True),
    "Cyndi Lauper": (r"lauper", False), "Sheila E.": (r"sheila e\b|sheila escovedo", False),
    "Frankie Goes to Hollywood": (r"frankie goes to hollywood|\bFGTH\b", False),
    "Corey Hart": (r"corey hart", False), "The Judds": (r"\bthe judds\b|naomi judd|wynonna", False),
    "Sade": (r"(?<![Dd]e )\b(?:Sade|SADE)\b", True), "a-ha": (r"\ba-ha\b", False),
    "Freddie Jackson": (r"freddie jackson", False),
    "Katrina and the Waves": (r"katrina (?:and|&) the waves", False),
    "Julian Lennon": (r"julian lennon", False), "Bruce Hornsby and the Range": (r"hornsby", False),
    "Glass Tiger": (r"glass tiger", False), "Nu Shooz": (r"nu[- ]?shooz", False),
    "Simply Red": (r"\b(?:Simply Red|SIMPLY RED)\b", True), "Timbuk3": (r"timbuk ?3", False),
    "Jody Watley": (r"jody watley", False),
    "Breakfast Club": (r"breakfast club", False),
    "Cutting Crew": (r"\b(?:Cutting Crew|CUTTING CREW)\b", True),
    "Terence Trent D'Arby": (r"terence trent d['’]?arby|trent d['’]arby|sananda maitreya", False),
    "Swing Out Sister": (r"swing out sister", False), "Tracy Chapman": (r"tracy chapman", False),
    "Rick Astley": (r"rick astley", False), "Toni Childs": (r"toni childs", False),
    "Take 6": (r"\b(?:Take (?:6|Six)|TAKE (?:6|SIX))\b", True),
    "Vanessa Williams": (r"vanessa (?:l\. |lynn )?williams", False),
    "Neneh Cherry": (r"neneh cherry", False), "Indigo Girls": (r"indigo girls", False),
    "Soul II Soul": (r"soul (?:ii|2) soul", False), "Tone Loc": (r"tone[- ]?l[oō]c", False),
}
ACT_RX = {a: re.compile(p, 0 if cs else re.I) for a, (p, cs) in ACT_NAME.items()}

NEEDS_ANCHOR = {
    "Asia": r"Heat of the Moment|Only Time Will Tell|Sole Survivor|Don['’]t Cry|The Smile Has Left|supergroup|"
            r"Wetton|Downes|Steve Howe|Carl Palmer|King Crimson|Buggles|\bELP\b|\bYes\b|\([Bb]and\)|[Pp]rog|"
            r"[Oo]riginal (?:[Ll]ine-?up|[Mm]embers)",
    "Breakfast Club": r"Gilroy|Stephen Bray|Steve Bray|Gary Burke|Paul Kauk|Right on Track|Rico Mambo|"
                      r"\(band\)|Madonna['’]s (?:first|early|old) band|Madonna (?:and|&) (?:[Tt]he )?Breakfast Club",
    "Big Country": r"Adamson|Brzezicki|Tony Butler|Bruce Watson|In a Big Country|Fields of Fire|The Crossing|"
                   r"Steeltown|Wonderland|Look Away|Mike Peters|Scottish (?:band|rock)|\([Bb]and\)",
}
NEEDS_ANCHOR = {a: re.compile(p) for a, p in NEEDS_ANCHOR.items()}

OTHER_SENSE = {
    "Breakfast Club": r"Charlamagne|DJ Envy|Angela Yee|Power 105|Jess Hilarious|Loren LoRosa|John Hughes|"
                      r"Molly Ringwald|Judd Nelson|Ally Sheedy|Emilio Estevez|Anthony Michael Hall|Brat Pack|"
                      r"Simple Minds|Forget About Me",
    "Asia": r"Asia Argento|Asia O['’]Hara|Asia Kate Dillon|Asia Monet|Sha-Asia|(?:South|East|Central|Southeast|"
            r"South-East|South East|West)(?:ern)? Asia|Asia[- ]?Pacific|\bAsian\b|\b(?:in|across|to|around|from) Asia\b|"
            r"Asia [Tt]our|Asia [Tt]rip|Asia [Cc]up",
    "Take 6": r"\(Take (?:6|Six)\)|\[Take (?:6|Six)\]|\btake (?:6|six)\b|Take (?:6|Six) (?:[Mm]inutes|[Mm]onths|"
              r"[Ww]eeks|[Dd]ays|[Yy]ears|[Hh]ours|[Ss]teps|[Tt]hings|[Ww]ays|[Tt]ips|[Ss]hots)",
    "Stray Cats": r"[Ff]eral|[Kk]itten|\bTNR\b|[Ss]pay|[Nn]euter|[Rr]escue|[Aa]dopt|[Ss]helter",
    "Culture Club": r"Pop Culture Club|[Yy]ogh?urt|[Kk]ombucha|[Ff]erment|[Cc]ompany [Cc]ulture|[Ww]orkplace [Cc]ulture",
    "Men at Work": r"[Cc]onstruction|[Rr]oad ?works?|Charlie Sheen|Emilio Estevez",
    "Sade": r"[Mm]arquis|\b[Dd]e Sade\b|Marat/Sade|Sade Sati|Sade Baderinwa",
    "a-ha": r"[Aa]-?ha moment|A-ha!|a-ha!",
    "The Knack": r"How to Get It|[Kk]nack (?:for|of)\b|\ba [Kk]nack\b",
    "Big Country": r"Bryant Reeves|Roy Nelson|[Bb]asketball|\bNBA\b|\bUFC\b|\bMMA\b|[Ff]ootball|Abilene|[Tt]rucking",
    "Eurythmics": r"[Dd]alcroze|[Ee]urh?ythmy|[Ee]urhythmics|[Ss]olf[eè]ge|[Mm]ovement class|[Mm]usic (?:education|educators?|teachers?)",
    "Corey Hart": r"Brewers|Pirates|Mariners|[Bb]aseball|\bMLB\b|[Hh]ome [Rr]un|[Oo]utfielder|[Ii]nning",
    "Vanessa Williams": r"Vanessa A\.? Williams|Vanessa Estelle Williams|Melrose Place|Soul Food",
    "Bruce Hornsby and the Range": r"Rogers Hornsby|Nick Hornsby|Hornsby,? (?:NSW|N\.S\.W|Sydney)|Hornsby (?:Shire|station|Westfield)",
    "Musical Youth": r"[Yy]outh [Oo]rchestra|[Yy]outh [Cc]hoir|[Yy]outh [Tt]heat(?:re|er)|[Yy]outh [Bb]and",
}
OTHER_SENSE = {a: re.compile(p) for a, p in OTHER_SENSE.items()}

# ═════════════════════════════════════════════════════════════════════════════
# BEST-KNOWN SONGS
# A song title with no interview word means a performance or music video. Song titles are also
# removed before looking for interview words, so "Talk of the Town" or "Baby Talks Dirty" does
# not count as a talk. One-word titles count only when set off by a dash, quotes or brackets.
# ═════════════════════════════════════════════════════════════════════════════
SONGS = {
    "Dire Straits": ["Sultans of Swing", "Money for Nothing", "Walk of Life", "Brothers in Arms", "Romeo and Juliet",
                     "Private Investigations", "Telegraph Road", "So Far Away", "Your Latest Trick", "Lady Writer",
                     "Twisting by the Pool", "Calling Elvis", "Down to the Waterline", "Water of Love",
                     "Tunnel of Love", "Skateaway", "Industrial Disease", "Heavy Fuel", "On Every Street"],
    "Rickie Lee Jones": ["Chuck E's in Love", "Young Blood", "We Belong Together", "Coolsville",
                         "Danny's All-Star Joint", "The Last Chance Texaco"],
    "The Knack": ["My Sharona", "Good Girls Don't", "Baby Talks Dirty", "Frustrated", "Pay the Devil"],
    "Christopher Cross": ["Sailing", "Ride Like the Wind", "Arthur's Theme", "Never Be the Same",
                          "Say You'll Be Mine", "Think of Laura", "No Time for Talk"],
    "Irene Cara": ["Fame", "Flashdance What a Feeling", "What a Feeling", "Out Here on My Own",
                   "Breakdance", "Hot Lunch Jam"],
    "Robbie Dupree": ["Steal Away", "Hot Rod Hearts"],
    "Amy Holland": ["How Do I Survive", "She's on Fire", "Shake Me Wake Me"],
    "The Pretenders": ["Brass in Pocket", "Back on the Chain Gang", "Don't Get Me Wrong", "Talk of the Town",
                       "Middle of the Road", "Hymn to Her", "I'll Stand by You", "Stop Your Sobbing",
                       "Message of Love", "My City Was Gone", "2000 Miles", "Night in My Veins", "Precious",
                       "Tattooed Love Boys", "I Go to Sleep", "Kid", "Thin Line Between Love and Hate"],
    "Sheena Easton": ["Morning Train", "Modern Girl", "For Your Eyes Only", "Sugar Walls", "Strut",
                      "U Got the Look", "We've Got Tonight", "Telefone", "Almost Over You", "The Lover in Me"],
    "Adam and the Ants": ["Stand and Deliver", "Prince Charming", "Antmusic", "Kings of the Wild Frontier",
                          "Goody Two Shoes", "Ant Rap", "Young Parisians",
                          "Desperate But Not Serious", "Vive le Rock", "Puss in Boots"],
    "The Go-Go's": ["We Got the Beat", "Our Lips Are Sealed", "Vacation", "Head over Heels", "Cool Jerk"],
    "James Ingram": ["Just Once", "One Hundred Ways", "Baby Come to Me", "Yah Mo B There", "Somewhere Out There",
                     "I Don't Have the Heart", "How Do You Keep the Music Playing"],
    "Luther Vandross": ["Never Too Much", "Dance with My Father", "Power of Love",
                        "Stop to Love", "Give Me the Reason", "Endless Love", "A House Is Not a Home",
                        "If Only for One Night", "Superstar", "The Best Things in Life Are Free"],
    "Men at Work": ["Down Under", "Who Can It Be Now", "Overkill", "It's a Mistake", "Be Good Johnny",
                    "Dr. Heckyll and Mr. Jive", "Everything I Need"],
    "Asia": ["Heat of the Moment", "Only Time Will Tell", "Sole Survivor", "Don't Cry",
             "The Smile Has Left Your Eyes", "Wildest Dreams"],
    "Jennifer Holliday": ["And I Am Telling You I'm Not Going", "And I Am Telling You", "I Am Love",
                          "Hard Times for Lovers"],
    "The Human League": ["Don't You Want Me", "Mirror Man", "Keep Feeling Fascination", "Together in Electric Dreams",
                         "Love Action", "Open Your Heart", "The Lebanon", "Louise", "Life on Your Own",
                         "Being Boiled", "Tell Me When", "The Sound of the Crowd"],
    "Stray Cats": ["Rock This Town", "Stray Cat Strut", "Sexy and 17", "Runaway Boys", "Rumble in Brighton"],
    "Culture Club": ["Do You Really Want to Hurt Me", "Karma Chameleon", "Church of the Poison Mind", "Miss Me Blind",
                     "Time Clock of the Heart", "I'll Tumble 4 Ya", "It's a Miracle", "The War Song", "Victims", "I Just Wanna Be Loved", "Mistake No. 3", "God Thank You Woman",
                     "Your Kisses Are Charity"],
    "Big Country": ["In a Big Country", "Fields of Fire", "Chance", "Wonderland", "Look Away", "Harvest Home",
                    "East of Eden", "Where the Rose Is Sown", "Just a Shadow", "King of Emotion", "Ships"],
    "Eurythmics": ["Sweet Dreams", "Here Comes the Rain Again", "Love Is a Stranger", "Who's That Girl",
                   "Right by Your Side", "Would I Lie to You", "There Must Be an Angel",
                   "Sisters Are Doin' It for Themselves", "Missionary Man", "Thorn in My Side", "When Tomorrow Comes",
                   "It's Alright Baby's Coming Back", "Sexcrime", "Julia", "I Saved the World Today", "17 Again"],
    "Men Without Hats": ["The Safety Dance", "Safety Dance", "Pop Goes the World"],
    "Musical Youth": ["Pass the Dutchie", "Never Gonna Give You Up", "Youth of Today"],
    "Cyndi Lauper": ["Girls Just Want to Have Fun", "Time After Time", "True Colors", "She Bop",
                     "All Through the Night", "Money Changes Everything", "Change of Heart", "I Drove All Night",
                     "The Goonies 'R' Good Enough"],
    "Sheila E.": ["The Glamorous Life", "A Love Bizarre", "The Belle of St. Mark", "Holly Rock", "Koo Koo"],
    "Frankie Goes to Hollywood": ["Relax", "Two Tribes", "The Power of Love", "Welcome to the Pleasuredome",
                                  "Rage Hard", "Warriors of the Wasteland", "Watching the Wildlife"],
    "Corey Hart": ["Sunglasses at Night", "Never Surrender", "It Ain't Enough", "Boy in the Box",
                   "Everything in My Heart", "Can't Help Falling in Love", "Eurasian Eyes", "In Your Soul"],
    "The Judds": ["Mama He's Crazy", "Why Not Me", "Grandpa", "Love Is Alive",
                  "Rockin' with the Rhythm of the Rain", "Love Can Build a Bridge", "Turn It Loose",
                  "Give a Little Love", "Let Me Tell You About Love"],
    "Sade": ["Smooth Operator", "Your Love Is King", "The Sweetest Taboo", "No Ordinary Love", "Kiss of Life",
             "By Your Side", "Hang On to Your Love", "Is It a Crime", "Paradise", "Love Is Stronger Than Pride",
             "Cherish the Day", "Soldier of Love", "Pearls", "Jezebel", "King of Sorrow", "Never as Good as the First Time"],
    "a-ha": ["Take On Me", "The Sun Always Shines on TV", "Hunting High and Low", "Stay on These Roads", "Cry Wolf",
             "The Living Daylights", "I've Been Losing You", "Manhattan Skyline", "Touchy", "Crying in the Rain", "Summer Moved On", "Analogue", "Foot of the Mountain"],
    "Freddie Jackson": ["Rock Me Tonight", "You Are My Lady", "Jam Tonight", "Have You Ever Loved Somebody",
                        "Nice 'n' Slow", "Tasty Love", "Do Me Again", "Love Is Just a Touch Away"],
    "Katrina and the Waves": ["Walking on Sunshine", "Love Shine a Light", "Do You Want Crying",
                              "Going Down to Liverpool"],
    "Julian Lennon": ["Too Late for Goodbyes", "Valotte", "Saltwater", "Say You're Wrong",
                      "Now You're in Heaven"],
    "Bruce Hornsby and the Range": ["The Way It Is", "Mandolin Rain", "Every Little Kiss", "The Valley Road",
                                    "Look Out Any Window", "Across the River", "On the Western Skyline",
                                    "The End of the Innocence", "Defenders of the Flag", "Harbor Lights"],
    "Glass Tiger": ["Don't Forget Me When I'm Gone", "Someday",
                    "Diamond Sun", "I'm Still Searching", "Animal Heart"],
    "Nu Shooz": ["I Can't Wait", "Point of No Return", "Should I Say Yes"],
    "Simply Red": ["Holding Back the Years", "Money's Too Tight to Mention", "Something Got Me Started",
                   "If You Don't Know Me by Now", "Stars", "Fairground", "A New Flame",
                   "Come to My Aid", "Jericho", "For Your Babies", "Say You Love Me",
                   "Sunrise", "Infidelity", "Ev'ry Time We Say Goodbye", "Open Up the Red Box", "Thrill Me"],
    "Timbuk3": ["The Future's So Bright I Gotta Wear Shades", "The Future's So Bright", "Hairstyles and Attitudes", "Just Another Movie"],
    "Jody Watley": ["Looking for a New Love", "Don't You Want Me", "Some Kind of Lover",
                    "Still a Thrill", "I'm the One You Need", "Precious Love", "Everything"],
    "Breakfast Club": ["Right on Track", "Kiss and Tell", "Rico Mambo"],
    "Cutting Crew": ["I Just Died in Your Arms", "Died in Your Arms", "I've Been in Love Before",
                     "One for the Mockingbird", "Everything but My Pride"],
    "Terence Trent D'Arby": ["Wishing Well", "Sign Your Name", "Dance Little Sister", "If You Let Me Stay",
                             "Delicate", "Holding On to You", "Do You Love Me Like You Say", "She Kissed Me",
                             "Let Her Down Easy"],
    "Swing Out Sister": ["Breakout", "Surrender", "Twilight World", "Am I the Same Girl",
                         "You on My Mind", "Notgonnachange", "La La Means I Love You"],
    "Tracy Chapman": ["Fast Car", "Talkin' 'bout a Revolution", "Baby Can I Hold You", "Give Me One Reason",
                      "Crossroads", "Across the Lines", "Mountains o' Things", "Smoke and Ashes", "Telling Stories", "Subcity", "She's Got Her Ticket"],
    "Rick Astley": ["Never Gonna Give You Up", "Together Forever", "Whenever You Need Somebody", "Cry for Help",
                    "She Wants to Dance with Me", "Take Me to Your Heart", "It Would Take a Strong Strong Man",
                    "Hold Me in Your Arms", "Angels on My Side", "Keep Singing",
                    "My Arms Keep Missing You", "When I Fall in Love"],
    "Toni Childs": ["Don't Walk Away", "Zimbabwe", "Stop Your Fussin'", "Many Rivers to Cross", "House of Hope",
                    "Walk and Talk Like Angels"],
    "Take 6": ["Spread Love", "Mary", "So Much 2 Say", "I L-O-V-E U", "Something Within Me", "Get Away Jordan"],
    "Vanessa Williams": ["Save the Best for Last", "Dreamin'", "Colors of the Wind",
                         "Running Back to You", "The Sweetest Days", "He's Got No Love"],
    "Neneh Cherry": ["Buffalo Stance", "Manchild", "Kisses on the Wind", "7 Seconds", "Woman",
                     "I've Got You Under My Skin", "Money Love", "Inna City Mamma"],
    "Indigo Girls": ["Closer to Fine", "Galileo", "Power of Two", "Least Complicated", "Kid Fears",
                     "Hammer and a Nail", "Land of Canaan", "Get Out the Map", "Watershed",
                     "Love's Recovery", "Ghost", "Joking", "Chickenman"],
    "Soul II Soul": ["Back to Life", "Keep On Movin'", "A Dream's a Dream", "Fairplay", "Jazzie's Groove"],
    "Tone Loc": ["Wild Thing", "Funky Cold Medina", "I Got It Goin' On", "Cheeba Cheeba"],
}

# ═════════════════════════════════════════════════════════════════════════════
# TYPES
# Checked on the title in this order (song titles removed first); then, if the title settles
# nothing, on the description.
# ═════════════════════════════════════════════════════════════════════════════
LINK_TYPES = ["Interview", "Q&A", "Panel discussion", "Workshop or masterclass", "Documentary", "Award or speech",
              "Unclear (check)"]
LEFT_TYPES = ["Performance or music video", "Cover, tribute, lesson or reaction", "Commentary by others",
              "Trailer or advert", "Not about the act", "Other-era member only", "Not a YouTube video"]
TYPE_ORDER = LINK_TYPES + LEFT_TYPES

# "(from the 2019 documentary)" names where a clip came from, not what the clip is
FROM_SOURCE = re.compile(r"[\(\[]\s*(?:taken |clip |excerpt )?from\b[^\)\]]*[\)\]]", re.I)
STRONG_THIRD = re.compile(
    r"\breaction\b|\breacts?\b|\breacting\b|\bfirst time (?:hearing|listening|watching|reacting)\b|\bvocal coach\b|"
    r"\bkaraoke\b|\blessons?\b|\btutorial\b|\bhow to play\b|\bchords?\b|\btabs?\b|\blicks\b|\btechniques\b|"
    r"\b(?:piano|guitar|bass|drum|vocal|singing) (?:style|technique|tone|sound|secrets?)\b|"
    r"\b(?:guitar|drum|bass|piano|vocal|ukulele|sax(?:ophone)?|violin|cello|harp|acoustic|instrumental|metal|punk|"
    r"jazz|reggae|a cappella|band|female|male|orchestral) cover\b|\bcover by\b|\bcovered by\b|\bplaythrough\b|"
    r"\bplay ?along\b|\bbacking track\b|\bisolated (?:vocals?|guitar|bass|drums|track)\b|\banaly[sz](?:is|es|ing)\b|"
    r"\bAI (?:cover|voice)\b|\bimpression\b|\bimpersonat\w*|\bparody\b", re.I)
STRONG_COMMENTARY = re.compile(
    r"\bnames (?:the |her |his |their )?(?:top )?\d+\b|"
    r"\b\d+ \w+ (?:she|he|they) (?:hates?|hated|resented|despised|can['’]?t stand|refused|regrets?)\b|"
    r"\btragic (?:story|life|death|end)\b|\bautopsy\b|\bnet worth\b|\bbefore (?:they|she|he) (?:were|was) famous\b|\bdisappeared\b|"
    r"\bfinally (?:admits|confirms|breaks|speaks|reveals|opens up|addresses)\b|"
    r"\bbreaks (?:her|his|their) silence\b", re.I)
TRAILER = re.compile(r"\btrailer\b|\bteaser\b|\bcommercial\b|\badvert\b|\bTV spot\b|\bTV ad\b", re.I)
QA = re.compile(
    r"\bq\s*(?:&|and|\+|n)\s*a\b|\bquestions? (?:and|&) answers?\b|\bfan questions?\b|\bask me anything\b|"
    r"\banswers? (?:fan |your |viewer )?questions\b|\btakes? (?:your |fan )?questions\b|\b\d+ questions\b", re.I)
PANEL = re.compile(r"\bpanel\b|\broundtable\b|\bround table\b|\bsymposium\b|\bforum\b|(?<!press )\bconference\b|"
                   r"\bkeynote\b|\bsummit\b|\bdebate\b", re.I)
WORKSHOP = re.compile(r"(?i:\bworkshops?\b|\bmaster ?class(?:es)?\b|\bclinic\b|\blecture\b|\bseminar\b|"
                      r"\btalk at\b|\bsongwriting (?:session|class|camp|tips|advice|workshop)\b|\bin residence\b)|"
                      r"\bTEDx?\b")
DOCUMENTARY = re.compile(
    r"\bdocumentar(?:y|ies|io)\b|\bdocu(?:series|drama|mentaire)?\b|\brockumentary\b|behind the music|"
    r"\b(?:the )?story of\b|(?<!on the )\bmaking of\b|\bretrospective\b|\bbiography\b|\bclassic albums\b|"
    r"\bbehind the scenes\b|\bportrait of\b|\bin (?:their|his|her) own words\b|\bunsung\b|true hollywood story|"
    r"\bVH-?1\b.{0,30}\bwhere are they now\b|\bwhere are they now\b.{0,30}\bVH-?1\b", re.I)
INTERVIEW = re.compile(
    r"\binterview(?:s|ed|ing)?\b|\binterviu|\bentrevista|\bintervista|\bgespr[äa]ch|\bintervju|\bentretien|"
    r"\btalks?\b|\btalking\b|\bchats?\b|\bchatting\b|\bconversation\b|\bspeaks\b|\bspeaking\b|\bsays\b|"
    r"\bdiscuss(?:es|ing|ion)?\b|\bbackstage\b|\bpress conference\b|\bmeets\b|\bmeet the\b|\bremembers\b|"
    r"\bremembering\b|\brecalls\b|\brecalling\b|\breflects?\b|\breflecting\b|\bexplains?\b|\breveals?\b|\btells\b|"
    r"\bresponds\b|\bopens up\b|\bsits? down\b|\bsit-down\b|\bone[- ]on[- ]one\b|\bpodcast\b|\bstory behind\b|"
    r"\bmeaning (?:of|behind)\b|\bon (?:writing|making|recording|the making of|the story of|the meaning of|how|why|"
    r"what)\b|\bhow (?:i|we|she|he|they) (?:wrote|made|recorded|met|got)\b|\bmessage (?:to|for|from)\b|"
    r"\bin their own words\b", re.I)
STRONG_INTERVIEW = re.compile(r"\binterview|\bconversation\b|\bpodcast\b|\bpress conference\b|"
                              r"\bq\s*(?:&|and)\s*a\b", re.I)
AWARD = re.compile(r"\bacceptance\b|\baccepts?\b|\bspeech\b|\bwins?\b|\bwinning\b|\bwon (?:the |a |an |best |grammy)|"
                   r"\binduct(?:s|ed|ion|ee|ees)?\b|\bhonou?red\b|\blifetime achievement\b", re.I)
TITLE_KEPT = [("Q&A", QA), ("Panel discussion", PANEL), ("Workshop or masterclass", WORKSHOP),
              ("Documentary", DOCUMENTARY), ("Interview", INTERVIEW), ("Award or speech", AWARD)]
WEAK_THIRD = re.compile(r"\bcovers?\b|\btribute\b|\bfan ?made\b|\bfan video\b|\bslideshow\b|\bphotos?\b|\bpictures\b|"
                        r"\bmash ?up\b|\bremake\b|\bcompilation\b", re.I)
WEAK_COMMENTARY = re.compile(
    r"\btragedy\b|\btragic\b|\bnever told\b|\buntold\b|\bthe truth about\b|\bwhat(?:ever)? happened to\b|"
    r"\brise and fall\b|\bdark side\b|\bshocking\b|\bdeath of\b|\blast days\b|\btop \d+\b|"
    r"\b\d+ (?:bands|songs|men|women|singers|artists|things|facts|reasons|secrets|times)\b|\bhates?\b|\bresented\b|"
    r"\bsecrets?\b|\bexposed\b|\bscandal\b|\bheartbreaking\b|\bsad (?:story|life|end)\b|"
    r"\blife (?:and (?:death|times) )?of\b|\bfacts\b|\branked\b|\branking\b|\bexplained\b|\bvideo essay\b|"
    r"\bremember\b|\bforgot(?:ten)?\b|\bwhy (?:did|does|do|is|was|were)\b|\bhow (?:did|does|do|is|was|were)\b|"
    r"\bwhere are they now\b|\bthen and now\b|\bdies\b|\bdied\b|\bdead at\b|\bpass(?:es|ed) away\b|\bfuneral\b|"
    r"\bobituary\b|\bR\.?I\.?P\b", re.I)
PERFORMANCE = re.compile(
    r"official (?:music )?video|music video|lyric video|\bvideo ?clip\b|\bclip officiel\b|\blyrics?\b|\bfull album\b|"
    r"\binstrumental\b|\bremix\b|\b(?:extended|12\"?|7\"?|single|radio|club|dub) (?:mix|version|edit)\b|"
    r"\blive (?:at|in|from|on|performance|version|concert|session|show|recording)\b|[\(\[]live[\)\]]|"
    r"[-–—|]\s*live\b|\blive\s*[-–—|]|\blive!|(?-i:\bLIVE\b)|\bin concert\b|\bconcert\b|\bperforms?\b|\bperformance\b|\bperforming\b|\bremaster(?:ed)?\b|"
    r"\bofficial audio\b|\baudio\b|\bunplugged\b|\bacoustic\b|\bmedley\b|\bsings?\b|\bsinging\b|\bduet\b|"
    r"\blip ?sync\b|\bplayback\b|\bmimed?\b|\brehearsal\b|\bsound ?check\b|\bpromo\b|\bofficial\b|\bgig\b|"
    r"\bfull show\b|\btop of the pops\b|\bTOTP\b|\bHQ\b|\bexcerpt\b|\bdemos?\b|\bunreleased\b|\bouttakes?\b|"
    r"\btapes\b|\bsnippet\b|\bsessions?\b", re.I)
# A music artefact in the title outweighs a weak talk word ("Simon Says (Gilroy Tapes excerpt)")
MUSIC_ARTIFACT = re.compile(r"official (?:music )?video|music video|lyric video|\blyrics?\b|official audio|"
                            r"full album|\bexcerpt\b|\btapes\b|\bremix\b|\bremaster(?:ed)?\b|\bsnippet\b", re.I)
TOPIC_OR_VEVO = re.compile(r"\s-\sTopic$|VEVO", re.I)
TALK_SHOWS = re.compile(
    r"\bWogan\b|Letterman|Johnny Carson|Tonight Show|\bCarson\b|Arsenio|Oprah|Donahue|Joan Rivers|Parkinson|Aspel|"
    r"Russell Harty|Des O['’]Connor|Pebble Mill|Late Show|Late Night|Today Show|TODAY Show|NBC TODAY|"
    r"Good Morning America|Entertainment Tonight|ET Canada|Access Hollywood|Inside Edition|E! News|MTV News|"
    r"Larry King|Charlie Rose|Graham Norton|Jonathan Ross|Loose Women|\bThe View\b|Sunday Morning|CBS Mornings|"
    r"CBS News|NBC News|ABC News|BBC News|BBC Breakfast|Sky News|\bCNN\b|Bloomberg|60 Minutes|20/20|Nightline|"
    r"Breakfast Time|\bTV-?am\b|\bGMTV\b|Good Morning Britain|This Morning|Lorraine|Richard (?:&|and) Judy|\bRove\b|"
    r"Studio 10|Today Extra|Midday Show|Ray Martin|Mike Walsh|Breakfast Television|Breakfast TV|CONAN|"
    r"Conan O['’]Brien|Team Coco|Howard Stern|Stern Show|Kelly Clarkson Show|Tamron Hall|Drew Barrymore|"
    r"BUILD Series|Rockonteurs|Sodajerker|Q with Tom Power|Studio Q|Absolute Radio|Virgin Radio|Celebs Up Close|"
    r"AM Northwest|Professor of Rock")
MUSIC_SHOWS = re.compile(
    r"Top of the Pops|\bTOTP\b|Solid Gold|American Bandstand|\bBandstand\b|Soul Train|Night Flight|\bCountdown\b|"
    r"\bThe Tube\b|Whistle Test|Razzmatazz|Tiswas|Rockpalast|Beat-?Club|Formel Eins|Musikladen|Friday Night Videos|"
    r"Hot Tracks|Dance Party USA|Video Hits|\bTop ?Pop\b|Saturday Night Live|\bSNL\b|Jools Holland|\bMTV\b|"
    r"\bVH-?1\b|Much ?Music|Live Aid|Band Aid|Prince['’]s Trust|\bGrammys?\b|Brit Awards|\bBRITs\b|BPI Awards|"
    r"American Music Awards|\bAMAs?\b|Video Music Awards|\bVMAs?\b|Juno Awards|ARIA Awards|Hall of Fame")
DESC_PERF = re.compile(
    r"music video by|\bperform(?:s|ing|ed|ances?)\b|\blyrics\b|\blive (?:at|in|from|on)\b|\bofficial (?:music )?video\b|"
    r"\bin concert\b|\bfull concert\b|\bset ?list\b|\btrack ?list\b|\bspotify\b|\bapple music\b|"
    r"\bstream (?:it|now|the)\b|provided to youtube by|auto-generated by youtube", re.I)
DESC_KEPT = [
    ("Q&A", QA), ("Panel discussion", PANEL), ("Workshop or masterclass", WORKSHOP),
    ("Documentary", re.compile(r"\bdocumentar(?:y|ies)\b|\brockumentary\b|behind the music", re.I)),
    ("Interview", re.compile(r"\binterview(?:s|ed|ing)?\b|\bconversation\b|\btalks?\s+(?:to|with|about)\b|"
                             r"\bchats?\s+(?:to|with|about)\b|\bspeaks?\s+(?:to|with|about)\b|\bsits? down with\b|"
                             r"\bjoins (?:\w+ ){1,3}(?:to (?:talk|discuss|chat)|for (?:a|an) (?:chat|interview|"
                             r"conversation|talk))\b|\bdiscuss(?:es|ing|ed)?\b|\bpress conference\b|\bpodcast\b", re.I)),
]
YEAR = re.compile(r"(?<!\d)(?<![Tt]op )(19[7-9]\d|20[0-2]\d)(?!\d)")
REC_YEAR = re.compile(r"(?:recorded|filmed|taped|broadcast|aired|televised|interviewed|interview (?:from|recorded|"
                      r"conducted|taken|was)|originally (?:shown|aired|broadcast)|shown on|first shown)\b[^.\n]{0,60}?"
                      r"(?<!\d)(19[7-9]\d|20[0-2]\d)(?!\d)", re.I)
YT_ID = re.compile(r"(?:youtube(?:-nocookie)?\.com/(?:watch\?(?:.*&)?v=|embed/|shorts/|v/|live/)|youtu\.be/)"
                   r"([A-Za-z0-9_-]{11})")


# ═════════════════════════════════════════════════════════════════════════════
# NAME AND SONG MATCHING
# ═════════════════════════════════════════════════════════════════════════════
def _name_pattern(n):
    p = re.escape(n).replace("\\'", "['’]").replace("'", "['’]").replace("\\ ", r"\s+")
    return p


def _compile_names(names):
    names = sorted(names, key=len, reverse=True)
    if not names:
        return None
    return re.compile(r"(?<!\w)(?<!\w-)(?:" + "|".join(_name_pattern(n) for n in names) + r")(?!\w)(?!-\w)", re.I)


def _key(s):
    return re.sub(r"\s+", " ", s.replace("’", "'")).strip().lower()


NAMES = {}
for _a, _v in ACTS.items():
    _lookup = {_key(n): (m, era, d) for n, m, era, d in _v["names"]}
    NAMES[_a] = (_compile_names([n for n, _, _, _ in _v["names"]]), _lookup)


def people_named(act, text):
    """Returns (nomination-era members named, of which named distinctively, other-era members named)."""
    rx, lookup = NAMES[act]
    nom, dist, oth = [], set(), []
    if rx is None:
        return nom, dist, oth
    for m in rx.finditer(text):
        member, era, d = lookup.get(_key(m.group(0)), (None, None, False))
        if member is None:
            continue
        if era == "nomination":
            if member not in nom:
                nom.append(member)
            if d:
                dist.add(member)
        elif member not in oth:
            oth.append(member)
    return nom, dist, oth


def _song_rx(song):
    toks = re.findall(r"[A-Za-z0-9]+", song)
    if len(toks) == 1:
        sep = r"[-–—:|\"“”‘’'«»()\[\]/]"
        return re.compile(r"(?:^|" + sep + r")\s*" + toks[0] + r"\s*(?=$|" + sep + r")", re.I)
    return re.compile(r"(?<![A-Za-z0-9])" + r"[\W_]*".join(toks) + r"(?![A-Za-z0-9])", re.I)


SONG_RX = {a: [(s, _song_rx(s)) for s in sorted(v, key=len, reverse=True)] for a, v in SONGS.items()}


def strip_songs(act, text):
    """Removes the act's song titles; returns (text without them, songs found)."""
    found = []
    for s, rx in SONG_RX.get(act, []):
        if rx.search(text):
            found.append(s)
            text = rx.sub(" | ", text)
    return text, found


# ═════════════════════════════════════════════════════════════════════════════
# SMALL HELPERS
# ═════════════════════════════════════════════════════════════════════════════
class SearchAborted(Exception):
    """Stop loudly rather than finish with an empty or truncated search."""


def clean_text(v):
    if v is None:
        return ""
    if isinstance(v, list):
        v = " ".join(str(x) for x in v)
    v = html.unescape(str(v))
    v = re.sub(r"<[^>]+>", " ", v)
    return re.sub(r"\s+", " ", v).strip()


def seconds_from_clock(s):
    s = str(s or "").strip()
    if not re.fullmatch(r"\d{1,2}(?::\d{2}){1,2}", s):
        return None
    total = 0
    for part in s.split(":"):
        total = total * 60 + int(part)
    return total


def seconds_from_iso(s):
    m = re.fullmatch(r"P(?:(\d+)D)?T?(?:(\d+)H)?(?:(\d+)M)?(?:(\d+)S)?", str(s or ""))
    if not m or not any(m.groups()):
        return None
    d, h, mi, se = (int(x) if x else 0 for x in m.groups())
    return d * 86400 + h * 3600 + mi * 60 + se


REL_DATE = re.compile(r"(\d+)\s+(year|month|week|day|hour|minute)s?\s+ago", re.I)


def parse_upload_date(s, fetched_iso):
    """Serper gives 'Mar 23, 2014' or '3 years ago'. Returns (ISO date, precision)."""
    s = str(s or "").strip()
    if not s:
        return "", ""
    m = REL_DATE.search(s)
    if m:
        n, unit = int(m.group(1)), m.group(2).lower()
        days = {"year": 365, "month": 30, "week": 7, "day": 1, "hour": 0, "minute": 0}[unit] * n
        base = datetime.fromisoformat(fetched_iso[:19]) if fetched_iso else datetime.now()
        return (base - timedelta(days=days)).date().isoformat(), f"approximate (Serper: '{s}')"
    try:
        return pd.to_datetime(s).date().isoformat(), "Serper"
    except Exception:
        return "", ""


def video_key(url):
    """('yt:ID', canonical URL, ID) for YouTube videos; ('url:...', URL, '') for anything else."""
    m = YT_ID.search(url or "")
    if m:
        return "yt:" + m.group(1), "https://www.youtube.com/watch?v=" + m.group(1), m.group(1)
    u = re.sub(r"[?&](?:utm_[^=&]+|si|feature|ref)=[^&#]*", "", str(url or "")).split("#")[0].rstrip("/")
    return "url:" + u.lower(), u, ""


def get_key(name, required):
    """Environment variable, then Colab secret, then ask (unless VIDEO_LINKS_NO_PROMPT=1)."""
    v = os.environ.get(name, "").strip()
    if not v:
        try:
            from google.colab import userdata
            v = (userdata.get(name) or "").strip()
        except Exception:
            v = ""
    if not v and os.environ.get("VIDEO_LINKS_NO_PROMPT") != "1":
        try:
            import getpass
            v = getpass.getpass(f"{name} was not found in Colab secrets. Paste it here"
                                f"{'' if required else ' (or press Enter to skip)'}: ").strip()
        except Exception:
            v = ""
    return v


# ═════════════════════════════════════════════════════════════════════════════
# CACHE: every raw response is kept, with the time it was fetched
# ═════════════════════════════════════════════════════════════════════════════
def load_cache():
    if os.path.exists(CACHE):
        with open(CACHE, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_cache(cache):
    tmp = CACHE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False)
    os.replace(tmp, CACHE)


SERPER_API_KEY = ""
YT_API_KEY = ""
QUERY_LOG = []
STATS = {"serper_calls": 0, "yt_search_calls": 0, "yt_detail_calls": 0, "cached": 0, "empty_in_a_row": 0,
         "since_save": 0}
YT_STOP = {"reason": ""}       # set when YouTube refuses (quota, key, API not enabled); Serper carries on


def _maybe_save(cache):
    STATS["since_save"] += 1
    if STATS["since_save"] >= SAVE_EVERY:
        save_cache(cache)
        STATS["since_save"] = 0


def serper_videos(cache, query, page):
    key = f"serper|{query}|{page}"
    if key in cache:
        STATS["cached"] += 1
        hit = cache[key]
        QUERY_LOG.append([query, "Serper videos", page, len(hit["data"].get("videos", [])), hit["fetched"], "yes"])
        return hit
    time.sleep(random.uniform(*DELAY))
    last = None
    for attempt in range(4):
        try:
            r = requests.post("https://google.serper.dev/videos",
                              headers={"X-API-KEY": SERPER_API_KEY, "Content-Type": "application/json"},
                              json={"q": query, "num": 10, "page": page, "gl": "us", "hl": "en"}, timeout=25)
        except Exception as e:
            last = f"{type(e).__name__}: {e}"
            time.sleep(3 * (attempt + 1))
            continue
        if r.status_code in (401, 403):
            raise SearchAborted(f"Serper refused the key ({r.status_code}): {r.text[:200]}")
        if r.status_code in (402, 429):
            if attempt == 3:
                raise SearchAborted(f"Serper returned {r.status_code} (out of credits or rate limited). The cache "
                                    "is saved; run the cell again later and it resumes where it stopped.")
            time.sleep(20 * (attempt + 1))
            continue
        if r.status_code != 200:
            last = f"HTTP {r.status_code}: {r.text[:200]}"
            time.sleep(3 * (attempt + 1))
            continue
        data = r.json()
        hit = {"fetched": datetime.now().isoformat(timespec="seconds"), "data": data}
        cache[key] = hit
        STATS["serper_calls"] += 1
        n = len(data.get("videos", []))
        STATS["empty_in_a_row"] = 0 if n else STATS["empty_in_a_row"] + 1
        if STATS["empty_in_a_row"] >= 25:
            raise SearchAborted("25 empty Serper responses in a row: the search is failing, not the queries.")
        QUERY_LOG.append([query, "Serper videos", page, n, hit["fetched"], "no"])
        return hit
    print(f"    [serper] gave up on {query!r} page {page}: {last}")
    QUERY_LOG.append([query, "Serper videos", page, 0, "", f"failed: {last}"])
    return {"fetched": "", "data": {}}


def _yt_refusal(r):
    """Stops further YouTube calls this run, with a reason she can act on."""
    try:
        err = r.json().get("error", {})
        reason = (err.get("errors") or [{}])[0].get("reason", "") or err.get("status", "")
        msg = err.get("message", "")
    except Exception:
        reason, msg = "", r.text[:200]
    if reason in ("quotaExceeded", "dailyLimitExceeded", "rateLimitExceeded"):
        why = ("the YouTube quota for today is used up. Run the cell again tomorrow: everything found so far is "
               "cached, and the rest is filled in")
    elif reason in ("keyInvalid", "badRequest") or "API key not valid" in msg:
        why = "YouTube rejected the API key. Check YT_API_KEY in Colab secrets"
    elif reason in ("accessNotConfigured", "SERVICE_DISABLED") or "has not been used" in msg or "disabled" in msg:
        why = ("the YouTube Data API v3 is not switched on for the key's project. In Google Cloud Console, open "
               "APIs & Services > Library, find YouTube Data API v3 and click Enable")
    else:
        why = f"YouTube refused the request (HTTP {r.status_code}, {reason or 'no reason given'}: {msg[:150]})"
    if not YT_STOP["reason"]:
        YT_STOP["reason"] = why
        print(f"\n    [youtube] Stopping YouTube calls for this run: {why}.\n")


def yt_search(cache, q):
    key = f"ytsearch|{q}|{YT_RESULTS}"
    if key in cache:
        STATS["cached"] += 1
        hit = cache[key]
        QUERY_LOG.append([q, "YouTube search", 1, len(hit["data"].get("items", [])), hit["fetched"], "yes"])
        return hit
    if YT_STOP["reason"] or not YT_API_KEY:
        return None
    time.sleep(random.uniform(*DELAY))
    params = {"part": "snippet", "q": q, "type": "video", "maxResults": YT_RESULTS, "order": "relevance",
              "relevanceLanguage": "en", "safeSearch": "none", "key": YT_API_KEY}
    for attempt in range(3):
        try:
            r = requests.get("https://www.googleapis.com/youtube/v3/search", params=params, timeout=30)
        except Exception as e:
            print(f"    [youtube search] {type(e).__name__}: {e}")
            time.sleep(3 * (attempt + 1))
            continue
        if r.status_code == 200:
            hit = {"fetched": datetime.now().isoformat(timespec="seconds"), "data": r.json()}
            cache[key] = hit
            STATS["yt_search_calls"] += 1
            QUERY_LOG.append([q, "YouTube search", 1, len(hit["data"].get("items", [])), hit["fetched"], "no"])
            return hit
        if r.status_code >= 500:
            time.sleep(3 * (attempt + 1))
            continue
        _yt_refusal(r)
        return None
    QUERY_LOG.append([q, "YouTube search", 1, 0, "", "failed"])
    return None


def youtube_details(cache, ids):
    """videos.list, 50 IDs per call (1 quota unit each). An ID missing from the answer is removed or private."""
    todo = [i for i in dict.fromkeys(ids) if f"yt|{i}" not in cache]
    if todo and not YT_API_KEY:
        print("No YouTube key: skipping YouTube details (exact upload date, description, captions).")
        todo = []
    calls = 0
    for k in range(0, len(todo), 50):
        if YT_STOP["reason"]:
            break
        chunk = todo[k:k + 50]
        try:
            r = requests.get("https://www.googleapis.com/youtube/v3/videos", timeout=30, params={
                "part": "snippet,contentDetails,statistics,status", "id": ",".join(chunk), "key": YT_API_KEY})
        except Exception as e:
            print(f"    [youtube details] {type(e).__name__}: {e}")
            continue
        if r.status_code != 200:
            _yt_refusal(r)
            break
        fetched = datetime.now().isoformat(timespec="seconds")
        items = {it["id"]: it for it in r.json().get("items", [])}
        for i in chunk:
            cache[f"yt|{i}"] = {"fetched": fetched, "item": items.get(i)}
        calls += 1
        STATS["yt_detail_calls"] += 1
        if calls % 20 == 0:
            save_cache(cache)
        time.sleep(0.2)
    if calls:
        QUERY_LOG.append([f"details for {len(todo)} videos", "YouTube details", "", len(todo),
                          datetime.now().isoformat(timespec="seconds"), "no"])
    return {i: cache[f"yt|{i}"]["item"] for i in ids if f"yt|{i}" in cache}


# ═════════════════════════════════════════════════════════════════════════════
# QUERIES
# ═════════════════════════════════════════════════════════════════════════════
def serper_queries(act):
    v = ACTS[act]
    s = v["search"]
    qs = [f"{s} interview", f"{s} interview {v['year'] - 1}", f"{s} interview {v['year']}",
          f"{s} documentary", f"{s} Q&A OR panel OR workshop OR masterclass"]
    for m in v["members"]:
        qs.append(f'"{m}" "{v["member_search"]}" interview')
    for m in v["distinctive"]:
        qs += [f'"{m}" interview', f'"{m}" Q&A OR panel OR workshop OR masterclass OR documentary']
    qs += v["extra_queries"]
    return list(dict.fromkeys(qs))


def youtube_query(act):
    return f"{ACTS[act]['search']} interview"


# ═════════════════════════════════════════════════════════════════════════════
# COLLECT
# ═════════════════════════════════════════════════════════════════════════════
def _add(rows, act, link, title, snippet, channel, duration_s, upload_raw, fetched, found_by, source, rank,
         platform=""):
    key, url, yid = video_key(link)
    row = rows.setdefault((act, key), {
        "act": act, "key": key, "url": url, "video_id": yid, "platform": platform or ("YouTube" if yid else ""),
        "title": title, "snippet": snippet, "channel": channel, "duration_s": duration_s,
        "upload_raw": upload_raw, "fetched": fetched, "found_by": [], "sources": set(), "best_rank": None})
    if found_by not in row["found_by"]:
        row["found_by"].append(found_by)
    row["sources"].add(source)
    if not row["title"] and title:
        row["title"] = title
    if len(snippet or "") > len(row["snippet"] or ""):
        row["snippet"] = snippet
    if rank is not None:
        row["best_rank"] = rank if row["best_rank"] is None else min(row["best_rank"], rank)
    return row


def collect(acts, cache):
    rows = {}
    for n, act in enumerate(acts, 1):
        qs = serper_queries(act)
        print(f"[{n}/{len(acts)}] {act}: {len(qs)} Serper queries" + (", 1 YouTube search" if YOUTUBE_SEARCH else ""))
        for q in qs:
            for page in range(1, SERPER_PAGES + 1):
                hit = serper_videos(cache, q, page)
                vids = hit["data"].get("videos", [])
                for it in vids:
                    rank = (page - 1) * 10 + int(it.get("position", 0) or 0)
                    _add(rows, act, it.get("link", ""), clean_text(it.get("title")), clean_text(it.get("snippet")),
                         clean_text(it.get("channel")), seconds_from_clock(it.get("duration")),
                         clean_text(it.get("date")), hit["fetched"], q, "Serper", rank,
                         platform=clean_text(it.get("source")))
                if len(vids) < 10:
                    break
            _maybe_save(cache)
        if YOUTUBE_SEARCH:
            q = youtube_query(act)
            hit = yt_search(cache, q)
            if hit:
                for i, it in enumerate(hit["data"].get("items", []), 1):
                    vid = (it.get("id") or {}).get("videoId")
                    if not vid:
                        continue
                    sn = it.get("snippet", {})
                    _add(rows, act, f"https://www.youtube.com/watch?v={vid}", clean_text(sn.get("title")),
                         clean_text(sn.get("description")), clean_text(sn.get("channelTitle")), None,
                         (sn.get("publishedAt") or "")[:10], hit["fetched"], f"YouTube search: {q}",
                         "YouTube search", i)
            _maybe_save(cache)
    save_cache(cache)
    return rows


def about_act(act, text):
    """Is the text about this act? Returns (about, act named, members, other-era members, note)."""
    act_m = ACT_RX[act].search(text)
    nom, dist, oth = people_named(act, text)
    note = ""
    about = bool(act_m) or bool(dist) or len(nom) >= 2
    anchored = bool(NEEDS_ANCHOR[act].search(text)) if act in NEEDS_ANCHOR else True
    if act in NEEDS_ANCHOR and act_m and not nom and not anchored:
        about, note = False, f"'{act}' with no member, song or band reference"
    other = OTHER_SENSE[act].search(text) if act in OTHER_SENSE else None
    if about and other and not nom and not (act in NEEDS_ANCHOR and anchored):
        about, note = False, f"the name used in another sense ('{other.group(0)}')"
    if not about and nom and not note:
        note = f"names {', '.join(nom)} (a name others share) but not the act"
    return about, bool(act_m), nom, oth, note


def cross_list(rows, acts):
    """A video found for one act whose title names another act is listed under that act too."""
    seen = {}
    for (act, key), r in rows.items():
        if r["video_id"] and key not in seen:
            seen[key] = (act, r)
    added = 0
    for key, (act, r) in seen.items():
        for other in acts:
            if other == act or (other, key) in rows or not ACT_RX[other].search(r["title"]):
                continue
            ok, act_named, _, _, _ = about_act(other, r["title"])
            if ok and act_named:
                row = dict(r, act=other, found_by=[f"named in the title of a video found for {act}"],
                           sources={"cross-listed"}, best_rank=None)
                rows[(other, key)] = row
                added += 1
    return added


def enrich(rows, cache):
    ids = sorted({r["video_id"] for r in rows.values() if r["video_id"]})
    need = [i for i in ids if f"yt|{i}" not in cache]
    print(f"\nYouTube details for {len(ids)} videos ({len(need)} not yet cached, "
          f"{(len(need) + 49) // 50} API calls of 1 unit each)...")
    info = youtube_details(cache, ids)
    save_cache(cache)
    for r in rows.values():
        if not r["video_id"] or r["video_id"] not in info:
            continue
        it = info[r["video_id"]]
        if it is None:
            r["available"] = "no (removed or private)"
            continue
        sn, cd = it.get("snippet", {}), it.get("contentDetails", {})
        st, stt = it.get("statistics", {}), it.get("status", {})
        r["available"] = "yes"
        r["yt_published"] = (sn.get("publishedAt") or "")[:10]
        r["yt_description"] = clean_text(sn.get("description"))
        r["channel"] = clean_text(sn.get("channelTitle")) or r["channel"]
        r["title"] = clean_text(sn.get("title")) or r["title"]
        r["duration_s"] = seconds_from_iso(cd.get("duration")) or r["duration_s"]
        r["captions"] = "yes" if str(cd.get("caption")).lower() == "true" else "no"
        r["language"] = sn.get("defaultAudioLanguage") or sn.get("defaultLanguage") or ""
        r["views"] = int(st["viewCount"]) if str(st.get("viewCount", "")).isdigit() else None
        r["privacy"] = stt.get("privacyStatus", "")


# ═════════════════════════════════════════════════════════════════════════════
# CLASSIFY
# ═════════════════════════════════════════════════════════════════════════════
def _speaker_before(act, text):
    """'Rick Astley reacts to ...' is the artist talking; 'Vocal coach reacts to Rick Astley' is not."""
    nom, _, _ = people_named(act, text)
    return bool(ACT_RX[act].search(text) or nom) and not re.search(r"\b(?:coach|teacher|producer|fan|fans|"
                                                                     r"musician|singer|guitarist|drummer|reacts?)\b",
                                                                     ACT_RX[act].sub(" ", text), re.I)


WEAK_VERB = re.compile(r"\b(?:talks?|talking|speaks|speaking|says|reveals?|recalls|recalling|remembers|remembering|"
                       r"reflects?|reflecting|explains?|tells|responds|opens up|discuss(?:es|ing)?|chats?|chatting)\b",
                       re.I)


def _act_is_speaker(act, t, m):
    """For a weak talk word, the act or a member should be the subject ('Boy George talks ...'), or the one
    talked with ('Talks with Boy George'). 'Drake Talks About Sade' is someone else talking."""
    before, after = t[:m.start()], t[m.end():]
    if ACT_RX[act].search(before) or people_named(act, before)[0]:
        return True
    w = re.match(r"(?:\s+[\w'’]+){0,4}?\s+(?:with|to)\s+(.{0,60})", after, re.I)
    return bool(w and (ACT_RX[act].match(w.group(1)) or people_named(act, w.group(1)[:40])[0]))


SPEAKER_ON_EXCLUDE = re.compile(r"(?:stage|tour|drums|guitar|piano|bass|keyboards|vocals|sax|tv|television|the radio|"
                                r"air|set|camera|film|record|vinyl|cd)\b", re.I)


def _speaker_on(act, t, music):
    """'Christopher Cross on surviving COVID', 'Roy Hay about Boy George': the act or a member talking about
    something. Not when what follows is a music show ('Culture Club on Top of the Pops') or a stage. When the
    title also has a song or a performance word, only talk about the song itself counts ('Mark Knopfler on
    Sultans of Swing'), not 'Tracy Chapman on Letterman - Fast Car'."""
    rx, _ = NAMES[act]
    found = [m for m in rx.finditer(t)] if rx else []
    act_m = ACT_RX[act].search(t)
    if act_m:
        found.append(act_m)
    for m in found:
        nxt = re.match(r"\s*(?:\([^)]*\)|\[[^\]]*\])?\s*(?:['’]s\s+)?(?:(?:lead\s+)?(?:front-?man|front-?woman|"
                       r"singer|vocalist|guitarist|drummer|bassist|keyboardist|saxophonist|flautist|members?|"
                       r"founder|co-founder)\s+)?(?:on|about)\s+(.{0,40})", t[m.end():], re.I)
        if nxt and not SPEAKER_ON_EXCLUDE.match(nxt.group(1)) and not MUSIC_SHOWS.search(nxt.group(1)):
            lead = re.match(r"\s*((?:[\w'’]+\s+){0,3})\|", nxt.group(1))      # "on 80s Classic Down Under"
            if not music or (lead and not TALK_SHOWS.search(lead.group(1)) and not MUSIC_SHOWS.search(lead.group(1))):
                return True
    return False


def classify(act, title, desc, channel, video_id):
    """Returns (type, evidence, act named, members, other-era members, notes)."""
    notes = []
    if not video_id:
        return "Not a YouTube video", "", "", [], [], notes
    about, act_named, nom, oth, why = about_act(act, f"{title}\n{desc}")
    in_title = bool(ACT_RX[act].search(title))
    named = "in title" if in_title else ("in description" if act_named else "no")
    if not about:
        return ("Other-era member only" if (oth and not nom) else "Not about the act"), why, named, nom, oth, notes
    if oth and not nom:
        notes.append(f"names only other-era members ({', '.join(oth)})")

    def res(kind, evidence):
        return kind, evidence, named, nom, oth, notes

    t, songs = strip_songs(act, FROM_SOURCE.sub(" ", title))
    m = STRONG_THIRD.search(t)
    if m and m.group(0).lower().startswith("react") and _speaker_before(act, t[:m.start()]):
        return res("Interview", f"title: the act or a member '{m.group(0)}'")
    if m:
        return res("Cover, tribute, lesson or reaction", f"title: '{m.group(0)}'")
    for rx, kind in ((STRONG_COMMENTARY, "Commentary by others"), (TRAILER, "Trailer or advert")):
        m = rx.search(t)
        if m:
            return res(kind, f"title: '{m.group(0)}'")
    for kind, rx in TITLE_KEPT:
        m = rx.search(t)
        if m:
            art = MUSIC_ARTIFACT.search(t)
            if art and kind in ("Interview", "Documentary", "Award or speech") and not STRONG_INTERVIEW.search(t):
                return res("Performance or music video", f"title: '{art.group(0)}'")
            if kind == "Interview" and WEAK_VERB.fullmatch(m.group(0)) and not STRONG_INTERVIEW.search(t) \
                    and not _act_is_speaker(act, t, m):
                notes.append(f"'{m.group(0)}': someone other than the act may be speaking")
                return res("Unclear (check)", f"title: '{m.group(0)}'")
            pm = PERFORMANCE.search(t)
            if pm:
                notes.append(f"the title also has '{pm.group(0)}': may include a performance")
            return res(kind, f"title: '{m.group(0)}'")
    if _speaker_on(act, t, bool(songs) or bool(PERFORMANCE.search(t))):
        return res("Interview", "title: the act or a member 'on' or 'about' something")
    for rx, kind in ((WEAK_THIRD, "Cover, tribute, lesson or reaction"), (WEAK_COMMENTARY, "Commentary by others"),
                     (PERFORMANCE, "Performance or music video")):
        m = rx.search(t)
        if m:
            return res(kind, f"title: '{m.group(0)}'")
    if TOPIC_OR_VEVO.search(channel or ""):
        return res("Performance or music video", f"channel: '{channel}'")
    d, _ = strip_songs(act, desc or "")
    dk = next(((k, rx.search(d)) for k, rx in DESC_KEPT if rx.search(d)), None)
    dp = DESC_PERF.search(desc or "")
    if dk and not dp:
        if songs:
            notes.append(f"song in title ('{songs[0]}')")
        return res(dk[0], f"description: '{dk[1].group(0)}'")
    if songs:
        return res("Performance or music video", f"song in title: '{songs[0]}'")
    if dp and not dk:
        return res("Performance or music video", f"description: '{dp.group(0)}'")
    m = TALK_SHOWS.search(f"{t}\n{channel}")
    if m and not dp:
        return res("Interview", f"talk show: '{m.group(0)}' (check)")
    if dk and dp:
        notes.append(f"description has both '{dk[1].group(0)}' and '{dp.group(0)}'")
    m = TALK_SHOWS.search(desc or "")
    if m and not dp:
        return res("Interview", f"talk show in description: '{m.group(0)}' (check)")
    m = MUSIC_SHOWS.search(f"{t}\n{channel}\n{desc or ''}")
    if m:
        notes.append(f"music or award show ('{m.group(0)}'): may be a performance")
    return res("Unclear (check)", "")


def best_year(title, desc, upload_iso):
    """(year, where it came from). Title years first, then a year beside a recording word in the
    description, then the upload year, which is only the latest the video can date from."""
    up = int(upload_iso[:4]) if upload_iso and upload_iso[:4].isdigit() else None
    top = up or datetime.now().year
    ty = sorted({int(y) for y in YEAR.findall(title) if 1970 <= int(y) <= top})
    if ty:
        return ty[0], ("year in title" if len(ty) == 1 else "earliest of several years in title")
    dy = sorted({int(m.group(1)) for m in REC_YEAR.finditer(desc or "") if 1970 <= int(m.group(1)) <= top})
    if dy:
        return dy[0], "year beside a recording word in the description"
    if up:
        return up, "upload year (the latest it can date from)"
    return None, ""


COLUMNS = ["ID", "Act", "Grammy year", "Type", "Type evidence", "Title", "Channel", "URL", "Minutes", "Year (best guess)",
           "Year from", "Years from nomination", "Period", "Upload date", "Show or event", "Act named",
           "Members named", "Other-era members named", "Captions by uploader", "Language", "Views", "Still online",
           "Found by", "Search source", "Best rank", "Notes", "Description", "Your check", "YouTube ID"]


def build_table(rows):
    out = []
    for r in rows.values():
        act = r["act"]
        desc = r.get("yt_description") or r.get("snippet") or ""
        kind, evidence, named, nom, oth, notes = classify(act, r["title"], desc, r.get("channel", ""), r["video_id"])
        dur = r.get("duration_s")
        if dur is not None and dur <= 60:
            notes.append("a minute or less")
        if dur is not None and dur > 7200:
            notes.append("over two hours")
        if r.get("available", "").startswith("no"):
            notes.append("no longer online")
        if r.get("yt_published"):
            up, up_prec = r["yt_published"], "YouTube"
        elif "YouTube search" in r["sources"] and re.fullmatch(r"\d{4}-\d{2}-\d{2}", r["upload_raw"] or ""):
            up, up_prec = r["upload_raw"], "YouTube search"
        else:
            up, up_prec = parse_upload_date(r["upload_raw"], r["fetched"])
        if up_prec.startswith("approximate"):
            notes.append(f"upload date {up_prec}")
        year, year_from = best_year(strip_songs(act, r["title"])[0], desc, up)
        gy = ACTS[act]["year"]
        dy = (year - gy) if year else None
        if year and not year_from.startswith("upload"):
            period = "before" if dy < 0 else ("nomination year" if dy == 0 else "after")
        else:
            period = "unknown"
        shows = sorted({m.group(0) for rx in (TALK_SHOWS, MUSIC_SHOWS) for m in rx.finditer(f"{r['title']}\n{desc}")},
                       key=str.lower)
        src = sorted(r["sources"])
        out.append({
            "ID": "", "Act": act, "Grammy year": gy, "Type": kind, "Type evidence": evidence, "Title": r["title"],
            "Channel": r.get("channel", ""), "URL": r["url"],
            "Minutes": round(dur / 60, 1) if dur else None,
            "Year (best guess)": year, "Year from": year_from, "Years from nomination": dy, "Period": period,
            "Upload date": up or None, "Show or event": "; ".join(dict.fromkeys(shows)), "Act named": named,
            "Members named": "; ".join(nom), "Other-era members named": "; ".join(oth),
            "Captions by uploader": r.get("captions", ""), "Language": r.get("language", ""),
            "Views": r.get("views"), "Still online": r.get("available", ""),
            "Found by": " | ".join(r["found_by"]),
            "Search source": " and ".join(src) if src != ["Serper", "YouTube search"] else "both",
            "Best rank": r.get("best_rank"), "Notes": "; ".join(dict.fromkeys(notes)),
            "Description": desc[:DESC_CHARS], "Your check": "", "YouTube ID": r["video_id"],
            "_sheet": "Links" if kind in LINK_TYPES else "Left out"})
    df = pd.DataFrame(out, columns=COLUMNS + ["_sheet"])
    if df.empty:
        return df
    order = {a: i for i, a in enumerate(ACTS)}
    df["_o"] = df["Act"].map(order)
    df["_t"] = df["Type"].map({t: i for i, t in enumerate(TYPE_ORDER)})
    df["_y"] = df["Year (best guess)"].fillna(9999)
    df = df.sort_values(["_o", "_t", "_y", "Best rank"], na_position="last", kind="mergesort")
    return df.drop(columns=["_o", "_t", "_y"]).reset_index(drop=True)


def assign_ids(df):
    """V1, V2 ... down the Links sheet, then on through Left out; one ID per video (a video listed under two
    acts keeps the same ID)."""
    df = df.copy()
    ids, n = {}, 0
    for sheet in ("Links", "Left out"):
        for i in df.index[df["_sheet"] == sheet]:
            key = df.at[i, "YouTube ID"] or df.at[i, "URL"]
            if key not in ids:
                n += 1
                ids[key] = f"V{n}"
            df.at[i, "ID"] = ids[key]
    return df


# ═════════════════════════════════════════════════════════════════════════════
# OUTPUT
# ═════════════════════════════════════════════════════════════════════════════
EXPLAIN = {
    "Act": "The Best New Artist nominee.",
    "Grammy year": "The year of the act's Grammy ceremony.",
    "Type": "Interview, Q&A, Panel discussion, Workshop or masterclass, Documentary, Award or speech, or Unclear "
            "(check). On Left out: Performance or music video; Cover, tribute, lesson or reaction; Commentary by "
            "others (listicles and video essays about the act); Trailer or advert; Not about the act; Other-era "
            "member only; Not a YouTube video.",
    "Type evidence": "The words that decided the type, and where they were (title, channel or description).",
    "Title": "The video's title on YouTube.", "Channel": "The YouTube channel that uploaded it.",
    "URL": "The video's page.", "Minutes": "Length in minutes.",
    "Year (best guess)": "When the video was probably recorded (see Year from).",
    "Year from": "'year in title'; 'year beside a recording word in the description' (for example 'recorded in "
                 "1985'); or 'upload year', which is only the latest the video can date from, since old TV "
                 "interviews are often uploaded decades later.",
    "Years from nomination": "Year (best guess) minus the Grammy year: negative before it, 0 in it, positive after.",
    "Period": "'before', 'nomination year' or 'after', only when the year comes from the title or description; "
              "otherwise 'unknown'.",
    "Upload date": "When the video was uploaded to YouTube. This is not the recording date.",
    "Show or event": "TV shows, award shows or events named in the title or description (for example Wogan, MTV, "
                     "Live Aid). A useful clue to the date.",
    "Act named": "Whether the act's name is in the title, only in the description, or not at all (the video names "
                 "a member instead, for example a solo-career interview).",
    "Members named": "Nomination-era members named in the title or description.",
    "Other-era members named": "Members from other eras named in the title or description.",
    "Captions by uploader": "Whether the uploader added captions. YouTube's automatic captions are not shown here.",
    "Language": "The language set by the uploader, when set.",
    "Views": "Views when the details were fetched.", "Still online": "Whether the video was still online.",
    "Found by": "The searches that found the video.",
    "Search source": "Serper (Google Videos), YouTube search, both, or cross-listed (the title names this act "
                     "although it was found for another).",
    "Best rank": "The video's best position in any search that found it (1 = top).",
    "Notes": "Anything worth checking.",
    "Description": f"The first {DESC_CHARS:,} characters of the video's description.",
    "Your check": "Blank, for notes from a check by hand.", "YouTube ID": "YouTube's own ID for the video.",
    "ID": "Our ID for the video (V1, V2 ...): numbered down the Links sheet, then on through Left out. A video "
          "listed under two acts keeps one ID. Transcripts and turns use it.",
}


def write_outputs(df, acts):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    links = df[df["_sheet"] == "Links"].drop(columns=["_sheet"])
    left = df[df["_sheet"] == "Left out"].drop(columns=["_sheet"])

    # By act
    rows = []
    for a in acts:
        L = links[links["Act"] == a]
        dated = L[L["Period"] != "unknown"]
        row = {"Act": a, "Grammy year": ACTS[a]["year"], "Type": ACTS[a]["type"], "Links": len(L)}
        for t in LINK_TYPES:
            row[t] = int((L["Type"] == t).sum())
        row["Hours (links)"] = round(L["Minutes"].fillna(0).sum() / 60, 1)
        row["Dated by title or description"] = len(dated)
        row["Before"] = int((dated["Period"] == "before").sum())
        row["Nomination year"] = int((dated["Period"] == "nomination year").sum())
        row["After"] = int((dated["Period"] == "after").sum())
        row["Left out"] = int((left["Act"] == a).sum())
        rows.append(row)
    by_act = pd.DataFrame(rows)

    counts = Counter(links["Type"])
    readme = [
        ["Video links", VERSION],
        ["What this is", "YouTube videos in which the acts may speak, for the transcription step. Nothing has been "
                         "downloaded. Links holds the interviews, Q&A, panel discussions, workshops, documentaries, "
                         "award speeches and unclear videos; Left out holds the rest, so no search result is lost."],
        ["How it searched", "As Harrison et al. (2023, p. 104) did for video, by the name of each act and of each "
                            "nomination-era member: Serper video search (Google Videos) with interview, documentary "
                            "and Q&A, panel, workshop or masterclass queries, plus YouTube's own search once per act. "
                            "The Query log sheet lists every search."],
        ["How types are decided", "From the title first, with the act's song titles removed so that a song such as "
                                  "'Talk of the Town' is not read as a talk. Reactions, lessons and covers by others, "
                                  "listicles and trailers come first; then Q&A, panel, workshop, documentary, "
                                  "interview, and award or speech words, or a title such as 'Christopher Cross on "
                                  "surviving COVID'; then tributes, video essays and performance words, and '- Topic' "
                                  "and VEVO channels. If the title settles nothing, the description decides, then the "
                                  "act's song titles (a performance), then a chat show named in the title (an "
                                  "interview, check). Otherwise the video is Unclear (check)."],
        ["Names", "An ordinary-word name such as Asia, Breakfast Club or Big Country counts only with a member, a "
                  "song or a band reference beside it. A member whose name others share counts only with the act's "
                  "name or a second member."],
        ["Dates", "Upload date is not the recording date. Year (best guess) uses a year in the title, then a year "
                  "beside a recording word in the description, then the upload year, which is only the latest "
                  "possible year."],
        ["Links", f"{len(links):,} rows ({links['URL'].nunique():,} distinct videos); "
                  + ", ".join(f"{t} {counts.get(t, 0):,}" for t in LINK_TYPES)
                  + f"; {links['Minutes'].fillna(0).sum() / 60:,.0f} hours in all."],
        ["Left out", f"{len(left):,} rows: " + ", ".join(f"{t} {int((left['Type'] == t).sum()):,}" for t in LEFT_TYPES)
                     + "."],
        ["Rows", "One row per video per act. A video that names two acts (a Live Aid backstage reel, say) is listed "
                 "under both."],
        ["IDs", f"Each video has our own ID, V1 to V{df['ID'].str[1:].astype(int).max()}: numbered down the Links "
                f"sheet, then on through Left out. A video listed under two acts keeps one ID."],
    ]
    if YT_STOP["reason"]:
        readme.append(["YouTube", f"Incomplete: {YT_STOP['reason']}."])
    readme.append(["", ""])
    readme.append(["Column", "Meaning"])
    readme += [[c, EXPLAIN[c]] for c in COLUMNS]

    payload = {"about": {"version": VERSION, "created": datetime.now().isoformat(timespec="seconds"),
                         "acts": len(acts), "links": len(links), "left_out": len(left),
                         "youtube_incomplete": YT_STOP["reason"]},
               "columns": COLUMNS + ["Sheet"],
               "data": json.loads(df.rename(columns={"_sheet": "Sheet"}).to_json(orient="records", force_ascii=False,
                                                                                 date_format="iso"))}
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)

    def as_dates(d):
        d = d.copy()
        d["Upload date"] = pd.to_datetime(d["Upload date"], errors="coerce")
        return d

    head = PatternFill("solid", fgColor="1F4E5F")
    widths = {"Act": 18, "Grammy year": 8, "Type": 16, "Type evidence": 24, "Title": 50, "Channel": 22, "URL": 34,
              "Minutes": 8, "Year (best guess)": 9, "Year from": 18, "Years from nomination": 10, "Period": 11,
              "Upload date": 11, "Show or event": 16, "Act named": 11, "Members named": 22,
              "Other-era members named": 16, "Captions by uploader": 9, "Language": 8, "Views": 10,
              "Still online": 9, "Found by": 34, "Search source": 12, "Best rank": 7, "Notes": 30,
              "Description": 60, "Your check": 20, "YouTube ID": 13, "ID": 8}
    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as xw:
        pd.DataFrame(readme).to_excel(xw, sheet_name="Read me", index=False, header=False)
        as_dates(links).to_excel(xw, sheet_name="Links", index=False)
        as_dates(left).to_excel(xw, sheet_name="Left out", index=False)
        by_act.to_excel(xw, sheet_name="By act", index=False)
        pd.DataFrame(QUERY_LOG, columns=["Query", "Source", "Page", "Results", "Fetched", "From cache"]) \
            .to_excel(xw, sheet_name="Query log", index=False)
        for name, ws in xw.sheets.items():
            if name == "Read me":
                ws.column_dimensions["A"].width = 24
                ws.column_dimensions["B"].width = 120
                for row in ws.iter_rows():
                    for c in row:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                    row[0].font = Font(bold=True)
                continue
            for j, c in enumerate(ws[1], start=1):
                c.fill, c.font = head, Font(bold=True, color="FFFFFF")
                c.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(j)].width = widths.get(c.value, 14)
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            if name in ("Links", "Left out"):
                heads = [c.value for c in ws[1]]
                col = {h: heads.index(h) + 1 for h in heads}
                wrap = {col[h] for h in ("Title", "Description", "Notes", "Found by", "Type evidence")}
                for row in ws.iter_rows(min_row=2):
                    for c in row:
                        c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap)
                    u = row[col["URL"] - 1]
                    if isinstance(u.value, str) and u.value.startswith("http"):
                        u.hyperlink, u.font = u.value, Font(color="1155CC", underline="single")
                    row[col["Upload date"] - 1].number_format = "yyyy-mm-dd"
                    row[col["Minutes"] - 1].number_format = "0.0"
    return links, left, by_act


# ═════════════════════════════════════════════════════════════════════════════
# MAIN
# ═════════════════════════════════════════════════════════════════════════════
def main():
    global SERPER_API_KEY, YT_API_KEY
    acts = [a for a in (PILOT_ACTS or list(ACTS)) if a in ACTS]
    missing = [a for a in (PILOT_ACTS or []) if a not in ACTS]
    if missing:
        raise ValueError(f"Not in the act list: {missing}. The names are: {list(ACTS)}")
    if not os.path.isdir(DIR):
        raise FileNotFoundError(f"{DIR} not found. Mount Drive first: from google.colab import drive; "
                                f"drive.mount('/content/drive')")
    YT_STOP["reason"] = ""
    QUERY_LOG.clear()
    STATS.update({k: 0 for k in STATS})
    SERPER_API_KEY = get_key("SERPER_API_KEY", required=True)
    YT_API_KEY = get_key("YT_API_KEY", required=False)
    if not SERPER_API_KEY:
        raise SearchAborted("No SERPER_API_KEY: add it to Colab secrets (key icon on the left) and run again.")
    if not YT_API_KEY:
        print("No YT_API_KEY: running with Serper only (no YouTube search, no exact dates or descriptions).")
    n_q = sum(len(serper_queries(a)) for a in acts)
    yt_on = YOUTUBE_SEARCH and bool(YT_API_KEY)
    print(f"{VERSION}\n{len(acts)} acts: {n_q} Serper queries (up to {n_q * SERPER_PAGES} calls)"
          + (f", {len(acts)} YouTube searches ({len(acts) * 100:,} quota units)" if yt_on else "")
          + ". Cached searches cost nothing.\n")
    cache = load_cache()
    try:
        rows = collect(acts, cache)
    except SearchAborted as e:
        save_cache(cache)
        print(f"\nSTOPPED: {e}")
        return None
    except KeyboardInterrupt:
        save_cache(cache)
        print("\nInterrupted. The cache is saved; run the cell again to resume.")
        return None
    added = cross_list(rows, acts)
    enrich(rows, cache)
    df = build_table(rows)
    if not df.empty:
        df = assign_ids(df)
    if df.empty:
        print("No videos found.")
        return df
    links, left, by_act = write_outputs(df, acts)
    print(f"\nSerper calls: {STATS['serper_calls']}, YouTube searches: {STATS['yt_search_calls']}, "
          f"YouTube detail calls: {STATS['yt_detail_calls']}, answered from the cache: {STATS['cached']}")
    print(f"{added} rows added because a title named another act.")
    print(f"\nLinks: {len(links):,} rows, {links['URL'].nunique():,} distinct videos, "
          f"{links['Minutes'].fillna(0).sum() / 60:,.0f} hours")
    print(links["Type"].value_counts().reindex(LINK_TYPES).fillna(0).astype(int).to_string())
    print(f"\nLeft out: {len(left):,} rows")
    print(left["Type"].value_counts().reindex(LEFT_TYPES).fillna(0).astype(int).to_string())
    few = by_act[by_act["Links"] < 5]["Act"].tolist()
    if few:
        print(f"\nActs with fewer than 5 links: {', '.join(few)}")
    if YT_STOP["reason"]:
        print(f"\nNOTE: YouTube is incomplete: {YT_STOP['reason']}.")
    print(f"\nWritten: {OUT_XLSX}\n         {OUT_JSON}")
    return df


if __name__ == "__main__":
    df = main()
